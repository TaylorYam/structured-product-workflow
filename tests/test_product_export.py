"""匯出資料庫的商品條件（#122）：`workflow.product` 整張寫成 Excel；商品總表（#291）。資料庫測試對真的 PostgreSQL；資料全部虛構。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from order_synth import CONFIG, A, B, order_row
from sp_workflow.lifecycle import KO, AutoEvent, replace_auto_events
from sp_workflow.product_export import (
    SUMMARY_COLUMNS,
    export_name,
    export_products,
    export_summary,
    summary_name,
    trade_columns,
)
from sp_workflow.product_terms import write_batch
from sp_workflow.window_actions import _drafts

UF_215 = {"Upfront / NotePrice (%)": Decimal("97.85")}


def import_order(pg, tdcc: str) -> None:
    receipt = write_batch(pg, _drafts([order_row(tdcc, UF_215)], CONFIG, pg), CONFIG, user="tester")
    assert receipt.lines[0].imported, receipt.lines[0].detail


def product_columns(pg) -> list[str]:
    rows = pg.execute(
        "SELECT column_name FROM information_schema.columns"
        " WHERE table_schema = 'workflow' AND table_name = 'product' ORDER BY ordinal_position"
    ).fetchall()
    return [name for (name,) in rows]


def read_back(path: Path) -> list[list[object]]:
    sheet = openpyxl.load_workbook(path).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_default_name_has_the_date_and_time() -> None:
    assert export_name(dt.datetime(2026, 10, 22, 15, 4, 5)) == "workflow_product_20261022_150405.xlsx"


@pytest.mark.db
def test_every_product_and_column_is_exported(pg, tmp_path: Path) -> None:
    import_order(pg, B)
    import_order(pg, A)
    pg.execute("UPDATE workflow.product SET locked_at = now() WHERE tdcc = %s", (A,))
    path = tmp_path / "商品條件.xlsx"

    receipt = export_products(pg, path)

    rows = read_back(path)
    columns = [*product_columns(pg), "存續狀態"]
    assert rows[0] == columns
    assert not [c for c in columns if c.startswith("status_")], "舊狀態欄已移除（#172），只有存續狀態"
    assert receipt.count == 2 and len(rows) == 3
    first = dict(zip(columns, rows[1], strict=True))
    assert [row[0] for row in rows[1:]] == [A, B], "依 TDCC 排序，前導零保留"
    stored = pg.execute("SELECT issuer, currency FROM workflow.product WHERE tdcc = %s", (A,)).fetchone()
    assert (first["issuer"], first["currency"]) == stored
    assert isinstance(first["trade_date"], dt.datetime), "日期照日期寫"
    assert isinstance(first["coupon_pct"], int | float), "數字照數字寫"
    assert isinstance(first["locked_at"], dt.datetime) and first["locked_at"].tzinfo is None, "時間換成本機時間"
    assert list(tmp_path.iterdir()) == [path], "不留暫存檔"


@pytest.mark.db
def test_an_empty_table_still_writes_the_header(pg, tmp_path: Path) -> None:
    receipt = export_products(pg, tmp_path / "x.xlsx")

    assert receipt.count == 0 and read_back(tmp_path / "x.xlsx") == [[*product_columns(pg), "存續狀態"]]


@pytest.mark.db
def test_lifecycle_status_is_exported(pg, tmp_path: Path) -> None:
    import_order(pg, A)
    import_order(pg, B)
    replace_auto_events(pg, A, [AutoEvent(KO, dt.date(2026, 11, 2))], user="tester")
    path = tmp_path / "商品條件.xlsx"

    export_products(pg, path)

    rows = read_back(path)
    assert [(row[0], row[-1]) for row in rows[1:]] == [(A, "KO"), (B, "在庫")]


# ---------------------------------------------------------------- 商品總表（#291）


def add_trades(pg, tdcc: str, rows: list[tuple[str, str, int]]) -> None:
    (run_id,) = pg.execute(
        "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES ('每日更新.xlsx', 'x', '{}')"
        " RETURNING run_id"
    ).fetchone()
    for seq, (branch, fc, amount) in enumerate(rows, start=1):
        pg.execute(
            "INSERT INTO workflow.trade_row (tdcc, row_seq, trade_date, product, currency, amount, fc_name, branch,"
            " first_run_id) VALUES (%s, %s, '2026-10-15', 'FCN', 'USD', %s, %s, %s, %s)",
            (tdcc, seq, amount, fc, branch, run_id),
        )


def as_dicts(path: Path) -> list[dict[str, object]]:
    header, *rows = read_back(path)
    return [dict(zip(header, row, strict=True)) for row in rows]


def test_summary_name_has_the_date_and_time() -> None:
    assert summary_name(dt.datetime(2026, 10, 27, 9, 8, 7)) == "商品總表_20261027_090807.xlsx"


@pytest.mark.db
def test_an_empty_summary_still_writes_the_header_with_two_trade_groups(pg, tmp_path: Path) -> None:
    receipt = export_summary(pg, tmp_path / "x.xlsx")

    assert receipt.count == 0
    assert read_back(tmp_path / "x.xlsx") == [[*SUMMARY_COLUMNS, *trade_columns(2)]]


@pytest.mark.db
def test_summary_rows_follow_the_new_summary_columns(pg, tmp_path: Path) -> None:
    import_order(pg, B)
    import_order(pg, A)
    pg.execute("INSERT INTO workflow.check_terms (tdcc, status) VALUES (%s, '完成')", (A,))
    pg.execute(
        "INSERT INTO workflow.underlying_price (tdcc, seq, code, initial_price, strike_price, ki_price, ko_price)"
        " VALUES (%s, 1, 'AAA UW', 100, 80, 65, 100)",
        (A,),
    )
    pg.execute(
        "INSERT INTO workflow.autocall_date (tdcc, slot, value) VALUES (%s, 1, '2026-11-16'), (%s, 2, NULL)", (A, A)
    )
    add_trades(pg, A, [("台北", "Amy", 100000), ("台中", "PB", 50000), ("新竹", "Ben", 30000)])
    replace_auto_events(pg, A, [AutoEvent(KO, dt.date(2026, 11, 16))], user="tester")
    path = tmp_path / "商品總表.xlsx"

    receipt = export_summary(pg, path)

    a, b = as_dicts(path)
    assert receipt.count == 2 and (a["TDCC Code"], b["TDCC Code"]) == (A, B), "依 TDCC 排序，前導零保留"
    assert list(a) == [*SUMMARY_COLUMNS, *trade_columns(3)], "交易組數是最多那一檔的筆數"
    assert (a["存續狀態"], a["存續事件日"], a["事件標的"]) == ("KO", dt.datetime(2026, 11, 16), "-")
    assert (b["存續狀態"], b["存續事件日"]) == ("在庫", "-")
    assert (a["發行機構"], a["承作幣別"], a["K(%)"], a["KI(Freq)"]) == ("Barclays", "USD", 80, "AM")
    assert (a["承作總額"], b["承作總額"]) == (180000, "-")
    assert [a[f"{n}_3"] for n in ("分公司", "FC", "金額")] == ["新竹", "Ben", 30000], "照交易明細順序"
    assert {b[c] for c in trade_columns(3)} == {"-"}
    assert (a["比價日_1"], a["比價日_2"], b["比價日_1"]) == (dt.datetime(2026, 11, 16), "-", "-")
    assert [a[c] for c in ("UL_1", "UL_1_進場價", "UL_1_執行價", "UL_1_下限價", "UL_1_KO價")] == [
        "AAA UW",
        100,
        80,
        65,
        100,
    ]
    assert (a["UL_2"], a["UL_2_進場價"], a["UL_3"]) == ("BBB UN", "-", "-"), "有代碼沒價格寫 -"
    assert {a[c] for c in ("私銀註記", "新加坡註記", "Lookback 期間")} == {"-"}
    assert all(value is not None for row in (a, b) for value in row.values()), "沒有空白格"


@pytest.mark.db
def test_summary_cells_are_centred_with_fonts_by_content(pg, tmp_path: Path) -> None:
    import_order(pg, A)
    add_trades(pg, A, [("台北", "Amy", 100000)])
    path = tmp_path / "商品總表.xlsx"

    export_summary(pg, path)

    sheet = openpyxl.load_workbook(path).active
    header = {cell.value: cell.column for cell in sheet[1]}
    cell = {name: sheet.cell(2, column) for name, column in header.items()}
    assert sheet.title == "商品總表" and sheet.freeze_panes == "G2" and sheet.auto_filter.ref == sheet.dimensions
    assert (sheet["A1"].font.name, sheet["A1"].font.b) == ("微軟正黑體", True)
    assert (sheet.cell(1, header["TDCC Code"]).font.name, sheet.cell(1, header["TDCC Code"]).font.b) == (
        "Calibri",
        True,
    )
    assert cell["分公司_1"].font.name == "微軟正黑體" and cell["FC_1"].font.name == "Calibri"
    assert cell["存續狀態"].font.name == "微軟正黑體" and cell["私銀註記"].font.name == "Calibri", "在庫是中文、- 不是"
    assert cell["TDCC Code"].number_format == "@" and cell["交易日"].number_format == "yyyy/m/d"
    assert {(c.alignment.horizontal, c.alignment.vertical) for row in sheet.iter_rows() for c in row} == {
        ("center", "center")
    }
