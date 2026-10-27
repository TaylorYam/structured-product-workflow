"""寄信格式範本：未設定 SPW_MAIL_TEMPLATE 時使用專案內建檔，設定了就以設定的檔案為準。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from sp_workflow.mail_drafts import BUNDLED_MAIL_TEMPLATE, mail_template_from_env


def _env(tmp_path: Path, extra: str = "") -> Path:
    env = tmp_path / ".env"
    env.write_text(f"PGHOST=localhost\n{extra}", encoding="utf-8")
    return env


def test_bundled_template_is_shipped_and_has_the_mail_format() -> None:
    assert BUNDLED_MAIL_TEMPLATE.is_file()
    ws = openpyxl.load_workbook(BUNDLED_MAIL_TEMPLATE).worksheets[0]
    assert ws["A1"].value == "TDCC"
    assert ws["A4"].value == "Product"
    assert ws["A5"].value is None  # 範例交易資料不隨專案發佈


def test_unset_template_uses_the_bundled_format(tmp_path: Path) -> None:
    assert mail_template_from_env(_env(tmp_path)) == BUNDLED_MAIL_TEMPLATE


def test_configured_template_overrides_the_bundled_format(tmp_path: Path) -> None:
    override = tmp_path / "other.xlsx"
    override.write_bytes(b"not used")
    assert mail_template_from_env(_env(tmp_path, f"SPW_MAIL_TEMPLATE={override}\n")) == override
