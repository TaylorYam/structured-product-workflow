"""放回資料夾：只透過 return_to_folders 驗證可觀察的結果（商品資料夾多了哪些檔、放回清單、批次資料夾不變）。"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from sp_workflow.preparation import prepare_batch
from sp_workflow.returning import (
    ALREADY_THERE,
    NOT_CHECKED_REASON,
    NOT_RETURNED,
    PENDING_ALREADY_THERE,
    PENDING_RETURNED,
    RETURNED,
    ReturnError,
    return_to_folders,
)
from synth import make_pdf, make_product_folder, snapshot, ts_lines

PREPARED_AT = dt.datetime(2026, 10, 19, 10, 30, 0)
RETURNED_AT = dt.datetime(2026, 10, 20, 9, 0, 0)
A = "029199990001"
B = "325199990002"
C = "147199990003"


def make_result_file(path: Path, passed: list[str | int], failed: list[str] = ()) -> Path:
    """合成核對結果檔：「回填後」第 1 列最新交易日、第 2 列表頭；「錯誤清單」列沒通過的 PDF。

    failed 的每一項是 PDF 檔名（例：`<TDCC>_IIS.pdf`），只寫 TDCC 時當成說明書沒通過。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "回填後"
    ws.append(["最新交易日", "2026-10-13"])
    ws.append(["TS/IIS完成檢核", "交易日期", "Product", "TDCC Code", "ISIN Code"])
    for tdcc in passed:
        ws.append(["V", "2026-10-13", "FCN", tdcc, "XS0000000000"])
    errors = wb.create_sheet("錯誤清單")
    errors.append(["TDCC Code", "PDF 檔名", "錯訊"])
    for item in failed:
        name = item if item.endswith(".pdf") else f"{item}_TS.pdf"
        errors.append([name[:12], name, "UL_1 進場價對不起來：參考條件表 123.4500／說明書 123.4000"])
    wb.save(path)
    return path


@pytest.fixture
def root(tmp_path: Path) -> Path:
    r = tmp_path / "FCN"
    r.mkdir()
    return r


@pytest.fixture
def batch(root: Path, tmp_path: Path) -> Path:
    """A、B 兩個商品已準備好的批次資料夾。"""
    for tdcc in (A, B):
        make_product_folder(root, tdcc)
    out = tmp_path / "out"
    out.mkdir()
    return prepare_batch(f"{A} {B}", root, out, now=PREPARED_AT).batch_dir


def run(batch: Path, result: Path):
    """每次放回都確認批次資料夾裡的 PDF 與準備清單沒被改動。"""
    before = {k: v for k, v in snapshot(batch).items() if not k.startswith("放回清單")}
    receipt = return_to_folders(batch, result, now=RETURNED_AT)
    assert {k: v for k, v in snapshot(batch).items() if not k.startswith("放回清單")} == before
    return receipt


def outcome(receipt, tdcc: str):
    [o] = [o for o in receipt.outcomes if o.tdcc == tdcc]
    return o


def read_list(path: Path) -> list[dict[str, object]]:
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    return [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


# ---- 主要流程 ----


def test_copies_passed_products_back_to_their_product_folders(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "核對結果.xlsx", [A], failed=[B])

    receipt = run(batch, result)

    o = outcome(receipt, A)
    assert (o.ts, o.iis) == (RETURNED, RETURNED)
    assert (root / A / f"{A}_TS.pdf").read_bytes() == (batch / f"{A}_TS.pdf").read_bytes()
    assert (root / A / f"{A}_IIS.pdf").read_bytes() == (batch / f"{A}_IIS.pdf").read_bytes()
    assert not (root / B / f"{B}_TS.pdf").exists()  # 沒通過的不用標準名


def test_only_rows_marked_v_are_returned(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "核對結果.xlsx", [A, B])
    wb = openpyxl.load_workbook(result)
    wb["回填後"]["A4"] = None  # B 那列沒有打 V
    wb.save(result)

    receipt = run(batch, result)

    assert outcome(receipt, B).reason == NOT_CHECKED_REASON  # 不在回填後、也不在錯誤清單：還沒核對
    assert not (root / B / f"{B}_TS.pdf").exists()
    assert not (root / B / f"{B}_TS_待確認.pdf").exists()


def test_tdcc_stored_as_number_still_matches(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "核對結果.xlsx", [int(A)])  # Excel 掉了前導零

    receipt = run(batch, result)

    assert outcome(receipt, A).ts == RETURNED
    assert (root / A / f"{A}_TS.pdf").exists()


def test_other_files_in_product_folder_are_untouched(root: Path, batch: Path, tmp_path: Path) -> None:
    before = snapshot(root / A)

    run(batch, make_result_file(tmp_path / "核對結果.xlsx", [A]))

    after = snapshot(root / A)
    assert set(after) - set(before) == {f"{A}_TS.pdf", f"{A}_IIS.pdf"}
    assert {k: after[k] for k in before} == before


# ---- 已在、衝突與重跑 ----


def test_identical_file_already_in_folder_counts_as_already_there(root: Path, batch: Path, tmp_path: Path) -> None:
    """UBS 上手本來就給標準檔名，內容和批次裡的一樣。"""
    (root / A / f"{A}_TS.pdf").write_bytes((batch / f"{A}_TS.pdf").read_bytes())

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [A])), A)

    assert (o.ts, o.iis) == (ALREADY_THERE, RETURNED)


def test_different_file_with_same_name_blocks_the_whole_product(root: Path, batch: Path, tmp_path: Path) -> None:
    make_pdf(root / A / f"{A}_IIS.pdf", ["同事放的另一份"])
    before = snapshot(root / A)

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [A])), A)

    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert f"商品資料夾已有內容不同的 {A}_IIS.pdf" in o.reason
    assert snapshot(root / A) == before  # TS 也不放


def test_returning_again_skips_what_is_already_there(root: Path, batch: Path, tmp_path: Path) -> None:
    run(batch, make_result_file(tmp_path / "first.xlsx", [A]))

    receipt = return_to_folders(
        batch, make_result_file(tmp_path / "second.xlsx", [A, B]), now=RETURNED_AT + dt.timedelta(hours=1)
    )

    assert (outcome(receipt, A).ts, outcome(receipt, A).iis) == (ALREADY_THERE, ALREADY_THERE)
    assert (outcome(receipt, B).ts, outcome(receipt, B).iis) == (RETURNED, RETURNED)


# ---- 不放回 ----


def test_tdcc_not_in_this_batch_is_not_returned(batch: Path, tmp_path: Path) -> None:
    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [C])), C)

    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert o.reason == "這批準備清單沒有這個商品"


def test_product_skipped_during_preparation_is_not_returned(root: Path, tmp_path: Path) -> None:
    make_pdf(root / C / "ts.pdf", ts_lines(C))  # 缺投資人須知，準備時跳過
    out = tmp_path / "out"
    out.mkdir()
    batch = prepare_batch(C, root, out, now=PREPARED_AT).batch_dir

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [C])), C)

    assert o.ts == NOT_RETURNED
    assert o.reason == "準備時已跳過這個商品"


def test_changed_batch_file_is_not_returned(root: Path, batch: Path, tmp_path: Path) -> None:
    make_pdf(batch / f"{A}_TS.pdf", ["被換掉的檔"])
    result = make_result_file(tmp_path / "r.xlsx", [A])

    receipt = return_to_folders(batch, result, now=RETURNED_AT)

    o = outcome(receipt, A)
    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert f"批次資料夾裡的 {A}_TS.pdf 和準備清單的指紋不符" in o.reason
    assert not (root / A / f"{A}_IIS.pdf").exists()


def test_missing_batch_file_is_not_returned(root: Path, batch: Path, tmp_path: Path) -> None:
    (batch / f"{A}_IIS.pdf").unlink()

    o = outcome(return_to_folders(batch, make_result_file(tmp_path / "r.xlsx", [A]), now=RETURNED_AT), A)

    assert o.iis == NOT_RETURNED
    assert f"批次資料夾裡找不到 {A}_IIS.pdf" in o.reason
    assert not (root / A / f"{A}_TS.pdf").exists()


def test_missing_product_folder_is_not_returned(root: Path, batch: Path, tmp_path: Path) -> None:
    for p in (root / A).iterdir():
        p.unlink()
    (root / A).rmdir()

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [A])), A)

    assert o.ts == NOT_RETURNED
    assert "商品資料夾不存在" in o.reason


def test_failure_halfway_through_a_copy_leaves_nothing_behind(
    root: Path, batch: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """公槽斷線：IIS 寫了一半就失敗，商品資料夾不能留下半個檔或只有 TS。"""
    from sp_workflow import returning

    real_copy = returning.shutil.copyfile

    def broken_copy(src, dst):
        if f"{A}_IIS.pdf" in str(dst):
            Path(dst).write_bytes(b"half")
            raise OSError("網路中斷")
        return real_copy(src, dst)

    monkeypatch.setattr(returning.shutil, "copyfile", broken_copy)
    before = snapshot(root / A)

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [A])), A)

    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert snapshot(root / A) == before


def test_copy_failure_does_not_leave_half_a_pair(
    root: Path, batch: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from sp_workflow import returning

    real_copy = returning.shutil.copyfile

    def flaky_copy(src, dst):
        if f"{A}_IIS.pdf" in str(dst):
            raise PermissionError("檔案被鎖住")
        return real_copy(src, dst)

    monkeypatch.setattr(returning.shutil, "copyfile", flaky_copy)

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [A, B])), A)

    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert "複製失敗" in o.reason
    assert not (root / A / f"{A}_TS.pdf").exists()


# ---- 整批停止 ----


def test_result_file_without_backfill_sheet_stops_everything(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "r.xlsx", [A])
    wb = openpyxl.load_workbook(result)
    wb["回填後"].title = "其他"
    wb.save(result)

    with pytest.raises(ReturnError, match="找不到「回填後」工作表"):
        return_to_folders(batch, result, now=RETURNED_AT)
    assert not (root / A / f"{A}_TS.pdf").exists()


def test_result_file_without_required_columns_stops_everything(batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "r.xlsx", [A])
    wb = openpyxl.load_workbook(result)
    wb["回填後"]["D2"] = "代號"
    wb.save(result)

    with pytest.raises(ReturnError, match="TDCC Code"):
        return_to_folders(batch, result, now=RETURNED_AT)


def test_folder_without_preparation_list_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ReturnError, match="準備清單"):
        return_to_folders(tmp_path, make_result_file(tmp_path / "r.xlsx", [A]), now=RETURNED_AT)


def test_broken_preparation_list_is_rejected_with_a_readable_message(batch: Path, tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(batch / "準備清單.xlsx")
    wb.active["B2"] = "XX"  # 種類
    wb.save(batch / "準備清單.xlsx")

    with pytest.raises(ReturnError, match="準備清單第 2 列的種類看不懂"):
        return_to_folders(batch, make_result_file(tmp_path / "r.xlsx", [A]), now=RETURNED_AT)


def test_unreadable_result_file_is_rejected(batch: Path, tmp_path: Path) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not excel")

    with pytest.raises(ReturnError, match="讀不到核對結果檔"):
        return_to_folders(batch, broken, now=RETURNED_AT)


def test_existing_return_list_is_never_overwritten(batch: Path, tmp_path: Path) -> None:
    (batch / "放回清單_20261020-090000.xlsx").write_bytes(b"old")

    with pytest.raises(ReturnError, match="已存在"):
        return_to_folders(batch, make_result_file(tmp_path / "r.xlsx", [A]), now=RETURNED_AT)
    assert (batch / "放回清單_20261020-090000.xlsx").read_bytes() == b"old"


# ---- 放回清單 ----


def test_return_list_records_each_product(batch: Path, tmp_path: Path) -> None:
    receipt = run(batch, make_result_file(tmp_path / "r.xlsx", [A, C]))

    assert receipt.list_path == batch / "放回清單_20261020-090000.xlsx"
    rows = read_list(receipt.list_path)
    assert list(rows[0]) == ["TDCC", "TS", "IIS", "原因"]
    assert rows == [
        {"TDCC": A, "TS": RETURNED, "IIS": RETURNED, "原因": None},
        {"TDCC": C, "TS": NOT_RETURNED, "IIS": NOT_RETURNED, "原因": "這批準備清單沒有這個商品"},
        {"TDCC": B, "TS": NOT_RETURNED, "IIS": NOT_RETURNED, "原因": NOT_CHECKED_REASON},  # 準備了但還沒核對
    ]


# ---- 沒通過的以「_待確認」放回（#17）----


def test_failed_document_is_returned_with_pending_suffix(root: Path, batch: Path, tmp_path: Path) -> None:
    receipt = run(batch, make_result_file(tmp_path / "r.xlsx", [A], failed=[B]))

    o = outcome(receipt, B)
    assert (o.ts, o.iis) == (PENDING_RETURNED, RETURNED)  # 只有沒過的那份加待確認
    assert (root / B / f"{B}_TS_待確認.pdf").read_bytes() == (batch / f"{B}_TS.pdf").read_bytes()
    assert (root / B / f"{B}_IIS.pdf").read_bytes() == (batch / f"{B}_IIS.pdf").read_bytes()
    assert not (root / B / f"{B}_TS.pdf").exists()
    assert o.reason == "核對沒通過（TS）：UL_1 進場價對不起來：參考條件表 123.4500／說明書 123.4000"


def test_both_failed_documents_get_the_suffix(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "r.xlsx", [], failed=[f"{B}_TS.pdf", f"{B}_IIS.pdf"])

    o = outcome(run(batch, result), B)

    assert (o.ts, o.iis) == (PENDING_RETURNED, PENDING_RETURNED)
    assert (root / B / f"{B}_IIS_待確認.pdf").exists()
    assert "核對沒通過（TS）" in o.reason and "核對沒通過（IIS）" in o.reason


def test_returning_a_pending_file_again_counts_as_already_there(root: Path, batch: Path, tmp_path: Path) -> None:
    run(batch, make_result_file(tmp_path / "first.xlsx", [], failed=[B]))

    receipt = return_to_folders(
        batch, make_result_file(tmp_path / "second.xlsx", [], failed=[B]), now=RETURNED_AT + dt.timedelta(hours=1)
    )

    assert (outcome(receipt, B).ts, outcome(receipt, B).iis) == (PENDING_ALREADY_THERE, ALREADY_THERE)


def test_different_pending_file_with_same_name_blocks_the_product(root: Path, batch: Path, tmp_path: Path) -> None:
    make_pdf(root / B / f"{B}_TS_待確認.pdf", ["上一版的說明書"])
    before = snapshot(root / B)

    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [], failed=[B])), B)

    assert (o.ts, o.iis) == (NOT_RETURNED, NOT_RETURNED)
    assert f"商品資料夾已有內容不同的 {B}_TS_待確認.pdf" in o.reason
    assert snapshot(root / B) == before


def test_old_pending_file_is_kept_and_reported_when_the_product_passes(root: Path, batch: Path, tmp_path: Path) -> None:
    run(batch, make_result_file(tmp_path / "first.xlsx", [], failed=[B]))

    receipt = return_to_folders(
        batch, make_result_file(tmp_path / "second.xlsx", [B]), now=RETURNED_AT + dt.timedelta(hours=1)
    )

    o = outcome(receipt, B)
    assert (o.ts, o.iis) == (RETURNED, ALREADY_THERE)
    assert (root / B / f"{B}_TS.pdf").exists()
    assert (root / B / f"{B}_TS_待確認.pdf").exists()  # 不自動刪
    assert o.reason == f"資料夾仍有舊的 {B}_TS_待確認.pdf，請確認後刪除"


def test_failed_product_not_in_this_batch_is_not_returned(batch: Path, tmp_path: Path) -> None:
    o = outcome(run(batch, make_result_file(tmp_path / "r.xlsx", [], failed=[C])), C)

    assert (o.ts, o.iis, o.reason) == (NOT_RETURNED, NOT_RETURNED, "這批準備清單沒有這個商品")


def test_result_file_without_error_sheet_stops_everything(root: Path, batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "r.xlsx", [A])
    wb = openpyxl.load_workbook(result)
    del wb["錯誤清單"]
    wb.save(result)

    with pytest.raises(ReturnError, match="找不到「錯誤清單」工作表"):
        return_to_folders(batch, result, now=RETURNED_AT)
    assert not (root / A / f"{A}_TS.pdf").exists()


def test_error_sheet_without_required_columns_stops_everything(batch: Path, tmp_path: Path) -> None:
    result = make_result_file(tmp_path / "r.xlsx", [A])
    wb = openpyxl.load_workbook(result)
    wb["錯誤清單"]["B1"] = "檔名"
    wb.save(result)

    with pytest.raises(ReturnError, match="PDF 檔名"):
        return_to_folders(batch, result, now=RETURNED_AT)
