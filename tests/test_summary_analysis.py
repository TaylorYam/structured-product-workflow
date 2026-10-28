"""商品總表的分析分頁（#293、#299）：公式寫法、動態陣列中繼資料；`excel` 標記的測試用真的 Excel 算一次。資料全部虛構。"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

from sp_workflow.file_tools import FileWriteError, save_export
from sp_workflow.summary_analysis import (
    BRANCH_SHEET,
    ISSUER_SHEET,
    TABLE_NAME,
    add_analysis_sheets,
    enable_spill,
    file_formula,
    summary_table,
)

HEADER = ["存續狀態", "發行機構", "承作幣別", "承作總額", "分公司_1", "FC_1", "金額_1", "分公司_2", "FC_2", "金額_2"]
ROWS = [
    ["在庫", "BNP", "USD", 300, "台北", "Amy", 100, "台中", "Ben", 200],
    ["在庫", "Barclays", "JPY", 5000, "台中", "Ben", 5000, "-", "-", "-"],
    ["KO", "BNP", "USD", 999, "台北", "Amy", 999, "-", "-", "-"],  # 已出場不算
    ["在庫", "BNP", "USD", 50, "新竹", "Cat", 50, "-", "-", "-"],
    ["在庫", "UBS", "USD", "-", "-", "-", "-", "-", "-", "-"],  # 沒有交易明細
]


def write_book(path: Path, rows: list[list[object]], groups: int = 2) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "商品總表"
    for row in [HEADER, *rows]:
        sheet.append(row)
    summary_table(sheet, len(HEADER), len(rows))
    add_analysis_sheets(book, groups)
    save_export(book, path, finish=enable_spill)
    return path


def test_formulas_use_the_excel_file_spelling() -> None:
    formula = file_formula(
        'LET(c,UNIQUE(FILTER(總表[承作幣別],總表[存續狀態]="在庫")),SORTBY(c,COUNTIFS(總表[承作幣別],c),-1))', ("c",)
    )

    assert formula == (
        '_xlfn.LET(_xlpm.c,_xlfn.UNIQUE(_xlfn._xlws.FILTER(總表[承作幣別],總表[存續狀態]="在庫")),'
        "_xlfn.SORTBY(_xlpm.c,COUNTIFS(總表[承作幣別],_xlpm.c),-1))"
    )


def test_the_summary_is_a_table_followed_by_two_analysis_sheets(tmp_path: Path) -> None:
    book = openpyxl.load_workbook(write_book(tmp_path / "x.xlsx", ROWS, groups=2))

    assert book.sheetnames == ["商品總表", ISSUER_SHEET, BRANCH_SHEET]
    table = book["商品總表"].tables[TABLE_NAME]
    assert table.ref == "A1:J6" and book["商品總表"].auto_filter.ref is None, "篩選按鈕由表格提供"
    for title in book.sheetnames[1:]:
        sheet = book[title]
        values = [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.value is not None]
        assert (values, sheet.freeze_panes) == (["A1"], "B2"), "只有 A1 一個公式，沒有說明文字"


def test_branch_formulas_cover_every_trade_group(tmp_path: Path) -> None:
    book = openpyxl.load_workbook(write_book(tmp_path / "x.xlsx", ROWS, groups=3))

    formula = book[BRANCH_SHEET]["A1"].value.text
    assert [f"總表[金額_{n}],總表[分公司_{n}]" in formula for n in (1, 2, 3)] == [True] * 3
    assert "總表[[分公司_1]:[金額_3]]" in formula


def test_saved_file_marks_every_formula_as_a_dynamic_array(tmp_path: Path) -> None:
    path = write_book(tmp_path / "x.xlsx", ROWS)

    with zipfile.ZipFile(path) as xlsx:
        names = xlsx.namelist()
        types = xlsx.read("[Content_Types].xml").decode()
        rels = xlsx.read("xl/_rels/workbook.xml.rels").decode()
        sheets = [xlsx.read(name).decode() for name in names if name.startswith("xl/worksheets/sheet")]
    assert "xl/metadata.xml" in names and "/xl/metadata.xml" in types and 'Target="metadata.xml"' in rels
    arrays = [cell for sheet in sheets for cell in re.findall(r"<c [^>]*><f t=\"array\"", sheet)]
    assert len(arrays) == 2 and all('cm="1"' in cell for cell in arrays), "兩個分頁各 1 個公式"


def test_a_failed_finish_leaves_no_temp_file(tmp_path: Path) -> None:
    def broken(_path: Path) -> None:
        raise PermissionError(13, "Permission denied")

    book = openpyxl.Workbook()
    with pytest.raises(FileWriteError, match="寫不進"):
        save_export(book, tmp_path / "x.xlsx", finish=broken)
    assert list(tmp_path.iterdir()) == []


# ---------------------------------------------------------------- 真的 Excel


@pytest.fixture(scope="module")
def excel():
    if sys.platform != "win32":
        pytest.skip("不是 Windows")
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        pytest.skip("沒有 pywin32")
    pythoncom.CoInitialize()
    try:
        app = win32com.client.DispatchEx("Excel.Application")
    except Exception:
        pytest.skip("無法啟動 Microsoft Excel")
    app.Visible = False
    app.DisplayAlerts = False
    yield app
    app.Quit()
    pythoncom.CoUninitialize()


def grid(book, title: str, rows: int, columns: int) -> list[list[str]]:
    sheet = book.Worksheets(title)
    return [[sheet.Cells(1 + r, 1 + c).Text for c in range(columns)] for r in range(rows)]


@pytest.mark.excel
def test_excel_spills_in_stock_amounts_and_follows_edits(excel, tmp_path: Path) -> None:
    book = excel.Workbooks.Open(str(write_book(tmp_path / "x.xlsx", ROWS)))
    try:
        excel.CalculateFull()
        assert grid(book, ISSUER_SHEET, 12, 3) == [
            ["發行機構", "USD", "JPY"],  # USD 在庫 3 檔、JPY 1 檔
            ["總計", "350", "5,000"],
            ["Barclays", "-", "5,000"],
            ["BNP", "350", "-"],
            ["UBS", "-", "-"],  # 在庫但沒有金額
            ["", "", ""],
            ["發行機構占比", "USD", "JPY"],
            ["總計", "100.0%", "100.0%"],
            ["Barclays", "-", "100.0%"],
            ["BNP", "100.0%", "-"],
            ["UBS", "-", "-"],
            ["", "", ""],
        ]
        branches = grid(book, BRANCH_SHEET, 11, 3)
        assert (branches[:2], sorted(branches[2:5])) == (
            [["分公司", "USD", "JPY"], ["總計", "350", "5,000"]],
            [
                ["台中", "200", "5,000"],  # 一檔兩家分公司，各分到自己的金額
                ["台北", "100", "-"],
                ["新竹", "50", "-"],
            ],
        )
        assert sorted(branches[8:]) == [["台中", "57.1%", "100.0%"], ["台北", "28.6%", "-"], ["新竹", "14.3%", "-"]]

        table = book.Worksheets("商品總表").ListObjects(TABLE_NAME)
        table.ListRows.Add().Range.Value = (("在庫", "HSBC", "EUR", 70, "台北", "Dan", 70, "-", "-", "-"),)
        excel.Calculate()
        issuers = grid(book, ISSUER_SHEET, 13, 4)
        assert issuers[:2] == [["發行機構", "USD", "JPY", "EUR"], ["總計", "350", "5,000", "70"]], (
            "表格下方新增一列：新的幣別、發行機構自動多出來（檔數一樣時照總表出現的順序）"
        )
        assert [row[0] for row in issuers] == [
            "發行機構", "總計", "Barclays", "BNP", "HSBC", "UBS", "",
            "發行機構占比", "總計", "Barclays", "BNP", "HSBC", "UBS",
        ], "上面多一列，占比表跟著往下，不會 #SPILL!"  # fmt: skip
    finally:
        book.Close(False)


@pytest.mark.excel
def test_excel_shows_no_error_values_without_products(excel, tmp_path: Path) -> None:
    book = excel.Workbooks.Open(str(write_book(tmp_path / "x.xlsx", [])))
    try:
        excel.CalculateFull()
        for title in (ISSUER_SHEET, BRANCH_SHEET):
            assert not any(text.startswith("#") for row in grid(book, title, 8, 3) for text in row), title
    finally:
        book.Close(False)
