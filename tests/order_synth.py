"""起單資料確認與視窗測試共用的虛構起單建檔與讀取結果。資料全部虛構。"""

from __future__ import annotations

import datetime as dt
from collections.abc import Mapping, Sequence
from decimal import Decimal
from pathlib import Path

import openpyxl

from fcn_checker.config import load_exchange_holidays
from sp_workflow.mailbox import Mailbox, Person
from sp_workflow.order_terms import NOT_FOUND, PENDING, READ, TDCC_FIELD, OrderRow
from sp_workflow.product_terms import Draft, ImportConfig, draft_from, read
from sp_workflow.status import OrderReview
from sp_workflow.trade_rows import TradeRowDraft, TradeRowsDraft, trade_product

A = "029199990001"  # BARC
B = "037199990002"  # SG
SHEET = Path("起單建檔_20261015.xlsm")
PREFIXES = {"029": "BARC", "037": "SG", "092": "UBS", "147": "MS"}
CONFIG = ImportConfig(
    {"USD": Decimal(10000), "JPY": Decimal(1000000)},
    {"BARC": "Barclays", "SG": "SG"},
    PREFIXES,
    load_exchange_holidays(Path(__file__).resolve().parents[1] / "config" / "exchange_holidays.toml"),
)

# 起單建檔的一列（畫面的欄名 → 值）
VALUES: dict[str, object] = {
    "Product": "FCN",
    "Currency": "USD",
    "Guaranteed Periods (m)": 1,
    "BBG Code 1": "AAA UW",
    "BBG Code 2": "BBB UN",
    "BBG Code 3": None,
    "BBG Code 4": None,
    "BBG Code 5": None,
    "Strike (%)": Decimal(80),
    "KO Type": "Daily Memory",
    "KO Barrier (%)": Decimal(100),
    "Coupon p.a. (%)": Decimal("12.5"),
    "Participation (%)": None,  # 起單建檔沒有這欄（只有 BEN 有，#373）
    "Lookback end date": None,  # 起單建檔沒有這欄（只有 lookback 的 BEN 有，#374）
    "Coupon Barrier (%)": None,  # 只有 DACN 有（#367）
    "Upfront / NotePrice (%)": Decimal("98.5"),
    "Tenor (m)": 6,
    "Barrier Type": "EKI",
    "KI Barrier (%)": Decimal(65),
    "Trade date": dt.date(2026, 10, 15),
    "Strike date": dt.date(2026, 10, 15),
    "Issue date": dt.date(2026, 10, 22),
    "Final valuation date": dt.date(2027, 4, 15),
    "Maturity date": dt.date(2027, 4, 22),
    "期初價格": "收盤價",
}


def order_row(
    tdcc: str = A,
    changes: Mapping[str, object] | None = None,
    *,
    notes: Mapping[str, str] | None = None,
    conflicts: Sequence[str] = (),
) -> OrderRow:
    """讀取結果：`changes` 改欄位的值；上手依 `PREFIXES`。"""
    values = {**VALUES, **(changes or {})}
    notes = dict(notes or {})
    status = PENDING if set(notes) - {TDCC_FIELD} or conflicts else READ
    return OrderRow(tdcc, status, PREFIXES.get(tdcc[:3], ""), values, notes, frozenset(conflicts), SHEET, (3,))


def missing_row(tdcc: str) -> OrderRow:
    return OrderRow(tdcc, NOT_FOUND, PREFIXES.get(tdcc[:3], ""), notes={TDCC_FIELD: f"{SHEET.name} 裡沒有這檔"})


def maintenance_draft(status: str = OrderReview.UNLOCKED, tdcc: str = A, issuer: str = "BARC", **changes) -> Draft:
    """人工維護打開的一檔（`draft_from(tdcc, …)` 的結果，不連資料庫）：`changes` 改欄位的值。"""
    values = {**VALUES, **changes}
    row = OrderRow(tdcc, status, issuer, dict(values))
    return Draft(row, dict(values), dict(values), dict(values), database_checked=True)


def open_terms(conn, tdcc: str, config: ImportConfig = CONFIG) -> Draft:
    """人工維護打開一檔（同 `window_actions.open_product`，不讀交易明細、不檢查兩邊都沒有）。"""
    return draft_from(tdcc, read(conn, tdcc), config, trade_product=trade_product(conn, tdcc))


# 資料庫裡一列交易明細（表頭 → 值），同 db_synth.trade 匯入後的樣子
TRADE: dict[str, object] = {
    "編號": 1,
    "交易日期": dt.date(2026, 10, 15),
    "幣別": "USD",
    "金額": Decimal("100000.0000"),
    "FC/PB": "Amy",
    "分公司": "台北",
    "前台主管/投專": None,
    "UF": Decimal("0.021500"),
    "期初價格": "收盤價",
    "備註": None,
    "募集": "N",
    "TS/IIS 核對": None,
    "核對TS/IIS人員": None,
    "起單人員": "A",
}
MAILBOX = Mailbox((Person("Amy", "amy@example.com", "FC", False),), {})


def maintenance_trades(*rows: Mapping[str, object], tdcc: str = A, mailbox: Mailbox = MAILBOX) -> TradeRowsDraft:
    """人工維護打開的交易明細（`load_trade_rows` 的結果，不連資料庫）：每列是畫面上對 `TRADE` 的修改。"""
    drafts = [TradeRowDraft(seq, {**TRADE, **changes}, dict(TRADE)) for seq, changes in enumerate(rows, 1)]
    return TradeRowsDraft(tdcc, drafts, mailbox)


# 起單建檔的表頭（第 2 列、從 B 欄開始；欄名後面的空白照真的檔）
HEADERS = (
    "No.", "TDCC", "Product", "Currency", "Guaranteed Periods (m) ", "BBG Code 1", "BBG Code 2", "BBG Code 3",
    "BBG Code 4", "BBG Code 5", "Strike (%) ", "KO Type", "KO Barrier (%)", "Coupon p.a. (%) ", "UF (%) ",
    "Tenor (m)", "Barrier Type", "KI Barrier (%) ", "Observation Frequency (m)", "OTC", "Funding Spread (bps)",
    "Effective Date Offset(Calendar Days) ", "Trade date", "Issue date", "Final valuation date", "Maturity date",
    "期初價格", "起單FC",
)  # fmt: skip
_FIELD_OF = {"UF (%)": "Upfront / NotePrice (%)"}


def sheet_line(tdcc: str = A, changes: Mapping[str, object] | None = None) -> dict[str, object]:
    """起單建檔的一列（表頭 → 格子），數字是 Excel 讀出來的 int／float、日期是 datetime。"""
    values = {**VALUES, **(changes or {})}
    line: dict[str, object] = {"TDCC": tdcc, "Observation Frequency (m)": 1, "OTC": "Note", "起單FC": "Amy"}
    for header in HEADERS:
        name = _FIELD_OF.get(header.strip(), header.strip())
        if name in values:
            value = values[name]
            if isinstance(value, Decimal):
                value = float(value)
            elif isinstance(value, dt.date) and not isinstance(value, dt.datetime):
                value = dt.datetime.combine(value, dt.time())
            line[header.strip()] = value
    return line


def write_sheet(
    path: Path, lines: Sequence[Mapping[str, object]], *, headers: Sequence[str] = HEADERS, sheets: int = 1
) -> Path:
    """寫一份虛構的起單建檔：第 2 列表頭、從 B 欄開始，No. 欄是公式（同真的檔）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"起單建檔_{path.stem[-8:]}"
    for i, header in enumerate(headers, 2):
        ws.cell(2, i, header)
    keys = [h.strip() for h in headers]
    for r, line in enumerate(lines, 3):
        ws.cell(r, 2, f'=IF(ISBLANK(C{r}),"",ROW()-2)')
        for key, value in line.items():
            if key in keys:
                ws.cell(r, 2 + keys.index(key), value)
    for n in range(1, sheets):
        wb.create_sheet(f"起單建檔_{n}")
    wb.save(path)
    return path
