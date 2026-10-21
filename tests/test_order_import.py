"""Order 條件確認與匯入（#86、ADR 0009）：讀信結果用 `template_row` 組出，資料庫測試對真的 PostgreSQL。資料全部虛構。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from db_synth import inventory, seed_mailbox, trade, write_daily
from fcn_checker.check_config import ConfigPaths
from sp_workflow.daily_import import DIFFERENT, FILLED, INVENTORY_MISSING, import_daily
from sp_workflow.mail_drafts import PENDING as MAIL_PENDING
from sp_workflow.mail_drafts import create_drafts
from sp_workflow.order_import import (
    AMOUNT_FIELD,
    FIELDS,
    LOCKED,
    NEW,
    OVERWRITTEN,
    SKIPPED,
    UNCHANGED,
    ImportConfig,
    OrderImportError,
    check_draft,
    drafts_from,
    export_drafts,
    format_field,
    import_drafts,
    load_import_config,
    parse_field,
)
from sp_workflow.order_template import COLUMNS, FLAG_FILL, HEADER_ROW, template_row
from sp_workflow.order_terms import NO_ORDER, PENDING, READ, OrderRow

A = "029199990001"  # BARC
B = "037199990002"  # SG
SENT = dt.datetime(2026, 10, 15, 14, 30)
CONFIG = ImportConfig(
    {"USD": Decimal(10000), "JPY": Decimal(1000000)},
    {"BARC": "Barclays", "SG": "SG"},
)
BUILTIN = ConfigPaths.with_fallback(Path(__file__).resolve().parents[1] / "config" / "review_standard.toml")


def order_row(tdcc: str = A, *, issuer: str = "BARC", flagged_fields=(), fields=None, **overrides) -> OrderRow:
    """讀信結果；`overrides` 改統一欄位（欄名有括號的用 `fields`）。"""
    values = {
        "商品": "FCN",
        "幣別": "USD",
        "承作金額": Decimal(500000),
        "價格(%)": Decimal("98.5"),
        "標的1": "AAA UW",
        "標的2": "BBB UN",
        "執行價(%)": Decimal(80),
        "KO(%)": Decimal(100),
        "KO觀察方式": "Daily Memory",
        "KI型態": "EKI",
        "KI(%)": Decimal(65),
        "年利率(%)": Decimal("12.5"),
        "天期(月)": 6,
        "Non-Call(月)": 1,
        "觀察頻率(月)": 1,
        "期初定價": "收盤價",
        "交易日": dt.date(2026, 10, 15),
        "發行日": dt.date(2026, 10, 22),
        "最終比價日": dt.date(2027, 4, 15),
        "到期日": dt.date(2027, 4, 22),
    }
    values.update(overrides)
    values.update(fields or {})
    values = {k: v for k, v in values.items() if v is not None}
    template, problems, flagged = template_row(tdcc, values, (), "", SENT, flagged_fields)
    return OrderRow(
        tdcc,
        PENDING if problems else READ,
        issuer,
        tuple(problems),
        values,
        (),
        (),
        Path(f"(Order) {tdcc}.msg"),
        SENT,
        template,
        flagged,
    )


def draft_of(row: OrderRow, conn=None):
    (draft,) = drafts_from([row], conn)
    return draft


# ---------- 確認資料（不需要資料庫） ----------


def test_draft_holds_the_template_fields_plus_the_amount() -> None:
    draft = draft_of(order_row())

    assert draft.values["KO Type"] == "Daily Memory"
    assert draft.values["Upfront / NotePrice (%)"] == Decimal("98.5")
    assert draft.values[AMOUNT_FIELD] == Decimal(500000)
    assert draft.values["Barrier Type"] == "EKI"
    assert set(draft.values) == set(FIELDS)
    assert check_draft(draft, CONFIG) == {}
    assert draft.edits() == []


def test_rows_without_an_order_have_no_draft() -> None:
    assert drafts_from([OrderRow(A, NO_ORDER, "BARC", ("商品資料夾裡沒有 (Order) 原信",))]) == []


def test_missing_final_valuation_and_maturity_dates_block_the_import() -> None:
    draft = draft_of(order_row(最終比價日=None, 到期日=None, 發行日=None, 標的2=None, 承作金額=None))

    assert check_draft(draft, CONFIG) == {"Final valuation date": "必填", "Maturity date": "必填"}
    assert set(draft.flagged) >= {"Final valuation date", "Maturity date"}  # 讀信時就標紅


def test_filling_the_red_cells_makes_the_product_importable() -> None:
    draft = draft_of(order_row(最終比價日=None, 到期日=None))
    draft.values["Final valuation date"] = parse_field("Final valuation date", "2027-04-15")
    draft.values["Maturity date"] = parse_field("Maturity date", "22-Apr-27")

    assert check_draft(draft, CONFIG) == {}
    assert draft.edits() == [
        ("Final valuation date", None, dt.date(2027, 4, 15)),
        ("Maturity date", None, dt.date(2027, 4, 22)),
    ]


@pytest.mark.parametrize(
    ("changes", "problems"),
    [
        ({"Barrier Type": "NONE"}, {"KI Barrier (%)": "Barrier Type 是 NONE 時要空白"}),
        ({"KI Barrier (%)": None}, {"KI Barrier (%)": "有 KI 時必填"}),
        ({"Currency": "TWD"}, {"Currency": "這個幣別沒有單位面額預設值"}),
        ({"Product": "BEN"}, {"Product": "目前只做 FCN"}),
        ({"Guaranteed Periods (m)": 0}, {"Guaranteed Periods (m)": "至少 1"}),
        ({"BBG Code 2": None, "BBG Code 3": "CCC UQ"}, {"BBG Code 2": "標的要從 BBG Code 1 依序填，中間不能空"}),
        ({"KO Type": "Weekly"}, {"KO Type": "要選 Daily／Daily Memory／Period End／Period End Memory"}),
        ({"Tenor (m)": "6"}, {"Tenor (m)": "要填整數"}),
        ({AMOUNT_FIELD: Decimal(0)}, {AMOUNT_FIELD: "要大於 0"}),
    ],
)
def test_field_checks(changes, problems) -> None:
    draft = draft_of(order_row())
    draft.values.update(changes)

    assert check_draft(draft, CONFIG) == problems


def test_more_than_five_underlyings_cannot_be_imported() -> None:
    draft = draft_of(order_row(標的6="FFF UN"))

    assert check_draft(draft, CONFIG) == {"BBG Code 5": "標的超過 5 檔，核對只支援 5 檔"}


def test_parse_and_format_fields() -> None:
    assert parse_field("Strike (%)", " 80% ") == Decimal(80)
    assert parse_field(AMOUNT_FIELD, "1,000,000") == Decimal(1000000)
    assert parse_field("Tenor (m)", "12") == 12
    assert parse_field("KO Type", "daily memory") == "Daily Memory"
    assert parse_field("Issue date", "") is None
    for name, text, message in (
        ("Strike (%)", "八十", "不是數字"),
        ("Tenor (m)", "6.5", "不是整數"),
        ("Trade date", "10/15", "日期要寫成 2026-10-21"),
        ("Barrier Type", "MKI", "要選 NONE／EKI／AKI"),
    ):
        with pytest.raises(ValueError, match=message):
            parse_field(name, text)
    assert format_field(Decimal("80.000000")) == "80"
    assert format_field(Decimal("2.150000")) == "2.15"
    assert format_field(dt.date(2026, 10, 15)) == "2026-10-15"
    assert format_field(None) == ""


def test_builtin_config_has_denominations_and_issuer_names() -> None:
    config = load_import_config(BUILTIN)

    assert config.denominations["USD"] == Decimal(10000)
    assert config.denominations["JPY"] == Decimal(1000000)
    assert len(config.denominations) == 8
    assert config.issuer_names["BARC"] == "Barclays"


# ---------- 匯出 Excel ----------


def sheet_values(path: Path, title: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path)
    return [tuple(r) for r in wb[title].iter_rows(values_only=True)]


def test_export_writes_confirmed_values_and_the_manual_edits(tmp_path: Path) -> None:
    fixed = order_row(到期日=None)
    other = order_row(B, issuer="SG")
    rows = [fixed, other]
    drafts = drafts_from(rows)
    drafts[0].values["Maturity date"] = dt.date(2027, 4, 22)
    drafts[1].values["Coupon p.a. (%)"] = Decimal("11.25")

    path = export_drafts(tmp_path, rows, drafts, tmp_path, (), CONFIG, now=dt.datetime(2026, 10, 21, 9, 0))

    wb = openpyxl.load_workbook(path)
    assert path.name == "Order條件清單_20261021-090000.xlsx"
    assert wb.sheetnames == ["Order原始資料", "Order條件清單", "原文對照", "資料來源", "人工修改"]
    ws = wb["Order原始資料"]
    headers = [c.header.strip() for c in COLUMNS]
    first = {h: ws.cell(HEADER_ROW + 1, 2 + i) for i, h in enumerate(headers)}
    second = {h: ws.cell(HEADER_ROW + 2, 2 + i) for i, h in enumerate(headers)}
    assert first["Maturity date"].value == dt.datetime(2027, 4, 22)
    assert second["Coupon p.a. (%)"].value == 11.25
    assert first["Effective Date Offset(Calendar Days)"].value == 7
    # 只匯出 Excel（沒有匯入）：讀信時標紅的格子保留淺紅底
    assert first["Maturity date"].fill.fgColor.rgb == FLAG_FILL
    assert second["Coupon p.a. (%)"].fill.fgColor.rgb != FLAG_FILL
    assert sheet_values(path, "人工修改") == [
        ("TDCC", "欄位", "信裡的值", "改成"),
        (A, "Maturity date", "（空白）", dt.datetime(2027, 4, 22)),
        (B, "Coupon p.a. (%)", 12.5, 11.25),
    ]
    # Order條件清單維持讀信結果
    listed = {r[0]: r for r in sheet_values(path, "Order條件清單")[1:]}
    assert listed[A][1] == PENDING


def test_export_refuses_an_existing_file(tmp_path: Path) -> None:
    rows = [order_row()]
    now = dt.datetime(2026, 10, 21, 9, 0)
    export_drafts(tmp_path, rows, drafts_from(rows), tmp_path, (), CONFIG, now=now)

    with pytest.raises(OrderImportError, match="清單已存在"):
        export_drafts(tmp_path, rows, drafts_from(rows), tmp_path, (), CONFIG, now=now)


# ---------- 匯入資料庫 ----------


def product(pg, tdcc: str) -> dict:
    cur = pg.execute("SELECT * FROM workflow.product WHERE tdcc = %s", (tdcc,))
    names = [d.name for d in cur.description]
    found = cur.fetchone()
    return None if found is None else dict(zip(names, found, strict=True))


@pytest.mark.db
def test_import_writes_a_new_product_from_the_confirmed_values(pg) -> None:
    draft = draft_of(order_row(到期日=None), pg)
    draft.values["Maturity date"] = dt.date(2027, 4, 22)

    receipt = import_drafts([draft], pg, CONFIG, user="tester")

    assert [(line.tdcc, line.result) for line in receipt.lines] == [(A, NEW)]
    p = product(pg, A)
    assert p["source"] == "Order" and p["order_import_id"] == receipt.import_id and p["first_run_id"] is None
    assert (p["product"], p["issuer"], p["currency"], p["isin"]) == ("FCN", "Barclays", "USD", None)
    assert (p["ko_freq"], p["ko_memory"], p["ko_type"]) == ("D", "Y", "Daily Memory")
    assert (p["ki_freq"], p["barrier_type"], p["ki_pct"]) == ("AM", "EKI", Decimal(65))
    assert p["upfront_pct"] == Decimal("1.5")
    assert p["denomination"] == Decimal(10000) and p["order_amount"] == Decimal(500000)
    assert p["maturity_date"] == dt.date(2027, 4, 22) and p["status_event"] == "在庫"
    assert (p["underlying_1"], p["underlying_2"], p["underlying_3"]) == ("AAA UW", "BBB UN", None)
    imported_by, counts = pg.execute(
        "SELECT imported_by, summary -> 'counts' FROM workflow.order_import WHERE import_id = %s",
        (receipt.import_id,),
    ).fetchone()
    assert imported_by == "tester" and counts[NEW] == 1
    result, source_mail, confirmed, edits = pg.execute(
        "SELECT result, source_mail, confirmed, edits FROM workflow.order_import_line WHERE import_id = %s",
        (receipt.import_id,),
    ).fetchone()
    assert (result, source_mail) == (NEW, f"(Order) {A}.msg")
    assert confirmed["Maturity date"] == "2027-04-22" and confirmed["Coupon p.a. (%)"] == "12.5"
    assert edits == [["Maturity date", None, "2027-04-22"]]


@pytest.mark.db
@pytest.mark.parametrize("currency", ["USD", "JPY", "CNH", "HKD", "AUD", "EUR", "GBP", "ZAR"])
def test_denomination_is_the_currency_default(pg, currency) -> None:
    config = load_import_config(BUILTIN)

    import_drafts([draft_of(order_row(幣別=currency), pg)], pg, config, user="tester")

    assert product(pg, A)["denomination"] == config.denominations[currency]


@pytest.mark.db
def test_products_with_problems_are_skipped_and_the_rest_imported(pg) -> None:
    bad = draft_of(order_row(到期日=None), pg)
    good = draft_of(order_row(B, issuer="SG"), pg)
    unknown = draft_of(order_row("029199990009", 幣別="TWD"), pg)

    receipt = import_drafts([bad, good, unknown], pg, CONFIG, user="tester")

    assert [(line.tdcc, line.result, line.detail) for line in receipt.lines] == [
        (A, SKIPPED, "Maturity date：必填"),
        (B, NEW, ""),
        ("029199990009", SKIPPED, "Currency：這個幣別沒有單位面額預設值"),
    ]
    assert product(pg, A) is None and product(pg, B)["issuer"] == "SG"


@pytest.mark.db
def test_reimport_overwrites_until_locked_and_keeps_the_previous_values(pg) -> None:
    import_drafts([draft_of(order_row(), pg)], pg, CONFIG, user="tester")

    same = import_drafts([draft_of(order_row(), pg)], pg, CONFIG, user="tester")
    changed = draft_of(order_row(), pg)
    changed.values["Coupon p.a. (%)"] = Decimal("11.25")
    overwritten = import_drafts([changed], pg, CONFIG, user="tester")

    assert same.lines[0].result == UNCHANGED
    assert (overwritten.lines[0].result, overwritten.lines[0].detail) == (OVERWRITTEN, "改了：Coupon p.a. (%)")
    assert product(pg, A)["coupon_pct"] == Decimal("11.25")
    assert product(pg, A)["order_import_id"] == overwritten.import_id
    previous = pg.execute(
        "SELECT previous FROM workflow.order_import_line WHERE import_id = %s", (overwritten.import_id,)
    ).fetchone()[0]
    assert previous["Coupon p.a. (%)"] == "12.5"

    pg.execute("UPDATE workflow.product SET locked_at = now() WHERE tdcc = %s", (A,))
    locked = draft_of(order_row(fields={"年利率(%)": Decimal(10)}), pg)
    receipt = import_drafts([locked], pg, CONFIG, user="tester")

    assert locked.locked.startswith("已核對通過")
    assert receipt.lines[0].result == LOCKED
    assert product(pg, A)["coupon_pct"] == Decimal("11.25")


@pytest.mark.db
def test_products_from_the_old_inventory_import_are_locked(pg, tmp_path: Path) -> None:
    seed_mailbox(pg)
    import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])]), pg)

    draft = draft_of(order_row(), pg)
    receipt = import_drafts([draft], pg, CONFIG, user="tester")

    assert draft.locked.startswith("以前清單更新匯入的商品")
    assert check_draft(draft, CONFIG) == {"TDCC": draft.locked}
    assert receipt.lines[0].result == LOCKED
    assert product(pg, A)["source"] == "清單更新"


@pytest.mark.db
def test_rereading_keeps_manual_fills_and_shows_database_differences(pg) -> None:
    first = draft_of(order_row(到期日=None), pg)
    first.values["Maturity date"] = dt.date(2027, 4, 22)
    import_drafts([first], pg, CONFIG, user="tester")

    again = draft_of(order_row(到期日=None, fields={"年利率(%)": Decimal(13)}), pg)

    assert again.values["Maturity date"] == dt.date(2027, 4, 22)  # 信裡空白 → 帶入資料庫的值
    assert again.values["Coupon p.a. (%)"] == Decimal(13)  # 兩邊不同 → 預設用信裡的值
    assert again.differences() == {"Coupon p.a. (%)": Decimal("12.5")}
    assert again.edits() == [("Maturity date", None, dt.date(2027, 4, 22))]


@pytest.mark.db
def test_export_after_import_clears_the_fill_of_imported_products(pg, tmp_path: Path) -> None:
    rows = [order_row(到期日=None), order_row(B, issuer="SG", 最終比價日=None)]
    drafts = drafts_from(rows, pg)
    drafts[0].values["Maturity date"] = dt.date(2027, 4, 22)
    receipt = import_drafts(drafts, pg, CONFIG, user="tester")

    path = export_drafts(tmp_path, rows, drafts, tmp_path, (), CONFIG, receipt=receipt)

    ws = openpyxl.load_workbook(path)["Order原始資料"]
    maturity = 2 + [c.header.strip() for c in COLUMNS].index("Maturity date")
    final = 2 + [c.header.strip() for c in COLUMNS].index("Final valuation date")
    assert ws.cell(HEADER_ROW + 1, maturity).fill.fgColor.rgb != FLAG_FILL  # 已匯入
    assert ws.cell(HEADER_ROW + 2, final).fill.fgColor.rgb == FLAG_FILL  # 沒匯入（缺最終比價日）


# ---------- 過渡期：清單更新只補空白（#87 上線前） ----------


@pytest.mark.db
def test_inventory_only_fills_blank_fields_of_order_products(pg, tmp_path: Path) -> None:
    seed_mailbox(pg)
    draft = draft_of(order_row(發行日=None), pg)
    import_drafts([draft], pg, CONFIG, user="tester")

    receipt = import_daily(
        write_daily(
            tmp_path / "d.xlsx",
            [trade(A, "Amy", "北區", 300000), trade(A, "Ben", "南區", 200000, number=None)],
            [inventory(A, [("北區", 300000), ("南區", 200000)], **{"Coupon p.a. (%)": 11})],
        ),
        pg,
    )

    lines = {(line.kind, line.detail) for line in receipt.lines if line.tdcc == A}
    assert (FILLED, "Order 匯入的商品，清單更新補上：ISIN、發行日、分公司配置") in lines
    assert any(kind == DIFFERENT and "coupon_pct" in detail for kind, detail in lines)
    p = product(pg, A)
    assert p["isin"] == "XS0000000001" and p["issue_date"] == dt.date(2026, 10, 22)
    assert p["coupon_pct"] == Decimal("12.5")  # 有值的不覆蓋
    assert pg.execute(
        "SELECT branch, amount FROM workflow.branch_allocation WHERE tdcc = %s ORDER BY seq", (A,)
    ).fetchall() == [
        ("北區", Decimal(300000)),
        ("南區", Decimal(200000)),
    ]


@pytest.mark.db
def test_order_product_without_branches_waits_for_the_inventory(pg, tmp_path: Path) -> None:
    seed_mailbox(pg)
    import_drafts([draft_of(order_row(), pg)], pg, CONFIG, user="tester")

    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], []), pg)

    assert [(line.kind, line.detail) for line in receipt.lines if line.tdcc == A and line.problem] == [
        (INVENTORY_MISSING, "Order 條件已匯入，等清單更新補上分公司配置後再對帳")
    ]


@pytest.mark.db
def test_mail_draft_waits_for_the_issue_date(pg, tmp_path: Path) -> None:
    seed_mailbox(pg)
    import_drafts([draft_of(order_row(發行日=None), pg)], pg, CONFIG, user="tester")
    import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], []), pg)
    root, out = tmp_path / "FCN", tmp_path / "out"
    folder = root / A
    folder.mkdir(parents=True)
    out.mkdir()
    (folder / f"{A}_TS.pdf").write_bytes(b"%PDF ts")
    (folder / f"{A}_IIS.pdf").write_bytes(b"%PDF iis")

    receipt = create_drafts(A, root, out, pg, drafter=lambda draft: None, now=dt.datetime(2026, 10, 21, 9, 0))

    (outcome,) = receipt.outcomes
    assert outcome.status == MAIL_PENDING and "商品資料還沒有發行日（等核對回填）" in outcome.reason
