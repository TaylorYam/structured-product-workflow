"""標的曝險分析（#345、#347、#352、#354、#356、#358、#360）：分類規則、狀態、匯出的「全部商品」。資料庫測試對真的 PostgreSQL；資料全部虛構。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from db_synth import inventory, seed_products
from sp_workflow.exposure import (
    DELIVER,
    FIXED,
    INTACT,
    MANUAL,
    NEAR,
    SAFE,
    UNKNOWN,
    WATCH,
    Holding,
    Leg,
    Position,
    assess,
    export_exposure,
    export_name,
    load_choices,
)
from sp_workflow.lifecycle import KO, AutoEvent, replace_auto_events
from sp_workflow.tabs.exposure import one_month_later

D = Decimal
DAY = dt.date(2026, 10, 28)
UNTIL = dt.date(2026, 11, 30)


def leg(seq: int, code: str, close: str | None, *, initial: str = "100", strike: str = "80", ki: str | None = "65"):
    return Leg(seq, code, D(initial), D(strike), None if ki is None else D(ki), None if close is None else D(close))


def position(barrier: str | None, *legs: Leg, ki_text: str = "", terms: str | None = "完成") -> Position:
    holdings = (Holding("台北", "Amy", D(100000)),)
    return Position("147199990001", "FCN", "XBANK", "USD", UNTIL, barrier, ki_text, terms, None, legs, holdings)


@pytest.mark.parametrize(
    ("barrier", "close", "ki_text", "category"),
    [
        ("NONE", "72", "-", DELIVER),  # 沒有 KI，低於執行價就接股
        ("EKI", "60", "EKI", DELIVER),  # 破下限價
        ("EKI", "70", "EKI", WATCH),  # 低於執行價、還在下限價之上
        ("AKI", "70", "AKI hit 2026-10-01 UL_1 AAA UW", DELIVER),  # 比價已記 AKI 觸及
        ("AKI", "70", "AKI", WATCH),
        ("MKI", "60", "MKI (manual)", MANUAL),
        ("NONE", "80", "-", NEAR),  # 等於執行價不接股，但距執行價 10% 以內
        ("NONE", "88", "-", NEAR),  # 剛好高 10% 還算以內
        ("NONE", "88.1", "-", SAFE),
    ],
)
def test_the_worst_underlying_decides(barrier: str, close: str, ki_text: str, category: str) -> None:
    result = assess(position(barrier, leg(1, "AAA UW", close), leg(2, "BBB UN", "95"), ki_text=ki_text))

    assert result.category == category
    assert result.worst is not None and result.worst.code == "AAA UW"
    assert result.notional == D(100000)


@pytest.mark.parametrize(
    ("barrier", "ki_text", "close", "status"),
    [
        ("EKI", "EKI", "60", "跌破KI"),
        ("EKI", "EKI", "70", "跌破K"),
        ("EKI", "EKI", "85", INTACT),
        ("AKI", "AKI hit 2026-10-01 UL_1 AAA UW", "85", "跌破KI"),  # 比價已記觸及，現在回到執行價之上也算
        ("MKI", "MKI (manual)", "60", "跌破KI"),
        ("NONE", "-", "60", "跌破K"),  # 沒有 KI，不會跌破KI
        ("NONE", "-", "80", INTACT),  # 等於執行價不算跌破
    ],
)
def test_status_is_what_the_worst_close_broke(barrier: str, ki_text: str, close: str, status: str) -> None:
    assert assess(position(barrier, leg(1, "AAA UW", close), ki_text=ki_text)).status == status


def test_status_keeps_fixed_and_unknown() -> None:
    assert INTACT == "", "未跌破留空（#358）"
    p = position("EKI", leg(1, "AAA UW", "60"), ki_text="EKI")
    assert assess(p, UNTIL + dt.timedelta(days=1)).status == FIXED
    assert assess(position("EKI", leg(1, "AAA UW", "60"), terms="待確認")).status == UNKNOWN
    assert assess(position("MKI", leg(1, "AAA UW", None))).status == UNKNOWN, "MKI 缺價格算不出"


def test_distance_to_strike() -> None:
    worst = assess(position("EKI", leg(1, "AAA UW", "72"), ki_text="EKI")).worst

    assert worst is not None and worst.to_strike == D("-0.1")


def test_products_past_the_final_valuation_are_fixed() -> None:
    p = position("NONE", leg(1, "AAA UW", "72"))

    fixed = assess(p, UNTIL + dt.timedelta(days=1))
    assert fixed.category == FIXED, "不用最新收盤價判斷會不會接股"
    assert fixed.worst is not None and fixed.worst.code == "AAA UW"
    assert assess(p, UNTIL).category == DELIVER, "最終比價日當天還用最新收盤價判斷"


def test_worst_is_lowest_performance_not_lowest_price() -> None:
    result = assess(position("NONE", leg(1, "AAA UW", "90"), leg(2, "BBB UN", "150", initial="200", strike="160")))

    assert result.worst is not None and result.worst.code == "BBB UN", "BBB −25% 比 AAA −10% 差"
    assert result.category == DELIVER


def test_missing_prices_cannot_be_assessed() -> None:
    assert assess(position("NONE")).category == UNKNOWN
    missing = assess(position("NONE", leg(1, "AAA UW", None), leg(2, "BBB UN", "95")))
    assert (missing.category, missing.reason) == (UNKNOWN, "缺進場價、執行價或收盤價：AAA UW")


@pytest.mark.parametrize(
    ("p", "reason"),
    [
        (position("NONE", leg(1, "AAA UW", "60"), terms="待確認"), "核對條件待確認"),  # 價格算得出來也不分類
        (position("NONE", leg(1, "AAA UW", "60"), terms=None), "核對條件還沒算"),
        (position(None, leg(1, "AAA UW", "60")), "沒有 KI 型態"),
        (position("EKI", leg(1, "AAA UW", "60", ki=None)), "EKI 沒有下限價：AAA UW"),
    ],
)
def test_incomplete_terms_cannot_be_assessed(p: Position, reason: str) -> None:
    assert (assess(p).category, assess(p).reason) == (UNKNOWN, reason)


def test_mki_is_listed_for_manual_review_even_without_prices() -> None:
    assert assess(position("MKI", leg(1, "AAA UW", None))).category == MANUAL


def test_one_month_later_clips_to_the_month_end() -> None:
    assert one_month_later(dt.date(2026, 10, 29)) == dt.date(2026, 11, 29)
    assert one_month_later(dt.date(2027, 1, 31)) == dt.date(2027, 2, 28)
    assert one_month_later(dt.date(2026, 12, 15)) == dt.date(2027, 1, 15)


def test_default_name_lists_the_tickers() -> None:
    name = export_name(["GOOG UW", "GOOG UQ", "GOOGL UW"], dt.datetime(2026, 10, 30, 9, 5))
    assert name == "標的曝險分析_GOOG_GOOGL_20261030_0905.xlsx"
    trial = export_name(["GOOG UW"], dt.datetime(2026, 10, 30, 9, 5), trial=True)
    assert trial == "標的曝險分析_GOOG_試算_20261030_0905.xlsx", "用了試算價，檔名看得出來"


# ---------------------------------------------------------------- 資料庫

A, B, C, LATE, GONE, OTHER, PENDING, OTHER_EXCHANGE = (f"14719999000{n}" for n in range(1, 9))


def seed(pg, tdcc: str, codes: list[str], *, barrier: str = "NONE", final=dt.date(2026, 11, 20), currency="USD"):
    seed_products(pg, [inventory(tdcc, [("台北", 1)], UL_1=codes[0], UL_2=codes[1], 承作幣別=currency)])
    pg.execute(
        "UPDATE workflow.product SET barrier_type = %s, final_valuation_date = %s, strike_date = '2026-05-20'"
        " WHERE tdcc = %s",
        (barrier, final, tdcc),
    )
    pg.execute("INSERT INTO workflow.check_terms (tdcc, status) VALUES (%s, '完成')", (tdcc,))
    for seq, code in enumerate(codes, start=1):
        pg.execute(
            "INSERT INTO workflow.underlying_price (tdcc, seq, code, initial_price, strike_price, ki_price, ko_price)"
            " VALUES (%s, %s, %s, 100, 80, %s, 100)",
            (tdcc, seq, code, None if barrier == "NONE" else 65),
        )


def trades(pg, tdcc: str, rows: list[tuple[str, str, int]]) -> None:
    (run_id,) = pg.execute(
        "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES ('每日更新.xlsx', 'x', '{}')"
        " RETURNING run_id"
    ).fetchone()
    for seq, (branch, fc, amount) in enumerate(rows, start=1):
        pg.execute(
            "INSERT INTO workflow.trade_row (tdcc, row_seq, trade_date, product, currency, amount, fc_name, branch,"
            " first_run_id) VALUES (%s, %s, '2026-05-20', 'FCN', 'USD', %s, %s, %s, %s)",
            (tdcc, seq, amount, fc, branch, run_id),
        )


def closes(pg, prices: dict[str, str]) -> None:
    (import_id,) = pg.execute(
        "INSERT INTO workflow.close_price_import (source, file_name, file_sha256, imported_by, summary)"
        " VALUES ('歷史收盤價', 'HP.xlsx', 'x', 'tester', '{}') RETURNING import_id"
    ).fetchone()
    for code, close in prices.items():
        for day, value in ((DAY - dt.timedelta(days=1), "100"), (DAY, close)):
            pg.execute(
                "INSERT INTO workflow.close_price (close_date, code, close, import_id) VALUES (%s, %s, %s, %s)",
                (day, code, value, import_id),
            )


def sheet_rows(path: Path, title: str) -> list[dict[str, object]]:
    header, *rows = openpyxl.load_workbook(path)[title].iter_rows(values_only=True)
    return [dict(zip(header, row, strict=True)) for row in rows]


@pytest.fixture
def seeded(pg):
    seed(pg, A, ["AAA UW", "BBB UN"])  # 沒有 KI：AAA −28% 低於執行價 → 會接股
    trades(pg, A, [("台北", "Amy", 100000), ("台中", "Amy", 50000), ("台中", "Ben", 30000)])
    seed(pg, B, ["CCC UN", "AAA UW"], barrier="EKI")  # EKI：AAA −28% 還在下限價之上 → 觀察
    trades(pg, B, [("台中", "Ben", 200000)])
    seed(pg, C, ["AAA UW", "DDD JT"], barrier="EKI", currency="JPY")  # DDD −40% 破下限 → 會接股，最差不是選定標的
    trades(pg, C, [("新竹", "Cara", 10000000)])
    seed(pg, LATE, ["AAA UW", "BBB UN"], final=dt.date(2026, 12, 31))  # 最終比價日晚於截止日
    seed(pg, GONE, ["AAA UW", "BBB UN"])  # 已出場
    replace_auto_events(pg, GONE, [AutoEvent(KO, dt.date(2026, 10, 1))], user="tester")
    seed(pg, OTHER, ["BBB UN", "CCC UN"])  # 沒有連結 AAA
    seed(pg, PENDING, ["AAA UW", "BBB UN"])  # 核對條件待確認：價格有，但不分類
    pg.execute("UPDATE workflow.check_terms SET status = '待確認', reason = '幣別不同' WHERE tdcc = %s", (PENDING,))
    closes(pg, {"AAA UW": "72", "BBB UN": "95", "CCC UN": "110", "DDD JT": "60", "USDJPY": "150"})
    return pg


@pytest.mark.db
def test_choices_merge_exchange_codes_by_ticker(seeded) -> None:
    seed(seeded, OTHER_EXCHANGE, ["AAA UQ", "BBB UN"])
    choices = {c.ticker: c for c in load_choices(seeded)}

    assert sorted(choices) == ["AAA", "BBB", "CCC", "DDD"]
    assert choices["AAA"].codes == ("AAA UW", "AAA UQ"), "連結多的代碼在前"
    assert choices["AAA"].products == 6, "兩個代碼的在庫檔數合計，已出場的不算"
    assert (choices["AAA"].close_date, choices["AAA"].close) == (DAY, D(72)), "收盤價取連結最多的代碼"


@pytest.mark.db
def test_every_exchange_code_of_a_ticker_is_analysed(seeded, tmp_path: Path) -> None:
    seed(seeded, OTHER_EXCHANGE, ["AAA UQ", "BBB UN"])
    closes(seeded, {"AAA UQ": "72"})
    path = tmp_path / "曝險.xlsx"

    export_exposure(seeded, ["AAA UW", "AAA UQ"], UNTIL, path, today=DAY)

    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    assert (every[OTHER_EXCHANGE]["狀態"], every[OTHER_EXCHANGE]["最差表現標的"]) == ("跌破K", "AAA UQ")
    assert {A, B, C, PENDING} < set(every)


@pytest.mark.db
def test_products_past_the_final_valuation_are_listed_apart(seeded, tmp_path: Path) -> None:
    seed(seeded, OTHER_EXCHANGE, ["AAA UW", "BBB UN"], final=DAY - dt.timedelta(days=3))  # 等股價
    path = tmp_path / "曝險.xlsx"

    receipt = export_exposure(seeded, ["AAA UW"], UNTIL, path, today=DAY)

    assert (receipt.counts[FIXED], receipt.counts[DELIVER]) == (1, 2), "AAA 72 也不算會接股"
    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    assert (every[OTHER_EXCHANGE]["狀態"], every[OTHER_EXCHANGE]["剩餘天數"]) == (FIXED, -3)


@pytest.mark.db
def test_export_is_one_sheet_of_every_product(seeded, tmp_path: Path) -> None:
    path = tmp_path / "曝險.xlsx"

    receipt = export_exposure(seeded, ["AAA UW"], UNTIL, path, today=DAY)

    assert receipt.count == 4, "截止日之後、已出場、沒連結的不列"
    assert (receipt.counts[DELIVER], receipt.counts[WATCH], receipt.counts[UNKNOWN]) == (2, 1, 1)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["全部商品"], "只留全部商品（#358）"
    sheet = book["全部商品"]
    assert not sheet._charts
    texts = {str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None}
    assert not [t for t in texts if any(word in t for word in ("估損", "配息", "淨損益"))], "只呈現現況（#352）"

    header, *_ = sheet.iter_rows(values_only=True, max_row=1)
    assert header == (
        *("狀態", "TDCC", "剩餘天數", "最終比價日", "商品", "最差表現標的", "最差表現"),
        *("K%", "KI%", "KI種類", "面額", "幣別", "發行機構"),
        *("分公司1", "FC1", "面額1", "分公司2", "FC2", "面額2", "分公司3", "FC3", "面額3"),
    ), "交易明細最多的 A 有三筆"
    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    a = every[A]
    assert (a["狀態"], a["剩餘天數"], a["商品"], a["KI種類"], a["幣別"], a["發行機構"]) == (
        "跌破K", 23, "FCN", None, "USD", "XBANK"
    )  # fmt: skip
    assert (a["最差表現標的"], a["最差表現"], a["K%"], a["KI%"], a["面額"]) == ("AAA UW", 72, 80, 65, 180000)
    assert [a[f"{name}{n}"] for n in (1, 2, 3) for name in ("分公司", "FC", "面額")] == [
        *("台北", "Amy", 100000),
        *("台中", "Amy", 50000),
        *("台中", "Ben", 30000),
    ], "依面額大到小"
    assert [every[B][name] for name in ("分公司1", "FC1", "面額1", "分公司2")] == ["台中", "Ben", 200000, None]
    assert (every[B]["狀態"], every[C]["狀態"], every[PENDING]["狀態"]) == ("跌破K", "跌破KI", UNKNOWN)
    assert every[C]["最差表現標的"] == "DDD JT", "最差不一定是選定標的"
    assert (every[B]["KI種類"], every[C]["KI種類"]) == ("EKI", "EKI"), "沒有 KI 的才留空（#360）"

    statuses = {sheet.cell(row, 2).value: sheet.cell(row, 1) for row in range(2, sheet.max_row + 1)}
    assert all(statuses[tdcc].fill.fgColor.rgb.endswith("FFC7CE") for tdcc in (A, B, C)), "跌破的狀態格淺紅色"
    assert statuses[PENDING].fill.fill_type is None
    assert sheet.cell(statuses[A].row, 3).fill.fill_type is None, "只填狀態那一欄"
    assert (sheet["D2"].number_format, sheet["G2"].number_format) == ("mm-dd-yy", "0.00"), (
        "簡短日期；最差表現不是百分比"
    )
    assert (sheet["H2"].number_format, sheet["K2"].number_format, sheet["P2"].number_format) == (
        "0.0",
        "#,##0",
        "#,##0",
    )

    colors = [cell.fill.fgColor.rgb[-6:] for cell in sheet[1]]
    blue, green, orange = "B8CCE4", "D7E4BC", "FCD5B4"
    assert colors == [blue] * 5 + [green] * 5 + [blue] * 3 + [orange] * 9, "基本資料、表現、交易明細三種底色"
    cells = [cell for row in sheet.iter_rows(max_row=sheet.max_row) for cell in row]
    assert {(c.alignment.horizontal, c.alignment.vertical) for c in cells} == {("center", "center")}, "全部置中"
    assert round(sheet.column_dimensions["F"].width, 1) == 18.4, "欄寬照範本"


@pytest.mark.db
def test_intact_status_is_blank_and_unfilled(seeded, tmp_path: Path) -> None:
    path = tmp_path / "曝險.xlsx"

    export_exposure(seeded, ["CCC UN"], UNTIL, path, today=DAY)

    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    assert (every[OTHER]["狀態"], every[OTHER]["最差表現標的"]) == (None, "BBB UN"), "BBB −5% 還在執行價之上"
    assert every[OTHER]["分公司1"] == "台北" and every[OTHER]["FC1"] == "-"
    sheet = openpyxl.load_workbook(path)["全部商品"]
    [row] = [r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 2).value == OTHER]
    assert sheet.cell(row, 1).fill.fill_type is None


@pytest.mark.db
def test_products_without_trade_rows_use_the_branch_allocation(seeded, tmp_path: Path) -> None:
    seeded.execute("DELETE FROM workflow.trade_row WHERE tdcc = %s", (B,))
    path = tmp_path / "曝險.xlsx"

    export_exposure(seeded, ["AAA UW"], UNTIL, path, today=DAY)

    b = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}[B]
    assert (b["分公司1"], b["FC1"], b["面額1"], b["分公司2"]) == ("台北", "-", 1, None)


@pytest.mark.db
def test_trial_prices_replace_the_latest_close_of_chosen_codes(seeded, tmp_path: Path) -> None:
    seed(seeded, OTHER_EXCHANGE, ["AAA UW", "BBB UN"], final=DAY - dt.timedelta(days=3))  # 已比價
    path = tmp_path / "曝險.xlsx"
    trial = {"AAA UW": D(50), "DDD JT": D(10)}  # DDD 不是選定標的，不理會

    receipt = export_exposure(seeded, ["AAA UW"], UNTIL, path, today=DAY, trial=trial)

    assert (receipt.counts[DELIVER], receipt.counts[WATCH], receipt.counts[FIXED]) == (3, 0, 1)
    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    assert (every[B]["狀態"], every[B]["最差表現"]) == ("跌破KI", 50), "破下限 65"
    assert (every[C]["最差表現標的"], every[C]["最差表現"]) == ("AAA UW", 50), "AAA −50% 比 DDD −40% 差"
    assert every[OTHER_EXCHANGE]["狀態"] == FIXED, "已比價的不受試算價影響"

    export_exposure(seeded, ["AAA UW", "CCC UN"], UNTIL, path, today=DAY, trial={"AAA UW": D(90)})
    every = {row["TDCC"]: row for row in sheet_rows(path, "全部商品")}
    assert every[A]["狀態"] is None, "AAA 90 還在執行價 80 之上"
    assert (every[C]["最差表現標的"], every[C]["最差表現"]) == ("DDD JT", 60), "DDD 照舊用收盤價"
