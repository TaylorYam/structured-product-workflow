"""本機真實樣本的位置：資料夾結構與檔名只定義在這裡（Issue #128）。

根目錄預設是專案的 `data/`（被 Git 忽略），可用環境變數 `FCN_TEST_DATA_DIR` 換到其他資料夾，底下結構不變：

    orders/核對表格式..xlsx          公司原始待建檔格式的參考條件表，涵蓋樣本內全部上手（回填欄位空白）
    term-sheets/IS/                  樣本內的說明書與投資人須知（<商品代號>_TS.pdf、<商品代號>_IIS.pdf；Issue #300）

樣本外（OOS）的 `term-sheets/OOS/<上手>/` 與 `orders/OOS/` 是驗證用，不在這裡讀、不放進測試斷言。
Super Bull 樣本（BNP Issue #311、CA Issue #312，說明書與投資人須知成對）在 `term-sheets/super-bull/`，或用環境變數 `FCN_SUPER_BULL_DIR`
指到私人雲端倉庫的資料夾。HSBC BEN 樣本（Issue #375，成對）在 `term-sheets/ben/`，或用環境變數 `FCN_BEN_DIR` 指到雲端倉庫。
DACN 樣本（Nomura workflow #368、SG workflow #369、DBS workflow #370、Natixis workflow #371，說明書與投資人須知成對）在 `term-sheets/dacn/`，或用環境變數 `FCN_DACN_DIR` 指到私人雲端倉庫的資料夾。
缺檔時測試略過，略過原因寫出缺的路徑；CI 沒有 data/，一律略過。
"""

from __future__ import annotations

import os
import re
import tomllib
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
DATA = Path(os.environ.get("FCN_TEST_DATA_DIR", ROOT / "data"))
ORDERS = DATA / "orders"
TERM_SHEETS = DATA / "term-sheets" / "IS"
CHECK_FORMAT = ORDERS / "核對表格式..xlsx"
SUPER_BULL = Path(os.environ.get("FCN_SUPER_BULL_DIR", DATA / "term-sheets" / "super-bull"))
BEN = Path(os.environ.get("FCN_BEN_DIR", DATA / "term-sheets" / "ben"))
DACN = Path(os.environ.get("FCN_DACN_DIR", DATA / "term-sheets" / "dacn"))
_REFERENCE_FORMAT = ROOT / "config" / "reference_sheet.toml"
_ISSUER_PREFIXES = ROOT / "config" / "issuer_prefixes.toml"
_SEED_RATINGS = ROOT / "src" / "sp_workflow" / "migrations" / "0024_credit_ratings.sql"


def pdfs(pattern: str = "*.pdf") -> list[Path]:
    return sorted(TERM_SHEETS.glob(pattern)) if TERM_SHEETS.is_dir() else []


def term_sheets(pattern: str = "*.pdf") -> list[Path]:
    """說明書（投資人須知 _IIS 不是說明書）。"""
    return [p for p in pdfs(pattern) if "_IIS" not in p.name]


def requires(*paths: Path) -> pytest.MarkDecorator:
    """缺任何一個就略過。檔名含 `*` 時，資料夾裡至少要有一個符合的檔案。"""
    gone = [p for p in paths if not _present(p)]
    shown = "、".join(_shown(p) for p in gone)
    return pytest.mark.skipif(bool(gone), reason=f"本機沒有真實樣本：缺 {shown}（data/ 被 Git 忽略）")


def _present(path: Path) -> bool:
    if "*" in path.name:
        return path.parent.is_dir() and any(path.parent.glob(path.name))
    return path.is_file()


def _shown(path: Path) -> str:
    try:
        return path.relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def rated_check_format(folder: Path) -> Path:
    """`CHECK_FORMAT` 的複本（寫在 `folder`，不在 data/ 產生檔案），加上信用評等欄（#240）並填初始評等表的值。

    `核對表格式..xlsx` 是舊格式、沒有評等欄（也沒有發行機構欄）；資料庫匯出的參考條件表才有。這裡依 TDCC Code
    前三碼（config/issuer_prefixes.toml）找上手，填 migration 0024 的初始評等，讓真實樣本的評等核對有值可比。
    """
    fmt = tomllib.loads(_REFERENCE_FORMAT.read_text(encoding="utf-8"))
    layout = fmt["layout"]
    prefixes = tomllib.loads(_ISSUER_PREFIXES.read_text(encoding="utf-8"))["prefixes"]
    header_of = {f: h for h, f in fmt["columns"].items() if f.startswith(("issuer_rating_", "guarantor_rating_"))}
    seed = {
        (code, role): dict(zip(("sp", "moodys", "fitch"), values, strict=True))
        for code, role, *values in (
            (m[0], m[1], *(None if v == "NULL" else v.strip("'") for v in m[2:]))
            for m in re.findall(
                r"\('(\w+)', '(發行機構|保證機構)', '1900-01-01', ('[^']*'|NULL), ('[^']*'|NULL), ('[^']*'|NULL)\)",
                _SEED_RATINGS.read_text(encoding="utf-8"),
            )
        )
    }
    wb = openpyxl.load_workbook(CHECK_FORMAT, data_only=True)  # 公式留下算好的值（存檔後不再重算）
    ws = wb.worksheets[0]
    header_row = next(
        r for r in range(1, layout["header_search_rows"] + 1) if any(c.value == layout["key_column"] for c in ws[r])
    )
    key_col = next(c.column for c in ws[header_row] if c.value == layout["key_column"])
    first = ws.max_column + 1
    columns = {std: first + i for i, std in enumerate(header_of)}
    for std, col in columns.items():
        ws.cell(header_row, col, header_of[std])
    for r in range(header_row + 1, ws.max_row + 1):
        code = prefixes.get(str(ws.cell(r, key_col).value or "").strip()[:3])
        for std, col in columns.items():
            role = "發行機構" if std.startswith("issuer_") else "保證機構"
            ws.cell(r, col, seed.get((code, role), {}).get(std.rsplit("_", 1)[1]))
    out = folder / CHECK_FORMAT.name
    wb.save(out)
    return out
