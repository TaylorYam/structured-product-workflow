"""審查標準依上手解析：固定警語、風險等級的期望值（#239）、商品名稱樣板、發行機構名稱、受託機構電話等價寫法。

測試切點是批量核對入口（預覽＋核對）；設定檔問題在載入核對設定時回報；BARC 與 HSBC 都用合成說明書。
"""

from __future__ import annotations

from pathlib import Path
from types import ModuleType

import pytest
from openpyxl.utils import get_column_letter

import barc_synth
import bnp_synth
import hsbc_synth
import ms_synth
import sg_synth
from barc_synth import Spec, build_pdf, check
from fcn_checker.ingestion import IngestionError
from fcn_checker.schema import ItemSource
from harness import (
    MISMATCH,
    PASS,
    REVIEW,
    REVIEW_STANDARD,
    STANDARD,
    check_all,
    check_sheet,
    load_config,
    results,
)
from pdf_writer import Edit
from reference_synth import REFERENCE_HEADERS, build_reference_sheet, reference_row

BARC_WARNING = STANDARD["risk"]["fixed_warning"]
HSBC_WARNING = STANDARD["issuer"]["hsbc"]["fixed_warning"]
PHONE = STANDARD["distributor"]["phone"]
# 設定檢查用的上手表寫法：每家上手一張表後同一個欄位名稱各家都有，取只有該上手才有的寫法（Issue #128）
MS_LEVELS = 'risk_level_formats = ["風險程度等級為{level}", "風險程度為'
MS_IGNORE = 'issuer_name_ignore = ["brackets", "'
MS_OPENING = 'fixed_warning_openings = ["本商品風險程度為{level}'
NOMURA_COUNT = "Issue #239）。\nfixed_warning_occurrences = "


def sub(tmp_path: Path, name: str) -> Path:
    (tmp_path / name).mkdir()
    return tmp_path / name


def standard_with(tmp_path: Path, old: str, new: str) -> Path:
    text = REVIEW_STANDARD.read_text(encoding="utf-8")
    assert old in text
    path = tmp_path / "review_standard.toml"
    path.write_text(text.replace(old, new), encoding="utf-8")
    return path


class _Filled(dict):
    """沒給的佔位符填空字串。"""

    def __missing__(self, key: str) -> str:
        return ""


def check_hsbc(
    tmp_path: Path,
    spec: hsbc_synth.Spec | None = None,
    standard: Path = REVIEW_STANDARD,
    edits: list[Edit] | None = None,
):
    s = spec or hsbc_synth.Spec()
    pdf = hsbc_synth.build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s, edits=edits or ())
    sheet = build_reference_sheet(tmp_path / "order.xlsx", [reference_row(s)])
    return check_sheet(pdf, sheet, load_config(review_standard=standard))


def check_barc(tmp_path: Path, spec: Spec, standard: Path):
    pdf = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    return check_sheet(pdf, sheet, load_config(review_standard=standard))


# ---------------------------------------------------------------- 固定警語


def test_each_issuer_uses_its_own_fixed_warning(tmp_path):
    assert results(check(sub(tmp_path, "barc")), "standard.fixed_warning")[0].status == PASS
    assert results(check_hsbc(sub(tmp_path, "hsbc")), "standard.fixed_warning")[0].status == PASS


def test_hsbc_document_with_barc_warning_is_mismatch(tmp_path):
    report = check_hsbc(tmp_path, edits=[Edit(HSBC_WARNING, BARC_WARNING, "ts")])
    r = results(report, "standard.fixed_warning")[0]
    assert (r.status, r.actual) == (MISMATCH, 0)


def test_barc_document_with_hsbc_warning_is_mismatch(tmp_path):
    r = results(check(tmp_path, edits=[Edit(BARC_WARNING, HSBC_WARNING, "ts")]), "standard.fixed_warning")[0]
    assert (r.status, r.actual) == (MISMATCH, 0)


# ---------------------------------------------------------------- 風險等級的期望值（#239）

RR_RULES = ("standard.risk_level", "standard.fixed_warning")
WITH_LEVEL = [*REFERENCE_HEADERS, "風險等級"]  # 資料庫匯出的參考條件表多一欄風險等級
# 合成文件（預設 RR4）改成 RR3：只換這檔商品的等級，警語裡「RR3、RR4為穩健型」的分類說明不動
RR3_EDITS = {
    "BARC": [Edit("為【RR4】", "為【RR3】"), Edit("程度：RR4", "程度：RR3")],  # 警語【】、投資人須知商品簡介
    "MS": [Edit("為RR4", "為RR3"), Edit("程度： RR4", "程度： RR3")],  # 無【】、開頭句「風險程度為」
    "BNP": [Edit("為[RR4]", "為[RR3]")],  # 半形方括號
    "SG": [Edit("為RR4", "為RR3"), Edit("為[RR4]", "為[RR3]")],  # 無【】，第一章第 2 項半形方括號（另一種警語寫法）
}
SYNTHS: dict[str, ModuleType] = {"BARC": barc_synth, "MS": ms_synth, "BNP": bnp_synth, "SG": sg_synth}


def rr_pair(tmp_path: Path, issuer: str = "BARC", *, level: str | None = None, edits=(), headers=WITH_LEVEL):
    """合成說明書與投資人須知（`edits` 改字）＋參考條件表一列（`level` 寫進風險等級欄；`headers` 不含那欄時是舊 Excel）；
    回傳兩份的 RR 相關結果 {(文件, rule_id, 欄位): 結果}。"""
    synth = SYNTHS[issuer]
    spec = synth.Spec()
    pdf = synth.build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec, edits=edits)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec, **{"風險等級": level})], headers)
    outcome = check_all(sheet, [pdf])  # 旁邊的投資人須知一起核對
    return {
        (r.document.value, r.rule_id, r.field): r
        for i in outcome.items
        for r in i.report.results
        if r.rule_id in RR_RULES
    }


@pytest.mark.parametrize("issuer", sorted(RR3_EDITS))
def test_rr3_documents_pass_against_an_rr3_reference_level_including_the_warnings(tmp_path, issuer):
    """自算 RR3 的商品：文件全文、固定警語（共用與各家版本、其他寫法）與投資人須知都寫 RR3 才通過。"""
    found = rr_pair(tmp_path, issuer, level="RR3", edits=RR3_EDITS[issuer])

    assert {
        ("說明書", "standard.risk_level", "risk_level"),
        ("說明書", "standard.fixed_warning", "fixed_warning"),
    } <= set(found)
    assert {k: r.status for k, r in found.items()} == dict.fromkeys(found, PASS)
    level = found[("說明書", "standard.risk_level", "risk_level")]
    cell = f"樣本清單!{get_column_letter(len(WITH_LEVEL))}4"
    assert (level.expected, level.item.source, level.order_source) == ("RR3", ItemSource.REFERENCE, [cell])


def test_a_reference_level_different_from_the_documents_is_a_mismatch(tmp_path):
    """文件寫 RR4、參考條件表（自算）是 RR3：全文風險等級、固定警語、商品簡介都不一致。"""
    found = rr_pair(tmp_path, level="RR3")

    assert {k: r.status for k, r in found.items()} == dict.fromkeys(found, MISMATCH)
    assert {k[0] for k in found} == {"說明書", "投資人須知"}
    level = found[("說明書", "standard.risk_level", "risk_level")]
    assert (level.expected, level.actual) == ("RR3", ["RR4"])


def test_a_pending_risk_level_requires_review_with_the_reason(tmp_path):
    found = rr_pair(tmp_path, level="待算")

    assert {k: (r.status, r.reason_code) for k, r in found.items()} == dict.fromkeys(
        found, (REVIEW, "risk_level_pending")
    )
    assert all("風險等級待算" in r.message for r in found.values())
    assert len(found) == 5  # 說明書：風險等級、固定警語；投資人須知：風險等級、商品簡介、固定警語


@pytest.mark.parametrize(("level", "reason"), [(None, "order_missing"), ("RR6", "order_invalid")])
def test_a_blank_or_unknown_risk_level_requires_review(tmp_path, level, reason):
    found = rr_pair(tmp_path, level=level)

    assert {(r.status, r.reason_code) for r in found.values()} == {(REVIEW, reason)}


@pytest.mark.parametrize(("edits", "status"), [((), PASS), (RR3_EDITS["BARC"], MISMATCH)])
def test_an_old_reference_sheet_without_the_column_uses_the_review_standard(tmp_path, edits, status):
    """舊 Excel 沒有風險等級欄：照審查標準 risk.level（RR4），同 #239 以前。"""
    found = rr_pair(tmp_path, edits=edits, headers=REFERENCE_HEADERS)

    assert {r.status for r in found.values()} == {status}
    level = found[("說明書", "standard.risk_level", "risk_level")]
    assert (level.expected, level.item.source, level.order_source) == ("RR4", ItemSource.STANDARD, [])


# ---------------------------------------------------------------- 商品名稱樣板


def test_unknown_product_name_placeholder_is_a_batch_config_error(tmp_path):
    standard = standard_with(
        tmp_path,
        'en = "{maxi_en}{daily_en}{memory_en}Autocallable',
        'en = "{maxi_en}{daily_en}{memory_en}{fancy_en}Autocallable',
    )
    with pytest.raises(IngestionError) as raised:
        load_config(review_standard=standard)
    err = raised.value
    assert err.reason_code == "config_invalid"
    assert "fancy_en" in str(err) and "issuer.hsbc.product_name" in str(err)


def test_product_name_placeholders_are_filled_from_the_document_for_any_issuer(tmp_path):
    """名稱樣板的 maxi／daily 佔位符依說明書欄位填入，不依上手：BARC 樣板加上 {daily_en} 也照填。"""
    text = standard_with(
        tmp_path,
        'en = "{tenor} Months {ccy} {memory_en}Autocallable',
        'en = "{tenor} Months {ccy} {daily_en}{memory_en}Autocallable',
    ).read_text(encoding="utf-8")
    path = tmp_path / "with_daily.toml"
    text = text.replace('memory_en = "Memory "\n', 'memory_en = "Memory "\ndaily_en = "Daily "\n', 1)
    path.write_text(text, encoding="utf-8")
    spec = Spec()  # D 型、記憶式
    report = check_barc(tmp_path, spec, path)
    r = results(report, "standard.product_name", "name_en")[0]
    assert r.status == MISMATCH
    assert r.expected.startswith("6 Months USD Daily Memory Autocallable")


@pytest.mark.parametrize(
    ("tenor", "written", "status"),
    [(1, None, MISMATCH), (1, "1 Month USD", PASS), (3, None, PASS), (3, "3 Month USD", MISMATCH)],
    ids=["1-Months", "1-Month", "3-Months", "3-Month"],
)
def test_english_name_is_singular_only_for_a_one_month_tenor(tmp_path, tenor, written, status):
    """天期 1 個月時英文名稱樣板的 Months 改成 Month、只接受單數；其他天期照舊 Months（Issue #343）。"""
    edits = [Edit(f"{tenor} Months USD", written, "ts")] if written else []
    report = check(tmp_path, Spec(tenor=tenor), edits=edits)
    assert results(report, "standard.product_name", "name_en")[0].status == status


@pytest.mark.parametrize(
    ("issuer", "one", "three"),
    [
        ("BARC", "1 Month USD", "3 Months USD"),
        ("MS", "issuance of 1 Month USD", "issuance of 3 Months USD"),
        ("BNP", "1 Month USD", "3 Months USD"),
        ("JPM", "JPMorgan 1Month USD", "JPMorgan 3Months USD"),
        ("GS", "1 Month USD", "3 Months USD"),
        ("CITI", "1-Month USD", "3-Month USD"),
    ],
)
def test_months_templates_take_the_singular_for_one_month(issuer, one, three):
    """所有用 `{tenor} Months`／`{tenor}Months` 的樣板天期 1 時改單數；寫 `-Month` 的上手不受影響（Issue #343）。"""
    tpl = load_config().review_standard.for_issuer(issuer).product_name
    assert one in tpl.template("en", 1).format_map(_Filled(tenor=1, ccy="USD"))
    assert three in tpl.template("en", 3).format_map(_Filled(tenor=3, ccy="USD"))


def test_missing_issuer_name_requires_review(tmp_path):
    standard = standard_with(tmp_path, '\nissuer_name = "香港商', '\n# issuer_name = "香港商')
    report = check_hsbc(tmp_path, standard=standard)
    for r in results(report, "standard.issuer_name"):
        assert (r.status, r.reason_code) == (REVIEW, "standard_missing")


# ---------------------------------------------------------------- 受託機構電話


def test_international_phone_listed_as_equivalent_passes_for_every_issuer(tmp_path):
    assert (
        results(check_hsbc(sub(tmp_path, "hsbc")), "standard.distributor", "distributor_phone_cover")[0].status == PASS
    )
    edits = [Edit(PHONE, "+886-2-5556-1313", "ts.cover.distributor")]
    r = results(check(sub(tmp_path, "barc"), edits=edits), "standard.distributor", "distributor_phone_cover")[0]
    assert (r.status, r.actual) == (PASS, "+886-2-5556-1313")


def test_phone_form_not_listed_in_standard_is_mismatch(tmp_path):
    edits = [Edit(PHONE, "02.5556.1313", "ts.cover.distributor")]
    r = results(check(sub(tmp_path, "barc"), edits=edits), "standard.distributor", "distributor_phone_cover")[0]
    assert r.status == MISMATCH

    standard = standard_with(
        tmp_path,
        'phone_equivalents = ["+886-2-5556-1313", "+886 2 5556 1313", "+886 (2) 5556 1313", "(+886)2-5556-1313", "(02)5556-1313"]',
        "phone_equivalents = []",
    )
    r = results(
        check_hsbc(sub(tmp_path, "hsbc"), standard=standard), "standard.distributor", "distributor_phone_cover"
    )[0]
    assert r.status == MISMATCH


# ---------------------------------------------------------------- MS 分節的設定檢查（Issue #135）


@pytest.mark.parametrize(
    "old,new,where",
    [
        ('basket_en = "Worst of Shares and/or ETFs"\n', "", "basket_en"),
        (f'{MS_LEVELS}{{level}}"', f'{MS_LEVELS}RRn"', "issuer.ms.risk_level_formats"),
        (f"{MS_IGNORE}trailing_period", f"{MS_IGNORE}commas", "issuer.ms.issuer_name_ignore"),
        (f'{MS_OPENING}。"]', f'{MS_OPENING}"]', "issuer.ms.fixed_warning_openings"),
        (
            f'{MS_OPENING}。"]',
            'fixed_warning_openings = ["本商品風險程度等級為{level}。"]',
            "issuer.ms.fixed_warning_openings",
        ),
        # SG 分節（Issue #194）：依幣別的倍數、固定警語其他寫法、英文名稱幣別代碼對照
        (
            "minimum_subscription = { USD = 10, JPY = 1, CNH = 1 }",
            "minimum_subscription = { USD = 10, JPY = 0, CNH = 1 }",
            "issuer.sg.min_amount_multiples",
        ),
        (
            'autocall_en = "Autocallable "\nccy_aliases = { CNH = "CNY" }',  # SG 才有這兩行相連（BNP、GS 也有 ccy_aliases）
            'autocall_en = "Autocallable "\nccy_aliases = { CNH = 1 }',
            "issuer.sg.product_name.ccy_aliases",
        ),
        # Nomura 分節（Issue #239）：上手別說明書固定警語次數須是正整數
        (f"{NOMURA_COUNT}2\n", f"{NOMURA_COUNT}0\n", "issuer.nomura.fixed_warning_occurrences"),
        (f"{NOMURA_COUNT}2\n", f'{NOMURA_COUNT}"2"\n', "issuer.nomura.fixed_warning_occurrences"),
        # UBS 分節（Issue #279）：名稱用了 {ko_observation_*} 就要有每日與每期兩種寫法
        ('periodic_zh = "定期觀察"\n', "", "periodic_zh"),
        # BNP 依商品種類的名稱樣板（#311）：鍵須是支援的商品種類寫法
        ('[issuer.bnp.product_names."Super Bull"]', '[issuer.bnp.product_names."Super Bear"]', "Super Bear"),
    ],
)
def test_invalid_ms_review_standard_is_a_batch_config_error(tmp_path, old, new, where):
    with pytest.raises(IngestionError) as raised:
        load_config(review_standard=standard_with(tmp_path, old, new))
    assert raised.value.reason_code == "config_invalid" and where in str(raised.value)


def test_issuer_name_ignores_apply_only_to_the_configured_issuer(tmp_path):
    """BARC 沒有設定 issuer_name_ignore：第二章少了括號仍是不一致。"""
    plain = STANDARD["issuer"]["barc"]["issuer_name"].replace("（", "").replace("）", "")
    r = results(
        check(tmp_path, edits=[Edit(STANDARD["issuer"]["barc"]["issuer_name"], plain, "ts.ch二.1")]),
        "standard.issuer_name",
        "issuer_name_ch2",
    )[0]
    assert r.status == MISMATCH


def test_address_form_not_listed_in_standard_is_mismatch(tmp_path):
    """MS 第二章沒有「松山區」的地址靠審查標準 address_equivalents 才算相符；清單拿掉就判不一致。"""
    variant = [
        Edit("◎ 營業所在地：台北市松山區民生東路三段158號6樓", "◎ 營業所在地：台北市民生東路三段158號6樓", "ts.ch二")
    ]
    r = results(
        ms_synth.check(sub(tmp_path, "listed"), edits=variant), "standard.distributor", "distributor_address_ch2"
    )[0]
    assert r.status == PASS and "address_equivalents" in r.tolerance
    standard = standard_with(
        tmp_path,
        'address_equivalents = ["台北市民生東路三段158號6樓", "臺北市105民生東路三段158號6樓", "台北市松山區民生東路3段158號6樓"]',
        "address_equivalents = []",
    )
    r = results(
        ms_synth.check(sub(tmp_path, "unlisted"), edits=variant, config=load_config(review_standard=standard)),
        "standard.distributor",
        "distributor_address_ch2",
    )[0]
    assert r.status == MISMATCH
