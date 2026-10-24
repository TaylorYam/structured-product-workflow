"""共用小工具（#139）：SHA-256 與簡單的 Excel 表、匯出用的 `save_rows`（#204 從 test_product_export 搬來）。資料全部虛構。"""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import openpyxl
import pytest

from sp_workflow.file_tools import (
    FileWriteError,
    save_book,
    save_rows,
    sha256_bytes,
    sha256_file,
    table_book,
    tdcc_as_text,
)


def test_sha256_of_a_file_and_of_bytes(tmp_path: Path) -> None:
    data = b"fictional" * 300_000  # 比讀檔的區塊大
    path = tmp_path / "a.bin"
    path.write_bytes(data)

    assert sha256_file(path) == sha256_bytes(data) == hashlib.sha256(data).hexdigest()


def test_table_book_writes_bold_headers_widths_and_tdcc_as_text(tmp_path: Path) -> None:
    path = tmp_path / "表.xlsx"
    book = table_book("清單", ["TDCC", "說明"], [["000123456789", "x"]], widths=(16, 40), style=tdcc_as_text)
    save_book(book, path)

    sheet = openpyxl.load_workbook(path).active
    assert sheet.title == "清單"
    assert [c.value for c in sheet[1]] == ["TDCC", "說明"]
    assert all(c.font.bold for c in sheet[1])
    assert sheet["A2"].value == "000123456789"
    assert sheet["A2"].number_format == "@"
    assert sheet["B2"].number_format == "General"
    assert (sheet.column_dimensions["A"].width, sheet.column_dimensions["B"].width) == (16, 40)
    assert sheet.freeze_panes == "A2"
    assert not list(tmp_path.glob("*.tmp"))


@pytest.mark.skipif(sys.platform != "win32", reason="不覆蓋靠 Windows 的 os.rename（目的地已存在時失敗）")
def test_save_without_overwrite_keeps_the_existing_file(tmp_path: Path) -> None:
    path = tmp_path / "表.xlsx"
    path.write_bytes(b"old")

    with pytest.raises(OSError):
        save_book(table_book("清單", ["TDCC"], []), path, overwrite=False)

    assert path.read_bytes() == b"old"
    assert not list(tmp_path.glob("*.tmp"))


def test_a_failed_save_leaves_no_temp_file_and_keeps_the_old_one(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "表.xlsx"
    path.write_bytes(b"old")

    def broken_replace(self, target):
        raise PermissionError("檔案開在 Excel")

    monkeypatch.setattr(Path, "replace", broken_replace)
    with pytest.raises(PermissionError):
        save_book(table_book("清單", ["TDCC"], [["1"]]), path)

    assert path.read_bytes() == b"old"
    assert not list(tmp_path.glob("*.tmp"))


# ---------------------------------------------------------------- save_rows（#204）


def read_back(path: Path) -> list[list[object]]:
    sheet = openpyxl.load_workbook(path).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_text_that_looks_like_a_formula_stays_text(tmp_path: Path) -> None:
    path = tmp_path / "x.xlsx"

    save_rows(path, "t", ("tdcc", "note"), [("029199990001", "=1+1")])

    assert read_back(path) == [["tdcc", "note"], ["029199990001", "=1+1"]]


def test_control_characters_are_dropped(tmp_path: Path) -> None:
    path = tmp_path / "x.xlsx"

    save_rows(path, "t", ("tdcc", "note"), [("029199990001", "KO\x01 2026")])

    assert read_back(path)[1] == ["029199990001", "KO 2026"]


def test_a_failed_write_keeps_the_old_file_and_leaves_no_temp_file(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "x.xlsx"
    path.write_bytes(b"old")

    def locked(_source, _target):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(Path, "replace", locked)

    with pytest.raises(FileWriteError, match="同名檔開在 Excel 時請先關掉"):
        save_rows(path, "t", ("tdcc",), [("029199990001",)])

    assert path.read_bytes() == b"old" and list(tmp_path.iterdir()) == [path]


def test_an_unexpected_error_leaves_no_temp_file(tmp_path: Path, monkeypatch) -> None:
    def broken(self, filename):
        Path(filename).write_bytes(b"half")
        raise ValueError("壞掉")

    monkeypatch.setattr(openpyxl.Workbook, "save", broken)

    with pytest.raises(ValueError):
        save_rows(tmp_path / "x.xlsx", "t", ("tdcc",), [])

    assert list(tmp_path.iterdir()) == []


def test_a_missing_folder_says_where(tmp_path: Path) -> None:
    with pytest.raises(FileWriteError, match="寫不進"):
        save_rows(tmp_path / "沒有這個資料夾" / "x.xlsx", "t", ("tdcc",), [])
