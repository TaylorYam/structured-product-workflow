"""收盤價（`sp_workflow.close_prices`，#153、ADR 0020）：代碼正規化、休市日、MarketData CSV 與歷史收盤價匯入。

價格與假日都是虛構的（代碼只借用常見的寫法）；假日表用測試自己組的 `TradingCalendar`。
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from db_synth import csv_price, price, write_daily, write_price_csv
from fcn_checker.trading_calendar import Market, TradingCalendar
from sp_workflow import close_prices
from sp_workflow.close_prices import (
    HISTORY_ERROR_BOOK,
    ClosePriceError,
    HistoryProblem,
    closes_from_csv,
    import_history,
    is_closed,
    is_uncovered,
)
from sp_workflow.daily_import import import_daily, import_price_csv
from sp_workflow.sheet_cells import bbg_code

YEAR = (dt.date(2026, 1, 1), dt.date(2026, 12, 31))
US_HOLIDAY = dt.date(2026, 10, 14)  # 虛構的休市日（週三）
JP_HOLIDAY = dt.date(2026, 10, 13)  # 虛構的休市日（週二）
CALENDAR = TradingCalendar(
    1,
    (
        Market("US", "美國", frozenset({"UN", "UW"}), *YEAR, frozenset({US_HOLIDAY})),
        Market("JP", "日本", frozenset({"JT"}), *YEAR, frozenset({JP_HOLIDAY})),
    ),
)
DAYS = [dt.date(2026, 10, d) for d in range(12, 17)]  # 週一到週五
NOW = dt.datetime(2026, 10, 22, 9, 30)


def count(pg, table: str) -> int:
    return pg.execute(f"SELECT count(*) FROM {table}").fetchone()[0]


def closes_in(pg) -> list[tuple]:
    return pg.execute("SELECT close_date, code, close FROM workflow.close_price ORDER BY code, close_date").fetchall()


# ---------------------------------------------------------------- 不碰資料庫


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        ("AAA UW Equity", "AAA UW"),
        (" 7203 JT equity ", "7203 JT"),
        ("USDJPY Curncy", "USDJPY"),
        ("SPX Index", "SPX"),
        ("BRK/B UN Equity", "BRK/B UN"),
        ("AAA UW", "AAA UW"),
        ("Equity", "Equity"),
        ("AAA  UW Equity", "AAA  UW"),
    ],
)
def test_codes_drop_only_the_trailing_class_word(raw: str, code: str) -> None:
    assert bbg_code(raw) == code


@pytest.mark.parametrize(
    ("code", "day", "closed"),
    [
        ("AAA UW", dt.date(2026, 10, 17), True),  # 週六
        ("AAA UW", US_HOLIDAY, True),
        ("7203 JT", US_HOLIDAY, False),
        ("AAA UW", JP_HOLIDAY, False),
        ("BMW GY", dt.date(2026, 10, 12), False),  # 交易所不在假日表：照存，比價時標待確認
        ("AAA UW", dt.date(2027, 10, 14), False),  # 超出涵蓋範圍：不能確認，另外處理（is_uncovered）
    ],
)
def test_only_confirmed_closed_days_are_closed(code: str, day: dt.date, closed: bool) -> None:
    assert is_closed(CALENDAR, code, day) is closed


@pytest.mark.parametrize(
    ("code", "day", "uncovered"),
    [
        ("AAA UW", dt.date(2027, 10, 14), True),
        ("AAA UW", dt.date(2027, 10, 16), False),  # 週六：本來就休市
        ("AAA UW", dt.date(2026, 10, 15), False),
        ("BMW GY", dt.date(2027, 10, 14), False),  # 交易所不在假日表：照存
    ],
)
def test_dates_outside_the_calendar_cannot_be_confirmed(code: str, day: dt.date, uncovered: bool) -> None:
    assert is_uncovered(CALENDAR, code, day) is uncovered


def test_csv_closes_outside_the_calendar_are_not_stored() -> None:
    day = dt.date(2027, 1, 4)

    found = closes_from_csv(day, {"AAA UW": Decimal("1"), "BMW GY": Decimal("2")}, CALENDAR)

    assert found.closes == {(day, "BMW GY"): Decimal("2")} and found.uncovered == ("AAA UW",)


def test_csv_closes_skip_closed_markets_and_conflicts() -> None:
    found = closes_from_csv(
        JP_HOLIDAY,
        {
            "AAA UW Equity": Decimal("101"),
            "7203 JT Equity": Decimal("2000"),
            "CCC UW": Decimal("5"),
            "CCC UW Equity": Decimal("6"),
            "DDD UN": None,
            "EEE UN": Decimal("7"),
        },
        CALENDAR,
        conflicted=("FFF UN Equity",),
    )

    assert found.closes == {(JP_HOLIDAY, "AAA UW"): Decimal("101"), (JP_HOLIDAY, "EEE UN"): Decimal("7")}
    assert found.closed == ("7203 JT",)
    assert found.conflicts == ("CCC UW",)


# ---------------------------------------------------------------- MarketData CSV


def test_a_price_csv_also_stores_closes_on_the_csv_date(pg, tmp_path: Path) -> None:
    source = write_price_csv(
        tmp_path / "MarketData_20261013.csv",
        [csv_price("AAA UW Equity", "101.25", day="2026-10-13"), csv_price("7203 JT Equity", "2000", day="2026-10-13")],
    )

    receipt = import_price_csv(source, pg, CALENDAR)

    assert closes_in(pg) == [(JP_HOLIDAY, "AAA UW", Decimal("101.25"))]
    assert receipt.closes is not None and (receipt.closes.stored, receipt.closes.codes) == (1, 1)
    assert receipt.closed == ("7203 JT",), "日本休市那天 CSV 帶的是前一天的價格，不存"
    assert pg.execute("SELECT DISTINCT price_date FROM workflow.daily_price").fetchall() == [(dt.date(2026, 10, 14),)]
    assert pg.execute("SELECT code FROM workflow.daily_price ORDER BY code").fetchall() == [
        ("7203 JT Equity",),
        ("AAA UW Equity",),
    ], "daily_price 照舊存原本的代碼"
    (import_row,) = pg.execute("SELECT source, file_name, run_id, summary FROM workflow.close_price_import").fetchall()
    assert import_row[:3] == ("MarketData CSV", "MarketData_20261013.csv", receipt.run_id)
    assert import_row[3]["closed"] == ["7203 JT"] and import_row[3]["stored"] == 1


def test_codes_that_clash_after_dropping_the_class_word_are_not_stored(pg, tmp_path: Path) -> None:
    source = write_price_csv(
        tmp_path / "MarketData_20261015.csv",
        [
            csv_price("AAA UW", "101"),
            csv_price("AAA UW Equity", "102"),
            csv_price("BBB UW", "1"),
            csv_price("BBB UW", "2"),
        ],
    )

    receipt = import_price_csv(source, pg, CALENDAR)

    assert closes_in(pg) == []
    assert receipt.close_conflicts == ("AAA UW",), "BBB UW 原本就是股價衝突，已經列在股價的問題裡"
    assert count(pg, "workflow.daily_price") == 3


def test_reimporting_a_csv_replaces_the_close_and_keeps_the_old_value(pg, tmp_path: Path) -> None:
    day = dt.date(2026, 10, 15)
    import_price_csv(
        write_price_csv(tmp_path / "a.csv", [csv_price("AAA UW", "101"), csv_price("BBB UW", "5")]), pg, CALENDAR
    )
    first = pg.execute("SELECT import_id FROM workflow.close_price_import").fetchone()[0]

    receipt = import_price_csv(
        write_price_csv(tmp_path / "b.csv", [csv_price("AAA UW", "102"), csv_price("BBB UW", "5")]), pg, CALENDAR
    )

    assert closes_in(pg) == [(day, "AAA UW", Decimal("102")), (day, "BBB UW", Decimal("5"))]
    assert receipt.closes is not None and (receipt.closes.replaced, receipt.closes.changed) == (2, 1)
    assert pg.execute(
        "SELECT close_date, code, previous_close, previous_import_id, import_id FROM workflow.close_price_change"
    ).fetchall() == [(day, "AAA UW", Decimal("101"), first, receipt.closes.import_id)]


def test_without_a_holiday_calendar_closes_are_not_stored(pg, tmp_path: Path) -> None:
    receipt = import_price_csv(write_price_csv(tmp_path / "a.csv", [csv_price("AAA UW")]), pg)

    assert receipt.closes is None
    assert count(pg, "workflow.close_price") == 0 and count(pg, "workflow.close_price_import") == 0
    assert count(pg, "workflow.daily_price") == 1


def test_bbg_data_never_writes_closes(pg, tmp_path: Path) -> None:
    import_daily(write_daily(tmp_path / "每日更新.xlsx", [], [price("AAA UW Equity")]), pg)

    assert count(pg, "workflow.daily_price") == 1
    assert count(pg, "workflow.close_price") == 0


# ---------------------------------------------------------------- 歷史收盤價（HP 格式）


def write_history(path: Path, rows: list[list], dates: list = DAYS) -> Path:
    """HP 格式：第 3 列從 E 欄起是收盤日，第 4 列起每列一檔（B 簡碼、C 帶類別字的彭博代碼）。"""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["虛構的歷史收盤價"])
    sheet.append([])
    sheet.append([None, "簡碼", "代碼", "名稱", *dates])
    for row in rows:
        sheet.append([None, *row])
    book.save(path)
    return path


HISTORY = [
    ["AAA", "AAA UW Equity", "虛構", 10, 11, "#N/A N/A", 13, "#N/A Requesting Data..."],
    ["7203", "7203 JT Equity", "虛構", 2000, 2010, 2020, "#N/A N/A", "-"],
    ["BMW", "BMW GY Equity", "虛構", 50, "#N/A N/A", 52, "亂碼ÿ", "#N/A Invalid Security"],
    ["AAA", "AAA UW", "重複", 10, 99, None, None, None],
    [None, None, None, None, None, None, None, None],
    ["USDJPY", "USDJPY Curncy", "虛構", 150.1, 150.2, 150.3, 150.4, 150.5],
    ["XYZ", "XYZ UN Equity", "失效公式", "=NA()", "#REF!", None, None, None],
]


def test_history_is_checked_as_a_whole_and_problems_are_listed(pg, tmp_path: Path) -> None:
    source = write_history(tmp_path / "HP.xlsx", HISTORY)

    receipt = import_history(source, pg, CALENDAR, user="tester", now=NOW)

    d12, d13, d14, d15, d16 = DAYS
    assert closes_in(pg) == [
        (d12, "7203 JT", Decimal("2000")),
        (d14, "7203 JT", Decimal("2020")),
        (d12, "AAA UW", Decimal("10")),
        (d15, "AAA UW", Decimal("13")),
        (d12, "BMW GY", Decimal("50")),
        (d14, "BMW GY", Decimal("52")),
        *((d, "USDJPY", Decimal(v)) for d, v in zip(DAYS, ("150.1", "150.2", "150.3", "150.4", "150.5"), strict=True)),
    ]
    assert (receipt.first_date, receipt.last_date, receipt.codes) == (d12, d16, 5)
    assert (receipt.summary.stored, receipt.missing, receipt.closed) == (11, 11, 1)
    assert {(p.code, p.reason, p.first, p.last, p.cells) for p in receipt.problems} == {
        ("7203 JT", "交易日沒有價格（#N/A N/A）", d15, d15, 1),
        ("7203 JT", "沒有價格（-）", d16, d16, 1),
        ("AAA UW", "代碼重複、同一天的值不同，沒有存", d13, d13, 1),
        ("AAA UW", "抓價未完成（#N/A Requesting Data...），請重抓再匯", d16, d16, 1),
        ("BMW GY", "亂碼或不是數字", d15, d15, 1),
        ("BMW GY", "代碼無效（#N/A Invalid Security）", d16, d16, 1),
        ("XYZ UN", "沒有價格（空白）", d12, d16, 3),
        ("XYZ UN", "錯誤值或失效公式（#REF!）", d13, d13, 1),
    }, "美國休市日的 #N/A N/A、交易所不在假日表的 #N/A N/A 不列"
    assert receipt.error_book_path == tmp_path / HISTORY_ERROR_BOOK.format(stamp="20261022-093000")
    sheet = openpyxl.load_workbook(receipt.error_book_path).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("代碼", "原因", "第一天", "最後一天", "格數") and len(rows) == 9
    (summary,) = pg.execute(
        "SELECT source, file_name, imported_by, summary FROM workflow.close_price_import"
    ).fetchall()
    assert summary[:3] == ("歷史收盤價", "HP.xlsx", "tester")
    assert (summary[3]["missing"], summary[3]["closed"], summary[3]["problems"]) == (11, 1, 8)


def test_a_row_without_a_code_is_listed_as_a_whole_row(pg, tmp_path: Path) -> None:
    rows = [HISTORY[5], ["QQQ", None, "沒有代碼", 1, 2, None, None, None]]

    receipt = import_history(write_history(tmp_path / "HP.xlsx", rows), pg, CALENDAR, now=NOW)

    assert receipt.problems == (HistoryProblem("第 5 列", "C 欄沒有彭博代碼，這列不匯入（B 欄：QQQ）", None, None, 2),)
    assert receipt.codes == 1 and receipt.summary.stored == 5


def test_a_dash_is_not_a_code(pg, tmp_path: Path) -> None:
    rows = [HISTORY[5], ["-", "-", "-", "-", "-", "-", "-", "-"]]

    receipt = import_history(write_history(tmp_path / "HP.xlsx", rows), pg, CALENDAR, now=NOW)

    assert receipt.codes == 1 and receipt.missing == 0
    assert [p.code for p in receipt.problems] == ["第 5 列"]


def test_history_outside_the_calendar_is_listed_not_stored(pg, tmp_path: Path) -> None:
    days = [dt.date(2025, 12, 24), dt.date(2025, 12, 25), dt.date(2026, 1, 2)]
    rows = [["AAA", "AAA UW Equity", "虛構", 10, 10, 11], ["BMW", "BMW GY Equity", "虛構", 5, 5, 6]]

    receipt = import_history(write_history(tmp_path / "HP.xlsx", rows, days), pg, CALENDAR, now=NOW)

    assert closes_in(pg) == [
        (dt.date(2026, 1, 2), "AAA UW", Decimal("11")),
        *((d, "BMW GY", Decimal(v)) for d, v in zip(days, ("5", "5", "6"), strict=True)),
    ], "2025 年不在虛構假日表的涵蓋範圍：美股不能確認是交易日；交易所不在假日表的照存"
    assert receipt.uncovered == 2
    assert [(p.code, p.first, p.last, p.cells) for p in receipt.problems] == [
        ("AAA UW", dt.date(2025, 12, 24), dt.date(2025, 12, 25), 2)
    ]


def test_a_failed_write_leaves_no_error_list(pg, tmp_path: Path, monkeypatch) -> None:
    def broken(*args, **kwargs):
        raise RuntimeError("資料庫寫入失敗")

    monkeypatch.setattr(close_prices, "write_closes", broken)
    source = write_history(tmp_path / "HP.xlsx", HISTORY)

    with pytest.raises(RuntimeError):
        import_history(source, pg, CALENDAR, now=NOW)

    assert list(tmp_path.iterdir()) == [source]


def test_a_clean_history_writes_no_error_list(pg, tmp_path: Path) -> None:
    receipt = import_history(write_history(tmp_path / "HP.xlsx", [HISTORY[5]]), pg, CALENDAR, now=NOW)

    assert receipt.problems == () and receipt.error_book_path is None
    assert list(tmp_path.iterdir()) == [tmp_path / "HP.xlsx"]


def test_reimporting_history_replaces_and_keeps_changed_values(pg, tmp_path: Path) -> None:
    import_history(write_history(tmp_path / "HP.xlsx", [HISTORY[5]]), pg, CALENDAR, now=NOW)
    changed = ["USDJPY", "USDJPY Curncy", "虛構", 150.1, 150.2, 150.3, 150.4, 151]

    receipt = import_history(write_history(tmp_path / "HP2.xlsx", [changed]), pg, CALENDAR, now=NOW)

    assert (receipt.summary.replaced, receipt.summary.changed) == (5, 1)
    assert pg.execute("SELECT code, previous_close FROM workflow.close_price_change").fetchall() == [
        ("USDJPY", Decimal("150.5"))
    ]
    assert closes_in(pg)[-1] == (DAYS[-1], "USDJPY", Decimal("151"))


@pytest.mark.parametrize(
    ("dates", "rows", "message"),
    [
        ([DAYS[0], "收盤日"], [HISTORY[5]], "F3 應為收盤日，實際是「收盤日」"),
        ([DAYS[0], DAYS[0]], [HISTORY[5]], "收盤日 2026-10-12 出現兩次"),
        ([], [HISTORY[5]], "第 3 列從 E 欄起沒有收盤日"),
        (DAYS, [], "沒有任何代碼"),
    ],
)
def test_a_badly_shaped_history_writes_nothing(pg, tmp_path: Path, dates, rows, message) -> None:
    source = write_history(tmp_path / "HP.xlsx", rows, dates)

    with pytest.raises(ClosePriceError, match=message):
        import_history(source, pg, CALENDAR, now=NOW)

    assert count(pg, "workflow.close_price_import") == 0
    assert list(tmp_path.iterdir()) == [source]


def test_a_history_that_cannot_be_opened_writes_nothing(pg, tmp_path: Path) -> None:
    source = tmp_path / "HP.xlsx"
    source.write_text("不是 Excel", encoding="utf-8")

    with pytest.raises(ClosePriceError, match="打不開歷史收盤價"):
        import_history(source, pg, CALENDAR, now=NOW)

    assert count(pg, "workflow.close_price_import") == 0
