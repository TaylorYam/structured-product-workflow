"""Order 條件讀取對照 10 月清單（#52）：只在本機設定 SPW_ORDER_CHECK_DIR 時執行，用真的 Outlook 開信，CI 自動略過。

把 2026-10-20 的人工比對變成回歸測試：重跑私人樣本倉庫裡每個 `Order問題信件_*` 回報包，
格式表逐欄對照作業人員確認過的 10 月清單，抓「無聲錯誤」——格子沒標淺紅底，但和清單不同（含已讀取列漏值）。
已標淺紅底的格子，以及整列有問題（TDCC 那格標紅）的待確認列裡留空的格子交給人工，不算錯。

真實資料不進 repo；這裡只有欄位對應與允許清單（TDCC 與欄名）。
"""

from __future__ import annotations

import datetime as dt
import os
import re
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from db_synth import BARRIER_TYPES, KO_TYPES  # ELN待建檔格式的代碼 → 格式表的寫法
from sp_workflow.order_template import COLUMNS as TEMPLATE_COLUMNS
from sp_workflow.order_terms import PENDING, read_orders

pytestmark = [
    pytest.mark.real_samples,
    pytest.mark.skipif(sys.platform != "win32", reason="需要 Windows 與傳統版 Outlook"),
]

ANSWER_FILE = "10月清單(用來核對order專案).xlsx"  # ELN待建檔格式，第 1 列表頭
PACK_GLOB = "Order問題信件_*"

# 格式表欄 → 10 月清單欄（值直接比）；KO Type、Barrier Type、價格另外換算
LIST_COLUMNS = {
    "Product": "Product",
    "Currency": "計價幣別",
    "Guaranteed Periods (m)": "Non-Call(月)",
    **{f"BBG Code {n}": f"UL_{n}" for n in range(1, 6)},
    "Strike (%)": "K(%)",
    "KO Barrier (%)": "KO(%)",
    "Coupon p.a. (%)": "Coupon p.a. (%)",
    "Tenor (m)": "天期(月)",
    "KI Barrier (%)": "KI(%)",
    "Trade date": "交易日",
    "Issue date": "發行日",
    "Final valuation date": "最終比價日",
    "Maturity date": "到期日",
    "期初價格": "期初定價",
}
PRICE_COLUMN = "Upfront / NotePrice (%)"  # = 100 − UF

_UBS_DATES = "UBS Order 原文的日期與最終資料不同（發行日、最終比價日、到期日三個一起）"
# Order 原文與最終資料本來就不同（清單是事後更正過的值），程式照原文讀是對的
# （BNP、MS 兩筆來自公司全量比對，目前樣本回報包沒有這兩檔，留著讓之後補進樣本時不誤報）
KNOWN = {
    ("082001154045", "BBG Code 3"): "BNP Order 原文的 BBG 交易所代碼與最終資料不同",
    ("147000029714", "BBG Code 3"): "MS Order 原文的 BBG 交易所代碼與最終資料不同",
    ("148000011430", "BBG Code 3"): "GS Order 原文的 BBG 交易所代碼與最終資料不同",
    ("148000011443", "BBG Code 4"): "GS Order 原文的 BBG 交易所代碼與最終資料不同",
    ("092001742900", "Issue date"): _UBS_DATES,
    ("092001742900", "Final valuation date"): _UBS_DATES,
    ("092001742900", "Maturity date"): _UBS_DATES,
}
# 樣本回報包裡本來就讀不出格式表列的商品；其他檔讀不出列就算退步
NO_TEMPLATE = {
    "328000195113": "樣本商品資料夾裡沒有 (Order) 原信",
    "328000195218": "樣本商品資料夾裡沒有 (Order) 原信",
}

_COLUMN_INDEX = {c.header.strip(): i for i, c in enumerate(TEMPLATE_COLUMNS)}
_TDCC_INDEX = 0  # 整列的問題標在 TDCC 那格


def _check_dir() -> Path:
    value = os.environ.get("SPW_ORDER_CHECK_DIR")
    if not value or not Path(value).is_dir():
        pytest.skip("沒有設定 SPW_ORDER_CHECK_DIR")
    return Path(value)


def _norm(value: object) -> object:
    """兩邊統一成可比較的值：`-` 與空白 = None、日期去掉時間、數字去掉浮點誤差。"""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, str):
        value = value.strip()
        return None if value in ("", "-") else value
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(round(float(value), 6))).normalize()
    return value


def _code(value: object) -> str:
    return str(value or "").strip()


def _expected(answer: dict[str, object]) -> dict[str, object]:
    want = {column: _norm(answer.get(name)) for column, name in LIST_COLUMNS.items()}
    want["KO Type"] = KO_TYPES.get((_code(answer.get("KO(Freq)")), _code(answer.get("KO(memo)"))))
    want["Barrier Type"] = BARRIER_TYPES.get(_code(answer.get("KI(Freq)")))
    uf = _norm(answer.get("UF"))
    want[PRICE_COLUMN] = _norm(100 - uf) if isinstance(uf, Decimal) else None
    return want


def _answers(path: Path) -> dict[str, list[dict[str, object]]]:
    """TDCC → 清單換算後的期望值（同一檔可能因通路不同出現多列，商品條件應該相同）。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        found: dict[str, list[dict[str, object]]] = {}
        for row in rows:
            answer = dict(zip(header, row, strict=False))
            tdcc = str(answer.get("TDCC Code") or "").strip()
            if tdcc:
                found.setdefault(tdcc, []).append(_expected(answer))
        return found
    finally:
        wb.close()


def test_order_terms_match_october_list(tmp_path: Path) -> None:
    root = _check_dir()
    answer_path = root / ANSWER_FILE
    assert answer_path.is_file(), f"{root} 底下沒有 {ANSWER_FILE}"
    answers = _answers(answer_path)
    packs = sorted(p for p in root.glob(PACK_GLOB) if p.is_dir())
    assert packs, f"{root} 底下沒有 {PACK_GLOB} 回報包"

    problems: list[str] = []
    used: set[tuple[str, str]] = set()
    templated: set[str] = set()
    for pack in packs:
        tdccs = sorted(p.name for p in pack.iterdir() if p.is_dir() and re.fullmatch(r"\d{12}", p.name))
        if not tdccs:
            continue
        out = tmp_path / pack.name
        out.mkdir()
        receipt = read_orders("\n".join(tdccs), pack, out)
        for row in receipt.rows:
            if not row.template:  # 找不到 Order、未支援上手：沒有格式表列
                if row.tdcc not in NO_TEMPLATE:
                    problems.append(f"{pack.name} {row.tdcc} [{row.status}] 讀不出格式表列：{row.reasons}")
                continue
            templated.add(row.tdcc)
            wants = answers.get(row.tdcc)
            if not wants:
                problems.append(f"{pack.name} {row.tdcc} 10 月清單沒有這檔，無法核對")
                continue
            want = wants[0]
            conflicts = [c for c in want if any(w[c] != want[c] for w in wants)]
            if conflicts:
                problems.append(f"{row.tdcc} 在 10 月清單有多列但條件不一致，無法核對：{'、'.join(conflicts)}")
                continue
            row_flagged = _TDCC_INDEX in row.flagged
            for column, expected in want.items():
                index = _COLUMN_INDEX[column]
                got = _norm(row.template[index])
                if got == expected or index in row.flagged:
                    continue
                if (row.tdcc, column) in KNOWN:
                    used.add((row.tdcc, column))
                    continue
                if got is None and row.status == PENDING and row_flagged:  # 整列待人工補
                    continue
                problems.append(
                    f"{pack.name} {row.tdcc} [{row.status}] {column}：程式 {got!r}，清單 {expected!r}（沒標淺紅底）"
                )

    # 允許清單過時（樣本裡有這檔，但那格已經對了或已標紅）要拿掉，免得之後的錯誤被放過
    problems += [f"KNOWN 過時：{t} {c}" for t, c in KNOWN if t in templated and (t, c) not in used]
    problems += [f"NO_TEMPLATE 過時：{t} 已讀出格式表列" for t in NO_TEMPLATE if t in templated]
    assert templated, "沒有任何一檔讀出格式表列（Outlook 是否能開信？）"
    assert not problems, f"有 {len(problems)} 個問題：\n" + "\n".join(problems)
