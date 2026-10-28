"""回補既有的 FCN 以外商品（#303、ADR 0027）：參考條件表格式的 Excel → 鎖住的商品條件、價格、比價日與分公司配置。

資料庫測試對真的 PostgreSQL；參考條件表、商品、價格與日期全部虛構。
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from fcn_checker.check_config import ConfigPaths
from fcn_checker.config import load_reference_format
from fcn_checker.orders.reference import headers_of
from order_synth import CONFIG
from sp_workflow.product_backfill import EXISTING, IMPORTED, SKIPPED, ProductBackfillError, backfill_products
from sp_workflow.product_terms import read
from sp_workflow.window_actions import import_other_products

CONFIG_DIR = Path(__file__).resolve().parents[1] / "config"
FMT = load_reference_format(CONFIG_DIR / "reference_sheet.toml")
PATHS = ConfigPaths.with_fallback(CONFIG_DIR / "review_standard.toml")
BRANCHES = [name for n in range(1, 8) for name in (f"分公司_{n}", f"分公司_{n}_承作金額")]
HEADERS = [*dict.fromkeys(headers_of(FMT).values()), "發行機構", "UF", *BRANCHES]  # UF 不是參考條件表的標準欄位
A, B, C, D = (f"02919999000{i}" for i in range(1, 5))  # BARC
E = "029199990005"  # 測試設定只有 BARC 等幾家；商品種類和上手無關
TRADE = dt.date(2026, 4, 28)
DAY_1 = dt.date(2026, 5, 6)  # 發行日後的第一個觀察日
FINAL = dt.date(2026, 10, 5)


def express(tdcc: str, **cells) -> dict:
    """舊總表的一檔 FCN_Express：Non-Call 1、比價日_1 是第一個觀察日、最後一期是最終比價日，兩檔標的、沒有 KI。"""
    row = {
        "Product": "FCN_Express",
        "發行機構": "Barclays",
        "TDCC Code": tdcc,
        "ISIN Code": "XS0000000001",
        "單位面額": 10000,
        "承作幣別": "USD",
        "分公司_1": "新莊",
        "分公司_1_承作金額": 1500000,
        "分公司_2": "台中",
        "分公司_2_承作金額": 200000,
        **{f"分公司_{n}{suffix}": "-" for n in range(3, 8) for suffix in ("", "_承作金額")},
        "交易日": TRADE,
        "發行日": dt.date(2026, 5, 5),
        **{f"比價日_{n}": "-" for n in range(1, 13)},
        "比價日_1": DAY_1,
        "比價日_5": FINAL,
        "最終比價日": FINAL,
        "到期日": dt.date(2026, 10, 8),
        "期初定價": "收盤價",
        "KO(%)": 100,
        "KO(Freq)": "D",
        "KO(memo)": "Y",
        "K(%)": 82.62,
        "KI(%)": "-",
        "KI(Freq)": "-",
        "UF": 1.57,
        "Coupon p.a. (%)": 10,
        "天期(月)": 5,
        "Non-Call(月)": 1,
        "UL_1": "AAA UW",
        "UL_1_進場價": 259.7,
        "UL_1_執行價": 214.56414,
        "UL_1_下限價": "-",
        "UL_1_KO價": 259.7,
        "UL_2": "BBB UN",
        "UL_2_進場價": 349.78,
        "UL_2_執行價": 288.988236,
        "UL_2_下限價": "-",
        "UL_2_KO價": 349.78,
        **{f"UL_{i}{suffix}": "-" for i in range(3, 6) for suffix in ("", "_進場價", "_執行價", "_下限價", "_KO價")},
    }
    row.update(cells)
    return row


def super_bull(tdcc: str) -> dict:
    """舊總表的一檔 Super Bull（#310）：每期觀察非記憶式、沒有 KI，第 3 個月比一次，比價日_6 是最終比價日。"""
    row = express(tdcc, Product="Super Bull", **{"KO(Freq)": "P", "KO(memo)": "N", "K(%)": 70})
    row.update({"KO(%)": 70, "天期(月)": 6, "Non-Call(月)": 3, "比價日_1": "-", "比價日_5": "-"})
    row.update({"比價日_3": DAY_1, "比價日_6": FINAL, "Coupon p.a. (%)": 8.5})
    return row


def write_sheet(path: Path, rows: list[dict], headers: list[str] = HEADERS) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    book.save(path)
    return path


def lines(receipt) -> dict[str, tuple[str, str]]:
    return {line.tdcc: (line.result, line.detail) for line in receipt.lines}


def other_product(pg, *tdccs: str) -> None:
    """以前每日匯入把 Express 當成其他商品略過、記下來的 TDCC。"""
    (run_id,) = pg.execute(
        "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES ('每日更新.xlsx', 'x', '{}')"
        " RETURNING run_id"
    ).fetchone()
    for tdcc in tdccs:
        pg.execute(
            "INSERT INTO workflow.other_product (tdcc, product, first_run_id) VALUES (%s, 'FCN_Express', %s)",
            (tdcc, run_id),
        )


@pytest.mark.db
def test_an_express_row_becomes_a_locked_product_with_prices_and_dates(pg, tmp_path: Path) -> None:
    other_product(pg, A)
    source = write_sheet(
        tmp_path / "FCN外商品數量.xlsx",
        [express(A), {**express(B), "Product": "FCN"}, {**express(C), "Product": "BEN"}],
    )

    receipt = backfill_products(pg, source, FMT, CONFIG)

    assert lines(receipt) == {A: (IMPORTED, "")}
    assert receipt.ignored == 2, "FCN 與不支援的商品種類不在範圍內，只計數"
    stored = read(pg, A)
    assert stored is not None and stored.locked_at is not None and stored.source == "歷史回補"
    columns = stored.columns
    assert (columns["product"], columns["non_call_months"], columns["issuer"]) == ("FCN_Express", 0, "Barclays")
    assert (stored.isin, columns["denomination"]) == ("XS0000000001", Decimal(10000))
    assert columns["strike_date"] == TRADE, "舊表沒有期初定價日：照舊商品存交易日"
    assert (columns["ko_type"], columns["ko_freq"], columns["ko_memory"]) == ("Daily Memory", "D", "Y")
    assert (columns["barrier_type"], columns["ki_freq"], columns["ki_pct"]) == ("NONE", "-", None)
    assert (columns["upfront_pct"], columns["strike_pct"], columns["tenor_months"]) == (
        Decimal("1.57"),
        Decimal("82.62"),
        5,
    )
    assert pg.execute(
        "SELECT seq, code, initial_price, strike_price, ki_price, ko_price FROM workflow.underlying_price"
        " WHERE tdcc = %s ORDER BY seq",
        (A,),
    ).fetchall() == [
        (1, "AAA UW", Decimal("259.7"), Decimal("214.56414"), None, Decimal("259.7")),
        (2, "BBB UN", Decimal("349.78"), Decimal("288.988236"), None, Decimal("349.78")),
    ]
    dates = dict(pg.execute("SELECT slot, value FROM workflow.autocall_date WHERE tdcc = %s", (A,)).fetchall())
    assert len(dates) == 12 and (dates[1], dates[5], dates[2]) == (DAY_1, FINAL, None)
    assert pg.execute(
        "SELECT branch, amount FROM workflow.branch_allocation WHERE tdcc = %s ORDER BY seq", (A,)
    ).fetchall() == [("新莊", Decimal(1500000)), ("台中", Decimal(200000))]
    assert pg.execute("SELECT status FROM workflow.check_terms WHERE tdcc = %s", (A,)).fetchone() == ("完成",)
    assert pg.execute("SELECT count(*) FROM workflow.other_product WHERE tdcc = %s", (A,)).fetchone() == (0,)
    assert receipt.error_book_path is None


@pytest.mark.db
def test_a_super_bull_row_keeps_its_non_call_and_dates(pg, tmp_path: Path) -> None:
    """Super Bull 照舊表匯入（#310）：Non-Call、比價日（含比價日_{天期} 的最終比價日）不轉換，出場紅利存 Coupon 欄。"""
    source = write_sheet(tmp_path / "x.xlsx", [super_bull(E)])

    receipt = backfill_products(pg, source, FMT, CONFIG)

    assert lines(receipt) == {E: (IMPORTED, "")}
    columns = read(pg, E).columns
    assert (columns["product"], columns["non_call_months"], columns["ko_type"]) == ("Super Bull", 3, "Period End")
    assert (columns["barrier_type"], columns["coupon_pct"]) == ("NONE", Decimal("8.5"))
    dates = dict(pg.execute("SELECT slot, value FROM workflow.autocall_date WHERE tdcc = %s", (E,)).fetchall())
    assert {slot: day for slot, day in dates.items() if day} == {3: DAY_1, 6: FINAL}


@pytest.mark.db
def test_a_blank_initial_pricing_is_imported_and_mentioned(pg, tmp_path: Path) -> None:
    source = write_sheet(tmp_path / "x.xlsx", [express(A, 期初定價=None)])

    receipt = backfill_products(pg, source, FMT, CONFIG)

    result, detail = lines(receipt)[A]
    assert result == IMPORTED and "期初定價空白" in detail
    assert read(pg, A).columns["initial_pricing"] is None


@pytest.mark.db
def test_existing_products_are_left_alone_and_a_rerun_changes_nothing(pg, tmp_path: Path) -> None:
    source = write_sheet(tmp_path / "x.xlsx", [express(A), express(B)])
    backfill_products(pg, source, FMT, CONFIG)
    pg.execute("UPDATE workflow.product SET coupon_pct = 11 WHERE tdcc = %s", (A,))

    receipt = backfill_products(pg, source, FMT, CONFIG)

    assert {line.result for line in receipt.lines} == {EXISTING}
    assert read(pg, A).columns["coupon_pct"] == Decimal(11), "資料庫已有的不覆蓋"
    assert pg.execute("SELECT count(*) FROM workflow.underlying_price WHERE tdcc = %s", (A,)).fetchone() == (2,)


@pytest.mark.parametrize(
    ("rows", "reason"),
    [
        ([express(A), express(A, 期初定價="開盤價")], "表上有 2 列"),
        ([express(A, **{"K(%)": "八十"})], "Strike (%)"),
        ([express(A, 比價日_3="2026/7/6")], "比價日_3 不是日期"),
        ([express(A, 比價日_1="-")], "比價日_1"),
        ([express(A, UL_2_進場價="-")], "UL_2 進場價是空白"),
        ([express(A, **{"KI(Freq)": "AM"})], "KI Barrier (%)"),
        ([express(A, 發行機構="SG")], "發行機構"),
        ([express("02919999001", **{"TDCC Code": "02919999001"})], "12 碼"),
        ([express(A, Product=None)], "Product 空白"),
    ],
)
@pytest.mark.db
def test_a_bad_row_is_skipped_and_listed_while_others_import(pg, tmp_path: Path, rows, reason) -> None:
    source = write_sheet(tmp_path / "FCN外商品數量.xlsx", [*rows, express(D)])

    receipt = backfill_products(pg, source, FMT, CONFIG, now=dt.datetime(2026, 10, 28, 9, 0, 0))

    bad = [line for line in receipt.lines if line.tdcc != D]
    assert len(bad) == 1 and bad[0].result == SKIPPED and reason in bad[0].detail, bad
    assert lines(receipt)[D] == (IMPORTED, "")
    if len(bad[0].tdcc) == 12:
        assert read(pg, bad[0].tdcc) is None
    assert receipt.error_book_path == tmp_path / "其他商品回補錯誤_20261028-090000.xlsx"
    book = openpyxl.load_workbook(receipt.error_book_path).active
    assert [c.value for c in book[1]] == ["TDCC", "原因", "表上的列"]
    assert book.cell(row=2, column=1).value == bad[0].tdcc


@pytest.mark.db
def test_a_file_without_required_columns_imports_nothing(pg, tmp_path: Path) -> None:
    source = write_sheet(tmp_path / "x.xlsx", [express(A)], [h for h in HEADERS if h != "K(%)"])

    with pytest.raises(ProductBackfillError, match="整份不匯入"):
        backfill_products(pg, source, FMT, CONFIG)

    assert read(pg, A) is None


@pytest.mark.db
def test_the_window_action_observes_after_importing(pg, pg_settings, tmp_path: Path) -> None:
    source = write_sheet(tmp_path / "x.xlsx", [express(A)])

    result = import_other_products(pg_settings, source, PATHS)

    assert result.receipt.count(IMPORTED) == 1
    assert result.observation.receipt is not None and result.observation.error == ""
