"""每日匯入（#32、#35、#87）：虛構的每日更新（交易明細＋BBG_Data），對真的 PostgreSQL 匯入；
信箱直接寫進資料庫。資料全部虛構。"""

from __future__ import annotations

import datetime as dt
import json
import re
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
import pytest

from db_synth import (
    PRICE_DATE,
    csv_price,
    inventory,
    price,
    seed_mailbox,
    seed_products,
    trade,
    write_daily,
    write_price_csv,
)
from sp_workflow import daily_import
from sp_workflow.daily_import import (
    BLOCKED,
    CANCELLED,
    DIFFERENT,
    MAILBOX_EMPTY,
    NAME_UNMATCHED,
    NEW_TRADES,
    PRICE_CONFLICT,
    PRICE_CSV_HEADERS,
    PRICE_HEADERS,
    PRICE_NOT_NUMBER,
    TRADE_HEADERS,
    UNCHANGED,
    DailyImportError,
    import_daily,
    import_price_csv,
    latest_import,
)
from sp_workflow.file_tools import FileWriteError
from sp_workflow.mailbox import is_pb

pytestmark = pytest.mark.db

A = "029199990001"
B = "037199990002"
C = "076199990003"
D = "325199990004"


@pytest.fixture
def mailbox(pg) -> None:
    seed_mailbox(pg)


def kinds(receipt, tdcc: str) -> list[str]:
    return [line.kind for line in receipt.lines if line.tdcc == tdcc]


def count(pg, table: str) -> int:
    return pg.execute(f"SELECT count(*) FROM {table}").fetchone()[0]


def test_new_trades_and_people_are_written(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(
        tmp_path / "每日更新.xlsx",
        [
            trade(A, "Amy", "北區", 300000),
            trade(A, "Ben", "南區", 100000, number=None),
            trade(B, "Ben", "南區", 500000),
        ],
        [price("AAA UW")],
    )

    receipt = import_daily(daily, pg)

    assert kinds(receipt, A) == [NEW_TRADES] and kinds(receipt, B) == [NEW_TRADES]
    assert not [line for line in receipt.lines if line.problem]
    assert count(pg, "workflow.product") == 0, "#87 起商品條件只由匯入起單資料建立"
    assert pg.execute(
        "SELECT row_seq, number, fc_name, amount FROM workflow.trade_row WHERE tdcc = %s ORDER BY row_seq", (A,)
    ).fetchall() == [(1, 1, "Amy", Decimal("300000")), (2, None, "Ben", Decimal("100000"))]
    summary = pg.execute("SELECT daily_file_name, summary FROM workflow.import_run").fetchone()
    assert summary[0] == "每日更新.xlsx" and summary[1]["counts"][NEW_TRADES] == 2
    assert (summary[1]["price_date"], summary[1]["price_count"]) == ("2026-10-16", 1)


def test_tdcc_loses_no_leading_zero(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(int(A), "Amy", "北區", 1)])

    import_daily(daily, pg)

    assert pg.execute("SELECT DISTINCT tdcc FROM workflow.trade_row").fetchall() == [(A,)]


def test_other_products_are_skipped_and_counted(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(
        tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1), trade(D, "Amy", "北區", 1, product="FCN_Stepdown")]
    )

    receipt = import_daily(daily, pg)

    assert receipt.skipped_trade_rows == 1
    assert count(pg, "workflow.trade_row") == 1 and kinds(receipt, D) == []
    assert pg.execute("SELECT tdcc, product FROM workflow.other_product").fetchall() == [(D, "FCN_Stepdown")]


def test_express_rows_are_imported_like_fcn(pg, tmp_path: Path, mailbox) -> None:
    """FCN_Express 和 FCN 一樣寫進交易明細列（#302）：商品種類照設定的寫法存，列給起單資料頁，不記成其他商品。"""
    daily = write_daily(
        tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1, product="fcn_express"), trade(B, "Amy", "北區", 1)]
    )

    receipt = import_daily(daily, pg)

    assert (receipt.skipped_trade_rows, receipt.fcn_trade_rows, receipt.fcn_tdccs) == (0, 2, (A, B))
    assert pg.execute("SELECT tdcc, product FROM workflow.trade_row ORDER BY tdcc").fetchall() == [
        (A, "FCN_Express"),
        (B, "FCN"),
    ]
    assert count(pg, "workflow.other_product") == 0


@pytest.mark.parametrize(("written", "stored"), [("SUPER BULL", "Super Bull"), (" ben ", "BEN")])
def test_super_bull_and_ben_rows_are_imported_like_fcn(pg, tmp_path: Path, mailbox, written, stored) -> None:
    """Super Bull（#310）、BEN（#373）和 FCN 一樣寫進交易明細列：大小寫、空白不同也照設定的寫法存，不記成其他商品。"""
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1, product=written)])

    receipt = import_daily(daily, pg)

    assert (receipt.skipped_trade_rows, receipt.fcn_tdccs) == (0, (A,))
    assert pg.execute("SELECT tdcc, product FROM workflow.trade_row").fetchall() == [(A, stored)]
    assert count(pg, "workflow.other_product") == 0


def test_latest_import_describes_the_newest_run(pg, tmp_path: Path, mailbox) -> None:
    assert latest_import(pg) is None
    import_daily(write_daily(tmp_path / "1.xlsx", []), pg)
    import_daily(write_daily(tmp_path / "2.xlsx", []), pg)

    info = latest_import(pg)

    assert (info.run_id, info.file_name, len(info.sha256)) == (2, "2.xlsx", 64)


@pytest.mark.parametrize("name", ["PB", " pb ", "Pb"])
def test_pb_is_recognized_however_it_is_typed(name) -> None:
    assert is_pb(name)


def test_importing_the_same_file_again_adds_nothing(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)])
    import_daily(daily, pg)

    receipt = import_daily(daily, pg)

    assert kinds(receipt, A) == [UNCHANGED]
    assert (count(pg, "workflow.trade_row"), count(pg, "workflow.import_run")) == (1, 2)


def test_changed_trades_are_reported_not_overwritten(pg, tmp_path: Path, mailbox) -> None:
    import_daily(write_daily(tmp_path / "1.xlsx", [trade(A, "Amy", "北區", 1)]), pg)

    receipt = import_daily(write_daily(tmp_path / "2.xlsx", [trade(A, "Amy", "北區", 2)]), pg)

    assert [(line.kind, line.detail) for line in receipt.lines if line.tdcc == A] == [
        (DIFFERENT, "交易列和資料庫不同（資料庫 1 列、這次 1 列）")
    ]
    assert pg.execute("SELECT amount FROM workflow.trade_row").fetchone() == (Decimal("1"),)


def test_a_bad_trade_row_blocks_all_rows_of_that_product(pg, tmp_path: Path, mailbox) -> None:
    bad = trade(A, "Ben", "南區", 1, number=None)
    bad["金額"] = None
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1), bad, trade(B, "Amy", "北區", 1)])

    receipt = import_daily(daily, pg)

    assert kinds(receipt, A) == [BLOCKED] and "金額 沒有值" in receipt.lines[0].detail
    assert pg.execute("SELECT DISTINCT tdcc FROM workflow.trade_row").fetchall() == [(B,)]


def test_cancel_note_is_skipped(pg, tmp_path: Path, mailbox) -> None:
    cancelled = trade(A, "Ben", "南區", 100000, number=None)
    cancelled["備註"] = "此筆取消"
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 300000), cancelled])

    receipt = import_daily(daily, pg)

    assert receipt.cancelled_trade_rows == 1
    assert CANCELLED in kinds(receipt, A)
    assert pg.execute(
        "SELECT row_seq, fc_name, amount FROM workflow.trade_row WHERE tdcc = %s ORDER BY row_seq", (A,)
    ).fetchall() == [(1, "Amy", Decimal("300000"))]


def test_fcn_tdccs_are_listed_for_the_order_tab_in_first_seen_order(pg, tmp_path: Path, mailbox) -> None:
    """#90：依第一次出現的順序、不重複；非 FCN、整檔都取消的不列；有一列沒取消就列。
    #93：只列剛好 12 碼的 TDCC；存成數字、前面的 0 掉了的照樣匯入資料庫，但不列。"""
    cancelled_b = trade(B, "Amy", "北區", 1)
    cancelled_b["備註"] = "Cancel"
    cancelled_c = trade(C, "Amy", "北區", 1, number=None)
    cancelled_c["備註"] = "此筆取消"
    daily = write_daily(
        tmp_path / "d.xlsx",
        [
            trade(C, "Amy", "北區", 1),
            trade(D, "Amy", "北區", 1, product="FCN_Stepdown"),
            trade(A, "Amy", "北區", 1),
            cancelled_b,
            cancelled_c,
            trade("14700000000X", "Amy", "北區", 1),
            trade(A, "Ben", "南區", 1, number=None),
            trade(int(B), "Ben", "南區", 1),
        ],
    )

    receipt = import_daily(daily, pg)

    assert receipt.fcn_tdccs == (C, A)
    assert receipt.fcn_trade_rows == 7, "非 FCN 不算，取消與 TDCC 不完整的列都算"
    assert NEW_TRADES in kinds(receipt, B), "前面的 0 掉了的 TDCC 匯入資料庫照舊"


def test_reimporting_the_same_file_still_lists_its_fcn_tdccs(pg, tmp_path: Path, mailbox) -> None:
    """#90：資料庫已有的（「已存在」）照樣列，因為讀 Order 不靠資料庫。"""
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)])
    import_daily(daily, pg)

    again = import_daily(daily, pg)

    assert UNCHANGED in kinds(again, A) and again.fcn_tdccs == (A,)


def test_no_fcn_trades_leaves_nothing_to_hand_over(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(D, "Amy", "北區", 1, product="FCN_Stepdown")])

    receipt = import_daily(daily, pg)

    assert (receipt.fcn_trade_rows, receipt.fcn_tdccs) == (0, ())


def test_unmatched_fc_names_are_listed_and_pb_is_not(pg, tmp_path: Path, mailbox) -> None:
    """匯入只列出信箱對不到的名字；比對規則本身（大小寫、別名）在 test_mailbox 測（#201）。"""
    daily = write_daily(
        tmp_path / "d.xlsx",
        [trade(A, "Amy", "北區", 1), trade(C, "Zoe", "北區", 1), trade(D, "PB", "北區", 1)],
    )

    receipt = import_daily(daily, pg)

    assert [(line.tdcc, line.detail) for line in receipt.lines if line.kind == NAME_UNMATCHED] == [(C, "Zoe")]


def test_import_does_not_change_the_mailbox(pg, tmp_path: Path, mailbox) -> None:
    before = pg.execute("SELECT name, email, unit, fixed_cc FROM mailbox.person ORDER BY name").fetchall()

    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)]), pg)

    assert pg.execute("SELECT name, email, unit, fixed_cc FROM mailbox.person ORDER BY name").fetchall() == before
    assert receipt.people == 3


def test_empty_mailbox_is_reported_once(pg, tmp_path: Path) -> None:
    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)]), pg)

    assert [line.kind for line in receipt.lines if line.problem] == [MAILBOX_EMPTY]
    assert count(pg, "workflow.trade_row") == 1


def test_mailbox_keys_are_generated_by_the_database(pg) -> None:
    seed_mailbox(pg, [(" Amy ", "amy@example.com", "FC", False)], aliases=[("Amie", " Amy ")])

    assert pg.execute("SELECT name_key FROM mailbox.person").fetchone() == ("amy",)
    assert pg.execute("SELECT alias_key FROM mailbox.alias").fetchone() == ("amie",)


@pytest.mark.parametrize(
    ("sql", "params"),
    [
        ("INSERT INTO mailbox.person (name, email) VALUES (%s, %s)", ("amy", "other@example.com")),
        ("INSERT INTO mailbox.person (name, email) VALUES (%s, %s)", ("Cat", "not-an-email")),
        ("INSERT INTO mailbox.alias (alias, person_name) VALUES (%s, %s)", ("Zed", "Nobody")),
        ("INSERT INTO mailbox.alias (alias, person_name) VALUES (%s, %s)", ("AMIE", "Ben")),
    ],
)
def test_database_rejects_bad_mailbox_rows(pg, sql, params) -> None:
    import psycopg

    seed_mailbox(pg, aliases=[("Amie", "Amy")])

    with pytest.raises(psycopg.errors.IntegrityError):
        pg.execute(sql, params)


def test_renaming_a_person_keeps_their_aliases(pg) -> None:
    seed_mailbox(pg, aliases=[("Amie", "Amy")])

    pg.execute("UPDATE mailbox.person SET name = 'Amelia' WHERE name = 'Amy'")

    assert pg.execute("SELECT alias, person_name FROM mailbox.alias").fetchall() == [("Amie", "Amelia")]


def test_wrong_header_rejects_the_whole_file(pg, tmp_path: Path, mailbox) -> None:
    headers = list(TRADE_HEADERS)
    headers[3], headers[4] = headers[4], headers[3]
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], trade_headers=headers)

    with pytest.raises(DailyImportError, match="D 欄應為「TDCC」，實際是「幣別」"):
        import_daily(daily, pg)

    assert count(pg, "workflow.import_run") == 0


@pytest.mark.parametrize("missing", [daily_import.TRADE_SHEET, daily_import.PRICE_SHEET])
def test_missing_sheet_rejects_the_whole_file(pg, tmp_path: Path, mailbox, missing) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")])
    wb = openpyxl.load_workbook(daily)
    del wb[missing]
    wb.save(daily)

    with pytest.raises(DailyImportError, match=f"缺少工作表「{missing}」"):
        import_daily(daily, pg)

    assert count(pg, "workflow.import_run") == 0


def test_failure_halfway_leaves_nothing_behind(pg, tmp_path: Path, mailbox, monkeypatch) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")])

    def broken(*_args):
        raise RuntimeError("中途失敗")

    monkeypatch.setattr(daily_import, "_unmatched_names", broken)
    with pytest.raises(RuntimeError):
        import_daily(daily, pg)

    tables = ("workflow.trade_row", "workflow.import_run", "workflow.price_import", "workflow.daily_price")
    assert [count(pg, t) for t in tables] == [0, 0, 0, 0]


# ---------------------------------------------------------------- 清單更新停用（#87）


def test_an_inventory_sheet_is_ignored(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")])
    wb = openpyxl.load_workbook(daily)
    wb.create_sheet("清單更新").append(["這張工作表不讀"])
    wb.save(daily)

    receipt = import_daily(daily, pg)

    assert kinds(receipt, A) == [NEW_TRADES] and count(pg, "workflow.product") == 0


def test_trades_of_products_already_in_the_database_are_imported_as_before(pg, tmp_path: Path, mailbox) -> None:
    """已在資料庫的商品（以前清單更新、或 Order 匯入建立的）：交易照樣匯入，商品條件與分公司配置不動。"""
    seed_products(pg, [inventory(A, [("北區", 300000)])])

    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "南區", 50000)]), pg)

    assert [line.kind for line in receipt.lines if line.tdcc == A] == [NEW_TRADES], "不再分公司對帳"
    assert pg.execute("SELECT branch, amount FROM workflow.branch_allocation").fetchall() == [
        ("北區", Decimal("300000.0000"))
    ]


def test_old_inventory_results_are_not_listed_as_errors(pg, tmp_path: Path, mailbox) -> None:
    """錯誤清單會重建歷史：以前清單更新的成功結果（新增商品、補上空白）不列成錯誤。"""
    old = {"lines": [[A, "新增商品", ""], [A, "補上空白", "ISIN"], [B, "庫存尚未更新", "x"]]}
    pg.execute(
        "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES ('舊.xlsx', %s, %s::jsonb)",
        ("0" * 64, json.dumps(old, ensure_ascii=False)),
    )

    import_daily(write_daily(tmp_path / "d.xlsx", []), pg)

    error_book = openpyxl.load_workbook(tmp_path / "每日匯入錯誤.xlsx", read_only=True, data_only=True)
    assert [row[3:5] for row in list(error_book.active.values)[1:]] == [(B, "庫存尚未更新")]


def test_an_error_book_that_cannot_be_written_stops_the_import_and_leaves_no_file(pg, tmp_path: Path, mailbox) -> None:
    """#207：錯誤清單寫不進去時整次不匯入（同以前），也不留下暫存檔。"""
    (tmp_path / "每日匯入錯誤.xlsx").mkdir()  # 同名的資料夾：換名換不過去

    with pytest.raises(FileWriteError, match="錯誤清單寫不進去，整份不匯入"):
        import_daily(write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)]), pg)

    assert count(pg, "workflow.import_run") == 0 and count(pg, "workflow.trade_row") == 0
    assert sorted(p.name for p in tmp_path.iterdir()) == ["d.xlsx", "每日匯入錯誤.xlsx"]
    assert (tmp_path / "每日匯入錯誤.xlsx").is_dir(), "原本的同名檔（這裡是資料夾）不動"


def failing_commit(conn, monkeypatch) -> None:
    """最外層的 transaction 在本體做完之後才失敗，像 commit 失敗一樣整個 rollback。"""
    real, depth = conn.transaction, [0]

    @contextmanager
    def transaction(*args, **kwargs):
        depth[0] += 1
        try:
            with real(*args, **kwargs) as tx:
                yield tx
                if depth[0] == 1:
                    raise psycopg.OperationalError("commit 失敗（測試）")
        finally:
            depth[0] -= 1

    monkeypatch.setattr(conn, "transaction", transaction)


def test_a_failed_commit_puts_the_previous_error_book_back(pg, tmp_path: Path, mailbox, monkeypatch) -> None:
    """#207：錯誤清單寫好之後 commit 失敗，放回上一次的清單，和資料庫一致。"""
    import_daily(write_daily(tmp_path / "d1.xlsx", [trade(A, "Amy", "北區", 1)]), pg)
    book = tmp_path / "每日匯入錯誤.xlsx"
    before = book.read_bytes()

    failing_commit(pg, monkeypatch)
    with pytest.raises(psycopg.OperationalError, match="commit 失敗"):
        import_daily(write_daily(tmp_path / "d2.xlsx", [trade(B, "Amy", "北區", 1)]), pg)

    assert count(pg, "workflow.import_run") == 1
    assert book.read_bytes() == before
    assert sorted(p.name for p in tmp_path.iterdir()) == ["d1.xlsx", "d2.xlsx", "每日匯入錯誤.xlsx"]


def test_a_failed_first_commit_leaves_no_error_book(pg, tmp_path: Path, mailbox, monkeypatch) -> None:
    failing_commit(pg, monkeypatch)
    with pytest.raises(psycopg.OperationalError, match="commit 失敗"):
        import_price_csv(write_price_csv(tmp_path / "MarketData_20261015.csv", [csv_price("AAA UW", "101.25")]), pg)

    assert count(pg, "workflow.import_run") == 0
    assert not (tmp_path / "每日匯入錯誤.xlsx").exists()


# ---------------------------------------------------------------- 每日股價 BBG_Data（#87）


def prices_in(pg) -> list[tuple]:
    return pg.execute(
        "SELECT price_date, code, company_name, last_close, last_open, previous_close, volatility_200d,"
        " market_cap_millions, conflict_rows IS NOT NULL FROM workflow.daily_price ORDER BY code, price_date"
    ).fetchall()


def test_prices_are_stored_by_price_date(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(
        tmp_path / "d.xlsx",
        [trade(A, "Amy", "北區", 1)],
        [price(" AAA UW ", 101.25), price("9999 JT", None, open_=None), price("USDJPY", 150.123456), ("", "空白代碼")],
    )

    receipt = import_daily(daily, pg)

    day = PRICE_DATE.date()
    assert prices_in(pg) == [
        (day, "9999 JT", "虛構公司", None, None, Decimal("98.0"), Decimal("31.5"), Decimal("12345.6"), False),
        (day, "AAA UW", "虛構公司", Decimal("101.25"), Decimal("99.25"), Decimal("98.0"), Decimal("31.5"),
         Decimal("12345.6"), False),
        (day, "USDJPY", "虛構公司", Decimal("150.123456"), Decimal("99.25"), Decimal("98.0"), Decimal("31.5"),
         Decimal("12345.6"), False),
    ]  # fmt: skip
    assert (receipt.prices.price_date, receipt.prices.count, receipt.prices.replaced) == (day, 3, 0)
    assert kinds(receipt, A) == [NEW_TRADES], "交易明細的結果和以前一樣"
    assert pg.execute("SELECT price_date, price_count, conflict_count FROM workflow.price_import").fetchall() == [
        (day, 3, 0)
    ]


def test_codes_keep_their_case_and_exchange(pg, tmp_path: Path, mailbox) -> None:
    import_daily(write_daily(tmp_path / "d.xlsx", [], [price("aaa uw"), price("AAA UQ")]), pg)

    assert [row[1] for row in prices_in(pg)] == ["AAA UQ", "aaa uw"]


def test_duplicate_codes_with_the_same_values_count_once(pg, tmp_path: Path, mailbox) -> None:
    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [], [price("AAA UW"), price("AAA UW")]), pg)

    assert receipt.prices.count == 1 and receipt.prices.conflicts == ()
    assert not [line for line in receipt.lines if line.problem]


def test_duplicate_codes_with_different_values_are_conflicts(pg, tmp_path: Path, mailbox) -> None:
    receipt = import_daily(
        write_daily(tmp_path / "d.xlsx", [], [price("AAA UW", 100), price("BBB UN"), price("AAA UW", 101)]), pg
    )

    assert receipt.prices.conflicts == ("AAA UW",)
    assert [(line.kind, line.detail) for line in receipt.lines if line.kind == PRICE_CONFLICT] == [
        (PRICE_CONFLICT, "AAA UW：同一個代碼有 2 列、值不同，價格沒有存，算價格時不能用")
    ]
    conflict = pg.execute("SELECT last_close, conflict_rows FROM workflow.daily_price WHERE code = 'AAA UW'").fetchone()
    assert conflict[0] is None and [row["last_close"] for row in conflict[1]] == ["100", "101"]
    assert pg.execute("SELECT conflict_count FROM workflow.price_import").fetchone() == (1,)


def test_a_conflict_counts_every_row_of_the_code(pg, tmp_path: Path, mailbox) -> None:
    """三列以上：訊息寫實際的列數，conflict_rows 只記不同的值。"""
    rows = [price("AAA UW", 100), price("AAA UW", 101), price("AAA UW", 100), price("AAA UW", 102)]

    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [], rows), pg)

    assert [line.detail for line in receipt.lines if line.kind == PRICE_CONFLICT] == [
        "AAA UW：同一個代碼有 4 列、值不同，價格沒有存，算價格時不能用"
    ]
    stored = pg.execute("SELECT conflict_rows FROM workflow.daily_price").fetchone()[0]
    assert [row["last_close"] for row in stored] == ["100", "101", "102"]


def test_a_price_that_is_not_a_number_is_stored_blank_and_listed(pg, tmp_path: Path, mailbox) -> None:
    receipt = import_daily(write_daily(tmp_path / "d.xlsx", [], [price("AAA UW", "#N/A N/A")]), pg)

    assert [(line.kind, line.detail) for line in receipt.lines if line.problem] == [
        (PRICE_NOT_NUMBER, "BBG_Data第 4 列 AAA UW：最新收盤價(值)「#N/A N/A」不是數字，存成空白")
    ]
    assert prices_in(pg)[0][3] is None


def test_reimporting_a_price_date_replaces_that_day(pg, tmp_path: Path, mailbox) -> None:
    import_daily(write_daily(tmp_path / "1.xlsx", [], [price("AAA UW", 100), price("BBB UN")]), pg)
    other_day = dt.datetime(2026, 10, 19)
    import_daily(write_daily(tmp_path / "2.xlsx", [], [price("AAA UW", 102)], price_date=other_day), pg)

    receipt = import_daily(write_daily(tmp_path / "3.xlsx", [], [price("AAA UW", 101)]), pg)

    assert receipt.prices.replaced == 2
    assert [(row[0], row[1], row[3]) for row in prices_in(pg)] == [
        (PRICE_DATE.date(), "AAA UW", Decimal("101")),
        (other_day.date(), "AAA UW", Decimal("102")),
    ]
    assert pg.execute("SELECT price_date FROM workflow.price_import ORDER BY run_id").fetchall() == [
        (PRICE_DATE.date(),),
        (other_day.date(),),
        (PRICE_DATE.date(),),
    ], "每次匯入都留紀錄"


@pytest.mark.parametrize("value", [None, "2026/10/16", 20261016])
def test_a_price_date_that_is_not_a_date_rejects_the_whole_file(pg, tmp_path: Path, mailbox, value) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")], price_date=value)

    with pytest.raises(DailyImportError, match="C1 應為取價日（日期）"):
        import_daily(daily, pg)

    assert count(pg, "workflow.import_run") == 0 and count(pg, "workflow.trade_row") == 0


def test_wrong_price_headers_reject_the_whole_file(pg, tmp_path: Path, mailbox) -> None:
    headers = list(PRICE_HEADERS)
    headers[2], headers[3] = headers[3], headers[2]
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")], price_headers=headers)

    with pytest.raises(DailyImportError, match="D2 應為「最新收盤價\\(值\\)」，實際是「最新開盤價\\(值\\)」"):
        import_daily(daily, pg)

    assert count(pg, "workflow.import_run") == 0


def test_a_price_sheet_without_headers_rejects_the_whole_file(pg, tmp_path: Path, mailbox) -> None:
    """只有 C1 的 BBG_Data：不能當成「沒有股價」把那一天清空。"""
    import_daily(write_daily(tmp_path / "1.xlsx", [], [price("AAA UW")]), pg)
    daily = write_daily(tmp_path / "2.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")])
    wb = openpyxl.load_workbook(daily)
    wb[daily_import.PRICE_SHEET].delete_rows(2, 10)
    wb.save(daily)

    with pytest.raises(DailyImportError, match="B2 應為「代碼\\(值\\)」，實際是「空白」"):
        import_daily(daily, pg)

    assert count(pg, "workflow.daily_price") == 1 and count(pg, "workflow.trade_row") == 0


def test_a_price_sheet_without_prices_rejects_the_whole_file(pg, tmp_path: Path, mailbox) -> None:
    """有欄名但沒有任何代碼：多半是取價沒有完成，不清空那一天已存的股價。"""
    import_daily(write_daily(tmp_path / "1.xlsx", [], [price("AAA UW")]), pg)

    with pytest.raises(DailyImportError, match="沒有任何股價"):
        import_daily(write_daily(tmp_path / "2.xlsx", [trade(A, "Amy", "北區", 1)], [("", "代碼空白")]), pg)

    assert count(pg, "workflow.daily_price") == 1 and count(pg, "workflow.import_run") == 1


# ---------------------------------------------------------------- 回補股價 CSV 與只匯交易明細（#112）


def test_a_price_csv_is_stored_as_the_next_morning(pg, tmp_path: Path, mailbox) -> None:
    """CSV 的取價日是收盤那一天，存成下一天：等同隔天早上用 BBG_Data 取到的價格。"""
    source = write_price_csv(
        tmp_path / "MarketData_20261015.csv",
        [csv_price(" AAA UW ", "101.25"), csv_price("USDJPY", "150.123456", open_=""), ["2026-10-15", "", "空白代碼"]],
    )

    receipt = import_price_csv(source, pg)

    day = dt.date(2026, 10, 16)
    assert receipt.close_date == dt.date(2026, 10, 15)
    assert (receipt.prices.price_date, receipt.prices.count, receipt.prices.replaced) == (day, 2, 0)
    assert prices_in(pg) == [
        (day, "AAA UW", "虛構公司", Decimal("101.25"), Decimal("99.25"), None, Decimal("31.5"), Decimal("12345.6"),
         False),
        (day, "USDJPY", "虛構公司", Decimal("150.123456"), None, None, Decimal("31.5"), Decimal("12345.6"), False),
    ]  # fmt: skip
    assert pg.execute("SELECT daily_file_name FROM workflow.import_run").fetchall() == [("MarketData_20261015.csv",)]
    assert pg.execute("SELECT price_date, price_count FROM workflow.price_import").fetchall() == [(day, 2)]
    assert count(pg, "workflow.trade_row") == 0
    assert latest_import(pg) is None, "回補股價不算每日匯入：寄信前要看的是交易明細"


def test_the_latest_import_skips_price_csvs(pg, tmp_path: Path, mailbox) -> None:
    """寄信前確認用的「最近一次每日匯入」：之後回補了股價 CSV，仍然是那份每日更新。"""
    import_daily(write_daily(tmp_path / "每日更新.xlsx", [trade(A, "Amy", "北區", 1)]), pg)
    import_price_csv(write_price_csv(tmp_path / "m.csv", [csv_price("AAA UW")]), pg)

    assert latest_import(pg).file_name == "每日更新.xlsx"


def test_a_quoted_line_break_stays_in_one_price_csv_row(pg, tmp_path: Path, mailbox) -> None:
    source = tmp_path / "m.csv"
    lines = [",".join(PRICE_CSV_HEADERS), '2026-10-15,AAA UW,"虛構\n公司",101,100,31.5,12345.6,USD,12345.6']
    source.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8-sig"))

    receipt = import_price_csv(source, pg)

    assert receipt.prices.count == 1 and receipt.lines == ()
    assert [(row[1], row[2], row[3]) for row in prices_in(pg)] == [("AAA UW", "虛構\n公司", Decimal("101"))]


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-inf"])
def test_nan_and_infinity_are_not_prices(pg, tmp_path: Path, mailbox, value) -> None:
    """NaN、Infinity 會讓算價格中斷：當成不是數字，存成空白並列出來（CSV 與 BBG_Data 都一樣）。"""
    other_day = csv_price("AAA UW", value, day="2026-10-18")  # 存成 10-19，不和 BBG_Data 的 10-16 同一天
    csv_receipt = import_price_csv(write_price_csv(tmp_path / "m.csv", [other_day]), pg)
    daily_receipt = import_daily(write_daily(tmp_path / "d.xlsx", [], [price("BBB UN", value)]), pg)

    assert [line.kind for line in csv_receipt.lines + daily_receipt.lines] == [PRICE_NOT_NUMBER, PRICE_NOT_NUMBER]
    assert [(row[1], row[3]) for row in prices_in(pg)] == [("AAA UW", None), ("BBB UN", None)]


def test_a_price_csv_replaces_the_day_it_is_stored_as(pg, tmp_path: Path, mailbox) -> None:
    """回補時先用 BBG_Data 匯進了錯的價格（取價日 10-16）：10-15 收盤的 CSV 整天取代它。"""
    import_daily(write_daily(tmp_path / "d.xlsx", [], [price("AAA UW", 200), price("BBB UN")]), pg)

    receipt = import_price_csv(write_price_csv(tmp_path / "m.csv", [csv_price("AAA UW", "101")]), pg)

    assert receipt.prices.replaced == 2
    assert [(row[0], row[1], row[3]) for row in prices_in(pg)] == [(PRICE_DATE.date(), "AAA UW", Decimal("101"))]


def test_price_csv_problems_are_listed_and_kept_in_the_error_book(pg, tmp_path: Path, mailbox) -> None:
    rows = [csv_price("AAA UW", "#N/A N/A"), csv_price("BBB UN", "1"), csv_price("BBB UN", "2")]

    receipt = import_price_csv(write_price_csv(tmp_path / "m.csv", rows), pg)

    assert [(line.kind, line.detail) for line in receipt.lines] == [
        (PRICE_NOT_NUMBER, "m.csv第 2 列 AAA UW：收盤價「#N/A N/A」不是數字，存成空白"),
        (PRICE_CONFLICT, "BBB UN：同一個代碼有 2 列、值不同，價格沒有存，算價格時不能用"),
    ]
    error_book = openpyxl.load_workbook(receipt.error_book_path)
    assert [row[2:5] for row in list(error_book.active.values)[1:]] == [
        ("m.csv", None, PRICE_NOT_NUMBER),
        ("m.csv", None, PRICE_CONFLICT),
    ]


@pytest.mark.parametrize(
    ("rows", "headers", "message"),
    [
        ([csv_price("AAA UW")], ("取價日", "代碼", "公司名稱", "開盤價"), "第 4 欄應為「收盤價」，實際是「開盤價」"),
        ([csv_price("AAA UW")], (*PRICE_CSV_HEADERS, "備註"), "之後還有其他欄"),
        ([csv_price("AAA UW", day="2026/10/32")], PRICE_CSV_HEADERS, "第 2 列的取價日「2026/10/32」不是日期"),
        ([csv_price("AAA UW", day="")], PRICE_CSV_HEADERS, "第 2 列的取價日「空白」不是日期"),
        (
            [csv_price("AAA UW"), csv_price("BBB UN", day="2026-10-16")],
            PRICE_CSV_HEADERS,
            "取價日不只一天：2026-10-15（第 2 列起）、2026-10-16（第 3 列起）",
        ),
        ([["2026-10-15", ""]], PRICE_CSV_HEADERS, "沒有任何股價"),
    ],
)
def test_a_bad_price_csv_changes_nothing(pg, tmp_path: Path, mailbox, rows, headers, message) -> None:
    import_daily(write_daily(tmp_path / "d.xlsx", [], [price("AAA UW")]), pg)
    source = write_price_csv(tmp_path / "m.csv", rows, headers=headers)

    with pytest.raises(DailyImportError, match=re.escape(message)):
        import_price_csv(source, pg)

    assert count(pg, "workflow.import_run") == 1 and count(pg, "workflow.daily_price") == 1


def test_a_price_csv_that_is_not_utf8_changes_nothing(pg, tmp_path: Path, mailbox) -> None:
    source = tmp_path / "m.csv"
    source.write_bytes(",".join(PRICE_CSV_HEADERS).encode("big5"))

    with pytest.raises(DailyImportError, match="不是 UTF-8"):
        import_price_csv(source, pg)

    assert count(pg, "workflow.import_run") == 0


def test_trades_only_leaves_prices_alone(pg, tmp_path: Path, mailbox) -> None:
    """只匯交易明細：不讀 BBG_Data（沒有這張工作表也可以），已存的股價不動。"""
    import_daily(write_daily(tmp_path / "1.xlsx", [], [price("AAA UW", 100)]), pg)
    daily = write_daily(tmp_path / "2.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW", 999)])
    wb = openpyxl.load_workbook(daily)
    del wb[daily_import.PRICE_SHEET]
    wb.save(daily)

    receipt = import_daily(daily, pg, prices=False)

    assert receipt.prices is None and kinds(receipt, A) == [NEW_TRADES]
    assert [(row[1], row[3]) for row in prices_in(pg)] == [("AAA UW", Decimal("100"))]
    assert count(pg, "workflow.price_import") == 1
    summary = pg.execute("SELECT summary FROM workflow.import_run ORDER BY run_id DESC LIMIT 1").fetchone()[0]
    assert (summary["price_date"], summary["price_count"]) == (None, None)


def test_trades_only_ignores_a_stale_price_sheet(pg, tmp_path: Path, mailbox) -> None:
    daily = write_daily(tmp_path / "d.xlsx", [trade(A, "Amy", "北區", 1)], [price("AAA UW")], price_date="壞掉的")

    receipt = import_daily(daily, pg, prices=False)

    assert kinds(receipt, A) == [NEW_TRADES] and count(pg, "workflow.daily_price") == 0
