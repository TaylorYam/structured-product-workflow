"""寄信草稿（#29）：資料庫用每日匯入寫入虛構資料，草稿用假的建立器，不碰真的信箱。資料全部虛構。"""

from __future__ import annotations

import datetime as dt
import os
import sys
import types
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from db_synth import inventory, seed_mailbox, seed_products, trade, write_daily
from sp_workflow import mail_drafts
from sp_workflow.daily_import import import_daily
from sp_workflow.mail_drafts import (
    DRAFTED,
    FAILED,
    NOT_NEEDED,
    OUTLOOK_FAILURE,
    PENDING,
    Draft,
    MailDraftError,
    OutlookDrafter,
    OutlookError,
    create_drafts,
)
from sp_workflow.mailbox import Person
from sp_workflow.risk_level import DONE, RiskDetail, RiskLevel
from sp_workflow.risk_level import PENDING as RISK_PENDING
from sp_workflow.risk_level import _write as write_risk  # 測試直接寫存好的風險等級（算法另有 test_risk_level）
from synth import snapshot

pytestmark = pytest.mark.db

A = "029199990001"
B = "037199990002"
NOW = dt.datetime(2026, 10, 20, 9, 0, 0)


class FakeDrafter:
    def __init__(self, fail: set[str] = frozenset()) -> None:
        self.drafts: list[Draft] = []
        self.fail = fail

    def __call__(self, draft: Draft) -> None:
        if draft.tdcc in self.fail:
            raise RuntimeError("Outlook 沒有回應")
        self.drafts.append(draft)


@pytest.fixture
def db(pg):
    seed_mailbox(pg)  # Amy、Ben 是 FC；Pat 是 PM、固定副本
    return pg


@pytest.fixture
def root(tmp_path: Path) -> Path:
    folder = tmp_path / "FCN"
    folder.mkdir()
    return folder


@pytest.fixture
def out(tmp_path: Path) -> Path:
    folder = tmp_path / "out"
    folder.mkdir()
    return folder


def load(db, tmp_path: Path, trades: list[dict], inventories: list[dict]) -> None:
    """每日匯入交易明細，再直接建商品（#87 起商品條件不從每日更新來）。"""
    import_daily(write_daily(tmp_path / f"d{len(list(tmp_path.glob('d*.xlsx')))}.xlsx", trades), db)
    seed_products(db, inventories)


def passed_folder(root: Path, tdcc: str, *, sub: str = "2026") -> Path:
    folder = root / sub / tdcc
    folder.mkdir(parents=True)
    (folder / f"{tdcc}_TS.pdf").write_bytes(b"%PDF ts")
    (folder / f"{tdcc}_IIS.pdf").write_bytes(b"%PDF iis")
    return folder


def risk(db, tdcc: str, level: str | None, reason: str = "缺 AAA UW 的 200 日波動率") -> None:
    """存好的自算風險等級（#238）：`level` 有值時是完成（明細虛構），None 時是待算。"""
    detail = RiskDetail(
        12, 1, 4, 1, 1, (("AAA UW", Decimal("20")),), Decimal("20"), 2, "發行機構", (("穆迪", "A1"),), 3
    )
    found = (
        RiskLevel(DONE, level, total=detail.total, detail=detail) if level else RiskLevel(RISK_PENDING, None, reason)
    )
    with db.transaction():
        write_risk(db, tdcc, found)


def run(db, root: Path, out: Path, text: str, drafter=None):
    drafter = drafter or FakeDrafter()
    return create_drafts(text, root, out, db, drafter=drafter, now=NOW), drafter


def test_ready_product_gets_one_draft_to_its_fcs_with_fixed_cc(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 300000), trade(A, "Ben", "南區", 100000, number=None)],
        [inventory(A, [("北區", 300000), ("南區", 100000)])],
    )
    folder = passed_folder(root, A)

    receipt, drafter = run(db, root, out, A)

    assert [(o.tdcc, o.status, o.to, o.cc) for o in receipt.outcomes] == [(A, DRAFTED, ("Amy", "Ben"), ("Pat",))]
    (draft,) = drafter.drafts
    assert draft.to == ("amy@example.com", "ben@example.com") and draft.cc == ("pat@example.com",)
    assert draft.subject == f"[中文TS/IIS] [FCN] - {A}"
    assert draft.attachments == (folder / f"{A}_TS.pdf", folder / f"{A}_IIS.pdf")
    assert draft.html_body.count("<table") == 3 and "<p>FYI</p>" in draft.html_body


def test_body_shows_the_terms_in_the_mail_format(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1), ("南區", 2)])])
    risk(db, A, "RR3")
    passed_folder(root, A)

    (receipt, drafter) = run(db, root, out, A)

    assert (
        receipt.outcomes[0].reason == ""
        and not receipt.outcomes[0].to_confirm
        and drafter.drafts[0].risk_level == "RR3"
    )
    body = drafter.drafts[0].html_body
    assert "RR4" not in body, "風險等級用自算的，不再寫死 RR4（#239）"
    for expected in (
        "10-15-26",  # Trade date
        "04-27-27",  # Maturity date
        "收盤價",
        "RR3",
        "AAA UW",
        "80.00",  # Strike
        "Daily Memory",
        "100.00",  # KO Barrier
        "12.50",  # Coupon
        "97.85",  # NotePrice = 100 − UF 2.15
        "EKI",
        "65.00",  # KI Barrier
        "分公司",
        "南區",
        'font-family: Calibri; mso-fareast-font-family: "Microsoft JhengHei";',
        '<font face="Calibri"',
        '<font face="Microsoft JhengHei"',
        "text-align: center;",
        'border="0"',
        'bgcolor="#002060"',
        'bgcolor="#F0F0F0"',
        'cellspacing="0"',
        'mso-fareast-font-family: "Microsoft JhengHei";',
        "#002060",
        "#00AB8E",
        "background: #C55A11; color: #FFFFFF; font-weight: bold;",
    ):
        assert expected in body, expected


@pytest.mark.parametrize(
    ("level", "why"),
    [(None, "風險等級待算（缺 AAA UW 的 200 日波動率）"), ("absent", "資料庫還沒有風險等級")],
)
def test_pending_or_missing_risk_level_is_marked_to_confirm(
    db, tmp_path: Path, root: Path, out: Path, level, why
) -> None:
    """#239：待算或還沒算過時不寫 RR4，草稿寫「待確認」照樣建立，結果清單提醒寄出前補上。"""
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    if level != "absent":
        risk(db, A, level)
    passed_folder(root, A)

    receipt, drafter = run(db, root, out, A)

    (outcome,) = receipt.outcomes
    assert outcome.status == DRAFTED and outcome.to_confirm
    assert outcome.reason == f"{why}，草稿寫「待確認」，寄出前請補上風險等級"
    (draft,) = drafter.drafts
    assert draft.risk_level == "待確認" and "待確認" in draft.html_body and "RR4" not in draft.html_body


def test_fixed_cc_order_prioritizes_max_bob_then_pm() -> None:
    people = (
        Person("PM2", "pm2@example.com", "PM", True),
        Person("Zed", "zed@example.com", "FC", True),
        Person("BOB", "bob@example.com", "FC", True),
        Person("MAX", "max@example.com", "FC", True),
        Person("PM1", "pm1@example.com", "PM", True),
    )

    assert [person.name for person in mail_drafts._order_fixed_cc(people)] == ["MAX", "BOB", "PM1", "PM2", "Zed"]


@pytest.mark.parametrize(
    ("text", "width"),
    [
        ("Guaranteed Periods (m)", 98),  # 範本 C 欄寬
        ("Upfront / NotePrice (%)", 98),  # Upfront /／NotePrice (%)：13 字
        ("Final valuation date", 105),  # Final／valuation date：14 字
        ("Daily Memory", 122),  # 範本 J 欄寬
    ],
)
def test_wide_columns_wrap_into_at_most_two_lines(db, tmp_path: Path, root: Path, out: Path, text, width) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)

    (_, drafter) = run(db, root, out, A)

    assert f' width="{width}" style="padding: 2px 6px; width: {width}px; white-space: normal;' in (
        _cell_style(drafter.drafts[0].html_body, text)
    )


def test_columns_no_wider_than_bbg_code_4_stay_on_one_line(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)

    (_, drafter) = run(db, root, out, A)

    for narrow in ("BBG Code 4", "Product", "期初價格", "分公司 1", "Strike (%)"):
        assert 'width="' in _cell_style(drafter.drafts[0].html_body, narrow), narrow


def _cell_style(body: str, text: str) -> str:
    """回傳內容是 text 的那一格（th 或 td）的開頭標籤。"""
    try:
        end = body.index(text)
    except ValueError:
        end = body.index(text.split(" ", 1)[0])
    return body[body.rindex("<t", 0, end) : end]


def test_no_ki_shows_none_and_a_dash(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)], **{"KI(Freq)": "-", "KI(%)": "-"})])
    passed_folder(root, A)

    (_, drafter) = run(db, root, out, A)

    assert ">NONE<" in drafter.drafts[0].html_body and ">-<" in drafter.drafts[0].html_body


def test_pb_is_left_out_and_a_pb_only_product_needs_no_mail(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 1), trade(A, "PB", "南區", 1, number=None), trade(B, "PB", "北區", 1)],
        [inventory(A, [("北區", 1), ("南區", 1)]), inventory(B, [("北區", 1)])],
    )
    passed_folder(root, A)
    passed_folder(root, B)

    receipt, drafter = run(db, root, out, f"{A}\n{B}")

    assert [(o.status, o.to) for o in receipt.outcomes] == [(DRAFTED, ("Amy",)), (NOT_NEEDED, ())]
    assert [d.tdcc for d in drafter.drafts] == [A]


def test_same_fc_twice_and_a_fixed_cc_who_is_a_recipient_appear_once(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 1), trade(A, "amy", "南區", 1, number=None), trade(A, "Pat", "北區", 1, number=None)],
        [inventory(A, [("北區", 2), ("南區", 1)])],
    )
    passed_folder(root, A)

    (_, drafter) = run(db, root, out, A)

    assert drafter.drafts[0].to == ("amy@example.com", "pat@example.com") and drafter.drafts[0].cc == ()


@pytest.mark.parametrize(
    ("setup", "reason"),
    [
        ("no_ts", f"商品資料夾沒有 {A}_TS.pdf（還沒核對通過或還沒放回）"),
        ("pending_only", f"商品資料夾沒有 {A}_TS.pdf、{A}_IIS.pdf"),
        ("no_folder", "根目錄底下找不到商品資料夾"),
        ("two_folders", "有 2 個同名商品資料夾"),
    ],
)
def test_attachments_must_be_the_two_standard_files_in_one_folder(
    db, tmp_path: Path, root: Path, out: Path, setup: str, reason: str
) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    if setup == "no_ts":
        (passed_folder(root, A) / f"{A}_TS.pdf").unlink()
    elif setup == "pending_only":
        folder = root / A
        folder.mkdir()
        (folder / f"{A}_TS_待確認.pdf").write_bytes(b"x")
        (folder / f"{A}_IIS_待確認.pdf").write_bytes(b"x")
    elif setup == "two_folders":
        passed_folder(root, A, sub="a")
        passed_folder(root, A, sub="b")

    receipt, drafter = run(db, root, out, A)

    assert receipt.outcomes[0].status == PENDING and reason in receipt.outcomes[0].reason
    assert drafter.drafts == []


def test_missing_database_data_and_unmatched_names_are_all_listed(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Zoe", "北區", 1)], [inventory(B, [("北區", 1)])])

    receipt, _ = run(db, root, out, f"{A} {B} 12345")

    reasons = {o.tdcc: o.reason for o in receipt.outcomes}
    assert all(o.status == PENDING for o in receipt.outcomes)
    assert reasons[A].startswith("資料庫沒有這檔的商品條件") and "名字對不到信箱：Zoe" in reasons[A]
    assert "資料庫沒有這檔的交易列" in reasons[B]
    assert reasons["12345"] == "不是 12 碼數字"


def test_blank_required_terms_wait_with_the_fields_named(db, tmp_path: Path, root: Path, out: Path) -> None:
    """#114：商品條件缺必填欄時不建草稿、寫出缺的欄，同批其他照常；期初價格空白照舊可以寄（以前清單更新的商品）。"""
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 1), trade(B, "Ben", "北區", 1)],
        [inventory(A, [("北區", 1)]), inventory(B, [("北區", 1)], **{"期初定價": None})],
    )
    passed_folder(root, A)
    passed_folder(root, B)
    db.execute("UPDATE workflow.product SET maturity_date = NULL, coupon_pct = NULL WHERE tdcc = %s", (A,))
    risk(db, B, "RR4")

    receipt, drafter = run(db, root, out, f"{A} {B}")

    assert [(o.tdcc, o.status, o.reason) for o in receipt.outcomes] == [
        (A, PENDING, "商品條件缺 Coupon p.a. (%)、Maturity date"),
        (B, DRAFTED, ""),
    ]
    assert [d.tdcc for d in drafter.drafts] == [B]


def test_one_failed_draft_does_not_stop_the_others(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 1), trade(B, "Ben", "北區", 1)],
        [inventory(A, [("北區", 1)]), inventory(B, [("北區", 1)])],
    )
    passed_folder(root, A)
    passed_folder(root, B)
    risk(db, B, "RR4")

    receipt, drafter = run(db, root, out, f"{A},{B}", FakeDrafter(fail={A}))

    assert [(o.status, o.reason) for o in receipt.outcomes] == [
        (FAILED, "建立草稿失敗：Outlook 沒有回應"),
        (DRAFTED, ""),
    ]
    assert [d.tdcc for d in drafter.drafts] == [B]


def test_running_again_creates_another_draft(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)
    drafter = FakeDrafter()

    create_drafts(A, root, out, db, drafter=drafter, now=NOW)
    create_drafts(A, root, out, db, drafter=drafter, now=NOW + dt.timedelta(seconds=1))

    assert len(drafter.drafts) == 2


def test_product_folders_are_read_only(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)
    before = snapshot(root)

    run(db, root, out, A)

    assert snapshot(root) == before


def test_result_list_records_outcomes_and_last_import(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)

    receipt, _ = run(db, root, out, f"{A} {B}")

    assert receipt.list_path == out / "寄信草稿清單_20261020-090000.xlsx" and receipt.last_import is not None
    wb = openpyxl.load_workbook(receipt.list_path)
    rows = list(wb["寄信草稿清單"].iter_rows(values_only=True))
    assert rows[0] == ("TDCC", "狀態", "收件人", "副本", "原因")
    assert rows[1][:4] == (A, DRAFTED, "Amy", "Pat") and rows[2][:2] == (B, PENDING)
    source = {row[0]: row[1] for row in wb["資料來源"].iter_rows(values_only=True)}
    assert source["匯入編號"] == receipt.last_import.run_id == 1
    assert source["每日更新檔"] == "d0.xlsx" and len(source["SHA-256"]) == 64
    assert source["讀不到的資料夾"] == 0


def test_full_width_digits_are_not_a_tdcc(db, root: Path, out: Path) -> None:
    receipt, _ = run(db, root, out, "０２９１９９９９０００１")

    assert [(o.status, o.reason) for o in receipt.outcomes] == [(PENDING, "不是 12 碼數字")]


def test_pb_typed_with_spaces_is_still_pb(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, " pb ", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)

    receipt, drafter = run(db, root, out, A)

    assert receipt.outcomes[0].status == NOT_NEEDED and drafter.drafts == []


def test_products_skipped_by_the_daily_import_are_unsupported(db, tmp_path: Path, root: Path, out: Path) -> None:
    load(
        db,
        tmp_path,
        [trade(B, "Amy", "北區", 1, product="FCN_Stepdown")],
        [inventory(B, [("北區", 1)], Product="FCN_Stepdown")],
    )

    receipt, _ = run(db, root, out, B)

    reason = "不支援的商品類型：FCN_Stepdown（目前只做 FCN、FCN_Express、Super Bull、BEN）"
    assert [(o.status, o.reason) for o in receipt.outcomes] == [(PENDING, reason)]


def test_express_gets_a_draft_like_fcn(db, tmp_path: Path, root: Path, out: Path, monkeypatch) -> None:
    """FCN_Express 和 FCN 一樣寄（#302）：主旨與表格的 Product 帶出商品種類，Guaranteed Periods 寫 0。"""
    express = inventory(A, [("北區", 1)], Product="FCN_Express", **{"Non-Call(月)": 0})
    load(db, tmp_path, [trade(A, "Amy", "北區", 1, product="FCN_Express")], [express])
    risk(db, A, "RR4")
    passed_folder(root, A)
    real_render, rendered = mail_drafts.render_body, []

    def render(tdcc, product, branches, risk_level):
        rendered.append((product["product"], product["non_call_months"]))
        return real_render(tdcc, product, branches, risk_level)

    monkeypatch.setattr(mail_drafts, "render_body", render)

    receipt, drafter = run(db, root, out, A)

    assert (receipt.outcomes[0].status, receipt.outcomes[0].reason) == (DRAFTED, "")
    assert drafter.drafts[0].subject == f"[中文TS/IIS] [FCN_Express] - {A}"
    assert rendered == [("FCN_Express", 0)] and ">FCN_Express<" in drafter.drafts[0].html_body


@pytest.mark.parametrize("product", ["Super Bull", "BEN"])
def test_super_bull_and_ben_get_a_draft_like_fcn(db, tmp_path: Path, root: Path, out: Path, product) -> None:
    """Super Bull（#310）、BEN（#373）和 FCN 一樣寄：主旨帶出商品種類，沒有 KI 的 KI Barrier 寫 -，不寫參與率。"""
    terms = inventory(A, [("北區", 1)], Product=product, **{"KI(Freq)": "-", "KI(%)": "-"})
    load(db, tmp_path, [trade(A, "Amy", "北區", 1, product=product)], [terms])
    risk(db, A, "RR4")
    passed_folder(root, A)

    receipt, drafter = run(db, root, out, A)

    assert (receipt.outcomes[0].status, receipt.outcomes[0].reason) == (DRAFTED, "")
    assert drafter.drafts[0].subject == f"[中文TS/IIS] [{product}] - {A}"
    assert f">{product}<" in drafter.drafts[0].html_body
    assert "Participation" not in drafter.drafts[0].html_body


def test_unreadable_folders_are_mentioned(db, tmp_path: Path, root: Path, out: Path, monkeypatch) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    locked = root / "鎖住的資料夾"
    monkeypatch.setattr(mail_drafts, "find_product_folders", lambda _root, _wanted: ({}, [locked]))

    receipt, _ = run(db, root, out, A)

    assert receipt.unreadable_dirs == (locked,)
    assert "（搜尋時有 1 個資料夾讀不到，可能在裡面）" in receipt.outcomes[0].reason
    source = {
        row[0]: row[1] for row in openpyxl.load_workbook(receipt.list_path)["資料來源"].iter_rows(values_only=True)
    }
    assert source["讀不到的資料夾"] == 1


def test_an_error_in_one_product_does_not_stop_the_batch(
    db, tmp_path: Path, root: Path, out: Path, monkeypatch
) -> None:
    load(
        db,
        tmp_path,
        [trade(A, "Amy", "北區", 1), trade(B, "Ben", "北區", 1)],
        [inventory(A, [("北區", 1)]), inventory(B, [("北區", 1)])],
    )
    passed_folder(root, A)
    passed_folder(root, B)
    risk(db, B, "RR4")
    real_render = mail_drafts.render_body

    def render(tdcc, product, branches, risk_level):
        if tdcc == A:
            raise ValueError("內文產生失敗")
        return real_render(tdcc, product, branches, risk_level)

    monkeypatch.setattr(mail_drafts, "render_body", render)

    receipt, drafter = run(db, root, out, f"{A} {B}")

    assert [(o.status, o.reason) for o in receipt.outcomes] == [(FAILED, "處理失敗：內文產生失敗"), (DRAFTED, "")]
    assert [d.tdcc for d in drafter.drafts] == [B] and receipt.list_path.exists()


def test_without_fixed_cc_nothing_is_drafted(pg, tmp_path: Path, root: Path, out: Path) -> None:
    seed_mailbox(pg, [("Amy", "amy@example.com", "FC", False)])

    with pytest.raises(MailDraftError, match="固定副本"):
        run(pg, root, out, A)

    assert list(out.iterdir()) == []


def test_bad_paths_and_empty_input_stop_before_anything(db, root: Path, out: Path, tmp_path: Path) -> None:
    with pytest.raises(MailDraftError, match="沒有輸入 TDCC"):
        run(db, root, out, " \n ")
    with pytest.raises(MailDraftError, match="根目錄不存在"):
        run(db, tmp_path / "none", out, A)
    with pytest.raises(MailDraftError, match="結果清單的位置不存在"):
        run(db, root, tmp_path / "none", A)


class FakeMail:
    """假的 Outlook 郵件：記下被設定的欄位；Send() 被呼叫就記下來，讓測試抓出來。"""

    def __init__(self) -> None:
        self.attached: list[str] = []
        self.Attachments = types.SimpleNamespace(Add=self.attached.append)
        self.saved = self.sent = False

    def Save(self) -> None:  # noqa: N802 — 對應 COM 方法名稱
        self.saved = True

    def Send(self) -> None:  # noqa: N802
        self.sent = True


@pytest.fixture
def fake_com(monkeypatch):
    """用假的 pythoncom／win32com 取代真的 Outlook；筆電與 CI 都能跑（#29）。"""

    class ComError(Exception):
        pass

    state = types.SimpleNamespace(mail=FakeMail(), error=None, com_depth=0, ComError=ComError)

    def dispatch(name: str):
        assert name == "Outlook.Application"
        if state.error is not None:
            raise state.error

        def create_item(kind: int) -> FakeMail:
            assert kind == 0  # olMailItem
            return state.mail

        return types.SimpleNamespace(CreateItem=create_item)

    pythoncom = types.ModuleType("pythoncom")
    pythoncom.com_error = ComError
    pythoncom.CoInitialize = lambda: setattr(state, "com_depth", state.com_depth + 1)
    pythoncom.CoUninitialize = lambda: setattr(state, "com_depth", state.com_depth - 1)
    client = types.ModuleType("win32com.client")
    client.Dispatch = dispatch
    package = types.ModuleType("win32com")
    package.client = client
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    return state


def sample_draft(tmp_path: Path) -> Draft:
    return Draft(
        A,
        ("amy@example.com", "ben@example.com"),
        ("pat@example.com",),
        f"[中文TS/IIS] [FCN] - {A}",
        "<p>FYI</p>",
        (tmp_path / "a.pdf", tmp_path / "b.pdf"),
    )


def test_outlook_drafter_fills_the_mail_and_only_saves_it(fake_com, tmp_path: Path) -> None:
    OutlookDrafter()(sample_draft(tmp_path))

    mail = fake_com.mail
    assert (mail.To, mail.CC, mail.Subject, mail.HTMLBody) == (
        "amy@example.com; ben@example.com",
        "pat@example.com",
        f"[中文TS/IIS] [FCN] - {A}",
        "<p>FYI</p>",
    )
    assert mail.attached == [str(tmp_path / "a.pdf"), str(tmp_path / "b.pdf")]
    assert mail.saved and not mail.sent
    assert fake_com.com_depth == 0  # CoInitialize／CoUninitialize 成對


def test_outlook_errors_are_explained_in_chinese(fake_com, tmp_path: Path) -> None:
    fake_com.error = fake_com.ComError(-2147467260, "亂碼", None, None)

    with pytest.raises(OutlookError) as info:
        OutlookDrafter()(sample_draft(tmp_path))

    assert str(info.value) == f"{OUTLOOK_FAILURE}，錯誤碼 0x80004004"
    assert fake_com.com_depth == 0


def test_outlook_failure_marks_only_that_product(fake_com, db, tmp_path: Path, root: Path, out: Path) -> None:
    load(db, tmp_path, [trade(A, "Amy", "北區", 1)], [inventory(A, [("北區", 1)])])
    passed_folder(root, A)
    fake_com.error = fake_com.ComError(-2147467260, "亂碼", None, None)

    receipt = create_drafts(A, root, out, db, drafter=OutlookDrafter(), now=NOW)

    assert receipt.outcomes[0].status == FAILED and OUTLOOK_FAILURE in receipt.outcomes[0].reason


@pytest.mark.skipif(
    sys.platform != "win32" or not os.environ.get("SPW_OUTLOOK_TEST"),
    reason="真的 Outlook 測試只在本機設定 SPW_OUTLOOK_TEST 時執行",
)
def test_outlook_creates_a_draft_without_sending(tmp_path: Path) -> None:
    import pythoncom
    import win32com.client

    attachment = tmp_path / "test.pdf"
    attachment.write_bytes(b"%PDF test")
    subject = f"SPW 測試草稿 {dt.datetime.now():%Y%m%d%H%M%S}（可刪除）"

    OutlookDrafter()(Draft("000000000000", ("nobody@example.com",), (), subject, "<p>FYI</p>", (attachment,)))

    pythoncom.CoInitialize()
    try:
        drafts = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI").GetDefaultFolder(16).Items
        found = [m for m in drafts if m.Subject == subject]
        assert len(found) == 1 and found[0].Attachments.Count == 1 and not found[0].Sent
    finally:
        for mail in locals().get("found", []):
            mail.Delete()
        pythoncom.CoUninitialize()
