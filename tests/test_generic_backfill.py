"""回填通用（#364）：ELN 通用回補格式以 TDCC 從資料庫回填 ISIN、比價日；期初定價 VWAP 的列另回填各標的價格。

資料庫測試對真的 PostgreSQL；表、商品、價格與日期全部虛構。
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from db_synth import check_terms, inventory, seed_products
from fcn_checker.check_config import ConfigPaths
from fcn_checker.config import load_reference_format
from sp_workflow.generic_backfill import NOT_FOUND, GenericBackfillError, backfill_generic
from sp_workflow.window_actions import fill_generic

CONFIG_DIR = Path(__file__).resolve().parents[1] / "config"
FMT = load_reference_format(CONFIG_DIR / "reference_sheet.toml")
PATHS = ConfigPaths.with_fallback(CONFIG_DIR / "review_standard.toml")
HEADERS = [
    "TS/IIS完成檢核", "交易日期", "Product", "TDCC Code", "ISIN Code", "私銀註記", "UF", "計價幣別", "單位面額",
    "交易日", "發行日", *(f"比價日_{n}" for n in range(1, 13)), "下一比價日", "最終比價日", "到期日", "期初定價",
    "KO(%)", "KO(Freq)", "KO(memo)", "K(%)", "KI(%)", "KI(Freq)", "Coupon p.a. (%)", "天期(月)", "Non-Call(月)",
    *(f"UL_{i}{s}" for i in range(1, 6) for s in ("", "_進場價", "_執行價", "_下限價", "_KO價", "_Memo")),
]  # fmt: skip
A, B, C, D = (f"14719999000{i}" for i in range(1, 5))
DAY_1, DAY_6 = dt.date(2026, 11, 23), dt.date(2027, 4, 22)
TRADE = dt.datetime(2026, 10, 15)


def sheet_row(tdcc: str, **cells) -> dict:
    """ELN 通用回補格式的一列：ISIN 與比價日空白，兩檔標的、收盤價（價格已填）。"""
    row = {
        "TS/IIS完成檢核": "待檢核",
        "交易日期": TRADE,
        "Product": "FCN",
        "TDCC Code": tdcc,
        "私銀註記": "#VALUE!",
        "交易日": TRADE,
        "期初定價": "收盤價",
        "UL_1": "AAA UW",
        "UL_1_進場價": 100,
        "UL_1_執行價": 80,
        "UL_1_下限價": 65,
        "UL_1_KO價": 100,
        "UL_2": "BBB UN",
        "UL_2_進場價": 50.5,
        "UL_2_執行價": 40.4,
        "UL_2_下限價": 32.825,
        "UL_2_KO價": 50.5,
        **{f"UL_{i}{s}": "-" for i in range(3, 6) for s in ("", "_進場價", "_執行價", "_下限價", "_KO價")},
    }
    row.update(cells)
    return row


def write_sheet(path: Path, rows: list[dict], headers: list[str] = HEADERS) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "工作表2"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    for r in range(2, sheet.max_row + 1):
        sheet.cell(row=r, column=2).number_format = "mm-dd-yy"
    book.save(path)
    return path


def filled(pg, tdcc: str, *, isin: str | None = "XS0000000001", vwap: bool = False, prices: bool = True) -> None:
    """資料庫一檔核對寫回過的商品：ISIN、比價日（1、6 有值，其他是不需要填的期別）、兩檔標的的價格。"""
    seed_products(
        pg, [inventory(tdcc, [("台北", 10000)], **{"ISIN Code": isin, "期初定價": "VWAP" if vwap else "收盤價"})]
    )
    for slot in range(1, 13):
        value = {1: DAY_1, 6: DAY_6}.get(slot)
        pg.execute("INSERT INTO workflow.autocall_date (tdcc, slot, value) VALUES (%s, %s, %s)", (tdcc, slot, value))
    check_terms(pg, tdcc)
    for seq, code, initial, ki in ((1, "AAA UW", "101.25", "65.8125"), (2, "BBB UN", "52", None)):
        values = (initial, "81", ki, initial) if prices else (None, None, None, None)
        pg.execute(
            "INSERT INTO workflow.underlying_price (tdcc, seq, code, initial_price, strike_price, ki_price, ko_price)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s)",
            (tdcc, seq, code, *values),
        )


def read_back(path: Path) -> dict[str, dict]:
    sheet = openpyxl.load_workbook(path).active
    headers = [c.value for c in sheet[1]]
    return {row[3]: dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=2, values_only=True)}


def lines(receipt) -> dict[str, tuple[int, tuple[str, ...]]]:
    return {line.tdcc: (line.filled, line.reasons) for line in receipt.lines}


@pytest.mark.db
def test_isin_and_compare_dates_are_filled_and_unknown_products_stay_blank(pg, tmp_path: Path) -> None:
    filled(pg, A)
    source = write_sheet(tmp_path / "通用.xlsx", [sheet_row(A), sheet_row(B, Product="BEN")])
    before = source.read_bytes()

    receipt = backfill_generic(pg, source, tmp_path / "通用_回填.xlsx", FMT)

    assert lines(receipt) == {A: (13, ()), B: (0, (NOT_FOUND,))}
    assert (receipt.cells, receipt.not_found, receipt.with_reasons) == (13, 1, 0)
    assert source.read_bytes() == before, "原檔不動"
    out = read_back(tmp_path / "通用_回填.xlsx")
    assert out[A]["ISIN Code"] == "XS0000000001"
    assert out[A]["比價日_1"] == dt.datetime(2026, 11, 23) and out[A]["比價日_6"] == dt.datetime(2027, 4, 22)
    assert {out[A][f"比價日_{n}"] for n in (2, 3, 4, 5, *range(7, 13))} == {"-"}, "不需要填的期別寫空值寫法"
    assert out[A]["UL_1_進場價"] == 100, "不是 VWAP 的價格欄不動"
    assert out[B]["ISIN Code"] is None and out[B]["比價日_1"] is None, "資料庫沒有這檔：維持空白"
    assert out[A]["私銀註記"] == "#VALUE!" and out[A]["TS/IIS完成檢核"] == "待檢核", "其他格不動"
    sheet = openpyxl.load_workbook(tmp_path / "通用_回填.xlsx").active
    assert sheet["L2"].number_format == "mm-dd-yy", "日期沿用表上既有日期格的顯示格式"


@pytest.mark.db
def test_existing_values_are_kept_unless_vwap_prices(pg, tmp_path: Path) -> None:
    filled(pg, A)
    filled(pg, B, vwap=True)
    source = write_sheet(
        tmp_path / "通用.xlsx",
        [
            sheet_row(A, **{"ISIN Code": "XS9999999999", "比價日_1": dt.datetime(2026, 11, 23)}),
            sheet_row(B, 期初定價="VWAP", UL_1="aaa uw", UL_2_進場價=None, UL_2_執行價=None),
        ],
    )

    receipt = backfill_generic(pg, source, tmp_path / "out.xlsx", FMT)

    count, reasons = lines(receipt)[A]
    assert count == 11 and reasons == ("ISIN Code 表上是 XS9999999999、資料庫是 XS0000000001，保留原值",)
    out = read_back(tmp_path / "out.xlsx")
    assert out[A]["ISIN Code"] == "XS9999999999", "不一致保留原值"
    assert lines(receipt)[B] == (13 + 8, ()), "VWAP：兩檔標的各四格（含覆寫）"
    assert [out[B][f"UL_1_{c}"] for c in ("進場價", "執行價", "下限價", "KO價")] == [101.25, 81, 65.8125, 101.25]
    assert [out[B][f"UL_2_{c}"] for c in ("進場價", "執行價", "下限價", "KO價")] == [52, 81, "-", 52], "資料庫空白寫 -"
    assert out[B]["UL_3_進場價"] == "-", "沒有的標的不動"


@pytest.mark.db
def test_missing_database_data_leaves_cells_blank(pg, tmp_path: Path) -> None:
    filled(pg, A, isin=None, vwap=True, prices=False)
    pg.execute("DELETE FROM workflow.autocall_date WHERE tdcc = %s", (A,))
    filled(pg, C, vwap=True)
    source = write_sheet(
        tmp_path / "通用.xlsx",
        [
            sheet_row(A, 期初定價="VWAP", UL_1_進場價=None),
            sheet_row(C, 期初定價="VWAP", UL_2="ZZZ UN", UL_2_進場價=None),
        ],
    )

    receipt = backfill_generic(pg, source, tmp_path / "out.xlsx", FMT)

    assert lines(receipt)[A] == (
        0,
        (
            "資料庫還沒有 ISIN",
            "資料庫還沒有比價日",
            "UL_1 AAA UW 資料庫還沒有 VWAP 價格",
            "UL_2 BBB UN 資料庫還沒有 VWAP 價格",
        ),
    )
    assert lines(receipt)[C] == (13 + 4, ("UL_2 ZZZ UN 資料庫這檔沒有這個標的",))
    assert receipt.with_reasons == 2
    out = read_back(tmp_path / "out.xlsx")
    assert out[A]["UL_1_進場價"] is None and out[C]["UL_2_進場價"] is None


@pytest.mark.db
def test_a_sheet_without_the_backfill_columns_is_refused(pg, tmp_path: Path) -> None:
    headers = [h for h in HEADERS if h not in ("ISIN Code", "比價日_3")]
    source = write_sheet(tmp_path / "通用.xlsx", [sheet_row(A)], headers)

    with pytest.raises(GenericBackfillError, match="缺少回填要用的欄：ISIN Code、比價日_3，整份不處理"):
        backfill_generic(pg, source, tmp_path / "out.xlsx", FMT)
    assert not (tmp_path / "out.xlsx").exists()

    (tmp_path / "壞.xlsx").write_text("不是 Excel", encoding="utf-8")
    with pytest.raises(GenericBackfillError, match="壞.xlsx 打不開"):
        backfill_generic(pg, tmp_path / "壞.xlsx", tmp_path / "out.xlsx", FMT)


@pytest.mark.db
def test_window_action_reads_the_reference_format(pg, pg_settings, tmp_path: Path) -> None:
    filled(pg, D)
    source = write_sheet(tmp_path / "通用.xlsx", [sheet_row(D)])

    receipt = fill_generic(pg_settings, source, tmp_path / "out.xlsx", PATHS)

    assert lines(receipt) == {D: (13, ())} and receipt.path == tmp_path / "out.xlsx"


@pytest.mark.db
def test_error_values_and_merged_cells_are_filled_and_the_source_may_be_overwritten(pg, tmp_path: Path) -> None:
    filled(pg, A)
    source = write_sheet(tmp_path / "通用.xlsx", [sheet_row(A), sheet_row(B)])
    book = openpyxl.load_workbook(source)
    sheet = book.active
    sheet["E2"] = "#VALUE!"  # 公式算錯留下的錯誤值
    sheet.merge_cells("L2:M2")  # 比價日_1、比價日_2 合併
    book.save(source)

    receipt = backfill_generic(pg, source, source, FMT)

    assert lines(receipt)[A] == (13, ())
    out = read_back(source)
    assert out[A]["ISIN Code"] == "XS0000000001", "錯誤值算空白"
    assert (out[A]["比價日_1"], out[A]["比價日_2"]) == (dt.datetime(2026, 11, 23), "-"), "合併儲存格拆開後各自回填"
    assert not openpyxl.load_workbook(source).active.merged_cells.ranges


@pytest.mark.db
def test_a_sheet_without_the_tdcc_header_is_refused(pg, tmp_path: Path) -> None:
    source = write_sheet(tmp_path / "通用.xlsx", [sheet_row(A)], [h for h in HEADERS if h != "TDCC Code"])

    with pytest.raises(GenericBackfillError, match=r"通用.xlsx 第一張工作表「工作表2」前 10 列找不到「TDCC Code」欄名"):
        backfill_generic(pg, source, tmp_path / "out.xlsx", FMT)
