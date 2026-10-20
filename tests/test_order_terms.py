"""Order 條件讀取（#40）：用虛構 HTML 仿造各上手的 Order 原信，不開 Outlook。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from sp_workflow.order_formats import ISSUERS, KNOWN_WITHOUT_FORMAT
from sp_workflow.order_template import COLUMNS as TEMPLATE_COLUMNS
from sp_workflow.order_template import FLAG_FILL, template_row
from sp_workflow.order_terms import (
    LIST_SHEET,
    NO_ORDER,
    PENDING,
    RAW_SHEET,
    READ,
    UNSUPPORTED,
    Mail,
    MailReadError,
    OrderTermsError,
    buy_segments,
    is_order_subject,
    parse_amount,
    parse_date,
    read_orders,
    split_embedded,
)

NOW = dt.datetime(2026, 10, 20, 9, 30, 0)
SENT = dt.datetime(2026, 10, 15, 16, 57, 36)
BARC = "029199990001"

PRICER_HEADER = (
    "Product",
    "Currency",
    "Guaranteed Periods (m)",
    "BBG Code 1",
    "BBG Code 2",
    "BBG Code 3",
    "BBG Code 4",
    "BBG Code 5",
    "Strike (%)",
    "KO Type",
    "KO Barrier (%)",
    "Coupon p.a. (%)",
    "Upfront / Note Price (%)",
    "Tenor (m)",
    "Barrier Type",
    "KI Barrier (%)",
    "Observation Frequency (m)",
    "OTC",
    "Funding Spread (bps)",
    "Effective Date offset",
    "Notional",
    "Trade Date",
    "Issue Date",
    "Final Valuation Date",
    "Maturity Date",
    "Quote ID",
)
PRICER_VALUES = (
    "FCN",
    "USD",
    "1",
    "AAA UW",
    "BBB UW",
    "CCC UN",
    "",
    "",
    "80.00",
    "Daily Memory",
    "100.00",
    "14.72",
    "98.50",
    "6",
    "NONE",
    "0",
    "1",
    "Note",
    "",
    "7",
    "10,000,000",
    "15-Oct-26",
    "22-Oct-26",
    "22-Apr-27",
    "27-Apr-27",
    "QUOTE001",
)


def table(*rows: tuple[str, ...]) -> str:
    body = "".join("<tr>" + "".join(f"<td><p>{c}</p></td>" for c in row) + "</tr>" for row in rows)
    return f"<table>{body}</table>"


def order_html(*parts: str, quoted: str = "") -> str:
    """`parts` 是依序的段落或表格（已是 HTML 的直接放入）。"""
    body = "".join(p if p.startswith("<table") else f"<p>{p}</p>" for p in parts)
    reply = f"<div><p><b>From:</b> someone@example.com</p></div>{quoted}" if quoted else ""
    return f"<html><head><style>p {{margin:0}}</style></head><body>{body}<p>玉山證券 E.SUN Securities</p><p>Tel: (02) 0000-0000</p>{reply}</body></html>"


def standard_order(
    tdcc: str = BARC, header: tuple[str, ...] = PRICER_HEADER, values: tuple[str, ...] = PRICER_VALUES
) -> str:
    return order_html(
        "Hi team,",
        f"We buy {tdcc} USD 200K @98.50",
        "Please settle thru Euroclear by acct DSU",
        "Risk level: RR4",
        "Initial price: Opening price",
        "PG rate: 0%",
        table(header, values),
    )


class FakeReader:
    """依檔名回傳虛構信件；記下讀過哪些檔。"""

    def __init__(self, mails: dict[str, Mail | Exception]) -> None:
        self.mails = mails
        self.read: list[str] = []

    def __call__(self, path: Path) -> Mail:
        self.read.append(path.name)
        mail = self.mails[path.name]
        if isinstance(mail, Exception):
            raise mail
        return mail


def product_folder(root: Path, tdcc: str, *names: str) -> Path:
    folder = root / "FCN" / tdcc
    folder.mkdir(parents=True)
    for name in names:
        (folder / name).write_bytes(b"msg")
    return folder


def run(tmp_path: Path, text: str, mails: dict[str, Mail | Exception]):
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    reader = FakeReader(mails)
    receipt = read_orders(text, tmp_path / "root", out, reader=reader, now=NOW)
    return receipt, reader


def term_reasons(row) -> tuple[str, ...]:
    """讀信本身的問題（不含格式表缺欄）；只放部分欄位的測試信用。"""
    return tuple(r for r in row.reasons if not r.startswith("格式表缺"))


@pytest.fixture
def root(tmp_path: Path) -> Path:
    path = tmp_path / "root"
    path.mkdir()
    return path


def test_reads_all_terms_of_a_standard_order(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    mail = Mail(f"(Order)_[Esun Sec] FCN_AAA/BBB/CCC_B--{BARC}", SENT, standard_order())

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    (row,) = receipt.rows
    assert (row.tdcc, row.status, row.issuer, row.reasons) == (BARC, READ, "BARC", ())
    assert row.values == {
        "交割說明": "Euroclear by acct DSU",
        "風險等級": "RR4",
        "期初定價": "開盤價",
        "PG rate(%)": Decimal("0"),
        "幣別": "USD",
        "承作金額": Decimal("200000"),
        "價格(%)": Decimal("98.50"),
        "商品": "FCN",
        "Non-Call(月)": 1,
        "標的1": "AAA UW",
        "標的2": "BBB UW",
        "標的3": "CCC UN",
        "執行價(%)": Decimal("80.00"),
        "KO觀察方式": "Daily Memory",
        "KO(%)": Decimal("100.00"),
        "年利率(%)": Decimal("14.72"),
        "天期(月)": 6,
        "KI型態": "NONE",
        "KI(%)": Decimal("0"),
        "觀察頻率(月)": 1,
        "發行日偏移": "7",
        "交易日": dt.date(2026, 10, 15),
        "發行日": dt.date(2026, 10, 22),
        "最終比價日": dt.date(2027, 4, 22),
        "到期日": dt.date(2027, 4, 27),
        "上手報價編號": "QUOTE001",
    }
    assert row.others == (("OTC", "Note"), ("Notional", "10,000,000"))  # 簽名檔的 Tel: 不算條件
    assert row.sent_on == SENT and row.source == root / "FCN" / BARC / "order.msg"


def test_list_workbook_keeps_tdcc_as_text_and_raw_values(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    mail = Mail(f"(Order)_FCN--{BARC}", SENT, standard_order())

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    assert receipt.list_path.name == "Order條件清單_20261020-093000.xlsx"
    wb = openpyxl.load_workbook(receipt.list_path)
    ws = wb[LIST_SHEET]
    header = [c.value for c in ws[1]]
    assert header[:4] == ["TDCC", "狀態", "上手", "待確認原因"]
    assert header[-3:] == ["其他欄位", "來源信件", "寄件時間"]
    row = {h: c for h, c in zip(header, ws[2], strict=True)}
    assert row["TDCC"].value == BARC and row["TDCC"].number_format == "@"
    assert row["發行日"].value == dt.datetime(2026, 10, 22) and row["發行日"].number_format == "yyyy-mm-dd"
    assert row["年利率(%)"].value == pytest.approx(14.72)
    assert row["其他欄位"].value == "OTC=Note\nNotional=10,000,000"
    raw = wb[RAW_SHEET]
    lines = {(r[2].value, r[3].value): (r[4].value, r[5].value) for r in raw.iter_rows(min_row=2)}
    assert lines[("We buy 行", "金額")] == ("200K", "承作金額")
    assert lines[("條件表 1", "BBG Code 1")] == ("AAA UW", "標的1")
    assert lines[("條件表 1", "OTC")] == ("Note", None)


def test_list_is_never_overwritten(tmp_path: Path, root: Path) -> None:
    out = tmp_path / "out"
    out.mkdir()
    (out / "Order條件清單_20261020-093000.xlsx").write_bytes(b"old")

    with pytest.raises(OrderTermsError, match="已存在"):
        read_orders(BARC, root, out, reader=FakeReader({}), now=NOW)


@pytest.mark.parametrize(
    ("subject", "expected"),
    [
        ("(Order)_[Esun Sec] FCN_AAA_B--029199990001", True),
        ("(ORder)[Esun Sec] FCN_AAA --029199990001", True),
        ("(Order)_EXTERNAL: [Esun Sec] FCN_AAA_M_029199990001", True),
        ("[External] RE: (Order)_[Esun Sec] FCN", False),
        ("RE: [EXTERNAL](ORder)[Esun Sec] FCN", False),
        ("[External] DONE: (Order)_[Esun Sec] FCN", False),
        ("[External] [TS/IIS] RE: (Order)_[Esun Sec] FCN", False),
        ("FW: (Order) FCN", False),
        ("[Esun Sec] FCN_AAA_B--TDCC", False),
    ],
)
def test_only_original_order_mails_are_read(subject: str, expected: bool) -> None:
    assert is_order_subject(subject) is expected


def test_replies_in_the_folder_are_skipped(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "a.msg", "b.msg", "c.msg", "note.txt")
    mails = {
        "a.msg": Mail(f"[External] RE: (Order)_FCN--{BARC}", SENT, standard_order()),
        "b.msg": Mail(f"(Order)_FCN--{BARC}", SENT, standard_order()),
        "c.msg": Mail("[External] BARX COMET Order EXECUTED", SENT, "<p>done</p>"),
    }

    receipt, reader = run(tmp_path, BARC, mails)

    assert receipt.rows[0].status == READ and receipt.rows[0].source.name == "b.msg"
    assert sorted(reader.read) == ["a.msg", "b.msg", "c.msg"]  # 只讀 .msg


def test_quoted_reply_below_the_order_is_ignored(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    other = table(("Product", "Currency", "BBG Code 1"), ("FCN", "EUR", "ZZZ UW"))
    html = order_html(
        f"We buy {BARC} USD 200K @98.50", table(PRICER_HEADER, PRICER_VALUES), quoted=f"<p>Please quote</p>{other}"
    )

    receipt, _ = run(tmp_path, BARC, {"order.msg": Mail(f"(Order)--{BARC}", SENT, html)})

    assert receipt.rows[0].values["標的1"] == "AAA UW" and receipt.rows[0].values["幣別"] == "USD"


def test_one_mail_with_two_tdccs_is_split_and_the_sibling_follows(tmp_path: Path, root: Path) -> None:
    first, second = "328199990001", "328199990002"
    header = ("Capital size", "Ref", "Product", "Currency", "BBG Code 1", "Upfront / NotePrice (%)", "Tenor (m)")
    html = order_html(
        f"We buy {first} USD 950K @96.15、{second} USD 1 mio @96.15",
        "Risk level: RR4",
        first,
        table(header, ("USD 950 K", "R1", "FCN", "USD", "AAA UW", "96.20%→96.15%", "12")),
        second,
        table(header, ("USD 1 mio", "R2", "FCN", "USD", "BBB UW", "96.15", "9")),
    )
    product_folder(root, first, "order.msg")

    receipt, _ = run(tmp_path, first, {"order.msg": Mail(f"(Order)--{first}、{second}", SENT, html)})

    assert [r.tdcc for r in receipt.rows] == [first, second]
    one, two = receipt.rows
    assert term_reasons(one) == ("價格(%) 有修改：96.20%→96.15%",)  # 這裡的表只有部分欄位，格式表缺的另外測
    assert one.values["價格(%)"] == Decimal("96.15") and one.values["承作金額"] == Decimal("950000")
    assert (one.values["標的1"], one.values["天期(月)"], one.values["上手報價編號"]) == ("AAA UW", 12, "R1")
    assert term_reasons(two) == ()
    assert (two.values["標的1"], two.values["承作金額"], two.values["風險等級"]) == (
        "BBB UW",
        Decimal("1000000"),
        "RR4",
    )


def test_sections_after_a_standalone_tdcc_belong_to_it(tmp_path: Path, root: Path) -> None:
    first, second = "156199990001", "156199990002"
    header = ("Nomura ID", "Product", "Currency", "BBG Code 1", "Upfront / NotePrice (%)")
    html = order_html(
        "HI team,",
        first,
        f"We buy {first} 600K USD @97.53. Please settle thru Euroclear by acct DSU",
        "Capital size: USD 600K",
        "Initial price: Closing price",
        table(header, ("SQ-1", "FCN", "USD", "AAA UW", "97.53")),
        second,
        f"We buy {second} 700K USD @97.25. Please settle thru Euroclear by acct OSU",
        "Capital size: USD 700K",
        "Initial price: VWAP",
        table(header, ("SQ-2", "FCN", "USD", "BBB UW", "97.25")),
    )
    product_folder(root, first, "order.msg")
    product_folder(root, second, "copy.msg")  # 同一封信也存在第二檔的資料夾
    mail = Mail(f"(Order) FCN--{first} {second}", SENT, html)

    receipt, _ = run(tmp_path, f"{first}\n{second}", {"order.msg": mail, "copy.msg": mail})

    one, two = receipt.rows
    assert (term_reasons(one), term_reasons(two)) == ((), ())
    assert (one.values["期初定價"], one.values["交割說明"], one.values["承作金額"]) == (
        "收盤價",
        "Euroclear by acct DSU",
        Decimal("600000"),
    )
    assert (two.values["期初定價"], two.values["交割說明"], two.values["價格(%)"]) == (
        "VWAP",
        "Euroclear by acct OSU",
        Decimal("97.25"),
    )


def test_ms_table_carries_its_tdcc_and_its_issue_date_column(tmp_path: Path, root: Path) -> None:
    tdcc = "147199990001"
    header = (
        "TDCC",
        "MS ID",
        "Product",
        "Issue Date (T + calendar days)",
        "Stock 1",
        "Curr",
        "Tenor (m)",
        "Note Price",
        "KO Memory",
    )
    html = order_html(
        f"We buy {tdcc} USD 740K @99.20",
        table(header, (tdcc, "M1", "FCN", "20-Oct-2026", "AAA UW", "USD", "1m", "99.20%", "No")),
    )
    product_folder(root, tdcc, "order.msg")

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, html)})

    (row,) = receipt.rows
    assert term_reasons(row) == ()
    assert (row.values["發行日"], row.values["天期(月)"], row.values["價格(%)"]) == (
        dt.date(2026, 10, 20),
        1,
        Decimal("99.20"),
    )
    assert row.others == (("KO Memory", "No"),)


def test_ubs_dates_inside_quote_id_are_split_out(tmp_path: Path, root: Path) -> None:
    tdcc = "092199990001"
    header = ("Product", "Currency", "BBG Code 1", "Cost (%)", "Quote Id")
    quote = (
        "Indicative Price Issue Date: 2026-10-20 Final Fixing Date: 2027-04-20 Final Maturity Date: 2027-04-23 "
        "1st KO Observation: 2026-11-20 UBS ID: abc-123 Cap Size: 5(mm USD)"
    )
    html = order_html(f"We buy {tdcc} USD 300K @98", table(header, ("FCN", "USD", "AAA UW", "98", quote)))
    product_folder(root, tdcc, "order.msg")

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, html)})

    (row,) = receipt.rows
    assert term_reasons(row) == ()
    assert (row.values["發行日"], row.values["最終比價日"], row.values["到期日"], row.values["上手報價編號"]) == (
        dt.date(2026, 10, 20),
        dt.date(2027, 4, 20),
        dt.date(2027, 4, 23),
        "abc-123",
    )
    assert row.others == (
        ("Quote Id", "Indicative Price"),
        ("Quote Id／1st KO Observation", "2026-11-20"),
        ("Quote Id／Cap Size", "5(mm USD)"),
    )


def test_split_embedded_handles_bracketed_parts() -> None:
    text = "[ Request ID: HK1 ] [RIC Quoted: AAA.OQ, BBB.N] [ Issue Date: 20Oct26 ] [ Issuer: GSI ]"

    leftover, parts = split_embedded(text, ISSUERS["148"].embedded["System Remark"])

    assert leftover == ""
    assert parts == [
        ("Request ID", "HK1"),
        ("RIC Quoted", "AAA.OQ, BBB.N"),
        ("Issue Date", "20Oct26"),
        ("Issuer", "GSI"),
    ]


@pytest.mark.parametrize(
    ("values", "reason"),
    [
        (PRICER_VALUES[:-1], "條件表 1 欄數（26）和值數（25）對不上"),
        ((*PRICER_VALUES[:21], "TBD", *PRICER_VALUES[22:]), "交易日 格式不認得：TBD（條件表 1「Trade Date」）"),
        (
            (*PRICER_VALUES[:12], "97.00", *PRICER_VALUES[13:]),
            "價格(%) 不一致：We buy 行「價格」98.50／條件表 1「Upfront / Note Price (%)」97.00",
        ),
        (("FCN", "", *PRICER_VALUES[2:3], "", *PRICER_VALUES[4:]), "讀不到標的1"),
    ],
)
def test_problems_mark_the_product_pending_without_guessing(tmp_path: Path, root: Path, values, reason: str) -> None:
    product_folder(root, BARC, "order.msg")
    mail = Mail(f"(Order)--{BARC}", SENT, standard_order(values=values))

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    (row,) = receipt.rows
    assert row.status == PENDING
    assert reason in row.reasons, row.reasons


def test_conflicting_values_are_left_blank(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    values = ("FCN", "EUR", *PRICER_VALUES[2:])
    mail = Mail(f"(Order)--{BARC}", SENT, standard_order(values=values))

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    (row,) = receipt.rows
    assert "幣別" not in row.values
    assert "幣別 不一致：We buy 行「幣別」USD／條件表 1「Currency」EUR" in row.reasons


def test_unknown_column_means_the_format_may_have_changed(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    mail = Mail(
        f"(Order)--{BARC}", SENT, standard_order(header=(*PRICER_HEADER, "New Field"), values=(*PRICER_VALUES, "x"))
    )

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    (row,) = receipt.rows
    assert row.status == PENDING
    assert row.reasons == ("條件表有沒列過的欄「New Field」（上手格式可能改版）",)
    assert ("New Field", "x") in row.others


def test_several_values_rows_cannot_be_chosen(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    html = order_html(f"We buy {BARC} USD 200K @98.50", table(PRICER_HEADER, PRICER_VALUES, PRICER_VALUES))

    receipt, _ = run(tmp_path, BARC, {"order.msg": Mail(f"(Order)--{BARC}", SENT, html)})

    assert "條件表 1 有 2 列值，無法判斷用哪一列" in receipt.rows[0].reasons


def test_blank_rows_and_warning_boxes_are_not_term_tables(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    warning = table(("", "This message was sent from outside of the company."))
    blank = ("",) * len(PRICER_HEADER)
    html = order_html(warning, f"We buy {BARC} USD 200K @98.50", table(PRICER_HEADER, PRICER_VALUES, blank))

    receipt, _ = run(tmp_path, BARC, {"order.msg": Mail(f"(Order)--{BARC}", SENT, html)})

    assert receipt.rows[0].values["標的1"] == "AAA UW"
    assert not any("列值" in r for r in receipt.rows[0].reasons)


def test_two_order_mails_need_a_person_to_choose(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "a.msg", "b.msg")
    mails = {
        "a.msg": Mail(f"(Order)--{BARC}", SENT, standard_order()),
        "b.msg": Mail(f"(Order) revised--{BARC}", SENT + dt.timedelta(hours=1), standard_order()),
    }

    receipt, _ = run(tmp_path, BARC, mails)

    (row,) = receipt.rows
    assert row.status == PENDING and row.values == {}
    assert row.reasons[0].startswith("有 2 封 (Order) 原信，請確認用哪一封：")


def test_missing_folder_mail_or_bad_input(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "reply.msg")
    product_folder(root, "029199990002", "order.msg")
    mails = {
        "reply.msg": Mail(f"RE: (Order)--{BARC}", SENT, standard_order()),
        "order.msg": Mail("(Order)--029199990009", SENT, standard_order("029199990009")),
    }

    receipt, _ = run(tmp_path, f"{BARC} 029199990002 029199990003 12345", mails)

    reasons = {r.tdcc: (r.status, r.reasons[0]) for r in receipt.rows}
    assert reasons[BARC] == (NO_ORDER, "商品資料夾裡沒有 (Order) 原信")
    assert reasons["029199990002"] == (NO_ORDER, "資料夾裡的 Order 信沒有提到這檔 TDCC：order.msg")
    assert reasons["029199990003"] == (NO_ORDER, "找不到商品資料夾")
    assert reasons["12345"] == (NO_ORDER, "不是 12 碼數字")
    assert "029199990009" not in reasons  # 別檔的信不會因為沒被貼上就加進來


def test_unreadable_mail_needs_checking(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "broken.msg")

    receipt, _ = run(tmp_path, BARC, {"broken.msg": MailReadError("Outlook 開不了這封信")})

    (row,) = receipt.rows
    assert row.status == PENDING
    assert row.reasons == ("商品資料夾裡沒有 (Order) 原信", "讀不到信件 broken.msg：Outlook 開不了這封信")


@pytest.mark.parametrize(
    ("tdcc", "reason"),
    [("998199990001", "這家上手還沒有 Order 讀法（需要樣本）"), ("999199990001", "TDCC 前三碼不是已知上手")],
)
def test_issuers_without_a_format_are_unsupported(
    tmp_path: Path, root: Path, monkeypatch: pytest.MonkeyPatch, tdcc: str, reason: str
) -> None:
    monkeypatch.setitem(KNOWN_WITHOUT_FORMAT, "998", "NEWBANK")
    product_folder(root, tdcc, "order.msg")

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, standard_order(tdcc))})

    (row,) = receipt.rows
    assert (row.status, row.reasons) == (UNSUPPORTED, (reason,))
    assert row.raw  # 原文仍保留，方便人工對照


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("200K", Decimal("200000")),
        ("1.01M", Decimal("1010000")),
        ("45 million", Decimal("45000000")),
        ("USD 950 K", Decimal("950000")),
        ("USD 1 mio", Decimal("1000000")),
        ("JPY 80M", Decimal("80000000")),
        ("10,000,000", Decimal("10000000")),
    ],
)
def test_parse_amount(text: str, expected: Decimal) -> None:
    assert parse_amount(text) == expected


@pytest.mark.parametrize("text", ["about 2M", "200 lots", ""])
def test_parse_amount_rejects_unknown_forms(text: str) -> None:
    with pytest.raises(ValueError):
        parse_amount(text)


@pytest.mark.parametrize("text", ["15-Oct-26", "15-Oct-2026", "15Oct26", "2026-10-15", "2026/10/15", "15 Oct 2026"])
def test_parse_date(text: str) -> None:
    assert parse_date(text) == dt.date(2026, 10, 15)


@pytest.mark.parametrize("text", ["2026-13-45", "Oct 15", "15-Okt-26"])
def test_parse_date_rejects_unknown_forms(text: str) -> None:
    with pytest.raises(ValueError):
        parse_date(text)


# ---------- 「Outlook orders 轉成 Excel」格式表 ----------

FULL_VALUES = {
    "商品": "FCN",
    "幣別": "USD",
    "Non-Call(月)": 1,
    "標的1": "AAA UW",
    "執行價(%)": Decimal("80"),
    "KO觀察方式": "Daily Memory",
    "KO(%)": Decimal("100"),
    "年利率(%)": Decimal("10"),
    "價格(%)": Decimal("98"),
    "天期(月)": 6,
    "KI型態": "EKI",
    "KI(%)": Decimal("65"),
    "觀察頻率(月)": 1,
    "交易日": dt.date(2026, 10, 19),
    "發行日": dt.date(2026, 10, 26),
    "最終比價日": dt.date(2027, 4, 26),
    "到期日": dt.date(2027, 4, 29),
    "期初定價": "收盤價",
}


def cells(values: dict, others=(), subject: str = "(Order)_[Esun Sec] FCN_AAA_B--x", sent_on=SENT):
    row, problems, _ = template_row(BARC, values, others, subject, sent_on)
    return dict(zip((c.header.strip() for c in TEMPLATE_COLUMNS), row, strict=True)), problems


def test_template_row_follows_the_example_columns() -> None:
    row, problems = cells(FULL_VALUES, (("OTC", "Note"),))

    assert problems == []
    assert row == {
        "TDCC": BARC,
        "Product": "FCN",
        "Currency": "USD",
        "Guaranteed Periods (m)": 1,
        "BBG Code 1": "AAA UW",
        "BBG Code 2": None,
        "BBG Code 3": None,
        "BBG Code 4": None,
        "BBG Code 5": None,
        "Strike (%)": Decimal("80"),
        "KO Type": "Daily Memory",
        "KO Barrier (%)": Decimal("100"),
        "Coupon p.a. (%)": Decimal("10"),
        "Upfront / NotePrice (%)": Decimal("98"),
        "Tenor (m)": 6,
        "Barrier Type": "EKI",
        "KI Barrier (%)": Decimal("65"),
        "Observation Frequency (m)": 1,
        "OTC": "Note",
        "Funding Spread (bps)": None,
        "Effective Date Offset(Calendar Days)": 7,
        "Trade date": dt.date(2026, 10, 19),
        "Issue date": dt.date(2026, 10, 26),
        "Final valuation date": dt.date(2027, 4, 26),
        "Maturity date": dt.date(2027, 4, 29),
        "期初價格": "收盤價",
        "起單FC": None,
    }


@pytest.mark.parametrize(
    ("ko_type", "memory", "expected"),
    [
        ("DAILY MEMORY", "", "Daily Memory"),
        ("Daily", "", "Daily"),
        ("OBSERVATION", "No", "Period End"),
        ("OBSERVATION", "Yes", "Period End Memory"),
    ],
)
def test_ko_type_is_normalized(ko_type: str, memory: str, expected: str) -> None:
    row, problems = cells({**FULL_VALUES, "KO觀察方式": ko_type}, (("KO Memory", memory),))

    assert (row["KO Type"], problems) == (expected, [])


@pytest.mark.parametrize(
    ("ki_type", "ki", "expected"),
    [
        ("NONE", Decimal("0"), ("NONE", None)),  # BARC 沒有 KI 時寫 0
        (None, None, ("NONE", None)),  # SG：KI Type、KI 都是 N/A
        ("EUROPEAN", Decimal("80"), ("EKI", Decimal("80"))),  # MS
        ("AKI", Decimal("55"), ("AKI", Decimal("55"))),
    ],
)
def test_barrier_type_is_normalized(ki_type, ki, expected) -> None:
    values = {k: v for k, v in FULL_VALUES.items() if k not in ("KI型態", "KI(%)")}
    values.update({k: v for k, v in (("KI型態", ki_type), ("KI(%)", ki)) if v is not None})

    row, problems = cells(values)

    assert ((row["Barrier Type"], row["KI Barrier (%)"]), problems) == (expected, [])


@pytest.mark.parametrize(
    ("change", "problem"),
    [
        ({"KO觀察方式": "Weekly"}, "KO Type 不認得：Weekly"),
        ({"KI型態": "XKI"}, "Barrier Type 不認得：XKI"),
        ({"KI型態": "NONE", "KI(%)": Decimal("60")}, "Barrier Type 是 NONE，但 KI Barrier 有值：60"),
        ({"標的6": "FFF UW"}, "標的超過 5 檔，格式表放不下"),
    ],
)
def test_values_the_template_cannot_hold_need_checking(change: dict, problem: str) -> None:
    _, problems = cells({**FULL_VALUES, **change})

    assert problem in problems


def test_values_missing_from_the_order_are_derived_by_fixed_rules() -> None:
    values = {k: v for k, v in FULL_VALUES.items() if k not in ("商品", "天期(月)", "觀察頻率(月)", "交易日", "發行日")}
    values["發行日偏移"] = "7"
    others = (("No. of Periods", "12"), ("Settlement Frequency", "Monthly"))

    row, problems = cells(values, others, subject="(Order)_[Esun Sec] FCN_AAA/BBB_Leo -- x")

    assert problems == []
    assert (row["Product"], row["Tenor (m)"], row["Observation Frequency (m)"], row["OTC"]) == ("FCN", 12, 1, "Note")
    # 交易日用 Order 寄出日期，發行日 = 交易日 + 日曆天數偏移
    assert (row["Trade date"], row["Issue date"], row["Effective Date Offset(Calendar Days)"]) == (
        SENT.date(),
        SENT.date() + dt.timedelta(days=7),
        7,
    )


def test_business_day_offset_is_not_used_to_guess_the_issue_date() -> None:
    values = {k: v for k, v in FULL_VALUES.items() if k != "發行日"}
    values["發行日偏移"] = "5BD"

    row, problems = cells(values)

    assert row["Issue date"] is None
    assert problems == ["格式表缺：Effective Date Offset(Calendar Days)、Issue date"]


def test_pending_products_stay_in_the_template_with_red_cells(tmp_path: Path, root: Path) -> None:
    product_folder(root, BARC, "order.msg")
    complete = "029199990002"
    product_folder(root, complete, "order2.msg")
    header = tuple(h for h in PRICER_HEADER if h not in ("Final Valuation Date", "Maturity Date"))
    values = tuple(v for h, v in zip(PRICER_HEADER, PRICER_VALUES, strict=True) if h in header)
    mails = {
        "order.msg": Mail(f"(Order)--{BARC}", SENT, standard_order(header=header, values=values)),
        "order2.msg": Mail(f"(Order)--{complete}", SENT, standard_order(complete)),
    }

    receipt, _ = run(tmp_path, f"{BARC} {complete}", mails)

    first, second = receipt.rows
    assert (first.status, first.reasons) == (PENDING, ("格式表缺：Final valuation date、Maturity date",))
    assert second.status == READ
    ws = openpyxl.load_workbook(receipt.list_path).worksheets[0]
    assert ws.title == "example"
    assert [c.value for c in ws[2]][1:5] == ["TDCC", "Product", "Currency", "Guaranteed Periods (m) "]
    assert ws.max_row == 4  # 待確認的也放進格式表
    # 待確認那列：能確定的照填，讀不到的兩格空白、淺紅底，其他格不標
    assert ws["B3"].value == BARC and ws["X3"].value == dt.datetime(2026, 10, 22)
    assert (ws["Y3"].value, ws["Z3"].value) == (None, None)
    assert {c.coordinate for c in ws[3] if c.fill.fgColor.rgb == FLAG_FILL} == {"Y3", "Z3"}
    # 已讀取那列：格式照範例、沒有紅底
    assert ws["B4"].value == complete and ws["B4"].number_format == "@"
    assert ws["X4"].value == dt.datetime(2026, 10, 22) and ws["X4"].number_format == "mm-dd-yy"
    assert ws["V4"].value == 7 and ws["AA4"].value == "開盤價"
    assert not [c.coordinate for c in ws[4] if c.fill.fgColor.rgb == FLAG_FILL]


@pytest.mark.parametrize(
    ("values", "header", "red"),
    [
        # 價格寫成 96.20→96.15：照填最後的值，價格那格紅底
        ((*PRICER_VALUES[:12], "98.60→98.50", *PRICER_VALUES[13:]), PRICER_HEADER, {"O3"}),
        # 條件表有沒列過的欄：不屬於某一格，標在 TDCC
        ((*PRICER_VALUES, "x"), (*PRICER_HEADER, "New Field"), {"B3"}),
    ],
)
def test_problem_cells_are_marked_red(tmp_path: Path, root: Path, values, header, red) -> None:
    product_folder(root, BARC, "order.msg")
    mail = Mail(f"(Order)--{BARC}", SENT, standard_order(header=header, values=values))

    receipt, _ = run(tmp_path, BARC, {"order.msg": mail})

    assert receipt.rows[0].status == PENDING
    ws = openpyxl.load_workbook(receipt.list_path).worksheets[0]
    assert ws["O3"].value == pytest.approx(98.5)
    assert {c.coordinate for c in ws[3] if c.fill.fgColor.rgb == FLAG_FILL} == red


GS_HEADER = ("Action", "GS Quote ID", "Product", "Currency", "BBG Code 1", "Cost (%)", "Strike Date", "System Remark")


def gs_order(tdcc: str) -> str:
    row = ("IOI", "Q1", "FCN", "USD", "AAA UW", "97.5", "13-Oct-26", "[ Issue Date: 20Oct26 ]")
    return order_html(f"We buy {tdcc} USD 1.06M @97.5", table(GS_HEADER, row))


def test_gs_dates_come_from_the_indicative_docs_mail(tmp_path: Path, root: Path) -> None:
    tdcc = "148199990001"
    product_folder(root, tdcc, "order.msg", "docs.msg", "other.msg")
    summary = table(
        ("TDCC Code", "Issue Date", "Final Valuation Date", "Maturity Date"),
        ("148199990009", "21-Oct-2026", "21-Oct-2027", "25-Oct-2027"),  # 別檔的列不取
        (tdcc, "20-Oct-2026", "20-Oct-2027", "22-Oct-2027"),
    )
    mails = {
        "order.msg": Mail(f"(Order)_[Esun Sec] FCN_AAA_B--{tdcc}", SENT, gs_order(tdcc)),
        "docs.msg": Mail(
            f"[External] Indicative Docs for {tdcc} - XS0000000000 - TD = 13Oct26",
            SENT,
            order_html("Order Summary", summary),
        ),
        "other.msg": Mail(
            "[External] Trade Executed", SENT, order_html(table(("TDCC Code", "Maturity Date"), (tdcc, "01-Jan-2030")))
        ),
    }

    receipt, _ = run(tmp_path, tdcc, mails)

    (row,) = receipt.rows
    assert (row.values["發行日"], row.values["最終比價日"], row.values["到期日"]) == (
        dt.date(2026, 10, 20),
        dt.date(2027, 10, 20),
        dt.date(2027, 10, 22),
    )
    assert not [r for r in term_reasons(row) if "日" in r]
    assert ("補日期：docs.msg", "Maturity Date", "22-Oct-2027") in {(e.source, e.header, e.value) for e, _ in row.raw}


def test_gs_without_the_indicative_docs_mail_says_where_to_look(tmp_path: Path, root: Path) -> None:
    tdcc = "148199990001"
    product_folder(root, tdcc, "order.msg")

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, gs_order(tdcc))})

    (row,) = receipt.rows
    assert row.status == PENDING
    assert any(r.startswith("沒找到補日期的上手信") for r in row.reasons)


def hsbc_order(tdcc: str, dates: tuple[str, str, str, str]) -> str:
    header = ("Product", "Currency", "BBG Code 1", "Upfront / NotePrice (%)")
    vertical = table(
        ("Strike Date", dates[0]), ("Issue Date", dates[1]), ("Final Valuation", dates[2]), ("Maturity Date", dates[3])
    )
    quoted = f"<p>Dear team,</p><p>TDCC code: {tdcc}</p>{vertical}{table(header, ('FCN', 'USD', 'AAA UW', '96.15'))}"
    return order_html(f"We buy {tdcc} USD 400K @96.15", table(header, ("FCN", "USD", "AAA UW", "96.15")), quoted=quoted)


def test_hsbc_dates_come_from_its_reply_below_the_order(tmp_path: Path, root: Path) -> None:
    tdcc = "325199990001"
    product_folder(root, tdcc, "order.msg")
    html = hsbc_order(tdcc, ("13/10/2026", "20/10/2026", "21/12/2026", "24/12/2026"))

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, html)})

    (row,) = receipt.rows
    # 日/月/年
    assert (row.values["交易日"], row.values["發行日"], row.values["最終比價日"], row.values["到期日"]) == (
        dt.date(2026, 10, 13),
        dt.date(2026, 10, 20),
        dt.date(2026, 12, 21),
        dt.date(2026, 12, 24),
    )
    assert not [r for r in term_reasons(row) if "日" in r]


def test_supplement_dates_never_override_the_order(tmp_path: Path, root: Path) -> None:
    tdcc = "325199990001"
    header = ("Product", "Currency", "BBG Code 1", "Upfront / NotePrice (%)", "Issue Date")
    vertical = table(
        ("Issue Date", "21/10/2026"),
    )
    quoted = f"<p>TDCC code: {tdcc}</p>{vertical}"
    html = order_html(
        f"We buy {tdcc} USD 400K @96.15", table(header, ("FCN", "USD", "AAA UW", "96.15", "20-Oct-26")), quoted=quoted
    )
    product_folder(root, tdcc, "order.msg")

    receipt, _ = run(tmp_path, tdcc, {"order.msg": Mail(f"(Order)--{tdcc}", SENT, html)})

    (row,) = receipt.rows
    assert "發行日" not in row.values
    assert any(r.startswith("發行日 不一致") for r in row.reasons)


def test_parse_date_day_first_only_when_asked() -> None:
    assert parse_date("05/06/2026", day_first=True) == dt.date(2026, 6, 5)
    with pytest.raises(ValueError):
        parse_date("05/06/2026")


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        (
            "We buy 029199990001 USD 200K @98.50",
            [("029199990001", [("幣別", "USD"), ("金額", "200K"), ("價格", "98.50")])],
        ),
        (
            "We buy 076199990001 45 million JPY @97.9. Please settle thru Euroclear by acct DSU.",
            [("076199990001", [("幣別", "JPY"), ("金額", "45 million"), ("價格", "97.9")])],
        ),
        (
            "We buy 328199990001 USD 950K @96.15、328199990002 USD 1 mio @96.15",
            [
                ("328199990001", [("幣別", "USD"), ("金額", "950K"), ("價格", "96.15")]),
                ("328199990002", [("幣別", "USD"), ("金額", "1 mio"), ("價格", "96.15")]),
            ],
        ),
        (
            "We buy 159199990001 USD 300K , UF=0.85%",
            [("159199990001", [("幣別", "USD"), ("金額", "300K"), ("UF", "0.85")])],
        ),
    ],
)
def test_buy_segments(line: str, expected) -> None:
    assert buy_segments(line) == expected


CITI_HEADER = (
    "Product",
    "Strike Date",
    "Currency",
    "Tenor (m)",
    "Issue T+",
    "BBG Code 1",
    "Strike (%)",
    "Barrier Type",
    "KI Barrier (%)",
    "Observation Frequency (m)",
    "Non Callable Periods",
    "KO Barrier (%)",
    "Memory Autocall",
    "Daily KO",
    "Coupon p.a. (%)",
    "Upfront (%)",
    "Notional Amount",
    "Issue Date",
    "Final Valuation Date",
    "Maturity Date",
    "PriceID",
)


def test_citi_gives_uf_instead_of_a_price(tmp_path: Path, root: Path) -> None:
    tdcc = "159199990001"
    values = (
        "FCA",
        "8-Oct-26",
        "USD",
        "7",
        "5",
        "AAA UN",
        "80",
        "European",
        "75",
        "1",
        "0",
        "105",
        "TRUE",
        "TRUE",
        "15.60",
        "0.85",
        "500,000",
        "16-Oct-26",
        "17-May-27",
        "20-May-27",
        "P1",
    )
    html = order_html(
        f"We buy {tdcc} USD 300K , UF=0.85%",
        "Initial price: Closing price",
        table(CITI_HEADER, values),
    )
    product_folder(root, tdcc, "order.msg")
    mail = Mail(f"(Order)_[Esun Sec] FCN_AAA_LU--{tdcc}", SENT, html)

    receipt, _ = run(tmp_path, tdcc, {"order.msg": mail})

    (row,) = receipt.rows
    assert (row.status, row.issuer, row.reasons) == (READ, "CITI", ())
    # 價格 = 100 − UF；Notional Amount 是報價名目本金，不是承作金額
    assert (row.values["價格(%)"], row.values["UF(%)"], row.values["承作金額"]) == (
        Decimal("99.15"),
        Decimal("0.85"),
        Decimal("300000"),
    )
    assert "商品" not in row.values and ("Product", "FCA") in row.others
    assert row.values["Non-Call(月)"] == 1  # Citi 的 Non Callable Periods 0 統一 +1
    template = dict(zip((c.header.strip() for c in TEMPLATE_COLUMNS), row.template, strict=True))
    assert (template["Product"], template["KO Type"], template["Barrier Type"]) == ("FCN", "Daily Memory", "EKI")
    assert (template["Trade date"], template["Effective Date Offset(Calendar Days)"]) == (dt.date(2026, 10, 8), 8)


@pytest.mark.parametrize(
    ("daily", "memory", "expected"),
    [
        ("TRUE", "TRUE", "Daily Memory"),
        ("TRUE", "FALSE", "Daily"),
        ("FALSE", "TRUE", "Period End Memory"),
        ("FALSE", "FALSE", "Period End"),
    ],
)
def test_citi_ko_flags(daily: str, memory: str, expected: str) -> None:
    values = {k: v for k, v in FULL_VALUES.items() if k != "KO觀察方式"}

    row, problems = cells(values, (("Daily KO", daily), ("Memory Autocall", memory)))

    assert (row["KO Type"], problems) == (expected, [])
