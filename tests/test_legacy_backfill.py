"""舊商品補價格與比價日（#154、ADR 0020「舊商品」）：只補已鎖住、還缺資料的商品，庫存狀態轉舊清單事件。

資料庫測試對真的 PostgreSQL；參考條件表、商品、價格與日期全部虛構。
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from db_synth import inventory, seed_products
from fcn_checker.check_config import ConfigPaths
from fcn_checker.config import load_reference_format
from fcn_checker.orders.reference import headers_of
from sp_workflow.legacy_backfill import (
    EXISTING,
    FILLED,
    SKIPPED,
    LegacyBackfillError,
    _LegacyEvent,
    _status_event,
    backfill_legacy,
)
from sp_workflow.lifecycle import KO, LEGACY, MATURITY, LifecycleError, current_events, record_legacy, record_redemption
from sp_workflow.window_actions import import_legacy_products

CONFIG_DIR = Path(__file__).resolve().parents[1] / "config"
FMT = load_reference_format(CONFIG_DIR / "reference_sheet.toml")
PATHS = ConfigPaths.with_fallback(CONFIG_DIR / "review_standard.toml")
HEADERS = [*dict.fromkeys(headers_of(FMT).values()), "發行機構", "庫存狀態"]
A, B, C, D, E, F, G, H = (f"14719999000{i}" for i in range(1, 9))
DAY_1 = dt.date(2026, 4, 22)


def sheet_row(tdcc: str, **cells) -> dict:
    """回填後庫存總表的一列：兩檔標的的四個價格、比價日 1 有值、2～12 是空值寫法。"""
    row = {
        "TDCC Code": tdcc,
        "發行機構": "XBANK",
        "UL_1": "AAA UW",
        "UL_1_進場價": 100,
        "UL_1_執行價": 80,
        "UL_1_下限價": 65,
        "UL_1_KO價": 100,
        "UL_2": "BBB UN",
        "UL_2_進場價": 50.5,
        "UL_2_執行價": 40.4,
        "UL_2_下限價": 32.825,
        "UL_2_KO價": 50.5,
        "比價日_1": DAY_1,
        **{f"比價日_{n}": "-" for n in range(2, 13)},
        **{f"UL_{i}{suffix}": "-" for i in range(3, 6) for suffix in ("", "_進場價", "_執行價", "_下限價", "_KO價")},
    }
    row.update(cells)
    return row


def write_sheet(path: Path, rows: list[dict], headers: list[str] = HEADERS) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    book.save(path)
    return path


def locked(pg, *tdccs: str) -> None:
    seed_products(pg, [inventory(t, [("台北", 10000)]) for t in tdccs])
    pg.execute("UPDATE workflow.product SET locked_at = now() WHERE tdcc = ANY(%s)", (list(tdccs),))


def prices(pg, tdcc: str) -> list[tuple]:
    return pg.execute(
        "SELECT seq, code, initial_price, strike_price, ki_price, ko_price FROM workflow.underlying_price"
        " WHERE tdcc = %s ORDER BY seq",
        (tdcc,),
    ).fetchall()


def dates(pg, tdcc: str) -> list[tuple]:
    return pg.execute(
        "SELECT slot, value FROM workflow.autocall_date WHERE tdcc = %s ORDER BY slot", (tdcc,)
    ).fetchall()


def lines(receipt) -> dict[str, tuple[str, str]]:
    return {line.tdcc: (line.result, line.detail) for line in receipt.lines}


@pytest.mark.db
def test_prices_dates_and_legacy_events_are_filled_once(pg, tmp_path: Path) -> None:
    locked(pg, A, B)
    source = write_sheet(
        tmp_path / "庫存總表.xlsx",
        [sheet_row(A, 庫存狀態="已於2026/10/1 KO"), sheet_row(B, 庫存狀態="到期還本"), sheet_row(C)],
    )
    seed_products(pg, [inventory(C, [("台北", 10000)])])  # 沒鎖住

    receipt = backfill_legacy(pg, source, FMT, user="tester")

    assert receipt.count(FILLED) == 2 and receipt.events == 2 and receipt.has_status
    assert prices(pg, A) == [
        (1, "AAA UW", Decimal(100), Decimal(80), Decimal(65), Decimal(100)),
        (2, "BBB UN", Decimal("50.5"), Decimal("40.4"), Decimal("32.825"), Decimal("50.5")),
    ]
    assert dates(pg, A) == [(1, DAY_1), *((n, None) for n in range(2, 13))], "空值寫法存空白"
    (terms,) = pg.execute("SELECT status, reason FROM workflow.check_terms WHERE tdcc = %s", (A,)).fetchall()
    assert terms[0] == "完成" and "庫存總表.xlsx" in terms[1], "價格掛在核對條件上，之後比價會比"
    (ko,) = current_events(pg, A)
    assert (ko.event, ko.event_date, ko.source, ko.recorded_by) == (KO, dt.date(2026, 10, 1), LEGACY, "tester")
    (maturity,) = current_events(pg, B)
    assert (maturity.event, maturity.event_date, maturity.source) == (MATURITY, None, LEGACY)
    assert lines(receipt)[C][0] == SKIPPED and "還沒鎖住" in lines(receipt)[C][1]
    assert prices(pg, C) == [] and dates(pg, C) == []
    (name,) = pg.execute("SELECT isin FROM workflow.product WHERE tdcc = %s", (A,)).fetchone()
    assert name == "XS0000000001", "其他欄不碰"

    again = backfill_legacy(pg, source, FMT, user="tester")

    assert again.count(EXISTING) == 2 and again.events == 0
    assert len(prices(pg, A)) == 2 and len(dates(pg, A)) == 12 and len(current_events(pg, A)) == 1


@pytest.mark.db
def test_a_bad_row_only_skips_that_product(pg, tmp_path: Path) -> None:
    locked(pg, B, C, D, F, G, H)
    source = write_sheet(
        tmp_path / "庫存總表.xlsx",
        [
            sheet_row(B, UL_1="ZZZ UW"),
            sheet_row(C, UL_1_進場價="abc"),
            sheet_row(D, 比價日_3="下週", UL_2_KO價="-"),
            sheet_row(E),
            sheet_row(G, UL_1="aaa  uw equity"),
            sheet_row(H),
            sheet_row(H),
        ],
    )

    receipt = backfill_legacy(pg, source, FMT, user="tester", now=dt.datetime(2026, 10, 23, 9, 0))

    found = lines(receipt)
    assert found[G][0] == FILLED, "代碼不分大小寫、空白與結尾的 Equity"
    assert "UL_1 代碼和資料庫不同（表上 ZZZ UW，資料庫 AAA UW）" in found[B][1]
    assert "UL_1 進場價不是數字（abc）" in found[C][1]
    assert "比價日_3 不是日期（下週）" in found[D][1] and "UL_2 KO 價是空白" in found[D][1]
    assert found[E][1] == "資料庫沒有這檔"
    assert found[F][1].startswith("表上沒有這檔")
    assert "表上有 2 列" in found[H][1]
    assert {t for t, (r, _) in found.items() if r == SKIPPED} == {B, C, D, E, F, H}
    for tdcc in (B, C, D, F, H):
        assert prices(pg, tdcc) == [] and dates(pg, tdcc) == [], "不匯入的那一檔什麼都不寫"
    assert receipt.error_book_path == tmp_path / "舊商品補資料錯誤_20261023-090000.xlsx"
    book = openpyxl.load_workbook(receipt.error_book_path)
    listed = [row[0] for row in book.active.iter_rows(min_row=2, values_only=True)]
    assert sorted(listed) == sorted([B, C, D, E, F, H])


@pytest.mark.db
def test_events_are_not_duplicated_or_contradicted(pg, tmp_path: Path) -> None:
    locked(pg, A, B)
    record_redemption(pg, A, dt.date(2026, 9, 1), "顧客來電", user="tester")
    record_legacy(pg, B, KO, dt.date(2026, 10, 1), "舊清單", user="tester")
    source = write_sheet(
        tmp_path / "庫存總表.xlsx",
        [sheet_row(A, 庫存狀態="已於2026/10/1 KO"), sheet_row(B, 庫存狀態="已於2026/10/2 KO")],
    )

    receipt = backfill_legacy(pg, source, FMT, user="tester")

    found = lines(receipt)
    assert (
        found[A][0] == SKIPPED and "已有存續事件「顧客贖回」（人工），和表上的「已於2026/10/1 KO」不同" in found[A][1]
    )
    assert found[B] == (
        FILLED,
        "補了價格、比價日 1～12；已有KO事件，沒有再寫（日期不同：資料庫 2026-10-01，表上 2026-10-02，請確認）",
    )
    assert len(current_events(pg, B)) == 1 and receipt.events == 0


@pytest.mark.db
def test_check_terms_that_are_not_done_are_kept_and_flagged(pg, tmp_path: Path) -> None:
    locked(pg, A)
    pg.execute("INSERT INTO workflow.check_terms (tdcc, status, reason) VALUES (%s, '待確認', '舊的原因')", (A,))
    source = write_sheet(tmp_path / "庫存總表.xlsx", [sheet_row(A)])

    receipt = backfill_legacy(pg, source, FMT, user="tester")

    assert lines(receipt)[A] == (
        FILLED,
        "補了價格、比價日 1～12；核對條件是「待確認」，比價不會比這檔，要先處理核對條件",
    )
    (terms,) = pg.execute("SELECT status, reason FROM workflow.check_terms WHERE tdcc = %s", (A,)).fetchall()
    assert terms == ("待確認", "舊的原因"), "已有的核對條件不動"


@pytest.mark.db
def test_a_failed_import_leaves_no_error_book(pg, tmp_path: Path, monkeypatch) -> None:
    import sp_workflow.legacy_backfill

    locked(pg, A)
    source = write_sheet(tmp_path / "庫存總表.xlsx", [sheet_row(A), sheet_row(E)])

    def broken(*args, **kwargs):
        raise RuntimeError("寫不進去")

    monkeypatch.setattr(sp_workflow.legacy_backfill, "_write_dates", broken)

    with pytest.raises(RuntimeError):
        backfill_legacy(pg, source, FMT, user="tester")
    assert list(tmp_path.glob("舊商品補資料錯誤_*.xlsx")) == [] and prices(pg, A) == []


@pytest.mark.db
def test_only_the_missing_part_is_filled(pg, tmp_path: Path) -> None:
    locked(pg, A)
    pg.execute("INSERT INTO workflow.autocall_date (tdcc, slot, value) VALUES (%s, 1, %s)", (A, dt.date(2026, 5, 1)))
    source = write_sheet(tmp_path / "庫存總表.xlsx", [sheet_row(A, 比價日_2="不是日期")])

    receipt = backfill_legacy(pg, source, FMT, user="tester")

    assert lines(receipt)[A] == (FILLED, "補了價格")
    assert dates(pg, A) == [(1, dt.date(2026, 5, 1))], "已經有的比價日不動，也不檢查"
    assert len(prices(pg, A)) == 2


@pytest.mark.db
def test_a_sheet_without_the_status_column_writes_no_events(pg, tmp_path: Path) -> None:
    locked(pg, A)
    source = write_sheet(tmp_path / "參考條件表.xlsx", [sheet_row(A)], [h for h in HEADERS if h != "庫存狀態"])

    receipt = backfill_legacy(pg, source, FMT, user="tester")

    assert receipt.count(FILLED) == 1 and not receipt.has_status and current_events(pg, A) == []


@pytest.mark.db
def test_a_sheet_missing_required_columns_is_refused(pg, tmp_path: Path) -> None:
    locked(pg, A)
    source = write_sheet(tmp_path / "參考條件表.xlsx", [sheet_row(A)], [h for h in HEADERS if h != "比價日_5"])

    with pytest.raises(LegacyBackfillError, match="缺少必要欄：比價日_5.*整份不匯入"):
        backfill_legacy(pg, source, FMT, user="tester")
    assert prices(pg, A) == []


@pytest.mark.db
def test_the_window_action_observes_afterwards(pg, pg_settings, tmp_path: Path) -> None:
    locked(pg, A)
    source = write_sheet(tmp_path / "庫存總表.xlsx", [sheet_row(A)])

    result = import_legacy_products(pg_settings, source, PATHS)

    assert result.receipt.count(FILLED) == 1
    assert result.observation.receipt is not None and result.observation.error == ""


@pytest.mark.parametrize(
    ("status", "found"),
    [
        ("已於2026/10/1 KO", (_LegacyEvent(KO, dt.date(2026, 10, 1)), "")),
        ("已於 2026 / 3 / 5 KO", (_LegacyEvent(KO, dt.date(2026, 3, 5)), "")),
        ("到期還本", (_LegacyEvent(MATURITY, None), "")),
        ("在庫", (None, "")),
        (None, (None, "")),
        ("已於2026/2/30 KO", (None, "庫存狀態「已於2026/2/30 KO」的 KO 日期不是有效的日期")),
        ("KO", (None, "庫存狀態「KO」看不出 KO 的日期（要寫成「已於 Y/M/D KO」）")),
    ],
)
def test_status_text_becomes_an_event(status, found) -> None:
    assert _status_event(status) == found


@pytest.mark.db
def test_legacy_events_need_a_product_without_events(pg) -> None:
    locked(pg, A)
    record_legacy(pg, A, MATURITY, None, "舊清單", user="tester")

    with pytest.raises(LifecycleError, match="已有存續事件「到期還本」"):
        record_legacy(pg, A, KO, dt.date(2026, 10, 1), "舊清單", user="tester")
    with pytest.raises(LifecycleError, match="只能是 KO 或到期還本"):
        record_legacy(pg, A, "顧客贖回", dt.date(2026, 10, 1), "舊清單", user="tester")
