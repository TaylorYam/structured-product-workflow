"""商品總表的分析分頁（#293、#299）：在庫部位 by 發行機構、分公司、幣別，全部用 Excel 公式從總表即時計算。

「商品總表」分頁是名為 `總表` 的 Excel 表格（Table），公式用結構化參照；在 Excel 裡改總表、或在表格下方新增一列，
分析分頁跟著重算。每頁只有 A1 一個動態陣列公式（`UNIQUE`、`FILTER`、`SORT`…），新的發行機構、分公司、幣別
自動多一列／一欄；需要 Excel 2021 或 Microsoft 365。

每頁兩張表上下並列：表頭（A 欄名稱、B 欄起幣別依在庫檔數由多到少）、總計、每個名稱一列；空一列後是同樣排法的占比
（同幣別內依金額）。上面那張的列數會變，分成兩個公式各自溢出會互撞成 `#SPILL!`，所以由同一個公式排出兩張表；
只用 Excel 2021 有的函式，不用 `VSTACK`。格式先套在一塊夠大的範圍，表頭底色、總計框線、占比的百分比格式
用條件式格式依 A 欄內容套上。金額是原幣、不換算；只算存續狀態「在庫」的列，`-` 與空白的名稱不列。
分公司照交易明細的 `分公司_n`／`金額_n`，一檔由幾家分公司承作就分到那幾家。

openpyxl 不支援寫動態陣列：公式以陣列公式寫入，存檔後由 `enable_spill` 在檔案裡補 Excel 的動態陣列中繼資料
（`xl/metadata.xml`、每格 `cm="1"`），否則 Excel 當成舊式陣列公式、不會溢出。檔案裡的公式要用 Excel 的存檔寫法：
新函式加 `_xlfn.` 前綴、`LET` 的變數加 `_xlpm.`（`file_formula`）。
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Callable
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

TABLE_NAME = "總表"
IN_STOCK = "在庫"
ISSUER_SHEET, BRANCH_SHEET = "發行機構x幣別", "分公司x幣別"
SHARE_SUFFIX = "占比"  # 下面那張表的表頭：發行機構占比、分公司占比
FORMAT_ROWS, FORMAT_COLUMNS = 205, 21  # 先套格式的範圍：兩張表各 100 個名稱、20 個幣別
AMOUNT_FORMAT = '#,##0;-#,##0;"-"'  # 0 顯示 -
SHARE_FORMAT = '0.0%;-0.0%;"-"'
CJK_FONT, LATIN_FONT = "微軟正黑體", "Calibri"
HEAD_COLOR = "1F4E79"

_NEW_FUNCTIONS = {"SORTBY": "_xlfn.", "SORT": "_xlfn._xlws.", "FILTER": "_xlfn._xlws.", "UNIQUE": "_xlfn.",
                  "SEQUENCE": "_xlfn.", "LET": "_xlfn."}  # fmt: skip
_NAMES = ("cs", "cur", "k", "n", "rw", "j", "b", "q", "cx", "kx", "t", "v")  # _stacked 的 LET 變數
_THIN = Side(style="thin", color="BFBFBF")
_BLACK = Side(style="thin", color="000000")


def summary_table(sheet: Worksheet, columns: int, rows: int) -> None:
    """把總表的範圍設成 Excel 表格（`TABLE_NAME`，有篩選按鈕）；沒有商品時多包一列空白，Excel 表格至少要有一列。"""
    ref = f"A1:{get_column_letter(columns)}{max(rows, 1) + 1}"
    style = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    sheet.add_table(Table(displayName=TABLE_NAME, ref=ref, tableStyleInfo=style))


def add_analysis_sheets(book: openpyxl.Workbook, groups: int) -> None:
    """總表後面接兩個分析分頁；`groups` 是總表的分公司／FC／金額組數。"""
    in_stock = f'{_col("存續狀態")},"{IN_STOCK}"'
    _stacked(
        book,
        ISSUER_SHEET,
        "發行機構",
        f'SORT(UNIQUE(FILTER({_col("發行機構")},{_in_stock()}*({_col("發行機構")}<>"-"),"")))',
        lambda name, currency: (
            f"SUMIFS({_col('承作總額')},{_col('發行機構')},{name},{_col('承作幣別')},{currency},{in_stock})"
        ),
        lambda currency: f"SUMIFS({_col('承作總額')},{_col('承作幣別')},{currency},{in_stock})",
    )

    block = f"{TABLE_NAME}[[分公司_1]:[金額_{groups}]]"  # 第 n 組的分公司在這塊的第 3n-2 欄
    _stacked(
        book,
        BRANCH_SHEET,
        "分公司",
        f'SORT(UNIQUE(FILTER(br,(INDEX({_col("存續狀態")},i)="{IN_STOCK}")*(br<>"-")*(br<>""),"")))',
        lambda name, currency: "+".join(
            f"SUMIFS({_col(f'金額_{n}')},{_col(f'分公司_{n}')},{name},{_col('承作幣別')},{currency},{in_stock})"
            for n in range(1, groups + 1)
        ),
        lambda currency: "+".join(
            f"SUMIFS({_col(f'金額_{n}')},{_col('承作幣別')},{currency},{in_stock})" for n in range(1, groups + 1)
        ),
        prelude=(
            f"s,SEQUENCE(ROWS({_col('存續狀態')})*{groups}),i,INT((s-1)/{groups})+1,"
            f"br,INDEX({block},i,MOD(s-1,{groups})*3+1),"
        ),
        prelude_names=("s", "i", "br"),
    )
    book.calculation.fullCalcOnLoad = True  # 公式沒有存值，打開時整本重算


def file_formula(formula: str, names: tuple[str, ...] = ()) -> str:
    """平常看到的公式 → Excel 檔案裡的寫法：新函式加前綴、`LET` 的變數（`names`）加 `_xlpm.`。"""
    for name in names:
        formula = re.sub(rf"(?<![\w.\[]){name}(?![\w\[(])", f"_xlpm.{name}", formula)
    for function, prefix in _NEW_FUNCTIONS.items():
        formula = re.sub(rf"(?<![\w.]){function}\(", f"{prefix}{function}(", formula)
    return formula


def enable_spill(path: Path) -> None:
    """存好的 xlsx 補上動態陣列中繼資料，陣列公式才會像 Excel 365 的公式一樣溢出。"""
    with zipfile.ZipFile(path) as source:
        parts = [(item, source.read(item.filename)) for item in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item, data in parts:
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", _CONTENT_TYPE + b"</Types>")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", _RELATIONSHIP + b"</Relationships>")
            elif item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb'<c (r="[A-Z]+\d+"[^>]*)><f t="array"', rb'<c \1 cm="1"><f t="array"', data)
            target.writestr(item, data)
        target.writestr("xl/metadata.xml", _METADATA)


def _col(name: str) -> str:
    return f"{TABLE_NAME}[{name}]"


def _in_stock() -> str:
    return f'({_col("存續狀態")}="{IN_STOCK}")'


def _stacked(
    book: openpyxl.Workbook,
    title: str,
    label: str,
    names: str,
    amount: Callable[[str, str], str],
    total: Callable[[str], str],
    *,
    prelude: str = "",
    prelude_names: tuple[str, ...] = (),
) -> None:
    """一個分析分頁：A1 的公式排出總額表、空一列、占比表。

    `names` 是直的名稱清單，可用 `prelude` 先定義的 `LET` 變數；`amount(名稱, 幣別)`、`total(幣別)` 是金額公式，
    拿到的名稱、幣別已攤成整個輸出的大小，SUMIFS 逐格算。輸出 2n+5 列：第 1 列表頭、第 2 列總計、
    第 3～n+2 列名稱、第 n+3 列空白，下面同樣排法是占比（`rw` 是第幾列、`j` 第幾欄，`b` 標出下面那張，`q` 是在自己那張的第幾列）。
    """
    currency = _col("承作幣別")
    formula = (
        f'LET(cs,UNIQUE(FILTER({currency},{_in_stock()}*({currency}<>"-"),"")),'
        f'cur,TRANSPOSE(SORTBY(cs,COUNTIFS({currency},cs,{_col("存續狀態")},"{IN_STOCK}"),-1)),'
        f"{prelude}k,{names},n,ROWS(k),rw,SEQUENCE(2*n+5),j,SEQUENCE(1,COLUMNS(cur)+1),"
        "b,rw>n+3,q,IF(b,rw-n-3,rw),cx,INDEX(cur,1,IF(j>1,j-1,1)),kx,INDEX(k,IF(q>2,q-2,1)),"
        f"t,{total('cx')},v,IF(q=2,t,{amount('kx', 'cx')}),"
        f'IF(rw=n+3,"",IF(q=1,IF(j=1,IF(b,"{label}{SHARE_SUFFIX}","{label}"),cx),'
        'IF(j=1,IF(q=2,"總計",kx),IF(b,IF(t=0,0,v/t),v)))))'
    )
    sheet = book.create_sheet(title)
    sheet["A1"] = ArrayFormula("A1", "=" + file_formula(formula, _NAMES + prelude_names))
    _layout(sheet, label)


def _layout(sheet: Worksheet, label: str) -> None:
    """格式先套在 `FORMAT_ROWS` × `FORMAT_COLUMNS`；表頭、總計、名稱列的框線與占比格式用條件式格式依 A 欄套上。"""
    sheet.sheet_view.showGridLines = False
    for row in range(1, FORMAT_ROWS + 1):
        sheet.cell(row, 1).font = Font(name=CJK_FONT)
        sheet.cell(row, 1).alignment = Alignment(vertical="center")
        for column in range(2, FORMAT_COLUMNS + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=LATIN_FONT)
            cell.number_format = AMOUNT_FORMAT
            cell.alignment = Alignment(horizontal="right", vertical="center")
    sheet.column_dimensions["A"].width = 16
    for column in range(2, FORMAT_COLUMNS + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 17
    end = f"{get_column_letter(FORMAT_COLUMNS)}{FORMAT_ROWS}"
    share = f"{label}{SHARE_SUFFIX}"
    used = 'A$1<>""'  # 第 1 列是表頭，有字的欄才在表格裡
    head = DifferentialStyle(font=Font(color="FFFFFF", bold=True), fill=PatternFill(bgColor=HEAD_COLOR))
    total = DifferentialStyle(font=Font(bold=True), border=Border(top=_BLACK, bottom=_BLACK))
    line = DifferentialStyle(border=Border(bottom=_THIN))
    percent = DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode=SHARE_FORMAT))
    sheet.conditional_formatting.add(f"A1:{end}", _rule(f'AND({used},OR($A1="{label}",$A1="{share}"))', head))
    sheet.conditional_formatting.add(f"A1:{end}", _rule(f'AND({used},$A1="總計")', total))
    sheet.conditional_formatting.add(f"A1:{end}", _rule(f'AND({used},$A1<>"")', line))
    sheet.conditional_formatting.add(f"B1:{end}", _rule(f'COUNTIF($A$1:$A1,"{share}")>0', percent))
    sheet.freeze_panes = "B2"


def _rule(formula: str, style: DifferentialStyle) -> Rule:
    return Rule(type="expression", dxf=style, formula=[formula])


_CONTENT_TYPE = (
    b'<Override PartName="/xl/metadata.xml"'
    b' ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
)
_RELATIONSHIP = (
    b'<Relationship Id="rIdSpill" Target="metadata.xml"'
    b' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata"/>'
)
_METADATA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    ' xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1"'
    ' pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1"'
    ' coerce="1" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
    '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)
