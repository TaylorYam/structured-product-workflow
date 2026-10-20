"""準備批次：依貼上的 TDCC 找商品資料夾，挑出中文說明書與投資人須知，改成標準檔名放進新的批次資料夾（#5、#7）。

對外只有 `prepare_batch`。商品資料夾只讀不寫；批次資料夾每次新開、已存在就不做。
同一種文件剛好一份候選才採用；說明書或投資人須知任一份找不到或有多份，整個商品跳過。
有上手 PDF 候選就只看 PDF；某種文件沒有 PDF 候選時，才把資料夾裡的 Word 複製到暫存區轉成 PDF 再判斷（#8、ADR 0002）。
Word 無法啟動或逾時，整個商品跳過；個別 Word 轉不了只影響它自己（需要它卻沒有其他候選時，原因寫出轉檔失敗）。
MS 說明書封面空白的「受託機構商品代號」在暫存複本補上 TDCC 再轉檔（#56）。

重新核對上手修正版（#21、#23）：放回過的 `_待確認` 檔不當候選，只標記「這個內容沒通過」；候選依內容指紋（全文）
只算一份；上手寄來新版時排除和待確認檔內容相同的舊版，還沒新版時沿用上次沒通過的那一版，同一版可以重複核對；
剩下不只一份時，上手檔名的版本優先：排除和上次放回的標準名檔內容相同的舊版，以及只以標準檔名存在的內容
（上手修正版沿用舊檔名、覆蓋了舊原檔時，舊版只剩工具放回的標準名複本，#27）。
"""

from __future__ import annotations

import datetime as dt
import hashlib
import os
import re
import shutil
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sp_workflow.conversion import ConversionError, Converter, default_converter
from sp_workflow.documents import DocKind, content_key, read_pdf
from sp_workflow.trustee_code import fill_trustee_code, needs_fill

PREPARED = "已準備"
PENDING_SUFFIX = "_待確認"  # 核對沒通過、放回時加在檔名最後（放回資料夾也用這個）
SKIPPED = "跳過"
LIST_NAME = "準備清單.xlsx"
# 準備清單的欄（欄名, 欄寬）
LIST_COLUMNS = (
    ("TDCC", 15),
    ("種類", 6),
    ("來源檔", 60),
    ("商品資料夾", 40),
    ("由 Word 轉檔", 12),
    ("輸出檔名", 22),
    ("SHA-256", 66),
    ("狀態", 8),
    ("原因", 60),
)

_SEPARATORS = re.compile(r"[\s,，、;；]+")
_TDCC = re.compile(r"[0-9]{12}")  # 只認半形數字
_PDF = ".pdf"
_WORD = (".docx", ".doc")


class PreparationError(Exception):
    """整批無法開始（輸入或資料夾有問題）；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class PreparedDocument:
    """一個商品的一種文件；跳過的商品若找到唯一候選也會記下來源檔，但沒有輸出。"""

    kind: DocKind
    source: Path | None = None
    output_name: str = ""
    sha256: str = ""
    from_word: bool = False


@dataclass(frozen=True)
class _Candidate:
    """判讀用的 PDF（上手 PDF 本身，或 Word 轉出的暫存 PDF）與它的來源檔。"""

    pdf: Path
    source: Path
    from_word: bool = False
    filled: bool = False  # 轉檔前補了受託機構商品代號（#56）


@dataclass(frozen=True)
class _Selection:
    """一種文件挑完版本之後：剩下的候選、排除掉的沒通過版本與已放回舊版，以及是否沿用上次沒通過的版本。"""

    chosen: list[_Candidate]
    failed: list[_Candidate]
    superseded: list[_Candidate]
    reused_failed: bool = False  # 上手還沒寄新版，沿用上次沒通過的那一版（chosen 就是它）


@dataclass(frozen=True)
class ProductOutcome:
    tdcc: str
    status: str
    reason: str = ""
    folder: Path | None = None
    documents: tuple[PreparedDocument, ...] = ()


@dataclass(frozen=True)
class PrepareReceipt:
    batch_dir: Path
    list_path: Path
    outcomes: tuple[ProductOutcome, ...]
    unreadable_dirs: tuple[Path, ...] = ()  # 搜尋商品資料夾時讀不到的資料夾（例如沒有權限）


def prepare_batch(
    text: str,
    root: Path,
    out_parent: Path,
    *,
    now: dt.datetime | None = None,
    converter: Converter | None = None,
) -> PrepareReceipt:
    """準備一個批次：回傳批次資料夾、準備清單路徑與逐商品結果（依輸入順序）。

    converter 預設用 Microsoft Word 轉檔（只有需要轉 Word 時才會啟動）。
    """
    convert = converter or default_converter()
    tokens = parse_tdccs(text)
    if not tokens:
        raise PreparationError("沒有輸入 TDCC")
    if not root.is_dir():
        raise PreparationError(f"根目錄不存在：{root}")
    if not out_parent.is_dir():
        raise PreparationError(f"輸出位置不存在：{out_parent}")
    batch_dir = out_parent / f"待核對_{(now or dt.datetime.now()):%Y%m%d-%H%M%S}"
    try:
        batch_dir.mkdir()
    except FileExistsError:
        raise PreparationError(f"批次資料夾已存在：{batch_dir}") from None

    valid = {t for t in tokens if is_tdcc(t)}
    folders, unreadable = find_product_folders(root, valid)
    outcomes = tuple(
        _prepare_product(t, folders.get(t, []), unreadable, batch_dir, convert)
        if t in valid
        else ProductOutcome(t, SKIPPED, "不是 12 碼數字")
        for t in tokens
    )
    list_path = batch_dir / LIST_NAME
    _write_list(list_path, outcomes)
    return PrepareReceipt(batch_dir, list_path, outcomes, tuple(unreadable))


def is_tdcc(text: str) -> bool:
    """剛好 12 個半形數字（全形數字不算）。"""
    return _TDCC.fullmatch(text) is not None


def parse_tdccs(text: str) -> list[str]:
    """換行、空白、逗號、頓號、分號都可以分隔；重複的只留第一次。"""
    return list(dict.fromkeys(t for t in _SEPARATORS.split(text) if t))


def find_product_folders(root: Path, wanted: set[str]) -> tuple[dict[str, list[Path]], list[Path]]:
    """根目錄底下任何深度、名稱剛好等於 TDCC 的資料夾（只走一次樹），以及讀不到的資料夾。"""
    found: dict[str, list[Path]] = {}
    unreadable: list[Path] = []
    if not wanted:
        return found, unreadable
    for dirpath, dirnames, _ in os.walk(root, onerror=lambda e: unreadable.append(Path(e.filename))):
        dirnames.sort()
        for name in dirnames:
            if name in wanted:
                found.setdefault(name, []).append(Path(dirpath) / name)
    return found, unreadable


def _prepare_product(
    tdcc: str, folders: list[Path], unreadable: list[Path], batch_dir: Path, convert: Converter
) -> ProductOutcome:
    if not folders:
        reason = "找不到商品資料夾"
        if unreadable:
            reason += f"（搜尋時有 {len(unreadable)} 個資料夾讀不到，例如 {unreadable[0]}）"
        return ProductOutcome(tdcc, SKIPPED, reason)
    if len(folders) > 1:
        paths = "、".join(str(f) for f in folders)
        return ProductOutcome(tdcc, SKIPPED, f"找到 {len(folders)} 個同名商品資料夾：{paths}")
    folder = folders[0]
    try:
        return _prepare_folder(tdcc, folder, batch_dir, convert)
    except OSError as exc:
        for kind in DocKind:  # 兩份要嘛都輸出、要嘛都不輸出
            (batch_dir / _output_name(tdcc, kind)).unlink(missing_ok=True)
        return ProductOutcome(tdcc, SKIPPED, f"讀取或複製失敗：{exc}", folder)


def _prepare_folder(tdcc: str, folder: Path, batch_dir: Path, convert: Converter) -> ProductOutcome:
    files = sorted((p for p in folder.iterdir() if p.is_file()), key=lambda p: p.name)
    words = [p for p in files if p.suffix.lower() in _WORD and not p.name.startswith("~$")]  # ~$ 是 Word 鎖定檔
    pdfs = [p for p in files if p.suffix.lower() == _PDF]
    pending = [p for p in pdfs if p.stem.endswith(PENDING_SUFFIX)]  # 上次核對沒通過、放回的檔：只當標記
    sources = [p for p in pdfs if p not in pending]
    versions = _Versions(pending, sources, {_output_name(tdcc, kind) for kind in DocKind})
    candidates: dict[DocKind, list[_Candidate]] = {kind: [] for kind in DocKind}
    mismatched: dict[DocKind, list[_Candidate]] = {kind: [] for kind in DocKind}

    def read(candidate: _Candidate, kinds: tuple[DocKind, ...]) -> None:
        reading = read_pdf(candidate.pdf, tdcc)
        if reading.kind in kinds:
            (candidates if reading.code_matches else mismatched)[reading.kind].append(candidate)

    for pdf in sources:
        read(_Candidate(pdf, pdf), tuple(DocKind))
    selections = {kind: versions.select(candidates[kind]) for kind in DocKind}

    # Word 逾時被強制結束時暫存檔可能還被鎖住，清不掉也不影響結果
    with tempfile.TemporaryDirectory(prefix="spw-prepare-", ignore_cleanup_errors=True) as tmp:
        # 只有本來就沒有 PDF 的那種文件才轉 Word（例：MS、BNP）；有 PDF 但都是沒通過的舊版時，不改用同一份的 Word 版
        need_word = tuple(kind for kind in DocKind if not candidates[kind])
        failures: dict[Path, str] = {}
        if need_word and words:
            try:
                converted, failures = _convert_words(tdcc, words, Path(tmp), convert)
            except ConversionError as exc:
                return _skipped(tdcc, [str(exc)], folder, selections)
            for candidate in converted:
                read(candidate, need_word)
            for kind in need_word:
                selections[kind] = versions.select(candidates[kind])

        problems = [_problem(kind, selections[kind], mismatched[kind], failures) for kind in DocKind]
        problems = [p for p in problems if p]
        if problems:
            return _skipped(tdcc, problems, folder, selections)
        # 沿用沒通過的版本時，資料夾裡其實可能有新檔、只是不能用：一起寫出來，不能只說「上手還沒寄新版」
        unusable = [
            note
            for kind in DocKind
            if selections[kind].reused_failed
            for note in _unusable_notes(kind, mismatched[kind], failures)
        ]

        docs = []
        for kind in DocKind:
            [picked] = selections[kind].chosen
            target = batch_dir / _output_name(tdcc, kind)
            shutil.copyfile(picked.pdf, target)
            docs.append(PreparedDocument(kind, picked.source, target.name, sha256(target), picked.from_word))
    filled = [c for kind in DocKind for c in selections[kind].chosen if c.filled]
    fill_note = f"已補受託機構商品代號（{_names(filled)}）" if filled else ""
    reason = "；".join(p for p in [fill_note, _version_note(selections), *unusable] if p)
    return ProductOutcome(tdcc, PREPARED, reason, folder, tuple(docs))


def _unusable_notes(kind: DocKind, mismatched: list[_Candidate], failures: dict[Path, str]) -> list[str]:
    """沿用舊版時，同一種文件不能用的新檔：封面商品代號不符、Word 轉檔失敗。

    有 PDF 時不轉 Word，所以不知道哪個 Word 是說明書；資料夾裡一定有公司表單的 Word，不逐一提醒（見 README）。
    """
    notes = []
    if mismatched:
        notes.append(f"{kind.label}另有 {_names(mismatched)}，封面商品代號不符")
    if failures:
        notes.append(f"{kind.label}另有 Word 轉檔失敗：{'；'.join(f'{w.name}：{r}' for w, r in failures.items())}")
    return notes


def _skipped(tdcc: str, problems: list[str], folder: Path, selections: dict[DocKind, _Selection]) -> ProductOutcome:
    """跳過的商品：原因之外也寫出排除了哪些舊版；有唯一候選的那種文件記下來源檔。"""
    reason = "；".join(p for p in [*problems, _version_note(selections)] if p)
    return ProductOutcome(tdcc, SKIPPED, reason, folder, _found_documents({k: s.chosen for k, s in selections.items()}))


def _version_note(selections: dict[DocKind, _Selection]) -> str:
    """寫進原因欄：沿用了哪些上次沒通過的版本、排除了哪些舊版。"""
    reused = [c for kind in DocKind if selections[kind].reused_failed for c in selections[kind].chosen]
    excluded = [c for kind in DocKind for c in selections[kind].failed + selections[kind].superseded]
    notes = []
    if reused:
        notes.append(f"沿用上次沒通過的版本（{_names(reused)}），上手還沒寄新版")
    if excluded:
        notes.append(f"已排除舊版：{_names(excluded)}")
    return "；".join(notes)


class _Versions:
    """判斷候選是哪個版本：上次沒通過的（和待確認檔內容相同）、上次已放回的（標準名複本），或新寄來的。

    內容指紋（全文）每個檔只算一次。只有一份候選時，標準名檔照用（例：UBS 上手本來就給標準檔名）；
    有多份內容不同時，上手檔名的版本優先，只以標準檔名存在的內容當成工具上次放回的舊版。
    """

    def __init__(self, pending: list[Path], sources: list[Path], standard_names: set[str]) -> None:
        self._pending = pending
        self._sources = sources
        self._standard_names = standard_names
        self._keys: dict[Path, str] = {}
        self._failed_keys: set[str] | None = None
        self._returned_keys: set[str] | None = None

    def key(self, path: Path) -> str:
        if path not in self._keys:
            self._keys[path] = content_key(path)
        return self._keys[path]

    @property
    def failed_keys(self) -> set[str]:
        if self._failed_keys is None:
            # 待確認檔的內容如果也有標準名檔（後來沿用同一版核對通過、放回了），那一版已經通過，不算沒通過
            passed = {self.key(p) for p in self._sources if p.name in self._standard_names}
            self._failed_keys = {self.key(p) for p in self._pending} - passed
        return self._failed_keys

    @property
    def returned_keys(self) -> set[str]:
        if self._returned_keys is None:
            standard = [p for p in self._sources if p.name in self._standard_names]
            originals = (
                {self.key(p) for p in self._sources if p.name not in self._standard_names} if standard else set()
            )
            self._returned_keys = {self.key(p) for p in standard} & originals
        return self._returned_keys

    def _as_pending_pdf(self, candidate: _Candidate) -> _Candidate:
        """沿用 Word 舊版時，改用資料夾裡那份待確認 PDF（內容相同）：同一份 Word 每次轉出的 PDF 不同，
        用待確認檔才能讓這次核對的、放回的、錯訊頁碼對應的都是同一份檔（#25）。"""
        if not candidate.from_word:
            return candidate
        key = self.key(candidate.pdf)
        for pending in self._pending:
            if self.key(pending) == key:
                return _Candidate(pending, candidate.source, from_word=True, filled=candidate.filled)
        return candidate

    def select(self, candidates: list[_Candidate]) -> _Selection:
        """內容相同只算一份（保留上手原檔名）；有新版時排除沒通過的版本，沒新版時沿用它；還有多份時排除已放回的舊版。"""
        if len(candidates) <= 1 and not self._pending:
            return _Selection(list(candidates), [], [])
        groups: dict[str, list[_Candidate]] = {}
        for candidate in candidates:
            groups.setdefault(self.key(candidate.pdf), []).append(candidate)
        representative = {
            key: min(group, key=lambda c: (c.source.name in self._standard_names, c.source.name))
            for key, group in groups.items()
        }
        failed = [representative[k] for k in groups if k in self.failed_keys]
        remaining = [k for k in groups if k not in self.failed_keys]
        if not remaining and failed:  # 上手還沒寄新版：沿用上次沒通過的版本，同一版可以重複核對
            return _Selection([self._as_pending_pdf(c) for c in failed], [], [], reused_failed=True)
        superseded: list[_Candidate] = []
        if len(remaining) > 1:
            # 上手檔名的版本優先：只以標準檔名存在的內容，或和上次放回的標準名檔相同的內容，都是舊版
            standard_only = {k for k in remaining if all(c.source.name in self._standard_names for c in groups[k])}
            newer = [k for k in remaining if k not in self.returned_keys and k not in standard_only]
            if newer:
                superseded = [representative[k] for k in remaining if k not in newer]
                remaining = newer
        return _Selection([representative[k] for k in remaining], failed, superseded)


def _convert_words(
    tdcc: str, words: list[Path], tmp: Path, convert: Converter
) -> tuple[list[_Candidate], dict[Path, str]]:
    """把 Word 複製到暫存區再轉 PDF（原始 Word 不動、Word 鎖定檔也不會留在商品資料夾）。

    MS 的暫存複本先補受託機構商品代號（只有說明書有這一格，其他 Word 不受影響）。

    回傳轉好的候選，以及轉不了的原始 Word 與原因。
    """
    jobs: dict[Path, tuple[Path, Path]] = {}
    filled: set[Path] = set()
    for i, word in enumerate(words):
        work = tmp / str(i)
        work.mkdir()
        copy = work / word.name
        shutil.copyfile(word, copy)
        if needs_fill(tdcc, word) and fill_trustee_code(copy, tdcc):
            filled.add(word)
        jobs[word] = (copy, work / "converted.pdf")
    reported = convert(list(jobs.values()))
    converted, failures = [], {}
    for word, (copy, pdf) in jobs.items():
        if copy in reported or not pdf.is_file():
            failures[word] = reported.get(copy, "沒有產生 PDF")
        else:
            converted.append(_Candidate(pdf, word, from_word=True, filled=word in filled))
    return converted, failures


def _found_documents(candidates: dict[DocKind, list[_Candidate]]) -> tuple[PreparedDocument, ...]:
    """跳過的商品：有唯一候選的那種文件記下來源檔，沒有輸出。"""
    return tuple(
        PreparedDocument(kind, c[0].source, from_word=c[0].from_word)
        if len(c := candidates[kind]) == 1
        else PreparedDocument(kind)
        for kind in DocKind
    )


def _output_name(tdcc: str, kind: DocKind) -> str:
    return f"{tdcc}_{kind.value}.pdf"


def _problem(kind: DocKind, selection: _Selection, mismatched: list[_Candidate], failures: dict[Path, str]) -> str:
    found = selection.chosen
    if len(found) == 1:
        return ""
    if found:
        return f"{kind.label}有 {len(found)} 份候選：{_names(found)}"
    if mismatched:
        return f"找不到{kind.label}（{_names(mismatched)} 封面商品代號不符）"
    if failures:
        details = "；".join(f"{word.name}：{reason}" for word, reason in failures.items())
        return f"找不到{kind.label}（有 Word 轉檔失敗：{details}）"
    return f"找不到{kind.label}"


def _names(candidates: list[_Candidate]) -> str:
    return "、".join(c.source.name for c in candidates)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _list_rows(outcomes: Iterable[ProductOutcome]) -> Iterable[tuple[str, ...]]:
    for o in outcomes:
        folder = str(o.folder) if o.folder else ""
        if not o.documents:
            yield (o.tdcc, "", "", folder, "", "", "", o.status, o.reason)
            continue
        for d in o.documents:
            source = str(d.source) if d.source else ""
            word = ("是" if d.from_word else "否") if d.source else ""
            yield (o.tdcc, d.kind.value, source, folder, word, d.output_name, d.sha256, o.status, o.reason)


def _write_list(path: Path, outcomes: Iterable[ProductOutcome]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "準備清單"
    ws.append([name for name, _ in LIST_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in _list_rows(outcomes):
        ws.append([v or None for v in row])
        ws.cell(ws.max_row, 1).number_format = "@"  # TDCC 是文字，保留前導零
    for i, (_, width) in enumerate(LIST_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
