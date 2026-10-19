"""放回資料夾：把批次資料夾的 `<TDCC>_TS.pdf`、`<TDCC>_IIS.pdf` 複製回各自的商品資料夾（#5、#9、#17）。

對外只有 `return_to_folders`。依據是核對工具存下的核對結果檔：
- 「回填後」工作表 `TS/IIS完成檢核` 為 `V` 的商品（說明書與投資人須知都通過或人工放行）用標準檔名放回；
- 「錯誤清單」上的 PDF（沒通過、也沒人工放行）檔名加 `_待確認` 放回，同商品通過的那份仍用標準檔名；
- 核對結果檔裡沒出現的商品（還沒核對）不放回。
放回標準檔名時，商品資料夾若還有同種的舊 `_待確認` 檔，只提醒、不刪。
目的地取準備清單記的商品資料夾；放回前重算指紋，和準備清單不符就不放。
商品資料夾已有同名檔：內容相同算已在，內容不同不覆蓋、整個商品不放回。一個商品的兩份要嘛都處理、要嘛都不動。
只複製；批次資料夾與準備清單不改。核對結果檔格式以 fcn-term-sheet-checker 現行版本為準（觀察基準 `34837cd`）。
"""

from __future__ import annotations

import datetime as dt
import os
import shutil
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sp_workflow.documents import DocKind
from sp_workflow.preparation import LIST_NAME, PREPARED, sha256

RETURNED = "已放回"
ALREADY_THERE = "已在"
NOT_RETURNED = "不放回"
PENDING_SUFFIX = "_待確認"  # 核對沒通過的檔，放回時加在檔名最後
PENDING_RETURNED = f"{RETURNED}（待確認）"
PENDING_ALREADY_THERE = f"{ALREADY_THERE}（待確認）"

_PREP_COLUMNS = ("TDCC", "種類", "商品資料夾", "輸出檔名", "SHA-256", "狀態")  # 讀準備清單用到的欄

RESULT_SHEET = "回填後"
RESULT_HEADER_ROW = 2
RESULT_TDCC_COLUMN = "TDCC Code"
RESULT_CHECKED_COLUMN = "TS/IIS完成檢核"
RESULT_CHECKED_VALUE = "V"
ERROR_SHEET = "錯誤清單"
ERROR_HEADER_ROW = 1
ERROR_PDF_COLUMN = "PDF 檔名"
ERROR_MESSAGE_COLUMN = "錯訊"
TDCC_DIGITS = 12

NOT_CHECKED_REASON = "核對結果檔沒有這個商品（還沒核對）"
_TEMP_SUFFIX = ".spw-tmp"  # 放回時先寫成暫存名，確認指紋後才改成正式檔名

# 放回清單的欄（欄名, 欄寬）
LIST_COLUMNS = (("TDCC", 15), ("TS", 10), ("IIS", 10), ("原因", 80))


class ReturnError(Exception):
    """整批無法放回（核對結果檔或準備清單有問題）；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class ReturnOutcome:
    tdcc: str
    ts: str
    iis: str
    reason: str = ""


@dataclass(frozen=True)
class ReturnReceipt:
    list_path: Path
    outcomes: tuple[ReturnOutcome, ...]


@dataclass(frozen=True)
class _Prepared:
    """準備清單上一個商品的一種文件。"""

    kind: DocKind
    folder: Path
    output_name: str
    sha256: str


def return_to_folders(batch_dir: Path, result_file: Path, *, now: dt.datetime | None = None) -> ReturnReceipt:
    """放回核對過的商品；回傳放回清單路徑與逐商品結果。

    依序列「回填後」的商品、錯誤清單的商品，再列這批已準備、但核對結果檔沒有的商品（還沒核對，不放回）。
    """
    prepared, skipped = _read_preparation_list(batch_dir / LIST_NAME)
    passed = _read_passed_tdccs(result_file)
    failed = {t: msgs for t, msgs in _read_failed_pdfs(result_file).items() if t not in passed}
    list_path = batch_dir / f"放回清單_{(now or dt.datetime.now()):%Y%m%d-%H%M%S}.xlsx"
    if list_path.exists():
        raise ReturnError(f"放回清單已存在：{list_path}")

    outcomes = []
    for tdcc in [*passed, *failed]:
        if tdcc in prepared:
            outcomes.append(_return_product(tdcc, prepared[tdcc], batch_dir, failed.get(tdcc, {})))
        else:
            reason = "準備時已跳過這個商品" if tdcc in skipped else "這批準備清單沒有這個商品"
            outcomes.append(ReturnOutcome(tdcc, NOT_RETURNED, NOT_RETURNED, reason))
    for tdcc in prepared:
        if tdcc not in passed and tdcc not in failed:
            outcomes.append(ReturnOutcome(tdcc, NOT_RETURNED, NOT_RETURNED, NOT_CHECKED_REASON))
    try:
        _write_list(list_path, outcomes)
    except OSError as exc:
        returned = sum(any(s.startswith(RETURNED) for s in (o.ts, o.iis)) for o in outcomes)
        raise ReturnError(f"已放回 {returned} 個商品，但放回清單寫不進去：{list_path}（{exc}）") from None
    return ReturnReceipt(list_path, tuple(outcomes))


def _read_preparation_list(path: Path) -> tuple[dict[str, dict[DocKind, _Prepared]], set[str]]:
    """準備清單上已準備的商品（兩種文件都在），以及準備時跳過的 TDCC。"""
    if not path.is_file():
        raise ReturnError(f"批次資料夾裡找不到準備清單：{path}")
    rows = _read_rows(path, None, "準備清單")
    if not rows:
        raise ReturnError(f"準備清單是空的：{path}")
    header = [str(h) if h is not None else "" for h in rows[0]]
    missing = [c for c in _PREP_COLUMNS if c not in header]
    if missing:
        raise ReturnError(f"準備清單少了欄：{'、'.join(missing)}（{path}）")
    prepared: dict[str, dict[DocKind, _Prepared]] = {}
    skipped: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        record = dict(zip(header, row, strict=False))
        tdcc = str(record["TDCC"] or "")
        if record["狀態"] != PREPARED:
            skipped.add(tdcc)
            continue
        try:
            kind = DocKind(record["種類"])
        except ValueError:
            raise ReturnError(f"準備清單第 {row_number} 列的種類看不懂：{record['種類']}（{path}）") from None
        if not all(record[c] for c in ("商品資料夾", "輸出檔名", "SHA-256")):
            raise ReturnError(f"準備清單第 {row_number} 列缺少商品資料夾、輸出檔名或 SHA-256（{path}）")
        prepared.setdefault(tdcc, {})[kind] = _Prepared(
            kind, Path(str(record["商品資料夾"])), str(record["輸出檔名"]), str(record["SHA-256"])
        )
    complete = {tdcc: docs for tdcc, docs in prepared.items() if set(docs) == set(DocKind)}
    return complete, skipped


def _read_passed_tdccs(path: Path) -> list[str]:
    """核對結果檔「回填後」上 TS/IIS完成檢核 為 V 的 TDCC（去重、保留順序）。"""
    rows = _read_rows(path, RESULT_SHEET, "核對結果檔")[RESULT_HEADER_ROW - 1 :]
    header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else ())]
    missing = [c for c in (RESULT_TDCC_COLUMN, RESULT_CHECKED_COLUMN) if c not in header]
    if missing:
        raise ReturnError(f"核對結果檔「{RESULT_SHEET}」第 {RESULT_HEADER_ROW} 列找不到欄：{'、'.join(missing)}")
    tdcc_col, checked_col = header.index(RESULT_TDCC_COLUMN), header.index(RESULT_CHECKED_COLUMN)
    passed = (
        _normalize_tdcc(row[tdcc_col])
        for row in rows[1:]
        if len(row) > max(tdcc_col, checked_col) and str(row[checked_col] or "").strip() == RESULT_CHECKED_VALUE
    )
    return list(dict.fromkeys(t for t in passed if t))


def _read_failed_pdfs(path: Path) -> dict[str, dict[str, str]]:
    """核對結果檔「錯誤清單」上的 PDF：{TDCC: {PDF 檔名: 錯訊}}（依清單順序）；PDF 檔名不是標準名的略過。"""
    rows = _read_rows(path, ERROR_SHEET, "核對結果檔")[ERROR_HEADER_ROW - 1 :]
    header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else ())]
    missing = [c for c in (ERROR_PDF_COLUMN, ERROR_MESSAGE_COLUMN) if c not in header]
    if missing:
        raise ReturnError(f"核對結果檔「{ERROR_SHEET}」第 {ERROR_HEADER_ROW} 列找不到欄：{'、'.join(missing)}")
    pdf_col, message_col = header.index(ERROR_PDF_COLUMN), header.index(ERROR_MESSAGE_COLUMN)
    failed: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        if len(row) <= max(pdf_col, message_col):
            continue
        name = str(row[pdf_col] or "").strip()
        tdcc, _, kind = name.removesuffix(".pdf").partition("_")
        if len(tdcc) == TDCC_DIGITS and tdcc.isdigit() and kind in {k.value for k in DocKind}:
            failed.setdefault(tdcc, {})[name] = str(row[message_col] or "").strip()
    return failed


def _pending_name(output_name: str) -> str:
    """`<TDCC>_TS.pdf` → `<TDCC>_TS_待確認.pdf`。"""
    path = Path(output_name)
    return f"{path.stem}{PENDING_SUFFIX}{path.suffix}"


def _read_rows(path: Path, sheet: str | None, label: str) -> list[tuple]:
    """讀一張工作表的所有列（sheet 為 None 時讀第一張），讀完就關檔，不鎖住檔案。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 對壞檔丟出的例外型別不固定
        raise ReturnError(f"讀不到{label}：{path}（{exc}）") from None
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise ReturnError(f"{label}裡找不到「{sheet}」工作表：{path}")
        return list((wb[sheet] if sheet else wb.worksheets[0]).iter_rows(values_only=True))
    finally:
        wb.close()


def _normalize_tdcc(value: object) -> str:
    """TDCC 以文字處理；Excel 存成數字、掉了前導零時補回 12 碼。"""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int | float):
        return str(int(value)).zfill(TDCC_DIGITS)
    return str(value).strip()


def _return_product(
    tdcc: str, docs: dict[DocKind, _Prepared], batch_dir: Path, failed: dict[str, str]
) -> ReturnOutcome:
    """failed 是這個商品在錯誤清單上的 PDF 與錯訊；它們以 `_待確認` 檔名放回，其餘用標準檔名。"""
    pending = {kind for kind, doc in docs.items() if doc.output_name in failed}
    names = {kind: _pending_name(doc.output_name) if kind in pending else doc.output_name for kind, doc in docs.items()}
    problems = []
    for doc in docs.values():
        batch_file = batch_dir / doc.output_name
        if not batch_file.is_file():
            problems.append(f"批次資料夾裡找不到 {doc.output_name}")
        elif sha256(batch_file) != doc.sha256:
            problems.append(f"批次資料夾裡的 {doc.output_name} 和準備清單的指紋不符")
    folder = docs[DocKind.TS].folder
    if not folder.is_dir():
        problems.append(f"商品資料夾不存在：{folder}")
    if problems:
        return ReturnOutcome(tdcc, NOT_RETURNED, NOT_RETURNED, "；".join(problems))

    copy_needed: dict[DocKind, bool] = {}
    for kind, doc in docs.items():
        target = folder / names[kind]
        if not target.exists():
            copy_needed[kind] = True
        elif target.is_file() and sha256(target) == doc.sha256:
            copy_needed[kind] = False
        else:
            problems.append(f"商品資料夾已有內容不同的 {names[kind]}，請確認要留哪一份")
    if problems:
        return ReturnOutcome(tdcc, NOT_RETURNED, NOT_RETURNED, "；".join(problems))

    to_copy = [kind for kind in docs if copy_needed[kind]]
    temps = {kind: folder / (names[kind] + _TEMP_SUFFIX) for kind in to_copy}
    placed: list[Path] = []
    try:
        # 先寫成暫存名並確認指紋，兩份都好了才改成正式檔名；複製到一半失敗也不會留下半個檔
        for kind in to_copy:
            shutil.copyfile(batch_dir / docs[kind].output_name, temps[kind])
            if sha256(temps[kind]) != docs[kind].sha256:
                raise OSError(f"{names[kind]} 複製後指紋不符")
        for kind in to_copy:
            target = folder / names[kind]
            os.rename(temps[kind], target)  # Windows 上目的地已存在時會失敗，不會覆蓋
            placed.append(target)
    except OSError as exc:
        for path in [*temps.values(), *placed]:  # 兩份要嘛都放、要嘛都不放
            path.unlink(missing_ok=True)
        return ReturnOutcome(tdcc, NOT_RETURNED, NOT_RETURNED, f"複製失敗：{exc}")

    def status(kind: DocKind) -> str:
        if kind in pending:
            return PENDING_RETURNED if copy_needed[kind] else PENDING_ALREADY_THERE
        return RETURNED if copy_needed[kind] else ALREADY_THERE

    notes = [f"核對沒通過（{kind.value}）：{failed[docs[kind].output_name]}" for kind in DocKind if kind in pending]
    notes += [
        f"資料夾仍有舊的 {_pending_name(docs[kind].output_name)}，請確認後刪除"
        for kind in DocKind
        if kind not in pending and (folder / _pending_name(docs[kind].output_name)).exists()
    ]
    return ReturnOutcome(tdcc, status(DocKind.TS), status(DocKind.IIS), "；".join(notes))


def _write_list(path: Path, outcomes: Iterable[ReturnOutcome]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "放回清單"
    ws.append([name for name, _ in LIST_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for o in outcomes:
        ws.append([o.tdcc, o.ts, o.iis, o.reason or None])
        ws.cell(ws.max_row, 1).number_format = "@"  # TDCC 是文字，保留前導零
    for i, (_, width) in enumerate(LIST_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
