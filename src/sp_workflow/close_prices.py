"""收盤價（#153、ADR 0020）：以「實際收盤日＋代碼」存進 `workflow.close_price`，給比價用。

兩個來源：
- 每天的 MarketData CSV：「1 每日匯入」的回補股價（`daily_import.import_price_csv`）在同一個 transaction 裡
  用 `closes_from_csv` 挑出 CSV 那一天的收盤價、`write_closes` 寫入；BBG_Data 的日期不明確，不寫。
- 歷史收盤價（HP 格式）：「維護」選單一次匯入（`import_history`）。橫向日期、每列一檔、代碼帶類別字；
  整份讀完、檢查完才寫，有問題的格子不存，依代碼與原因整理成錯誤清單放在來源檔同一個資料夾。
代碼去掉彭博最後的類別字（`normalize_code`），其餘一字不差，和 `workflow.product.underlying_n` 一致。
只存該代碼市場的交易日：週末、假日表上的休市日不存（`is_closed`）；交易所在假日表、但日期超出涵蓋範圍的
不能確認是不是交易日（休市日常帶著前一天的價格），也不存、列出來，補了假日表再匯（`is_uncovered`）。
交易所不在假日表的代碼照存，由比價標待確認。同一（收盤日、代碼）再匯入整筆取代，值改了的舊值留在
`workflow.close_price_change`，每次匯入一筆 `workflow.close_price_import`。
"""

from __future__ import annotations

import datetime as dt
import getpass
import hashlib
import json
from collections import defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    import psycopg

    from fcn_checker.trading_calendar import TradingCalendar

CSV_SOURCE = "MarketData CSV"
HISTORY_SOURCE = "歷史收盤價"
CLASS_WORDS = frozenset(("EQUITY", "CURNCY", "INDEX", "COMDTY"))  # 彭博代碼最後的類別字

# 歷史收盤價（HP 格式）：第 3 列從 E 欄起是收盤日，第 4 列起每列一檔，C 欄是帶類別字的彭博代碼（B 欄是簡碼）
HISTORY_DATE_ROW = 3
HISTORY_FIRST_ROW = 4
HISTORY_SHORT_CODE_COLUMN = 2
HISTORY_CODE_COLUMN = 3
HISTORY_FIRST_DATE_COLUMN = 5
HISTORY_ERROR_BOOK = "歷史收盤價錯誤_{stamp}.xlsx"
ERROR_BOOK_COLUMNS = (("代碼", 22), ("原因", 50), ("第一天", 12), ("最後一天", 12), ("格數", 8))

HOLIDAY_MARK = "#N/A N/A"  # 彭博在休市日填的值
UNCOVERED_REASON = "日期超出交易所假日表的涵蓋範圍，不能確認是交易日，沒有存；補了假日表再匯"


class ClosePriceError(Exception):
    """整份不能匯入（格式不對、寫不出錯誤清單）；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class CloseSummary:
    """一次寫入收盤價的結果。"""

    import_id: int
    first: dt.date | None  # 存進去的收盤日範圍
    last: dt.date | None
    codes: int  # 存了幾個代碼
    stored: int  # 存了幾筆（含取代的）
    replaced: int  # 取代了資料庫原本的幾筆
    changed: int  # 其中值不一樣的幾筆（舊值記在 close_price_change）


@dataclass(frozen=True)
class CsvCloses:
    """MarketData CSV 那一天要存的收盤價，與沒存的原因。"""

    closes: dict[tuple[dt.date, str], Decimal]
    closed: tuple[str, ...]  # 那一天休市、沒存的代碼
    uncovered: tuple[str, ...]  # 那一天超出假日表涵蓋範圍、沒存的代碼
    conflicts: tuple[str, ...]  # 去掉類別字後才重複、值不同、沒存的代碼（原本就股價衝突的已另外列出）


@dataclass(frozen=True)
class HistoryProblem:
    """錯誤清單的一列：一個代碼、一種原因，從哪天到哪天共幾格；整列不匯入的（C 欄沒有代碼）沒有日期。"""

    code: str  # 代碼；整列不匯入時是「第 N 列」
    reason: str
    first: dt.date | None
    last: dt.date | None
    cells: int


@dataclass(frozen=True)
class HistoryReceipt:
    summary: CloseSummary
    first_date: dt.date  # 檔案的日期範圍
    last_date: dt.date
    codes: int  # 檔案裡的代碼數
    missing: int  # 沒有價格的格數（休市日、週末不算）
    closed: int  # 休市日有值、沒有存的格數
    uncovered: int  # 超出假日表涵蓋範圍、有值但沒有存的格數
    problems: tuple[HistoryProblem, ...]
    error_book_path: Path | None  # 沒有問題時不寫


def normalize_code(raw: str) -> str:
    """去掉彭博代碼最後的類別字（例：`AA UN Equity` → `AA UN`），其餘一字不差；只去頭尾空白。"""
    code = raw.strip()
    head, _, last = code.rpartition(" ")
    return head.rstrip() if head.strip() and last.upper() in CLASS_WORDS else code


def is_closed(calendar: TradingCalendar, code: str, day: dt.date) -> bool:
    """這個代碼的市場在這一天確定休市：週末，或交易所在假日表、日期在涵蓋範圍內且是休市日。"""
    if day.weekday() >= 5:
        return True
    market = calendar.market_of(code)
    return market is not None and market.covers(day, day) and day in market.holidays


def is_uncovered(calendar: TradingCalendar, code: str, day: dt.date) -> bool:
    """平日、交易所在假日表，但日期超出涵蓋範圍：不能確認是不是交易日。"""
    market = calendar.market_of(code)
    return day.weekday() < 5 and market is not None and not market.covers(day, day)


def is_confirmed_trading_day(calendar: TradingCalendar, code: str, day: dt.date) -> bool:
    """確定是交易日：平日、交易所在假日表、日期在涵蓋範圍內且不是休市日。"""
    market = calendar.market_of(code)
    return day.weekday() < 5 and market is not None and market.covers(day, day) and day not in market.holidays


# ---------- MarketData CSV ----------


def closes_from_csv(
    close_date: dt.date,
    closes: Mapping[str, Decimal | None],
    calendar: TradingCalendar,
    conflicted: Iterable[str] = (),
) -> CsvCloses:
    """CSV 那一天要存的收盤價：`closes` 是原始代碼 → 收盤價，`conflicted` 是股價衝突的原始代碼。

    去掉類別字後同一個代碼有不同的值、或原本就股價衝突的不存；休市、超出假日表涵蓋範圍的不存。"""
    found: dict[str, set[Decimal]] = defaultdict(set)
    for raw, close in closes.items():
        if close is not None and close > 0:
            found[normalize_code(raw)].add(close)
    bad = {normalize_code(raw) for raw in conflicted}
    conflicts = tuple(sorted(code for code, values in found.items() if len(values) > 1 and code not in bad))
    kept, closed, uncovered = {}, [], []
    for code, values in sorted(found.items()):
        if len(values) > 1 or code in bad:
            continue
        if is_closed(calendar, code, close_date):
            closed.append(code)
        elif is_uncovered(calendar, code, close_date):
            uncovered.append(code)
        else:
            kept[(close_date, code)] = next(iter(values))
    return CsvCloses(kept, tuple(closed), tuple(uncovered), conflicts)


# ---------- 寫入 ----------


def write_closes(
    conn: psycopg.Connection,
    closes: Mapping[tuple[dt.date, str], Decimal],
    *,
    source: str,
    file_name: str,
    sha256: str,
    details: Mapping[str, Any],
    run_id: int | None = None,
    user: str | None = None,
) -> CloseSummary:
    """在呼叫端的 transaction 裡寫入一次收盤價：記一筆匯入、值改了的留舊值、整筆取代。"""
    days = [day for day, _ in closes]
    first, last = (min(days), max(days)) if days else (None, None)
    codes = len({code for _, code in closes})
    (import_id,) = conn.execute(
        "INSERT INTO workflow.close_price_import (source, file_name, file_sha256, run_id, imported_by, summary)"
        " VALUES (%s, %s, %s, %s, %s, '{}') RETURNING import_id",
        (source, file_name, sha256, run_id, user or getpass.getuser()),
    ).fetchone()
    conn.execute(
        "CREATE TEMP TABLE incoming_close (close_date date, code text, close numeric,"
        " PRIMARY KEY (close_date, code)) ON COMMIT DROP"
    )
    with conn.cursor() as cur, cur.copy("COPY incoming_close (close_date, code, close) FROM STDIN") as copy:
        for (day, code), close in closes.items():
            copy.write_row((day, code, close))
    (replaced,) = conn.execute(
        "SELECT count(*) FROM incoming_close i JOIN workflow.close_price c USING (close_date, code)"
    ).fetchone()
    changed = conn.execute(
        "INSERT INTO workflow.close_price_change (close_date, code, previous_close, previous_import_id, import_id)"
        " SELECT c.close_date, c.code, c.close, c.import_id, %s"
        " FROM incoming_close i JOIN workflow.close_price c USING (close_date, code) WHERE c.close <> i.close",
        (import_id,),
    ).rowcount
    conn.execute(
        "INSERT INTO workflow.close_price (close_date, code, close, import_id)"
        " SELECT close_date, code, close, %s FROM incoming_close"
        " ON CONFLICT (close_date, code) DO UPDATE SET close = EXCLUDED.close, import_id = EXCLUDED.import_id",
        (import_id,),
    )
    conn.execute("DROP TABLE incoming_close")
    summary = CloseSummary(import_id, first, last, codes, len(closes), replaced, changed)
    record = {
        "first": first.isoformat() if first else None,
        "last": last.isoformat() if last else None,
        "codes": codes,
        "stored": summary.stored,
        "replaced": replaced,
        "changed": changed,
        **details,
    }
    conn.execute(
        "UPDATE workflow.close_price_import SET summary = %s::jsonb WHERE import_id = %s",
        (json.dumps(record, ensure_ascii=False), import_id),
    )
    return summary


# ---------- 歷史收盤價（HP 格式） ----------


@dataclass
class _History:
    dates: list[dt.date]
    closes: dict[tuple[dt.date, str], Decimal]
    codes: set[str]
    missing: int = 0
    closed: int = 0
    uncovered: int = 0
    cells: dict[tuple[str, str], list[dt.date]] = field(
        default_factory=lambda: defaultdict(list)
    )  # (代碼, 原因) → 哪幾天
    rows: list[HistoryProblem] = field(default_factory=list)  # 整列不匯入的列


def import_history(
    path: Path,
    conn: psycopg.Connection,
    calendar: TradingCalendar,
    *,
    user: str | None = None,
    now: dt.datetime | None = None,
) -> HistoryReceipt:
    """匯入一份歷史收盤價：整份讀完、檢查完，先寫錯誤清單，再在一個 transaction 裡寫入。

    錯誤清單寫不出來就整份不匯入；寫入資料庫失敗時刪掉這次的錯誤清單，不留下沒有匯入的清單。"""
    history = _read_history(path, calendar)
    problems = (
        *history.rows,
        *(
            HistoryProblem(code, reason, min(days), max(days), len(days))
            for (code, reason), days in sorted(history.cells.items())
        ),
    )
    error_book = None
    if problems:
        error_book = path.parent / HISTORY_ERROR_BOOK.format(stamp=f"{(now or dt.datetime.now()):%Y%m%d-%H%M%S}")
        _write_error_book(error_book, problems)
    details = {
        "file_first": history.dates[0].isoformat(),
        "file_last": history.dates[-1].isoformat(),
        "file_codes": len(history.codes),
        "missing": history.missing,
        "closed": history.closed,
        "uncovered": history.uncovered,
        "problems": len(problems),
    }
    try:
        with conn.transaction():
            summary = write_closes(
                conn,
                history.closes,
                source=HISTORY_SOURCE,
                file_name=path.name,
                sha256=_sha256(path),
                details=details,
                user=user,
            )
    except BaseException:
        if error_book is not None:
            error_book.unlink(missing_ok=True)
        raise
    return HistoryReceipt(
        summary,
        history.dates[0],
        history.dates[-1],
        len(history.codes),
        history.missing,
        history.closed,
        history.uncovered,
        problems,
        error_book,
    )


def _read_history(path: Path, calendar: TradingCalendar) -> _History:
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 對壞檔丟的例外種類很多
        raise ClosePriceError(f"打不開歷史收盤價：{path}（{exc}）") from exc
    try:
        rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()
    columns = _date_columns(rows[HISTORY_DATE_ROW - 1] if len(rows) >= HISTORY_DATE_ROW else ())
    history = _History(sorted(columns.values()), {}, set())
    values: dict[tuple[dt.date, str], set[Decimal]] = defaultdict(set)
    no_price: dict[tuple[dt.date, str], str] = {}
    for number, row in enumerate(rows[HISTORY_FIRST_ROW - 1 :], HISTORY_FIRST_ROW):
        raw = _cell(row, HISTORY_CODE_COLUMN)
        code = normalize_code(raw) if isinstance(raw, str) else ""
        if not any(ch.isalnum() for ch in code):  # 空白、「-」這類不是代碼
            short = _cell(row, HISTORY_SHORT_CODE_COLUMN)
            filled = sum(_cell(row, c) not in (None, "") for c in columns)
            if short not in (None, "") or filled:
                reason = f"C 欄沒有彭博代碼，這列不匯入（B 欄：{short if short not in (None, '') else '空白'}）"
                history.rows.append(HistoryProblem(f"第 {number} 列", reason, None, None, filled))
            continue
        history.codes.add(code)
        for column, day in columns.items():
            price, reason = _price(_cell(row, column))
            if price is not None:
                values[(day, code)].add(price)
            else:
                no_price.setdefault((day, code), reason)
    for (day, code), found in values.items():
        if is_closed(calendar, code, day):
            history.closed += 1
        elif is_uncovered(calendar, code, day):
            history.uncovered += 1
            history.cells[(code, UNCOVERED_REASON)].append(day)
        elif len(found) > 1:
            history.missing += 1
            history.cells[(code, "代碼重複、同一天的值不同，沒有存")].append(day)
        else:
            history.closes[(day, code)] = next(iter(found))
    for (day, code), reason in no_price.items():
        if (day, code) in values or is_closed(calendar, code, day):
            continue
        history.missing += 1
        if reason != HOLIDAY_MARK or is_confirmed_trading_day(calendar, code, day):
            history.cells[(code, _reason_text(reason))].append(day)
    if not history.codes:
        raise ClosePriceError(f"歷史收盤價第 {HISTORY_FIRST_ROW} 列起沒有任何代碼（C 欄），整份不匯入：{path}")
    return history


def _date_columns(header: tuple[Any, ...]) -> dict[int, dt.date]:
    """第 3 列從 E 欄起的收盤日（欄號 → 日期）；不是日期、重複或一個都沒有就整份不匯入。"""
    columns: dict[int, dt.date] = {}
    for column in range(HISTORY_FIRST_DATE_COLUMN, len(header) + 1):
        value = header[column - 1]
        if value in (None, ""):
            continue
        where = f"{get_column_letter(column)}{HISTORY_DATE_ROW}"
        if isinstance(value, dt.datetime):
            value = value.date()
        if not isinstance(value, dt.date):
            raise ClosePriceError(f"歷史收盤價的 {where} 應為收盤日，實際是「{value}」，整份不匯入")
        if value in columns.values():
            raise ClosePriceError(f"歷史收盤價的收盤日 {value} 出現兩次（{where}），整份不匯入")
        columns[column] = value
    if not columns:
        raise ClosePriceError(
            f"歷史收盤價第 {HISTORY_DATE_ROW} 列從 {get_column_letter(HISTORY_FIRST_DATE_COLUMN)} 欄起沒有收盤日，整份不匯入"
        )
    return columns


def _cell(row: tuple[Any, ...], column: int) -> Any:
    return row[column - 1] if len(row) >= column else None


def _price(value: Any) -> tuple[Decimal | None, str]:
    """一格的收盤價；不是正數時回傳 (None, 原因)。錯誤值（#N/A…）、`-`、亂碼、失效公式一律沒有價格。"""
    if isinstance(value, bool):
        return None, str(value)
    if isinstance(value, int | float | Decimal):
        number = Decimal(str(value))
        return (number, "") if number > 0 else (None, f"不是正數（{value}）")
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "空白"
    text = str(value).strip()
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None, text
    return (number, "") if number.is_finite() and number > 0 else (None, text)


def _reason_text(value: str) -> str:
    if value.startswith(HOLIDAY_MARK):
        return f"交易日沒有價格（{HOLIDAY_MARK}）"
    if value.startswith("#N/A Requesting Data"):
        return "抓價未完成（#N/A Requesting Data...），請重抓再匯"
    if value.startswith("#N/A Invalid Security"):
        return "代碼無效（#N/A Invalid Security）"
    if value == "空白" or value.startswith("不是正數"):
        return f"沒有價格（{value}）"
    if value == "-":
        return "沒有價格（-）"
    if value.startswith("#"):
        return f"錯誤值或失效公式（{value[:30]}）"
    return "亂碼或不是數字"


def _write_error_book(path: Path, problems: tuple[HistoryProblem, ...]) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "歷史收盤價錯誤"
    sheet.append([name for name, _ in ERROR_BOOK_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for p in problems:
        sheet.append([p.code, p.reason, p.first, p.last, p.cells])
        sheet.cell(sheet.max_row, 3).number_format = sheet.cell(sheet.max_row, 4).number_format = "yyyy-mm-dd"
    for i, (_, width) in enumerate(ERROR_BOOK_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    try:
        book.save(path)
    except OSError as exc:
        raise ClosePriceError(f"錯誤清單寫不進去，整份不匯入：{path}（{exc.strerror or exc}）") from exc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()
