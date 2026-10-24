"""各步驟共用的小工具（#139）：檔案與位元組的 SHA-256、寫一張簡單的 Excel 表。

簡單的表：第 1 列粗體表頭、凍結第一列，欄寬與格子格式由呼叫端決定；存檔時先寫暫存檔再換名，
寫不進去時丟 `OSError`、暫存檔刪掉、原本的同名檔不動。核對結果檔與回報包的 Excel 有自己的版面，不用這裡。
給作業人員看的「檔案寫不進去」是 `FileWriteError`（#207，例：`import_run.error_book`）。
`save_rows`：值照資料庫型別寫成一張有篩選的表，商品匯出、進度與出場清單匯出共用（#204 從 `product_export` 搬來）。
"""

from __future__ import annotations

import datetime as dt
import hashlib
import os
from collections.abc import Callable, Iterable, Sequence
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sp_workflow.database import OperatorError

_MAX_WIDTH = 40  # `save_rows` 的欄寬上限


class FileWriteError(OperatorError):
    """檔案寫不進去（#207）；訊息是給作業人員看的中文，`path` 是那個檔，`error` 是原本的 `OSError`。"""

    def __init__(self, message: str, path: Path, error: OSError) -> None:
        super().__init__(message)
        self.path = path
        self.error = error


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    with path.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def tdcc_as_text(cell: Cell) -> None:
    """第 1 欄（TDCC）存成文字格式，保留前導零。"""
    if cell.column == 1:
        cell.number_format = "@"


def table_book(
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    *,
    widths: Sequence[float] = (),
    style: Callable[[Cell], None] | None = None,
    auto_filter: bool = False,
) -> openpyxl.Workbook:
    """一張工作表（`title`）的活頁簿：粗體表頭、每列寫完交給 `style` 設格式、依序設欄寬、凍結第一列。"""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(list(row))
        if style is not None:
            for cell in sheet[sheet.max_row]:
                style(cell)
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A2"
    if auto_filter:
        sheet.auto_filter.ref = sheet.dimensions
    return book


def save_book(book: openpyxl.Workbook, path: Path, *, overwrite: bool = True) -> None:
    """先寫 `<檔名>.tmp` 再換名。`overwrite=False` 時同名檔已存在就失敗（Windows 的 `os.rename`），不覆蓋。"""
    temp = path.with_name(path.name + ".tmp")
    try:
        book.save(temp)
        if overwrite:
            temp.replace(path)
        else:
            os.rename(temp, path)
    finally:
        temp.unlink(missing_ok=True)  # 換名成功時已經不在


def _excel(value: object) -> object:
    """資料庫的值 → 寫進 Excel 的值。"""
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        return value.astimezone().replace(tzinfo=None)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def save_rows(path: Path, title: str, columns: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    """一張工作表（`title`）：第 1 列粗體欄名、凍結窗格、自動篩選、欄寬依內容；值照型別寫入（有時區的時間換成本機
    時間、Excel 不收的控制字元拿掉、文字不當公式）。先寫暫存檔再換名；寫不進去時丟 `FileWriteError`。
    商品匯出、進度與出場清單匯出共用（#156、#204）。
    """
    values = [[_excel(value) for value in row] for row in rows]
    widths = [len(name) for name in columns]
    for row in values:
        for i, value in enumerate(row):
            widths[i] = max(widths[i], len(str(value or "")))
    book = table_book(
        title,
        columns,
        values,
        widths=[min(width + 2, _MAX_WIDTH) for width in widths],
        style=_text_cell,
        auto_filter=True,
    )
    try:
        save_book(book, path)
    except OSError as exc:
        raise FileWriteError(
            f"寫不進 {path}：{exc.strerror or exc}。同名檔開在 Excel 時請先關掉，或換一個檔名再匯出。", path, exc
        ) from exc


def _text_cell(cell: Cell) -> None:
    if isinstance(cell.value, str):
        cell.data_type = "s"  # 「=」開頭的文字照文字寫，不當公式
        cell.number_format = "@"  # TDCC 等文字保留前導零
