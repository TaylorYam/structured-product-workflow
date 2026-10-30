"""標的曝險分析（#345、#347、#352、#356）：選定的標的大跌時，近期到期、連結它們的在庫商品目前有哪些會接股、面額多少。

只呈現目前的狀態（檔數、面額、距觸發點多遠），不做估損、配息、淨損益或情境模擬等任何預估金額（#352）。

- 選標的照股票代號（彭博代碼的第一段，例 `GOOG`）：一個代號包含它所有交易所的代碼（`GOOG UW`、`GOOG UQ`），
  作業人員只打代號就好，分析與依分公司、理專的加總都算進全部代碼（`Choice.codes`）。
- 範圍：存續狀態在庫（沒有還有效的存續事件）、連結任一選定標的、最終比價日不晚於截止日的商品，商品種類不限
  （FCN、FCN_Express、Super Bull）。最終比價日早於分析日、還在庫的（例：等股價）是「已比價」，
  結果以最終比價日為準，不用最新收盤價判斷，另列一類。
- 價格：每個標的用收盤價表最新的一筆，依拆股登記換回價格基準（同比價，`stock_splits.SplitBook.factor`）。
  總覽最上面寫「以 X/X 收盤價判斷」。
- 試算價（#356）：盤後大跌等不及收盤價時，作業人員可替選定標的填價格（`trial`，代碼 → 價格），
  取代最新收盤價、當作分析日的價格（拆股照分析日換算）；其他標的照舊用最新收盤價。試算價不存進資料庫，
  總覽的判斷基準、明細的「價格依據」都標出來，檔名加「_試算」。
- 核對條件不是完成（同比價只比完成的）、缺 KI 型態、有 KI 卻沒有下限價、缺價格的商品無法計算；MKI 不比、人工確認。
- 其他看最差標的（收盤價 ÷ 進場價最低）：收盤價 < 執行價，且沒有 KI、收盤價 < 下限價或比價已記 AKI 觸及 → 會接股；
  < 執行價但還在下限價之上（EKI 或還沒觸及的 AKI）→ KI 未破下限（觀察）；還在執行價之上、距執行價 10% 以內
  → 接近執行價；其他是高於執行價。
- 緩衝（#352）：最差標的收盤價 ÷ 觸發價 − 1。還有下限價要看的（EKI、還沒觸及的 AKI、MKI）觸發價是下限價，
  沒有 KI（含 Super Bull）或已觸及 AKI 的是執行價；分組已跌破、<3%、3–10%、>10%，已比價的另成一組。
- 選定標的／其他標的：最差標的是選定的代碼之一時算「選定標的造成」，否則算「其他標的造成」（#347），
  總覽與依分公司、依理專分開列，避免把籃子裡別的股票造成的狀況算到選定標的頭上。
- 金額照交易明細逐筆（分公司、理專、金額），沒有交易明細的用庫存的分公司配置（理專寫 `-`）。
  USD 以收盤價表最新的匯率換算；查不到匯率的 USD 欄空白、不算進 USD 合計。
- 排序：依分公司、依理專照會接股面額、再看 KI 未破下限面額；通知明細照剩餘天數、再看緩衝小的在前。
- 全部商品（#354）：一檔一列，固定欄位給作業人員看；狀態簡化成跌破KI、跌破K、未跌破（已比價、無法計算照分類），
  表現用 100 為基數（100 + 表現 %，例 −26.1% → 73.9），各標的由差到好拆成標的N、表現N。
- 圖表：總覽右邊放各分類面額（選定標的／其他標的）、緩衝分組、會接股面額依最差標的三張圖與它們的資料表；
  依分公司、依理專右邊放會接股面額與面額組成的橫條圖（排前面的 `CHART_ROWS` 名）。圖是 Excel 原生圖表。
- 只讀資料庫，不改任何資料。

`assess` 只判斷一檔、不碰資料庫；`load_choices`、`load_positions` 讀資料庫；`export_exposure` 寫 Excel。
傳入的 `conn` 由 `database.connect` 建立，資料表版本已確認（#200），這裡不再檢查。
"""

from __future__ import annotations

import datetime as dt
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sp_workflow.check_terms import NO_BARRIER
from sp_workflow.file_tools import excel_value, save_export, table_book, text_cell
from sp_workflow.lifecycle import IN_STOCK, lifecycle_status
from sp_workflow.status import CheckTerms
from sp_workflow.stock_splits import SplitBook, load_splits

if TYPE_CHECKING:
    import psycopg

DELIVER = "會接股"
WATCH = "KI 未破下限"
NEAR = "接近執行價"
MANUAL = "MKI 人工確認"
SAFE = "高於執行價"
UNKNOWN = "無法計算"
FIXED = "已比價"  # 最終比價日早於分析日、還在庫（#352）
TRIAL = "試算價"  # 作業人員填的價格，取代最新收盤價（#356）
CATEGORIES = (DELIVER, WATCH, NEAR, MANUAL, SAFE, FIXED, UNKNOWN)
NOTIFY = (DELIVER, WATCH, NEAR, MANUAL)  # 通知明細列的分類
NEAR_BUFFER = Decimal("0.10")  # 接近執行價的門檻（距執行價，#347），也是緩衝分組 3–10% 的上限（#352）
TIGHT_BUFFER = Decimal("0.03")  # 緩衝分組 <3% 的上限
CROSSED = "已跌破"  # 最差標的已低於觸發價（下限價或執行價）
BUFFER_GROUPS = (CROSSED, "<3%", "3–10%", ">10%", FIXED)
KI_BASIS = "下限價"
STRIKE_BASIS = "執行價"
MKI = "MKI"
AKI_HIT = "AKI hit"  # 比價寫的 KI 欄：`AKI hit 日期 標的`（observation._ki_text）
NO_FC = "-"
MAX_LEGS = 5  # 一檔最多連結幾個標的
_UNDERLYING_COLUMNS = ", ".join(f"underlying_{n}" for n in range(1, MAX_LEGS + 1))  # workflow.product 的連結標的
KI_BROKEN = "跌破KI"  # 全部商品的狀態（#354）
K_BROKEN = "跌破K"
INTACT = "未跌破"


@dataclass(frozen=True)
class Choice:
    """視窗上可選的一個股票代號：它在各交易所的代碼、連結任一代碼的在庫檔數，與連結最多那個代碼的最新收盤價。"""

    ticker: str
    codes: tuple[str, ...]  # 依連結檔數多到少
    products: int
    close_date: dt.date | None
    close: Decimal | None


def ticker_of(code: str) -> str:
    """股票代號：彭博代碼的第一段（`GOOG UW` → `GOOG`）。"""
    return code.split()[0] if code.split() else code


@dataclass(frozen=True)
class Leg:
    """一檔商品的一個標的：條件價格與換回價格基準的最新收盤價。"""

    seq: int
    code: str
    initial: Decimal | None
    strike: Decimal | None
    ki: Decimal | None
    close: Decimal | None  # 換回價格基準的收盤價或試算價
    close_date: dt.date | None
    trial: bool = False  # `close` 是作業人員填的試算價（#356）

    @property
    def price_source(self) -> str | None:
        """價格依據：`試算價` 或 `10/29 收盤`；缺價格時 None。"""
        if self.close is None or self.close_date is None:
            return None
        return TRIAL if self.trial else f"{self.close_date.month}/{self.close_date.day} 收盤"

    @property
    def performance(self) -> Decimal | None:
        """收盤價 ÷ 進場價 − 1；缺價格時 None。"""
        if self.close is None or not self.initial:
            return None
        return self.close / self.initial - 1

    @property
    def to_strike(self) -> Decimal | None:
        """距執行價：收盤價 ÷ 執行價 − 1（負的是已低於執行價）；缺價格時 None。"""
        return None if self.close is None or not self.strike else self.close / self.strike - 1

    @property
    def to_ki(self) -> Decimal | None:
        """距下限價：收盤價 ÷ 下限價 − 1；沒有下限價或缺價格時 None。"""
        return None if self.close is None or not self.ki else self.close / self.ki - 1


@dataclass(frozen=True)
class Holding:
    """一筆交易明細（沒有交易明細時是一筆分公司配置）。"""

    branch: str
    fc: str
    amount: Decimal


@dataclass(frozen=True)
class Position:
    """一檔在庫商品的分析資料。"""

    tdcc: str
    product: str
    issuer: str
    currency: str
    final_date: dt.date | None
    barrier_type: str | None
    ki_text: str  # 比價狀態的 KI 欄
    terms_status: str | None  # 核對條件狀態；還沒算過是 None
    terms_reason: str | None
    legs: tuple[Leg, ...]
    holdings: tuple[Holding, ...]
    usd_rate: Decimal | None  # 1 單位幣別換幾 USD
    strike_pct: Decimal | None = None  # 執行價 (%)
    ki_pct: Decimal | None = None  # 下限價 (%)

    @property
    def watches_ki(self) -> bool:
        """還有下限價要看：EKI、還沒觸及的 AKI、MKI；沒有 KI 或已觸及 AKI 的只看執行價。"""
        return self.barrier_type not in (None, NO_BARRIER) and not self.ki_text.startswith(AKI_HIT)


@dataclass(frozen=True)
class Assessed:
    """一檔的分類與最差標的。"""

    position: Position
    category: str
    worst: Leg | None
    reason: str = ""

    @property
    def notional(self) -> Decimal:
        return sum((h.amount for h in self.position.holdings), Decimal(0))

    @property
    def status(self) -> str:
        """全部商品的狀態（#354）：最差標的的最新收盤價跌破下限價（或比價已記 AKI 觸及）、跌破執行價或都沒有；
        已比價與算不出的照分類。"""
        if self.category in (FIXED, UNKNOWN):
            return self.category
        w, p = self.worst, self.position
        if w is None or w.close is None or w.to_strike is None:
            return UNKNOWN
        if p.ki_text.startswith(AKI_HIT) or (p.barrier_type != NO_BARRIER and w.ki is not None and w.close < w.ki):
            return KI_BROKEN
        return K_BROKEN if w.to_strike < 0 else INTACT

    @property
    def buffer_basis(self) -> str | None:
        """緩衝看哪個價格（#352）：還有下限價要看的是下限價，其他是執行價；已比價或算不出最差標的時 None。"""
        if self.worst is None or self.category == FIXED:  # 已比價的不用最新收盤價判斷
            return None
        return KI_BASIS if self.position.watches_ki and self.worst.ki is not None else STRIKE_BASIS

    @property
    def buffer(self) -> Decimal | None:
        """最差標的距 `buffer_basis`：收盤價 ÷ 下限價或執行價 − 1，負的是已跌破；沒有 `buffer_basis` 時 None。"""
        if self.worst is None or self.buffer_basis is None:
            return None
        return self.worst.to_ki if self.buffer_basis == KI_BASIS else self.worst.to_strike

    @property
    def buffer_group(self) -> str | None:
        """依 `buffer` 分組；已比價的不分距離；算不出時 None。"""
        if self.category == FIXED:
            return FIXED
        gap = self.buffer
        if gap is None:
            return None
        if gap < 0:
            return CROSSED
        return BUFFER_GROUPS[1] if gap < TIGHT_BUFFER else BUFFER_GROUPS[2] if gap <= NEAR_BUFFER else BUFFER_GROUPS[3]


def assess(position: Position, today: dt.date | None = None) -> Assessed:
    """依最差標的判斷一檔的分類（見 module 說明）；給 `today` 時，最終比價日早於它的是已比價。"""

    def unknown(reason: str) -> Assessed:
        return Assessed(position, UNKNOWN, None, reason)

    priced = all(leg.performance is not None and leg.strike is not None for leg in position.legs)
    worst = min(position.legs, key=lambda leg: (leg.performance, leg.seq)) if position.legs and priced else None
    final = position.final_date
    if today is not None and final is not None and final < today:
        return Assessed(position, FIXED, worst, f"{final} 已比價，結果以最終比價日為準，不用最新收盤價判斷")
    if position.terms_status != CheckTerms.DONE:
        status = position.terms_status or "還沒算"
        return unknown(f"核對條件{status}" + (f"：{position.terms_reason}" if position.terms_reason else ""))
    if position.barrier_type == MKI:
        return Assessed(position, MANUAL, worst, "MKI 每期觀察，系統不比，請人工確認")
    if position.barrier_type is None:
        return unknown("沒有 KI 型態")
    if worst is None:
        missing = [leg.code for leg in position.legs if leg.performance is None or leg.strike is None]
        return unknown(f"缺進場價、執行價或收盤價：{'、'.join(missing)}" if missing else "沒有標的價格")
    if position.barrier_type != NO_BARRIER and worst.ki is None:
        return unknown(f"{position.barrier_type} 沒有下限價：{worst.code}")
    to_strike = worst.to_strike
    assert to_strike is not None and worst.close is not None
    if to_strike >= 0:
        near = to_strike <= NEAR_BUFFER
        return Assessed(position, NEAR if near else SAFE, worst, "距執行價 10% 以內" if near else "")
    knocked = (
        position.barrier_type == NO_BARRIER
        or (worst.ki is not None and worst.close < worst.ki)
        or position.ki_text.startswith(AKI_HIT)
    )
    if not knocked:
        return Assessed(position, WATCH, worst, "低於執行價，但還在下限價之上")
    return Assessed(position, DELIVER, worst)


# ---------------------------------------------------------------- 讀資料庫


def load_choices(conn: psycopg.Connection) -> list[Choice]:
    """在庫商品連結的標的，照股票代號合併各交易所的代碼，依代號排序。"""
    rows = conn.execute(
        f"SELECT DISTINCT tdcc, code FROM workflow.product, unnest(ARRAY[{_UNDERLYING_COLUMNS}]) AS code"
        " WHERE code IS NOT NULL"
    ).fetchall()
    statuses = lifecycle_status(conn, [tdcc for tdcc, _ in rows])
    by_code: dict[str, set[str]] = defaultdict(set)
    for tdcc, code in rows:
        if statuses[tdcc].status == IN_STOCK:
            by_code[code].add(tdcc)
    latest = _latest_closes(conn, by_code)
    by_ticker: dict[str, list[str]] = defaultdict(list)
    for code in by_code:
        by_ticker[ticker_of(code)].append(code)
    choices = []
    for ticker in sorted(by_ticker):
        codes = tuple(sorted(by_ticker[ticker], key=lambda c: (-len(by_code[c]), c)))
        products = len(set().union(*(by_code[c] for c in codes)))
        main = next((c for c in codes if c in latest), None)
        choices.append(Choice(ticker, codes, products, *(latest[main] if main else (None, None))))
    return choices


def load_positions(
    conn: psycopg.Connection,
    codes: Iterable[str],
    until: dt.date,
    *,
    trial: Mapping[str, Decimal] | None = None,
    today: dt.date | None = None,
) -> list[Position]:
    """在庫、連結任一 `codes`、最終比價日不晚於 `until` 的商品，依最終比價日、TDCC 排序。

    `trial` 裡的代碼用試算價、當作 `today`（預設今天）的價格，其他用最新收盤價（#356）。
    """
    trial = trial or {}
    today = today or dt.date.today()
    wanted = sorted(set(codes))
    rows = conn.execute(
        "SELECT p.tdcc, p.product, p.issuer, p.currency, p.final_valuation_date, p.barrier_type, p.strike_date,"
        " s.ki, c.status, c.reason, p.strike_pct, p.ki_pct FROM workflow.product p"
        " LEFT JOIN workflow.pricing_status s ON s.tdcc = p.tdcc LEFT JOIN workflow.check_terms c ON c.tdcc = p.tdcc"
        f" WHERE p.final_valuation_date <= %s AND ARRAY[{_UNDERLYING_COLUMNS}] && %s::text[]"
        " ORDER BY p.final_valuation_date, p.tdcc",
        (until, wanted),
    ).fetchall()
    statuses = lifecycle_status(conn, [row[0] for row in rows])
    rows = [row for row in rows if statuses[row[0]].status == IN_STOCK]
    tdccs = [row[0] for row in rows]
    prices: dict[str, list[tuple]] = defaultdict(list)
    for tdcc, *values in conn.execute(
        "SELECT tdcc, seq, code, initial_price, strike_price, ki_price, price_basis_date"
        " FROM workflow.underlying_price WHERE tdcc = ANY(%s) ORDER BY tdcc, seq",
        (tdccs,),
    ):
        prices[tdcc].append(tuple(values))
    holdings = _holdings(conn, tdccs)
    latest = _latest_closes(conn, {code for values in prices.values() for _, code, *_ in values})
    currencies = {row[3] for row in rows}
    fx = _latest_closes(conn, _fx_codes(currencies))
    rates = {ccy: _usd_rate(ccy, fx) for ccy in currencies}
    book = SplitBook(load_splits(conn))
    positions = []
    for row in rows:
        tdcc, product, issuer, currency, final, barrier, strike_date, ki_text, terms_status, terms_reason = row[:10]
        legs = []
        for seq, code, initial, strike, ki, basis in prices.get(tdcc, ()):
            day, close = (today, trial[code]) if code in trial else latest.get(code, (None, None))
            base = basis or strike_date
            if close is not None and day is not None and base is not None:
                close = close * book.factor(code, base, day)
            legs.append(Leg(seq, code, initial, strike, ki, close, day, code in trial))
        positions.append(
            Position(
                tdcc,
                product,
                issuer,
                currency,
                final,
                barrier,
                ki_text or "",
                terms_status,
                terms_reason,
                tuple(legs),
                tuple(holdings.get(tdcc, ())),
                rates[currency],
                *row[10:],
            )
        )
    return positions


def _holdings(conn: psycopg.Connection, tdccs: Sequence[str]) -> dict[str, list[Holding]]:
    """交易明細逐筆；沒有交易明細的商品用庫存的分公司配置。"""
    found: dict[str, list[Holding]] = defaultdict(list)
    for tdcc, branch, fc, amount in conn.execute(
        "SELECT tdcc, branch, fc_name, amount FROM workflow.trade_row WHERE tdcc = ANY(%s) ORDER BY tdcc, row_seq",
        (list(tdccs),),
    ):
        found[tdcc].append(Holding(branch, fc, amount))
    traded = set(found)
    for tdcc, branch, amount in conn.execute(
        "SELECT tdcc, branch, amount FROM workflow.branch_allocation WHERE tdcc = ANY(%s) ORDER BY tdcc, seq",
        (list(tdccs),),
    ):
        if tdcc not in traded:
            found[tdcc].append(Holding(branch, NO_FC, amount))
    return found


def _latest_closes(conn: psycopg.Connection, codes: Iterable[str]) -> dict[str, tuple[dt.date, Decimal]]:
    rows = conn.execute(
        "SELECT DISTINCT ON (code) code, close_date, close FROM workflow.close_price"
        " WHERE code = ANY(%s) ORDER BY code, close_date DESC",
        (sorted(set(codes)),),
    ).fetchall()
    return {code: (day, close) for code, day, close in rows}


def _fx_codes(currencies: Iterable[str]) -> set[str]:
    return {code for ccy in currencies if ccy != "USD" for code in (f"{ccy}USD", f"USD{ccy}")}


def _usd_rate(currency: str, closes: Mapping[str, tuple[dt.date, Decimal]]) -> Decimal | None:
    """1 單位幣別換幾 USD：收盤價表的 `XXXUSD` 直接用，只有 `USDXXX` 時取倒數；都沒有時 None。"""
    if currency == "USD":
        return Decimal(1)
    if f"{currency}USD" in closes:
        return closes[f"{currency}USD"][1]
    if f"USD{currency}" in closes:
        return 1 / closes[f"USD{currency}"][1]
    return None


# ---------------------------------------------------------------- 匯出


@dataclass(frozen=True)
class ExposureReceipt:
    path: Path
    count: int  # 列入的商品檔數
    counts: Mapping[str, int]  # 各分類檔數


def export_name(codes: Sequence[str], now: dt.datetime | None = None, *, trial: bool = False) -> str:
    """存檔視窗預設的檔名；用了試算價時加「_試算」（#356）。"""
    tickers = "_".join(dict.fromkeys(ticker_of(code) for code in codes))[:40]
    return f"標的曝險分析_{tickers}{'_試算' if trial else ''}_{now or dt.datetime.now():%Y%m%d_%H%M}.xlsx"


def export_exposure(
    conn: psycopg.Connection,
    codes: Sequence[str],
    until: dt.date,
    path: Path,
    *,
    today: dt.date | None = None,
    trial: Mapping[str, Decimal] | None = None,
) -> ExposureReceipt:
    """分析並寫成一份 Excel：總覽、依分公司、依理專、通知明細、全部商品。

    `trial` 是選定代碼的試算價（#356）；不在 `codes` 裡的代碼不理會，只有選定標的能用試算價。
    """
    today = today or dt.date.today()
    trial = {code: price for code, price in (trial or {}).items() if code in codes}
    with conn.transaction():  # 每個 SELECT 看同一個時間點
        conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ")
        positions = load_positions(conn, codes, until, trial=trial, today=today)
        chosen = _latest_closes(conn, codes)
    results = [assess(p, today) for p in positions]
    report = _Report(results, frozenset(codes), today)
    counts = {category: sum(r.category == category for r in results) for category in CATEGORIES}
    book = table_book(OVERVIEW, ("項目", "值"), _overview(codes, chosen, until, report, trial))
    _style(book.active, widths=(30, 70))
    _overview_charts(book.active, report)
    for title, headers, rows, widths in (
        (BY_BRANCH, BRANCH_HEADERS, _rollup(report, lambda h: h.branch, lambda h: h.fc), ROLLUP_WIDTHS),
        (BY_FC, FC_HEADERS, _rollup(report, lambda h: h.fc, lambda h: h.branch), ROLLUP_WIDTHS),
        (DETAIL, DETAIL_HEADERS, _detail(report), DETAIL_WIDTHS),
        (ALL, ALL_HEADERS, [_all_row(report, r) for r in results], ALL_WIDTHS),
    ):
        sheet = book.create_sheet(title)
        sheet.append(list(headers))
        for row in rows:
            sheet.append([excel_value(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
        _style(sheet, widths=widths, index=ALL_INDEX if title == ALL else ())
        if title in (BY_BRANCH, BY_FC):
            _rollup_charts(sheet, title, rows)
    save_export(book, path)
    return ExposureReceipt(path, len(results), counts)


@dataclass(frozen=True)
class _Report:
    """匯出用：分析結果、選定的代碼與分析日。"""

    results: Sequence[Assessed]
    chosen: frozenset[str]
    today: dt.date

    def caused_by_chosen(self, r: Assessed) -> bool:
        return r.worst is not None and r.worst.code in self.chosen

    def days_left(self, r: Assessed) -> int | None:
        final = r.position.final_date
        return None if final is None else (final - self.today).days


OVERVIEW = "總覽"
BY_BRANCH = "依分公司"
BY_FC = "依理專"
DETAIL = "通知明細"
ALL = "全部商品"
CHOSEN = "選定標的造成"
OTHERS = "其他標的造成"
_ROLLUP = (
    "會接股檔數",
    "會接股面額(USD)",
    f"會接股面額：{CHOSEN}(USD)",
    f"會接股面額：{OTHERS}(USD)",
    "KI 未破下限檔數",
    "KI 未破下限面額(USD)",
    "接近執行價檔數",
    "接近執行價面額(USD)",
    "MKI 檔數",
    "連結檔數",
    "連結面額(USD)",
)
BRANCH_HEADERS = ("分公司", *_ROLLUP, "理專")
FC_HEADERS = ("理專", *_ROLLUP, "服務的分公司")
ROLLUP_WIDTHS = (12, 12, 18, 28, 28, 16, 20, 14, 20, 10, 10, 16, 40)
_POSITION = ("最終比價日", "剩餘天數", "分類", "緩衝分組", "緩衝基準", "緩衝")
_WORST = (
    "KI",
    "最差標的",
    "最差是選定標的",
    "最差表現",
    "執行價(%)",
    "下限價(%)",
    "距執行價",
    "距下限價",
    "最差標的價格",
    "價格依據",
    "執行價",
    "下限價",
)
_POSITION_WIDTHS = (12, 9, 12, 10, 10, 9)
_WORST_WIDTHS = (24, 12, 14, 10, 10, 10, 10, 10, 14, 12, 10, 10)
DETAIL_HEADERS = (*_POSITION, "分公司", "理專", "TDCC", "商品", "發行機構", "幣別", *_WORST, "金額", "說明")
DETAIL_WIDTHS = (*_POSITION_WIDTHS, 10, 12, 15, 12, 10, 6, *_WORST_WIDTHS, 14, 30)
_LEG_HEADERS = tuple(f"{name}{n}" for n in range(1, MAX_LEGS + 1) for name in ("標的", "表現"))
ALL_HEADERS = (
    "狀態",
    "TDCC",
    "剩餘天數",
    "最終比價日",
    "商品",
    "最差表現標的",
    "最差表現",
    "K%",
    "KI%",
    "KI種類",
    "面額",
    "發行機構",
    "幣別",
    *_LEG_HEADERS,
)
ALL_WIDTHS = (10, 15, 9, 12, 12, 14, 10, 7, 7, 8, 14, 10, 6, *(11, 8) * MAX_LEGS)
ALL_INDEX = ("最差表現", *_LEG_HEADERS[1::2])  # 全部商品的表現用 100 為基數，不是百分比
SPLIT = (DELIVER, WATCH, NEAR)  # 總覽與圖分開選定標的、其他標的的分類


@dataclass
class _Sum:
    """一群商品的檔數與面額 USD 合計；查不到匯率的商品算檔數、不算面額。"""

    tdccs: set[str] = field(default_factory=set)
    notional: Decimal = Decimal(0)
    missing_rate: bool = False

    def add(self, r: Assessed, amount: Decimal) -> None:
        self.tdccs.add(r.position.tdcc)
        if r.position.usd_rate is None:
            self.missing_rate = True
        else:
            self.notional += amount * r.position.usd_rate

    def text(self) -> str:
        notes = "；部分幣別查不到匯率，未計入" if self.missing_rate else ""
        return f"{len(self.tdccs)} 檔，面額 USD {self.notional:,.0f}{notes}"


def _overview(
    codes: Sequence[str],
    chosen: Mapping[str, tuple[dt.date, Decimal]],
    until: dt.date,
    report: _Report,
    trial: Mapping[str, Decimal],
) -> list[tuple[object, ...]]:
    results = report.results
    rows: list[tuple[object, ...]] = [
        ("判斷基準", _close_basis(codes, chosen, trial)),
        ("分析日", report.today),
        ("最終比價日截止", until),
    ]
    for code in codes:
        day, close = chosen.get(code, (None, None))
        latest = f"最新收盤 {close}（{day}）" if close is not None else "收盤價表沒有這個代碼"
        rows.append((f"選定標的 {code}", f"{TRIAL} {trial[code]}；{latest}" if code in trial else latest))
    rows.append(("列入商品", f"{len(results)} 檔（在庫、連結選定標的、最終比價日不晚於截止日）"))
    for category in CATEGORIES:
        sel = [r for r in results if r.category == category]
        if category in SPLIT:  # 分開最差標的是不是選定標的（#347）
            for label, wanted in ((CHOSEN, True), (OTHERS, False)):
                total = _Sum()
                for r in sel:
                    if report.caused_by_chosen(r) == wanted:
                        total.add(r, r.notional)
                rows.append((f"{category}｜{label}", total.text()))
            continue
        total = _Sum()
        for r in sel:
            total.add(r, r.notional)
        rows.append((category, total.text()))
    rows.append(("說明", "只呈現目前的狀態（檔數、面額、距觸發點多遠），不做任何預估金額。"))
    rows.append(("", "會接股：最差標的收盤價 < 執行價，且沒有 KI、收盤價 < 下限價或已觸及 AKI。"))
    rows.append(("", "選定標的造成：最差標的是選定的代碼；其他標的造成：籃子裡別的股票最差。"))
    rows.append(("", "接近執行價：還在執行價之上、距執行價 10% 以內。"))
    rows.append(
        ("", "緩衝 = 最差標的收盤價 ÷ 觸發價 − 1：EKI、還沒觸及的 AKI、MKI 看下限價，沒有 KI 或已觸及 AKI 看執行價。")
    )
    rows.append(("", "已比價：最終比價日早於分析日、還在庫的商品，結果以最終比價日為準，不用最新收盤價判斷。"))
    rows.append(("", "USD 以收盤價表最新匯率換算；查不到匯率的幣別，USD 欄空白、不算進各表的 USD 合計。"))
    if trial:
        rows.append(
            ("", f"{TRIAL}：作業人員填的價格（盤後、盤前或盤中），不是收盤價、不存進資料庫；其他標的用最新收盤價。")
        )
    return rows


def _close_basis(
    codes: Sequence[str], chosen: Mapping[str, tuple[dt.date, Decimal]], trial: Mapping[str, Decimal]
) -> str:
    """「以 10/23 收盤價判斷」；選定標的最新收盤日不同時全部列出；有試算價的寫在最前面（#356）。"""
    days = sorted({chosen[code][0] for code in codes if code not in trial and code in chosen})
    text = f"以 {'、'.join(f'{day.month}/{day.day}' for day in days)} 收盤價判斷" if days else ""
    text += "（選定標的最新收盤日不同）" if len(days) > 1 else ""
    if trial:
        prices = "、".join(f"{code} {price}" for code, price in trial.items())
        return f"{prices} 以{TRIAL}判斷" + (f"；其他選定標的{text}" if text else "")
    return text or "收盤價表沒有選定標的的收盤價"


def _rollup(
    report: _Report, key: Callable[[Holding], str], other: Callable[[Holding], str]
) -> list[tuple[object, ...]]:
    """依 `key`（分公司或理專）加總、`other` 是最後一欄列出的另一邊；一檔分給同一個人好幾筆時只算一檔。

    依會接股面額排序，相同時看 KI 未破下限面額（#352）。
    """
    sums: dict[str, dict[str, _Sum]] = defaultdict(lambda: defaultdict(_Sum))
    others: dict[str, set[str]] = defaultdict(set)
    for r in report.results:
        cause = CHOSEN if report.caused_by_chosen(r) else OTHERS
        for h in r.position.holdings:
            group = sums[key(h)]
            others[key(h)].add(other(h))
            group["all"].add(r, h.amount)
            group[r.category].add(r, h.amount)
            if r.category == DELIVER:
                group[cause].add(r, h.amount)
    rows = []
    for name, group in sums.items():
        deliver, watch = group[DELIVER], group[WATCH]
        rows.append(
            (
                name,
                len(deliver.tdccs),
                deliver.notional,
                group[CHOSEN].notional,
                group[OTHERS].notional,
                len(watch.tdccs),
                watch.notional,
                len(group[NEAR].tdccs),
                group[NEAR].notional,
                len(group[MANUAL].tdccs),
                len(group["all"].tdccs),
                group["all"].notional,
                "、".join(sorted(others[name])),
            )
        )
    rows.sort(key=lambda row: (-row[2], -row[6], str(row[0])))  # 會接股面額、KI 未破下限面額
    return rows


def _position_cells(report: _Report, r: Assessed) -> tuple[object, ...]:
    """最終比價日到緩衝那幾欄（`_POSITION`）。"""
    return (r.position.final_date, report.days_left(r), r.category, r.buffer_group, r.buffer_basis, r.buffer)


def _worst_cells(report: _Report, r: Assessed) -> tuple[object, ...]:
    """KI 到下限價那幾欄（`_WORST`）。"""
    p, w = r.position, r.worst
    return (
        p.ki_text,
        w and w.code,
        None if w is None else ("是" if report.caused_by_chosen(r) else "否"),
        w and w.performance,
        _ratio(p.strike_pct),
        _ratio(p.ki_pct),
        w and w.to_strike,
        w and w.to_ki,
        w and w.close,
        w and w.price_source,
        w and w.strike,
        w and w.ki,
    )


def _detail(report: _Report) -> list[tuple[object, ...]]:
    """通知明細：依剩餘天數排序，天數相同時緩衝小的在前（#352）。"""

    def order(r: Assessed) -> tuple[object, ...]:
        days, gap = report.days_left(r), r.buffer
        return (days is None, days or 0, gap is None, gap or 0, r.position.tdcc)

    rows = []
    for r in sorted((r for r in report.results if r.category in NOTIFY), key=order):
        p = r.position
        for h in p.holdings:
            rows.append(
                (
                    *_position_cells(report, r),
                    h.branch,
                    h.fc,
                    p.tdcc,
                    p.product,
                    p.issuer,
                    p.currency,
                    *_worst_cells(report, r),
                    h.amount,
                    r.reason,
                )
            )
    return rows


def _all_row(report: _Report, r: Assessed) -> tuple[object, ...]:
    """全部商品的一列（#354）：各標的依表現由差到好，表現用 100 為基數。"""
    p, w = r.position, r.worst
    legs = sorted(p.legs, key=lambda leg: (leg.performance is None, leg.performance or 0, leg.seq))
    return (
        r.status,
        p.tdcc,
        report.days_left(r),
        p.final_date,
        p.product,
        w and w.code,
        w and _index(w.performance),
        p.strike_pct,
        p.ki_pct,
        p.barrier_type,
        r.notional,
        p.issuer,
        p.currency,
        *(value for leg in legs for value in (leg.code, _index(leg.performance))),
    )


def _index(performance: Decimal | None) -> Decimal | None:
    """表現換成 100 為基數：100 + 表現 %（−26.1% → 73.9）。"""
    return None if performance is None else 100 + performance * 100


def _ratio(pct: Decimal | None) -> Decimal | None:
    """資料庫存的百分點（80）→ 比例（0.8），Excel 用百分比格式顯示。"""
    return None if pct is None else pct / 100


MONEY_WORDS = ("面額", "金額")  # 表頭有這些字的欄是金額，顯示到整數、千分位
PERCENT_HEADERS = ("緩衝", "最差表現", "執行價(%)", "下限價(%)", "距執行價", "距下限價")
WIDEST = 70  # 欄寬上限（說明較長）


def _style(sheet, *, widths: Sequence[float], index: Sequence[str] = ()) -> None:
    """日期、百分比、金額的格式與欄寬；`index` 的欄是 100 為基數的表現，顯示一位小數。"""
    headers = {cell.column: str(cell.value) for cell in sheet[1]}
    money = {col for col, name in headers.items() if any(word in name for word in MONEY_WORDS)}
    indexed = {col for col, name in headers.items() if name in index}
    percent = {col for col, name in headers.items() if name in PERCENT_HEADERS} - indexed
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            text_cell(cell)
            if isinstance(cell.value, dt.date):
                cell.number_format = "yyyy/m/d"
            elif isinstance(cell.value, int | float | Decimal) and not isinstance(cell.value, bool):
                if cell.column in indexed:
                    cell.number_format = "0.0"
                elif cell.column in percent:
                    cell.number_format = "0.0%"
                elif cell.column in money:
                    cell.number_format = "#,##0"
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = min(width, WIDEST)


# ---------------------------------------------------------------- 圖表（#347）

CHART_ROWS = 15  # 圖只畫排前面的幾名，太多條看不清楚
_CHART_COLUMN = 4  # 總覽的圖表資料從 D 欄開始（A、B 是文字）
_CHART_ANCHOR = "H"  # 圖放在資料表右邊


def _overview_charts(sheet, report: _Report) -> None:
    """總覽右邊：各分類面額、緩衝分組、會接股面額依最差標的，三張圖與資料表。"""
    results = report.results
    by_category = []
    for category in SPLIT:
        split = {True: _Sum(), False: _Sum()}
        for r in results:
            if r.category == category:
                split[report.caused_by_chosen(r)].add(r, r.notional)
        by_category.append((category, split[True].notional, split[False].notional))
    by_buffer = []
    for group in BUFFER_GROUPS:
        total = _Sum()
        for r in results:
            if r.buffer_group == group:
                total.add(r, r.notional)
        by_buffer.append((group, len(total.tdccs), total.notional))
    worst: dict[str, _Sum] = defaultdict(_Sum)
    for r in results:
        if r.category == DELIVER and r.worst is not None:
            worst[r.worst.code].add(r, r.notional)
    ranked = sorted(((code, total.notional) for code, total in worst.items()), key=lambda item: (-item[1], item[0]))
    by_worst: list[tuple[str, Decimal | int]] = list(ranked[:CHART_ROWS]) or [("（沒有會接股）", 0)]
    if rest := ranked[CHART_ROWS:]:
        by_worst.append((f"其他 {len(rest)} 檔標的", sum((usd for _, usd in rest), Decimal(0))))

    row = 1
    row = _chart_table(
        sheet, row, ("分類", f"面額：{CHOSEN}(USD)", f"面額：{OTHERS}(USD)"), by_category,
        _bar("各分類面額（USD）", stacked=True), anchor=1,
    )  # fmt: skip
    row = _chart_table(
        sheet, row + 2, ("緩衝分組（距下限價或執行價）", "檔數", "面額(USD)"), by_buffer,
        _bar("緩衝分組（面額 USD）"), anchor=18, series=(2,),
    )  # fmt: skip
    _chart_table(
        sheet, row + 2, ("會接股的最差標的", "會接股面額(USD)"), by_worst,
        _bar("會接股面額：最差標的（USD）", horizontal=True), anchor=35,
    )  # fmt: skip
    for col, width in zip("DEFG", (28, 24, 24, 4), strict=True):
        sheet.column_dimensions[col].width = width


def _chart_table(
    sheet,
    top: int,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    chart: BarChart,
    *,
    anchor: int,
    series: Sequence[int] = (),
) -> int:
    """在 `top` 列、D 欄起寫一張小資料表，畫成 `chart` 放在 H 欄 `anchor` 列；回傳資料表最後一列。

    `series` 是要畫的欄（`headers` 的位置，0 是分類）；沒給就畫分類以外的每一欄。
    """
    first = _CHART_COLUMN
    for i, header in enumerate(headers):
        cell = sheet.cell(row=top, column=first + i, value=header)
        cell.font = Font(bold=True)
    for r, values in enumerate(rows, start=top + 1):
        for i, value in enumerate(values):
            cell = sheet.cell(row=r, column=first + i, value=excel_value(value))
            if isinstance(value, Decimal | int | float) and not isinstance(value, bool) and "(USD)" in headers[i]:
                cell.number_format = "#,##0"
    last = top + len(rows)
    for i in series or range(1, len(headers)):
        col = first + i
        chart.add_data(Reference(sheet, min_col=col, min_row=top, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=first, min_row=top + 1, max_row=last))
    sheet.add_chart(chart, f"{_CHART_ANCHOR}{anchor}")
    return last


_COMPOSITION = ("會接股面額(USD)", "KI 未破下限面額(USD)", "接近執行價面額(USD)")


def _rollup_charts(sheet, title: str, rows: Sequence[Sequence[object]]) -> None:
    """依分公司、依理專右邊：會接股面額（選定標的／其他標的）與面額組成，各畫排前面的 `CHART_ROWS` 名。

    `rows` 是 `_rollup` 的列（已依會接股面額排序）；會接股面額圖只畫有會接股的，沒有就不畫。
    """
    headers = [str(cell.value) for cell in sheet[1]]
    delivering = sum(1 for row in rows[:CHART_ROWS] if row[headers.index("會接股面額(USD)")])
    who = "分公司" if title == BY_BRANCH else "理專"
    shown = min(len(rows), CHART_ROWS)
    charts = (
        (
            f"會接股面額（USD，前 {delivering} 名{who}）",
            delivering,
            (f"會接股面額：{CHOSEN}(USD)", f"會接股面額：{OTHERS}(USD)"),
        ),
        (f"面額組成（USD，前 {shown} 名）", shown, _COMPOSITION),
    )
    anchor = get_column_letter(len(headers) + 2)
    top = 1
    for name, count, columns in charts:
        if not count:
            continue
        chart = _bar(name, stacked=True, horizontal=True)
        for column in columns:
            col = headers.index(column) + 1
            chart.add_data(Reference(sheet, min_col=col, min_row=1, max_row=1 + count), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=1 + count))
        chart.height = max(7.5, 0.6 * count + 3)
        sheet.add_chart(chart, f"{anchor}{top}")
        top += int(chart.height * 2) + 2  # 一列約 0.5 公分


def _bar(title: str, *, stacked: bool = False, horizontal: bool = False) -> BarChart:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = _title(title)
    chart.legend.position = "b"  # 圖例放下面，不擠繪圖區
    chart.width = 18
    chart.height = 7.5
    if stacked:
        chart.grouping = "stacked"
        chart.overlap = 100
    if horizontal:
        chart.y_axis.majorGridlines = None
        chart.x_axis.scaling.orientation = "maxMin"  # 第一名在最上面
        chart.y_axis.crosses = "max"  # 分類反過來後數值軸會跑到上面蓋到標題，放回下面（#350）
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.delete = False  # openpyxl 3.1 預設不寫，Excel 會把座標軸藏起來
    chart.y_axis.delete = False
    return chart


def _title(text: str) -> Title:
    """圖表標題：11 pt 粗體、放在繪圖區外面，不蓋到圖（#350）。"""
    font = CharacterProperties(sz=1100, b=True)
    paragraph = Paragraph(pPr=ParagraphProperties(defRPr=font), r=[RegularTextRun(rPr=font, t=text)])
    return Title(tx=Text(rich=RichText(p=[paragraph])), overlay=False)
