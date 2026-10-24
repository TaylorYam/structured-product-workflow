"""每日匯入：把每日更新（`交易明細`＋`BBG_Data`）匯入最小 FCN 資料庫（#32、#87、ADR 0003、0008）。

對外是 `import_daily` 與回補股價用的 `import_price_csv`（#112）；交易明細的欄與解讀規則（`TRADE_FIELDS`、
`TRADE_REQUIRED`、`trade_value`、`is_cancel_note`）也給人工維護修改交易明細用（#116），兩邊存進去的值一樣。規則：
- 需要 `交易明細`、`BBG_Data` 兩張工作表，少一張就整份不匯入；`清單更新` 已停用，檔案裡有也不讀（#87）。
  「只匯交易明細」（`prices=False`，回補時用）不讀 `BBG_Data`、不動股價，沒有這張工作表也可以。
  商品條件改由「匯入起單資料」建立（#86），每日匯入不再寫商品條件、也不做分公司對帳。
- `交易明細` 第 1 列是欄名，必須和交易清單月份工作表的欄名與順序完全相同；不同就整份不匯入。
  只收 Product 為 FCN 的列，其他略過並計數；略過的 TDCC 與 Product 記在 `workflow.other_product`。
  以 TDCC 為單位、第一次寫入為準：資料庫已有的交易不重複寫入，內容不同只列出差異、不覆蓋。
- `BBG_Data`：C1 是取價日（要是日期），第 2 列 B～H 是欄名（要完全相同），第 4 列起每列一檔。
  代碼只去頭尾空白；代碼空白的列略過；價格格子空白就存空白，不是數字也存空白並列出來。
  同一個代碼出現多列：值都一樣算一筆，值不一樣記成衝突（價格留空），算價格時（#88）不能用。
  同一個取價日重新匯入時，整份取代那一天的股價；每次匯入都留紀錄（`workflow.price_import`）。
- 回補股價 CSV（`MarketData_YYYYMMDD.csv`，#112）：第 1 列欄名要完全相同，「取價日」是收盤那一天、整份同一天；
  存成取價日＋1 天，等同隔天早上用 BBG_Data 取到的價格。沒有 1日前收盤價；幣別、市值(百萬美元) 不存。
  代碼、重複代碼、不是數字、整天取代的處理同 BBG_Data；也留一筆 `workflow.import_run`。
  同一個 transaction 也把 CSV 那一天的收盤價寫進 `workflow.close_price`（#153、ADR 0020，見 `close_prices`）；
  讀不到交易所假日表時不能判斷休市日，收盤價這次不存。BBG_Data 不寫收盤價。
- 信箱由資料庫直接維護（#35、ADR 0004）：匯入不修改 `mailbox`，只用它比對 FC 名字。
- 結果另外列出交易明細的 FCN TDCC（依第一次出現的順序、不重複，不是剛好 12 碼與整檔都取消的不列），
  給視窗帶到起單資料頁（#90）。
- 整次匯入在一個 transaction 內完成，任何錯誤都不留下部分寫入。
傳入的 `conn` 由 `database.connect` 建立，資料表版本已確認（#200），這裡不再檢查。
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import TYPE_CHECKING, Any, NamedTuple

import openpyxl
from openpyxl.utils import get_column_letter

from sp_workflow.close_prices import CSV_SOURCE, CloseSummary, closes_from_csv, write_closes
from sp_workflow.database import OperatorError
from sp_workflow.file_tools import save_book, sha256_file, table_book
from sp_workflow.mailbox import Mailbox, is_pb, load_mailbox
from sp_workflow.product_record import FCN

if TYPE_CHECKING:
    import psycopg

    from fcn_checker.trading_calendar import TradingCalendar

TRADE_SHEET = "交易明細"
PRICE_SHEET = "BBG_Data"
TDCC_DIGITS = 12

TRADE_HEADERS = (
    "編號", "交易日期", "Product", "TDCC", "幣別", "金額", "FC/PB", "分公司", "前台主管/投專",
    "UF", "期初價格", "備註", "募集", "TS/IIS 核對", "核對TS/IIS人員", "起單人員",
)  # fmt: skip
# BBG_Data（ADR 0008）：C1 是取價日；第 2 列從 B 欄起是欄名；第 4 列起每列一檔
PRICE_DATE_ROW, PRICE_DATE_COLUMN = 1, 3
PRICE_HEADER_ROW, PRICE_FIRST_ROW = 2, 4
PRICE_FIRST_COLUMN = 2  # B 欄是代碼
PRICE_HEADERS = (
    "代碼(值)", "公司名稱", "最新收盤價(值)", "最新開盤價(值)", "1日前收盤價(值)", "200日波動率(值)",
    "市值(原幣/百萬)(值)",
)  # fmt: skip
# 回補股價 CSV（#112）：取價日是收盤那一天，存成下一天（同 BBG_Data 的取價日：內容是前一個交易日的價格）
PRICE_CSV_HEADERS = (
    "取價日", "代碼", "公司名稱", "收盤價", "開盤價", "200日波動率(%)", "市值(原幣/百萬)", "幣別", "市值(百萬美元)",
)  # fmt: skip
PRICE_CSV_DATE_SHIFT = dt.timedelta(days=1)


class Price(NamedTuple):
    """BBG_Data 一個代碼的值，依 `PRICE_HEADERS` 代碼之後的順序；欄名同 `workflow.daily_price`。"""

    company_name: str | None
    last_close: Decimal | None
    last_open: Decimal | None
    previous_close: Decimal | None
    volatility_200d: Decimal | None
    market_cap_millions: Decimal | None


# 結果的種類
NEW_TRADES = "新增交易"
UNCHANGED = "已存在"
DIFFERENT = "內容不同（未覆蓋）"
BLOCKED = "擋下"
NAME_UNMATCHED = "名字對不到信箱"
MAILBOX_EMPTY = "信箱尚未建立"
CANCELLED = "交易已取消"
PRICE_CONFLICT = "股價衝突"
PRICE_NOT_NUMBER = "股價不是數字"
IMPORT_ERROR_BOOK_NAME = "每日匯入錯誤.xlsx"
_SUCCESS_KINDS = frozenset((NEW_TRADES, UNCHANGED, CANCELLED))
# 清單更新停用前（#87）的成功結果「新增商品」「補上空白」：錯誤清單重建歷史時照樣不列
_RETIRED_SUCCESS_KINDS = frozenset(("新增商品", "補上空白"))

_PCT = Decimal("0.000001")
_AMOUNT = Decimal("0.0001")
_FULL_TDCC = re.compile(rf"[0-9]{{{TDCC_DIGITS}}}")

# 交易明細的欄 → `workflow.trade_row` 的欄（Product 只收 FCN、TDCC 是哪一檔，不在這裡）；人工維護可以改這些欄（#116）
TRADE_FIELDS = {
    "編號": "number", "交易日期": "trade_date", "幣別": "currency", "金額": "amount", "FC/PB": "fc_name",
    "分公司": "branch", "前台主管/投專": "supervisor", "UF": "uf", "期初價格": "initial_pricing", "備註": "note",
    "募集": "raised", "TS/IIS 核對": "ts_iis_checked", "核對TS/IIS人員": "checker", "起單人員": "originator",
}  # fmt: skip
TRADE_REQUIRED = ("交易日期", "幣別", "金額", "FC/PB", "分公司")
_TRADE_COLUMNS = ("product", *TRADE_FIELDS.values())  # 寫入、比對資料庫的交易列用的欄


class DailyImportError(OperatorError):
    """整份無法匯入（檔案或欄名有問題）；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class ImportLine:
    tdcc: str
    kind: str
    detail: str = ""

    @property
    def problem(self) -> bool:
        return self.kind not in _SUCCESS_KINDS


@dataclass(frozen=True)
class PriceSummary:
    """這次匯入的股價（#87）。"""

    price_date: dt.date  # 取價日
    count: int  # 存了幾個代碼（含衝突的）
    replaced: int  # 取代了這個取價日原本的幾筆（第一次匯入這一天是 0）
    conflicts: tuple[str, ...]  # 值不一樣的重複代碼


@dataclass(frozen=True)
class ImportReceipt:
    run_id: int
    lines: tuple[ImportLine, ...]
    people: int
    skipped_trade_rows: int
    cancelled_trade_rows: int
    error_book_path: Path
    prices: PriceSummary | None  # 只匯交易明細時是 None
    fcn_trade_rows: int  # 交易明細裡 Product 為 FCN 的列數，含取消的
    fcn_tdccs: tuple[str, ...]  # 帶到起單資料頁的 TDCC（#90）

    def count(self, kind: str) -> int:
        return sum(line.kind == kind for line in self.lines)


@dataclass(frozen=True)
class PriceCsvReceipt:
    """回補股價 CSV 的結果（#112）。"""

    run_id: int
    close_date: dt.date  # CSV 的取價日（收盤那一天）；存進資料庫的取價日是 `prices.price_date`
    prices: PriceSummary
    lines: tuple[ImportLine, ...]  # 股價不是數字、股價衝突
    error_book_path: Path
    closes: CloseSummary | None = None  # 收盤價（#153）；讀不到假日表時沒有存
    closed: tuple[str, ...] = ()  # 那一天休市、收盤價沒有存的代碼
    close_uncovered: tuple[str, ...] = ()  # 那一天超出假日表涵蓋範圍、收盤價沒有存的代碼
    close_conflicts: tuple[str, ...] = ()  # 去掉類別字後才重複、值不同、收盤價沒有存的代碼


@dataclass(frozen=True)
class _SheetRow:
    number: int
    values: dict[str, Any]


@dataclass
class _Prices:
    """BBG_Data 讀出來的股價：值一樣的重複列只留一筆；值不一樣的代碼記成衝突（各組不同的值）。"""

    price_date: dt.date
    rows: dict[str, Price] = field(default_factory=dict)
    conflicts: dict[str, list[Price]] = field(default_factory=dict)
    row_counts: Counter[str] = field(default_factory=Counter)  # 每個代碼在檔案裡有幾列
    problems: list[ImportLine] = field(default_factory=list)  # 股價不是數字、股價衝突


@dataclass
class _Parsed:
    prices: _Prices | None  # 只匯交易明細時是 None
    trades: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    blocked: list[ImportLine] = field(default_factory=list)
    cancelled_lines: list[ImportLine] = field(default_factory=list)
    other_products: dict[str, str] = field(default_factory=dict)  # 略過的非 FCN：TDCC → Product
    fcn_tdccs: dict[str, bool] = field(default_factory=dict)  # 交易明細的 FCN TDCC（出現順序）→ 有沒有沒取消的列
    skipped_trade_rows: int = 0
    cancelled_trade_rows: int = 0
    fcn_trade_rows: int = 0


def import_daily(daily_book: Path, conn: psycopg.Connection, *, prices: bool = True) -> ImportReceipt:
    """匯入一份每日更新；回傳每檔 TDCC 的結果。整份有問題時丟 `DailyImportError`、不寫入任何資料。

    `prices=False`：只匯交易明細，不讀 BBG_Data、不動股價（回補時股價改用 `import_price_csv`）。
    """
    parsed = _read_daily(daily_book, prices=prices)

    with conn.transaction():
        mailbox = load_mailbox(conn)
        run_id = conn.execute(
            "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES (%s, %s, '{}') RETURNING run_id",
            (daily_book.name, sha256_file(daily_book)),
        ).fetchone()[0]
        lines = [*parsed.blocked, *parsed.cancelled_lines, *(parsed.prices.problems if parsed.prices else ())]
        lines += _write_trades(conn, parsed, run_id)
        _write_other_products(conn, parsed, run_id)
        summary = None if parsed.prices is None else _write_prices(conn, parsed.prices, run_id)
        lines += _unmatched_names(parsed, mailbox)
        error_book_path = daily_book.parent / IMPORT_ERROR_BOOK_NAME
        receipt = ImportReceipt(
            run_id=run_id,
            lines=tuple(sorted(lines, key=lambda line: line.tdcc)),
            people=len(mailbox.people),
            skipped_trade_rows=parsed.skipped_trade_rows,
            cancelled_trade_rows=parsed.cancelled_trade_rows,
            error_book_path=error_book_path,
            prices=summary,
            fcn_trade_rows=parsed.fcn_trade_rows,
            fcn_tdccs=tuple(tdcc for tdcc, live in parsed.fcn_tdccs.items() if live),
        )
        conn.execute(
            "UPDATE workflow.import_run SET summary = %s::jsonb WHERE run_id = %s",
            (json.dumps(_summary(receipt), ensure_ascii=False), run_id),
        )
        _write_error_book(conn, error_book_path)
    return receipt


def import_price_csv(path: Path, conn: psycopg.Connection, calendar: TradingCalendar | None = None) -> PriceCsvReceipt:
    """回補股價（#112）：匯入一份 `MarketData_YYYYMMDD.csv`。整份有問題時丟 `DailyImportError`、不寫入任何資料。

    有 `calendar`（交易所假日表）時同時寫 CSV 那一天的收盤價（#153）；沒有時收盤價不存。"""
    close_date, prices = _read_price_csv(path)
    sha256 = sha256_file(path)
    with conn.transaction():
        run_id = conn.execute(
            "INSERT INTO workflow.import_run (daily_file_name, daily_sha256, summary) VALUES (%s, %s, '{}') RETURNING run_id",
            (path.name, sha256),
        ).fetchone()[0]
        summary = _write_prices(conn, prices, run_id)
        closes, closed, uncovered, close_conflicts = None, (), (), ()
        if calendar is not None:
            found = closes_from_csv(
                close_date, {code: price.last_close for code, price in prices.rows.items()}, calendar, prices.conflicts
            )
            closed, uncovered, close_conflicts = found.closed, found.uncovered, found.conflicts
            details = {"closed": list(closed), "uncovered": list(uncovered), "conflicts": list(close_conflicts)}
            closes = write_closes(
                conn,
                found.closes,
                source=CSV_SOURCE,
                file_name=path.name,
                sha256=sha256,
                details=details,
                run_id=run_id,
            )
        error_book_path = path.parent / IMPORT_ERROR_BOOK_NAME
        receipt = PriceCsvReceipt(
            run_id,
            close_date,
            summary,
            tuple(prices.problems),
            error_book_path,
            closes,
            closed,
            uncovered,
            close_conflicts,
        )
        record = {
            "counts": dict(Counter(line.kind for line in receipt.lines)),
            "close_date": close_date.isoformat(),
            "price_date": summary.price_date.isoformat(),
            "price_count": summary.count,
            "lines": [[line.tdcc, line.kind, line.detail] for line in receipt.lines],
        }
        conn.execute(
            "UPDATE workflow.import_run SET summary = %s::jsonb WHERE run_id = %s",
            (json.dumps(record, ensure_ascii=False), run_id),
        )
        _write_error_book(conn, error_book_path)
    return receipt


@dataclass(frozen=True)
class ImportInfo:
    """一次每日匯入的紀錄（寄信前確認資料是不是最新的）。"""

    run_id: int
    imported_at: dt.datetime
    file_name: str
    sha256: str


def latest_import(conn: psycopg.Connection) -> ImportInfo | None:
    """最近一次每日匯入；還沒匯入過時回傳 None。回補股價 CSV（summary 有 close_date）不算：寄信前要看的是交易明細。"""
    with conn.transaction():
        row = conn.execute(
            "SELECT run_id, imported_at, daily_file_name, daily_sha256 FROM workflow.import_run"
            " WHERE NOT summary ? 'close_date' ORDER BY run_id DESC LIMIT 1"
        ).fetchone()
    return None if row is None else ImportInfo(*row)


# ---------- 每日更新 ----------


def _read_daily(path: Path, *, prices: bool) -> _Parsed:
    wb = _open(path)
    try:
        for sheet in (TRADE_SHEET, PRICE_SHEET) if prices else (TRADE_SHEET,):
            if sheet not in wb.sheetnames:
                raise DailyImportError(f"每日更新缺少工作表「{sheet}」")
        trade_rows = _rows_with_header(wb[TRADE_SHEET], TRADE_HEADERS)
        parsed = _Parsed(_read_prices(wb[PRICE_SHEET]) if prices else None)
    finally:
        wb.close()
    _parse_trades(trade_rows, parsed)
    return parsed


def _rows_with_header(sheet: Any, expected: Sequence[str]) -> list[_SheetRow]:
    """檢查第 1 列欄名和順序完全相同（後面多出的空白欄不算），回傳 (列號, {欄名: 值})。"""
    rows = sheet.iter_rows(values_only=True)
    header = [_text(v) for v in next(rows, ())]
    while header and header[-1] is None:
        header.pop()
    if tuple(header) != tuple(expected):
        for i, want in enumerate(expected):
            got = header[i] if i < len(header) else None
            if got != want:
                where = openpyxl.utils.get_column_letter(i + 1)
                raise DailyImportError(
                    f"「{sheet.title}」的欄名和來源表不同，整份不匯入：{where} 欄應為「{want}」，實際是「{got or '空白'}」"
                )
        raise DailyImportError(f"「{sheet.title}」在「{expected[-1]}」之後還有其他欄，整份不匯入")
    width = len(expected)
    result = []
    for number, row in enumerate(rows, 2):
        values = (tuple(row) + (None,) * width)[:width]
        if any(_text(v) is not None for v in values):
            result.append(_SheetRow(number, dict(zip(expected, values, strict=True))))
    return result


def _read_prices(sheet: Any) -> _Prices:
    """BBG_Data：取價日、欄名不對，或沒有任何股價，就整份不匯入（ADR 0008、#87）。"""
    width = PRICE_FIRST_COLUMN - 1 + len(PRICE_HEADERS)
    rows = [(tuple(row) + (None,) * width)[:width] for row in sheet.iter_rows(max_col=width, values_only=True)]
    blank = (None,) * width
    prices = _Prices(_price_date(rows[PRICE_DATE_ROW - 1] if len(rows) >= PRICE_DATE_ROW else blank))
    _check_price_headers(rows[PRICE_HEADER_ROW - 1] if len(rows) >= PRICE_HEADER_ROW else blank)
    for number, row in enumerate(rows[PRICE_FIRST_ROW - 1 :], PRICE_FIRST_ROW):
        _add_price(prices, number, row[PRICE_FIRST_COLUMN - 1 :], PRICE_SHEET, PRICE_HEADERS[2:])
    if not prices.rows and not prices.conflicts:
        raise DailyImportError(
            f"「{PRICE_SHEET}」沒有任何股價（第 {PRICE_FIRST_ROW} 列起都沒有代碼），整份不匯入；請確認取價是否完成"
        )
    _list_conflicts(prices)
    return prices


def _list_conflicts(prices: _Prices) -> None:
    prices.problems += [
        ImportLine(
            "", PRICE_CONFLICT, f"{code}：同一個代碼有 {prices.row_counts[code]} 列、值不同，價格沒有存，算價格時不能用"
        )
        for code in prices.conflicts
    ]


def _read_price_csv(path: Path) -> tuple[dt.date, _Prices]:
    """回補股價 CSV（#112）：回傳 (收盤那一天, 股價)；股價的取價日是下一天。欄名、取價日不對或沒有股價就整份不匯入。"""
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DailyImportError(f"股價 CSV 不是 UTF-8 編碼，整份不匯入：{path}") from exc
    except OSError as exc:
        raise DailyImportError(f"打不開股價 CSV：{path}（{exc}）") from exc
    rows = list(csv.reader(io.StringIO(text, newline="")))  # 引號裡的換行留在同一格
    header = [_text(v) for v in (rows[0] if rows else [])]
    while header and header[-1] is None:
        header.pop()
    width = len(PRICE_CSV_HEADERS)
    for i, want in enumerate(PRICE_CSV_HEADERS):
        got = header[i] if i < len(header) else None
        if got != want:
            raise DailyImportError(
                f"股價 CSV 的欄名和範本不同，整份不匯入：第 {i + 1} 欄應為「{want}」，實際是「{got or '空白'}」"
            )
    if len(header) > width:
        raise DailyImportError(f"股價 CSV 在「{PRICE_CSV_HEADERS[-1]}」之後還有其他欄，整份不匯入")

    found: list[tuple[int, list[str]]] = []
    dates: dict[dt.date, int] = {}  # 取價日 → 第一次出現的列號
    for number, raw in enumerate(rows[1:], 2):
        row = (raw + [""] * width)[:width]
        if _text(row[1], dash_is_none=False) is None:
            continue  # 代碼空白的列略過
        try:
            close_date = _date(row[0])
        except ValueError:
            close_date = None
        if close_date is None:
            raise DailyImportError(f"股價 CSV 第 {number} 列的取價日「{row[0] or '空白'}」不是日期，整份不匯入")
        dates.setdefault(close_date, number)
        found.append((number, row))
    if not found:
        raise DailyImportError(f"股價 CSV 沒有任何股價（第 2 列起都沒有代碼），整份不匯入：{path}")
    if len(dates) > 1:
        shown = "、".join(f"{d:%Y-%m-%d}（第 {n} 列起）" for d, n in sorted(dates.items()))
        raise DailyImportError(f"股價 CSV 的取價日不只一天：{shown}，整份不匯入；一個 CSV 只放一天")

    close_date = next(iter(dates))
    prices = _Prices(close_date + PRICE_CSV_DATE_SHIFT)
    labels = (*PRICE_CSV_HEADERS[3:5], "1日前收盤價", *PRICE_CSV_HEADERS[5:7])
    for number, row in found:
        code, name, close, open_, volatility, market_cap = row[1:7]
        _add_price(prices, number, (code, name, close, open_, None, volatility, market_cap), path.name, labels)
    _list_conflicts(prices)
    return close_date, prices


def _price_date(row: Sequence[Any]) -> dt.date:
    value = row[PRICE_DATE_COLUMN - 1]
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    where = f"{get_column_letter(PRICE_DATE_COLUMN)}{PRICE_DATE_ROW}"
    shown = _text(value, dash_is_none=False) or "空白"
    raise DailyImportError(f"「{PRICE_SHEET}」的 {where} 應為取價日（日期），實際是「{shown}」，整份不匯入")


def _check_price_headers(row: Sequence[Any]) -> None:
    for i, want in enumerate(PRICE_HEADERS):
        got = _text(row[PRICE_FIRST_COLUMN - 1 + i])
        if got != want:
            where = f"{get_column_letter(PRICE_FIRST_COLUMN + i)}{PRICE_HEADER_ROW}"
            raise DailyImportError(
                f"「{PRICE_SHEET}」的欄名和範本不同，整份不匯入：{where} 應為「{want}」，實際是「{got or '空白'}」"
            )


def _add_price(prices: _Prices, number: int, values: Sequence[Any], source: str, labels: Sequence[str]) -> None:
    """一列股價（代碼、公司名稱，再依 `Price` 價格欄的順序）；代碼空白的列略過。
    同一個代碼再出現時，值一樣不重複算，值不一樣記成衝突。`source`、`labels` 是訊息裡的來源與價格欄名。"""
    code = _text(values[0], dash_is_none=False)
    if code is None:
        return
    prices.row_counts[code] += 1
    numbers = []
    for header, value in zip(labels, values[2:], strict=True):
        try:
            numbers.append(_decimal(value))
        except ValueError:
            prices.problems.append(
                ImportLine(
                    "", PRICE_NOT_NUMBER, f"{source}第 {number} 列 {code}：{header}「{value}」不是數字，存成空白"
                )
            )
            numbers.append(None)
    found = Price(_text(values[1], dash_is_none=False), *numbers)
    if code in prices.conflicts:
        if found not in prices.conflicts[code]:
            prices.conflicts[code].append(found)
    elif code not in prices.rows:
        prices.rows[code] = found
    elif prices.rows[code] != found:
        prices.conflicts[code] = [prices.rows.pop(code), found]


def _parse_trades(rows: Iterable[_SheetRow], parsed: _Parsed) -> None:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    errors: dict[str, list[str]] = defaultdict(list)
    for sheet_row in rows:
        number, row = sheet_row.number, sheet_row.values
        if (_text(row["Product"]) or "").upper() != FCN:
            parsed.skipped_trade_rows += 1
            _remember_other(parsed, row["TDCC"], row["Product"])
            continue
        parsed.fcn_trade_rows += 1
        _remember_fcn(parsed, row)
        where = f"交易明細第 {number} 列"
        if is_cancel_note(row["備註"]):
            parsed.cancelled_trade_rows += 1
            parsed.cancelled_lines.append(
                ImportLine(
                    _tdcc_or_raw(row["TDCC"]),
                    CANCELLED,
                    f"{where}：備註為 Cancel，未匯入資料庫",
                )
            )
            continue
        try:
            tdcc = _tdcc(row["TDCC"])
        except ValueError as exc:
            parsed.blocked.append(ImportLine(_tdcc_or_raw(row["TDCC"]), BLOCKED, f"{where}：{exc}"))
            continue
        try:
            trade = _trade(row)
            groups[tdcc].append(trade)
        except ValueError as exc:
            errors[tdcc].append(f"{where}：{exc}")
    for tdcc, trades in groups.items():
        if errors[tdcc]:
            parsed.blocked.append(ImportLine(tdcc, BLOCKED, "；".join(errors[tdcc]) + "（這檔的交易列都不匯入）"))
        else:
            parsed.trades[tdcc] = trades
    for tdcc in errors.keys() - groups.keys():
        parsed.blocked.append(ImportLine(tdcc, BLOCKED, "；".join(errors[tdcc])))


def _remember_fcn(parsed: _Parsed, row: Mapping[str, Any]) -> None:
    """記下這列的 FCN TDCC（#90）；同一檔只要有一列沒取消就帶。

    只收剛好 12 個半形數字（#93，同「3 準備」）：Excel 存成數字、前面的 0 掉了的不補 0，不帶。
    """
    tdcc = _text(row["TDCC"])
    if tdcc is None or not _FULL_TDCC.fullmatch(tdcc):
        return
    parsed.fcn_tdccs[tdcc] = parsed.fcn_tdccs.get(tdcc, False) or not is_cancel_note(row["備註"])


def _trade(row: Mapping[str, Any]) -> dict[str, Any]:
    trade: dict[str, Any] = {"product": FCN}
    for header, column in TRADE_FIELDS.items():
        try:
            value = trade_value(header, row[header])
        except ValueError as exc:
            raise ValueError(f"{header}：{exc}") from exc
        if value is None and header in TRADE_REQUIRED:
            raise ValueError(f"{header} 沒有值")
        trade[column] = value
    return trade


def trade_value(header: str, value: Any) -> Any:
    """交易明細的一格（Excel 的值或畫面打的字）→ 存進資料庫的值；每日匯入與人工維護（#116）共用，
    存進去的值才會和之後匯入的一樣。空白（含 `-`）回傳 None，必填由呼叫端擋；格式錯丟 ValueError（不含欄名）。"""
    return _TRADE_CONVERT.get(header, _text)(value)


# ---------- 寫入資料庫 ----------


def _write_prices(conn: psycopg.Connection, prices: _Prices, run_id: int) -> PriceSummary:
    """股價依取價日存入：同一個取價日整天取代（#87）；衝突的代碼價格留空，各組不同的值記在 conflict_rows。"""
    count = len(prices.rows) + len(prices.conflicts)
    conn.execute(
        "INSERT INTO workflow.price_import (run_id, price_date, price_count, conflict_count) VALUES (%s, %s, %s, %s)",
        (run_id, prices.price_date, count, len(prices.conflicts)),
    )
    replaced = conn.execute("DELETE FROM workflow.daily_price WHERE price_date = %s", (prices.price_date,)).rowcount
    columns = ("price_date", "code", *Price._fields, "conflict_rows", "run_id")
    blank = (None,) * len(Price._fields)
    rows = [(prices.price_date, code, *price, None, run_id) for code, price in prices.rows.items()]
    rows += [
        (prices.price_date, code, *blank, json.dumps([_jsonable(p) for p in found], ensure_ascii=False), run_id)
        for code, found in prices.conflicts.items()
    ]
    with conn.cursor() as cur:
        cur.executemany(
            f"INSERT INTO workflow.daily_price ({', '.join(columns)}) VALUES ({', '.join(['%s'] * len(columns))})", rows
        )
    return PriceSummary(prices.price_date, count, replaced, tuple(prices.conflicts))


def _jsonable(price: Price) -> dict[str, Any]:
    return {name: str(v) if isinstance(v, Decimal) else v for name, v in price._asdict().items()}


def _write_trades(conn: psycopg.Connection, parsed: _Parsed, run_id: int) -> list[ImportLine]:
    lines = []
    for tdcc, trades in parsed.trades.items():
        stored = conn.execute(
            f"SELECT {', '.join(_TRADE_COLUMNS)} FROM workflow.trade_row WHERE tdcc = %s ORDER BY row_seq", (tdcc,)
        ).fetchall()
        if not stored:
            columns = ("tdcc", "row_seq", *_TRADE_COLUMNS, "first_run_id")
            with conn.cursor() as cur:
                cur.executemany(
                    f"INSERT INTO workflow.trade_row ({', '.join(columns)}) VALUES ({', '.join(['%s'] * len(columns))})",
                    [(tdcc, seq, *(t[c] for c in _TRADE_COLUMNS), run_id) for seq, t in enumerate(trades, 1)],
                )
            lines.append(ImportLine(tdcc, NEW_TRADES, f"{len(trades)} 列"))
        elif len(stored) != len(trades) or any(
            (s != t[c])
            for row, t in zip(stored, trades, strict=False)
            for s, c in zip(row, _TRADE_COLUMNS, strict=True)
        ):
            lines.append(
                ImportLine(tdcc, DIFFERENT, f"交易列和資料庫不同（資料庫 {len(stored)} 列、這次 {len(trades)} 列）")
            )
        else:
            lines.append(ImportLine(tdcc, UNCHANGED, "交易列"))
    return lines


def _remember_other(parsed: _Parsed, tdcc: Any, product: Any) -> None:
    try:
        parsed.other_products.setdefault(_tdcc(tdcc), _text(product) or "（空白）")
    except ValueError:
        pass  # TDCC 不完整的列只計數


def _write_other_products(conn: psycopg.Connection, parsed: _Parsed, run_id: int) -> None:
    with conn.cursor() as cur:
        cur.executemany(
            "INSERT INTO workflow.other_product (tdcc, product, first_run_id) VALUES (%s, %s, %s)"
            " ON CONFLICT (tdcc) DO NOTHING",
            [(tdcc, product, run_id) for tdcc, product in parsed.other_products.items()],
        )


def _unmatched_names(parsed: _Parsed, mailbox: Mailbox) -> list[ImportLine]:
    if not mailbox.people:
        return [ImportLine("", MAILBOX_EMPTY, "mailbox.person 沒有任何人員，FC 名字無法比對")] if parsed.trades else []
    lines = []
    for tdcc, trades in parsed.trades.items():
        names = sorted(
            {t["fc_name"] for t in trades if not is_pb(t["fc_name"]) and mailbox.match(t["fc_name"]) is None}
        )
        if names:
            lines.append(ImportLine(tdcc, NAME_UNMATCHED, "、".join(names)))
    return lines


def _summary(receipt: ImportReceipt) -> dict[str, Any]:
    counts: dict[str, int] = defaultdict(int)
    for line in receipt.lines:
        counts[line.kind] += 1
    return {
        "counts": dict(counts),
        "people": receipt.people,
        "skipped_trade_rows": receipt.skipped_trade_rows,
        "price_date": None if receipt.prices is None else receipt.prices.price_date.isoformat(),
        "price_count": None if receipt.prices is None else receipt.prices.count,
        "lines": [[line.tdcc, line.kind, line.detail] for line in receipt.lines],
    }


def _write_error_book(conn: psycopg.Connection, path: Path) -> None:
    rows = []
    runs = conn.execute(
        "SELECT run_id, imported_at, daily_file_name, summary FROM workflow.import_run ORDER BY run_id"
    ).fetchall()
    for run_id, imported_at, file_name, summary in runs:
        for tdcc, kind, detail in summary.get("lines", []):
            if kind in _SUCCESS_KINDS or kind in _RETIRED_SUCCESS_KINDS:
                continue
            rows.append(
                (
                    imported_at.strftime("%Y-%m-%d %H:%M:%S%z"),
                    run_id,
                    file_name,
                    tdcc,
                    kind,
                    detail,
                )
            )

    book = table_book(
        "匯入錯誤",
        ("匯入時間", "批次", "來源檔案", "TDCC", "錯誤類型", "說明"),
        rows,
        widths=(22, 10, 30, 16, 24, 80),
        auto_filter=True,
    )
    try:
        save_book(book, path)
    except OSError as exc:
        raise DailyImportError(f"錯誤清單寫入失敗：{path}（{exc}）") from exc


# ---------- 值的轉換 ----------


def _open(path: Path) -> Any:
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise DailyImportError(f"打不開 Excel：{path}（{exc}）") from exc


def _text(value: Any, *, dash_is_none: bool = True) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text or (dash_is_none and text == "-"):
        return None
    return text


def is_cancel_note(value: Any) -> bool:
    """交易列的備註是 Cancel 或含「取消」：每日匯入不存進資料庫（人工維護也不能改成這樣，#116）。"""
    text = (_text(value, dash_is_none=False) or "").casefold()
    return text == "cancel" or "取消" in text


def _tdcc(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _text(value)
    if text is None or not text.isdigit() or len(text) > TDCC_DIGITS:
        raise ValueError(f"TDCC「{text or '空白'}」不是 12 碼以內的數字")
    return text.zfill(TDCC_DIGITS)


def _tdcc_or_raw(value: Any) -> str:
    try:
        return _tdcc(value)
    except ValueError:
        return _text(value) or "（空白）"


def _decimal(value: Any) -> Decimal | None:
    if _text(value) is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"「{value}」不是數字")
    try:
        number = Decimal(str(value).strip().replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"「{value}」不是數字") from exc
    if not number.is_finite():  # NaN、Infinity 算價格時會中斷
        raise ValueError(f"「{value}」不是數字")
    return number


def _whole_number(value: Any) -> int | None:
    number = _decimal(value)
    if number is not None and number != number.to_integral_value():
        raise ValueError(f"「{number}」不是整數")
    return None if number is None else int(number)


def _pct(value: Any) -> Decimal | None:
    number = _decimal(value)
    return None if number is None else number.quantize(_PCT)


def _amount(value: Any) -> Decimal | None:
    number = _decimal(value)
    return None if number is None else number.quantize(_AMOUNT)


def _date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    if text is None:
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"「{text}」不是日期")


# 交易明細各欄的轉換（`trade_value`）；沒列的是文字
_TRADE_CONVERT = {"編號": _whole_number, "交易日期": _date, "金額": _amount, "UF": _pct}
