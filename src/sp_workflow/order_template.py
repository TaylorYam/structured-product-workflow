"""「Outlook orders 轉成 Excel」格式表（#40 追加）：把讀好的 Order 條件排成作業人員既有的格式。

版面照作業人員提供的範例檔 `example` 工作表：第 2 列表頭、從 B 欄開始，一檔一列。
讀到 Order 的商品都放（含待確認）；讀不到、換算不了、有修改或不一致的格子照樣填能確定的值，
並標淺紅底，該檔在條件清單標「待確認」並寫原因，不猜值。整列的問題（例如條件表有沒列過的欄）標在 TDCC 那格。

信裡沒有直接寫、但依固定規則換算的欄位：
- Product：條件表沒有時取主旨 `] FCN_…` 的商品別。
- KO Type：統一成 Daily／Daily Memory／Period End／Period End Memory；MS 的 `OBSERVATION` 依 KO Memory 判斷，
  Citi 依 Daily KO、Memory Autocall 兩欄（TRUE／FALSE）判斷。
- Barrier Type：統一成 NONE／EKI／AKI（MS `EUROPEAN` = EKI、`AMERICAN` = AKI）；沒有 KI 時 KI Barrier 留空。
- SG：Tenor = No. of Periods、Observation Frequency = 1（Settlement Frequency 是 Monthly 時）。
- OTC：信裡沒寫時填 `Note`。
- Trade date：條件表沒有時用 Order 寄出日期。
- Issue date：條件表沒有、但有日曆天數的發行日偏移時 = Trade date + 偏移天數。
- Effective Date Offset：Issue date − Trade date 的日曆天數。
"""

from __future__ import annotations

import datetime as dt
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from decimal import Decimal

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

SHEET_TITLE = "example"  # 與範例檔相同，方便直接複製貼上
HEADER_ROW = 2
FIRST_COLUMN = 2  # B 欄

_TEAL = "FF00AB8E"
_NAVY = "FF002060"
_BLUE = "FF0B76A0"  # 範例檔佈景主題「輔色 4、較深 25%」
_PCT = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
_DATE = "mm-dd-yy"


@dataclass(frozen=True)
class Column:
    header: str
    width: float
    number_format: str = "General"
    fill: str = _BLUE
    align: str = "center"


COLUMNS: tuple[Column, ...] = (
    Column("TDCC", 14, "@", align="left"),
    Column("Product", 9.5),
    Column("Currency", 10, fill=_TEAL),
    Column("Guaranteed Periods (m) ", 11),
    *(Column(f"BBG Code {n}", 12, align="left") for n in range(1, 6)),
    Column("Strike (%) ", 9.5, _PCT),
    Column("KO Type", 16, _PCT, align="left"),
    Column("KO Barrier (%)", 10, _PCT),
    Column("Coupon p.a. (%) ", 10, "0.00_);[Red]\\(0.00\\)"),
    Column("Upfront / NotePrice (%) ", 11, "0.00"),
    Column("Tenor (m)", 8),
    Column("Barrier Type", 10),
    Column("KI Barrier (%) ", 10, _PCT),
    Column("Observation Frequency (m)", 12, fill=_TEAL),
    Column("OTC", 8, fill=_TEAL),
    Column("Funding Spread (bps)", 10, fill=_TEAL),
    Column("Effective Date Offset(Calendar Days) ", 12, fill=_TEAL),
    Column("Trade date", 11, _DATE, _NAVY),
    Column("Issue date", 11, _DATE, _NAVY),
    Column("Final valuation date", 11, _DATE, _NAVY),
    Column("Maturity date", 11, _DATE, _NAVY),
    Column("期初價格", 10, fill=_NAVY),
    Column("起單FC", 10, fill=_NAVY),
)

_COLUMN_INDEX = {c.header.strip(): i for i, c in enumerate(COLUMNS)}
# 統一欄位 → 格式表的欄（讀信時這個欄位有問題，就把這一格標淺紅底）
FIELD_COLUMNS = {
    "商品": "Product",
    "幣別": "Currency",
    "Non-Call(月)": "Guaranteed Periods (m)",
    **{f"標的{n}": f"BBG Code {n}" for n in range(1, 6)},
    "執行價(%)": "Strike (%)",
    "KO觀察方式": "KO Type",
    "KO(%)": "KO Barrier (%)",
    "年利率(%)": "Coupon p.a. (%)",
    "價格(%)": "Upfront / NotePrice (%)",
    "UF(%)": "Upfront / NotePrice (%)",
    "天期(月)": "Tenor (m)",
    "KI型態": "Barrier Type",
    "KI(%)": "KI Barrier (%)",
    "觀察頻率(月)": "Observation Frequency (m)",
    "發行日偏移": "Effective Date Offset(Calendar Days)",
    "交易日": "Trade date",
    "發行日": "Issue date",
    "最終比價日": "Final valuation date",
    "到期日": "Maturity date",
    "期初定價": "期初價格",
}
FLAG_FILL = "FFFFC7CE"  # Excel 內建「淺紅色填滿」

KO_TYPES = {k.lower(): k for k in ("Daily", "Daily Memory", "Period End", "Period End Memory")}
_BARRIER_TYPES = {"none": "NONE", "eki": "EKI", "aki": "AKI", "european": "EKI", "american": "AKI", "no": "NONE"}
_PRODUCT_IN_SUBJECT = re.compile(r"\]\s*([A-Za-z]{2,10})_")
DEFAULT_OTC = "Note"


def _other(others: Sequence[tuple[str, str]], name: str) -> str:
    key = name.lower()
    return next((v for k, v in others if k.lower() == key), "")


def template_row(
    tdcc: str,
    values: dict[str, object],
    others: Sequence[tuple[str, str]],
    subject: str,
    sent_on: dt.datetime | None,
    flagged_fields: Iterable[str] = (),
) -> tuple[tuple[object, ...], list[str], frozenset[int]]:
    """一檔商品在格式表的一列（依 COLUMNS 順序）、換算不了的原因，以及要標淺紅底的欄。

    flagged_fields 是讀信時有問題的統一欄位；對不到格式表某一欄的（或空字串＝整列的問題）標在 TDCC 欄。
    """
    problems: list[str] = []
    flagged: set[int] = set()
    for name in flagged_fields:
        flagged.add(_COLUMN_INDEX.get(FIELD_COLUMNS.get(name, ""), 0))

    product = values.get("商品")
    if product is None:
        found = _PRODUCT_IN_SUBJECT.search(subject)
        product = found.group(1).upper() if found else None

    ko_type = _ko_type(values.get("KO觀察方式"), _other(others, "KO Memory"), problems)
    if values.get("KO觀察方式") is None:
        ko_type = _ko_from_flags(_other(others, "Daily KO"), _other(others, "Memory Autocall"), problems)
    barrier, ki = _barrier(values.get("KI型態"), values.get("KI(%)"), problems)

    tenor, frequency = values.get("天期(月)"), values.get("觀察頻率(月)")
    periods, settlement = _other(others, "No. of Periods"), _other(others, "Settlement Frequency")
    if tenor is None and periods.isdigit() and settlement.lower() == "monthly":
        tenor, frequency = int(periods), frequency or 1

    trade = values.get("交易日") or (sent_on.date() if sent_on else None)
    issue = values.get("發行日")
    offset_text = str(values.get("發行日偏移") or "")
    if issue is None and trade is not None and offset_text.isdigit():
        issue = trade + dt.timedelta(days=int(offset_text))
    offset = (issue - trade).days if isinstance(issue, dt.date) and isinstance(trade, dt.date) else None

    if values.get("標的6"):
        problems.append("標的超過 5 檔，格式表放不下")
        flagged.add(_COLUMN_INDEX["BBG Code 5"])
    cells = (
        tdcc,
        product,
        values.get("幣別"),
        values.get("Non-Call(月)"),
        *(values.get(f"標的{n}") for n in range(1, 6)),
        values.get("執行價(%)"),
        ko_type,
        values.get("KO(%)"),
        values.get("年利率(%)"),
        values.get("價格(%)"),
        tenor,
        barrier,
        ki,
        frequency,
        _other(others, "OTC") or DEFAULT_OTC,
        _other(others, "Funding Spread (bps)") or None,
        offset,
        trade,
        issue,
        values.get("最終比價日"),
        values.get("到期日"),
        values.get("期初定價"),
        None,  # 起單FC：使用者決定留空，由人補
    )
    optional = {"BBG Code 2", "BBG Code 3", "BBG Code 4", "BBG Code 5", "Funding Spread (bps)", "起單FC"}
    if barrier == "NONE":
        optional.add("KI Barrier (%) ")
    missing = [c.header.strip() for c, v in zip(COLUMNS, cells, strict=True) if v is None and c.header not in optional]
    if missing:
        problems.append(f"格式表缺：{'、'.join(missing)}")
        flagged.update(_COLUMN_INDEX[name] for name in missing)
    if any(p.startswith("Barrier Type 是 NONE") for p in problems):
        flagged.add(_COLUMN_INDEX["KI Barrier (%)"])
    return cells, problems, frozenset(flagged)


def _ko_type(raw: object, memory: str, problems: list[str]) -> str | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if text.lower() == "observation":  # MS：每期觀察，記憶式另寫在 KO Memory
        if memory.lower() in ("yes", "y"):
            return "Period End Memory"
        if memory.lower() in ("no", "n"):
            return "Period End"
        problems.append(f"KO Type 是 OBSERVATION，但 KO Memory 不認得：{memory or '（空白）'}")
        return None
    if text.lower() in KO_TYPES:
        return KO_TYPES[text.lower()]
    problems.append(f"KO Type 不認得：{text}")
    return None


def _ko_from_flags(daily: str, memory: str, problems: list[str]) -> str | None:
    """Citi 用兩欄 TRUE／FALSE：Daily KO（每日觀察）、Memory Autocall（記憶式）。"""
    if not daily and not memory:
        return None
    flags = {"true": True, "false": False}
    if daily.lower() not in flags or memory.lower() not in flags:
        problems.append(f"Daily KO／Memory Autocall 不認得：{daily or '（空白）'}／{memory or '（空白）'}")
        return None
    kind = "Daily" if flags[daily.lower()] else "Period End"
    return f"{kind} Memory" if flags[memory.lower()] else kind


def _barrier(raw: object, ki: object, problems: list[str]) -> tuple[str | None, object]:
    if raw is None:
        if ki is None:
            return "NONE", None  # 例：SG 的 KI Type、KI 都是 N/A
        problems.append("有 KI 價格但沒有 Barrier Type")
        return None, ki
    barrier = _BARRIER_TYPES.get(str(raw).strip().lower())
    if barrier is None:
        problems.append(f"Barrier Type 不認得：{raw}")
        return None, ki
    if barrier == "NONE":
        if ki not in (None, Decimal(0)):
            problems.append(f"Barrier Type 是 NONE，但 KI Barrier 有值：{ki}")
        return barrier, None
    return barrier, ki


def write_sheet(ws: Worksheet, rows: Sequence[tuple[tuple[object, ...], frozenset[int]]]) -> None:
    """照範例檔版面寫表頭與資料列；rows 是 (一列的值, 要標淺紅底的欄位置)。"""
    flag = PatternFill("solid", fgColor=FLAG_FILL)
    ws.title = SHEET_TITLE
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[HEADER_ROW].height = 75
    ws.column_dimensions["A"].width = 3
    for i, column in enumerate(COLUMNS):
        cell = ws.cell(HEADER_ROW, FIRST_COLUMN + i, column.header)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=column.fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = column.width
    for r, (values, flagged) in enumerate(rows, start=HEADER_ROW + 1):
        for i, (column, value) in enumerate(zip(COLUMNS, values, strict=True)):
            cell = ws.cell(r, FIRST_COLUMN + i, float(value) if isinstance(value, Decimal) else value)
            cell.number_format = column.number_format
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal=column.align)
            cell.border = border
            if i in flagged:
                cell.fill = flag
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, FIRST_COLUMN + 1)
