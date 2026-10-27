"""可連結標的白名單（#273）：作業人員大約每月用「維護」選單匯入「ELN可連結標的白名單.xlsm」，整份取代存進資料庫。

可以連結的標的只有九個指數的成分股、ADR、ETF（#252）；風險等級的標的種類分把白名單上的 ETF 算 3 分（`risk_level`）。
- 讀「總表」工作表（第 2 列表頭，第 3 列起每檔一列）：`TICKER` 之後到「所屬指數數」之前的每一欄是一個分類
  （`GROUPS`：SP、NQ…、ADR、ETF），有值就是屬於那一類；抓取日是「說明」工作表的 B3。白名單格式固定：
  工作表、表頭（含不認得的分類欄）、抓取日與代碼有任何不對都整份不寫入，說明哪裡不對。
- 每次匯入整份取代 `workflow.eln_whitelist`，並在 `workflow.eln_whitelist_import` 留一筆紀錄（抓取日、檔名、檔數、
  和上一次匯入比新增、移除、分類變了幾檔、誰）；同一份重複匯入內容不變。
- `current_import`：最近一次匯入（畫面在匯入前先給作業人員看）；`stored_etfs`：白名單上的 ETF，還沒匯入過時 None
  （風險等級改用設定檔的備用清單）。
傳入的 `conn` 由 `database.connect` 建立。
"""

from __future__ import annotations

import datetime as dt
import getpass
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from sp_workflow.database import OperatorError

if TYPE_CHECKING:
    import psycopg

SUMMARY_SHEET, NOTES_SHEET = "總表", "說明"
CODE_HEADER, COUNT_HEADER = "TICKER", "所屬指數數"  # 分類欄在這兩欄之間
ETF_GROUP = "ETF"
GROUPS = (
    "SP",
    "NQ",
    "SOX",
    "R1000",
    "R2000",
    "R3000",
    "NKY",
    "TPX500",
    "HSI",
    "ADR",
    ETF_GROUP,
)  # 白名單「設定」的簡稱
_HEADER_ROW, _FETCHED_CELL = 2, "B3"


class WhitelistError(OperatorError):
    """白名單檔案打不開或格式不對；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class Listed:
    """白名單上的一檔：代碼（白名單寫法，例 `NVDA UW`、`AAAU US`）與所屬分類（照總表欄位順序）。"""

    code: str
    groups: tuple[str, ...]


@dataclass(frozen=True)
class Whitelist:
    fetched_on: dt.date  # 說明!B3 的抓取日期
    rows: tuple[Listed, ...]

    @property
    def etfs(self) -> frozenset[str]:
        return frozenset(row.code for row in self.rows if ETF_GROUP in row.groups)


@dataclass(frozen=True)
class WhitelistImport:
    """一次匯入的紀錄（`workflow.eln_whitelist_import` 的一列）。"""

    fetched_on: dt.date
    source_file: str
    code_count: int
    etf_count: int
    imported_by: str
    imported_at: dt.datetime


@dataclass(frozen=True)
class WhitelistReceipt:
    """一次匯入的結果：這一份的抓取日、檔數與 ETF 檔數，和上一次匯入的那份比新增、移除、分類變了的檔
    （(之前, 之後)）；`previous` 是上一次匯入那份的抓取日，第一次是 None。"""

    fetched_on: dt.date
    code_count: int
    etf_count: int
    added: tuple[Listed, ...]
    removed: tuple[Listed, ...]
    regrouped: tuple[tuple[Listed, Listed], ...]
    previous: dt.date | None


def read_whitelist(path: Path) -> Whitelist:
    """讀白名單的總表與抓取日；打不開或格式不對時丟 `WhitelistError`，說明哪裡不對。"""
    import openpyxl

    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # 檔案不在、不是 Excel、被別的程式鎖住
        raise WhitelistError(f"白名單打不開（{path.name}）：{exc}") from exc
    try:
        for name in (SUMMARY_SHEET, NOTES_SHEET):
            if name not in book.sheetnames:
                raise WhitelistError(f"白名單沒有「{name}」工作表，請選「ELN可連結標的白名單」原檔。")
        cells = book[NOTES_SHEET].iter_rows(min_row=3, max_row=3, min_col=2, max_col=2, values_only=True)
        fetched = next(cells, (None,))[0]
        rows = list(book[SUMMARY_SHEET].iter_rows(min_row=_HEADER_ROW, values_only=True))
    finally:
        book.close()
    if not isinstance(fetched, dt.datetime | dt.date):
        raise WhitelistError(f"「{NOTES_SHEET}」{_FETCHED_CELL} 不是抓取日期，請先按「更新成分股」再存檔。")
    fetched_on = fetched.date() if isinstance(fetched, dt.datetime) else fetched
    return Whitelist(fetched_on, _listed(rows))


def _listed(rows: list[tuple[object, ...]]) -> tuple[Listed, ...]:
    header = [str(cell).strip() if cell is not None else "" for cell in (rows[0] if rows else ())]
    if not header or header[0] != CODE_HEADER or COUNT_HEADER not in header:
        raise WhitelistError(
            f"「{SUMMARY_SHEET}」第 {_HEADER_ROW} 列的表頭不對：要從 {CODE_HEADER} 開始、有「{COUNT_HEADER}」。"
        )
    groups = header[1 : header.index(COUNT_HEADER)]
    unknown = [g or "（空白）" for g in groups if g not in GROUPS]
    if unknown:
        raise WhitelistError(
            f"「{SUMMARY_SHEET}」的表頭有不認得的分類欄：{'、'.join(unknown)}。白名單新增或改名指數時要先更新程式。"
        )
    if ETF_GROUP not in groups:
        raise WhitelistError(f"「{SUMMARY_SHEET}」的表頭沒有 {ETF_GROUP} 欄。")
    found: list[Listed] = []
    seen: set[str] = set()
    for row in rows[1:]:
        code = str(row[0]).strip() if row and row[0] is not None else ""
        if not code:
            continue
        if code in seen:
            raise WhitelistError(f"「{SUMMARY_SHEET}」的 {code} 出現兩次。")
        seen.add(code)
        member = tuple(g for g, cell in zip(groups, row[1:], strict=False) if cell not in (None, ""))
        if not member:
            raise WhitelistError(f"「{SUMMARY_SHEET}」的 {code} 沒有所屬的指數、ADR 或 ETF。")
        found.append(Listed(code, member))
    if not found:
        raise WhitelistError(f"「{SUMMARY_SHEET}」沒有任何代碼；請在 Excel 開啟、確認總表有算出來後存檔再匯入。")
    if not any(ETF_GROUP in row.groups for row in found):
        raise WhitelistError(f"「{SUMMARY_SHEET}」沒有任何 {ETF_GROUP}，請確認是完整的白名單。")
    return tuple(found)


def import_whitelist(
    conn: psycopg.Connection, whitelist: Whitelist, source: Path, *, user: str | None = None
) -> WhitelistReceipt:
    """整份取代資料庫的白名單並留紀錄（自己一個 transaction）。"""
    with conn.transaction():
        conn.execute("LOCK TABLE workflow.eln_whitelist IN EXCLUSIVE MODE")  # 兩個人同時匯入時一個等另一個
        before = {
            code: Listed(code, tuple(groups))
            for code, groups in conn.execute("SELECT code, groups FROM workflow.eln_whitelist")
        }
        previous = current_import(conn)
        now = {row.code: row for row in whitelist.rows}
        receipt = WhitelistReceipt(
            whitelist.fetched_on,
            len(whitelist.rows),
            len(whitelist.etfs),
            added=tuple(row for code, row in now.items() if code not in before),
            removed=tuple(row for code, row in before.items() if code not in now),
            regrouped=tuple((before[code], row) for code, row in now.items() if code in before and before[code] != row),
            previous=None if previous is None else previous.fetched_on,
        )
        conn.execute("DELETE FROM workflow.eln_whitelist")
        with conn.cursor().copy("COPY workflow.eln_whitelist (code, groups) FROM STDIN") as copy:
            for row in whitelist.rows:
                copy.write_row((row.code, list(row.groups)))
        conn.execute(
            "INSERT INTO workflow.eln_whitelist_import (fetched_on, source_file, code_count, etf_count, added_count,"
            " removed_count, regrouped_count, imported_by) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
            (
                receipt.fetched_on,
                source.name,
                receipt.code_count,
                receipt.etf_count,
                len(receipt.added),
                len(receipt.removed),
                len(receipt.regrouped),
                user or getpass.getuser(),
            ),
        )
    return receipt


def current_import(conn: psycopg.Connection) -> WhitelistImport | None:
    """最近一次匯入（就是資料庫目前這一份）；還沒匯入過時 None。"""
    row = conn.execute(
        "SELECT fetched_on, source_file, code_count, etf_count, imported_by, imported_at"
        " FROM workflow.eln_whitelist_import ORDER BY id DESC LIMIT 1"
    ).fetchone()
    return None if row is None else WhitelistImport(*row)


def stored_etfs(conn: psycopg.Connection) -> frozenset[str] | None:
    """資料庫白名單上的 ETF（白名單寫法）；還沒匯入過時 None。"""
    if current_import(conn) is None:
        return None
    return frozenset(
        code
        for (code,) in conn.execute("SELECT code FROM workflow.eln_whitelist WHERE %s = ANY (groups)", (ETF_GROUP,))
    )
