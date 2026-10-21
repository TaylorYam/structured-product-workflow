"""寄信草稿：為核對通過的 FCN 在 Outlook 建立 TS／IIS 草稿給承作的 FC（#29，需求更新見 #29 留言）。

對外只有 `create_drafts`。資料全部讀最小 FCN 資料庫（#32、#35），不讀 Excel。每檔 TDCC 都要符合：
1. 資料庫有這檔的商品條件與交易列；
2. 根目錄底下找得到剛好一個商品資料夾，裡面有 `<TDCC>_TS.pdf` 與 `<TDCC>_IIS.pdf`
   （準備與放回工具只會以標準檔名放回核對通過的文件，所以這就代表核對通過）；
3. 交易列上的 FC（略過 PB）都對得到信箱（不分大小寫、別名）。
只有 PB 承作時「不需寄送」；每日匯入略過的非 FCN 商品標「不支援的商品類型」。符合的一檔建一封草稿：收件人是 FC，副本是信箱的固定副本（已在收件人的不重複）。
只建草稿、不寄出；同一檔重跑會再建一封；某一檔出錯只影響那一檔。結果寫成 `寄信草稿清單_<日期時間>.xlsx`，
並記錄用的是哪一次每日匯入。商品資料夾只讀。
"""

from __future__ import annotations

import datetime as dt
import html
import math
import shutil
import tempfile
import time
import unicodedata
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sp_workflow.daily_import import ImportInfo, Mailbox, is_pb, latest_import, load_mailbox
from sp_workflow.database import DatabaseError, load_env_file, require_current_schema
from sp_workflow.preparation import find_product_folders, is_tdcc, parse_tdccs

if TYPE_CHECKING:
    import psycopg

DRAFTED = "已建草稿"
PENDING = "待處理"
NOT_NEEDED = "不需寄送"
FAILED = "建立失敗"

SUBJECT_FORMAT = "{product} {tdcc} 中文TS/IIS"  # 例：FCN 029199990001 中文TS/IIS
BODY_INTRO = "FYI"
RISK_LEVEL = "RR4"  # 使用者確認：風險等級固定 RR4
NAVY = "#002060"
GREEN = "#00AB8E"
ORANGE = "#C55A11"  # 格式檔分公司表頭：佈景主題「輔色 2、較深 25%」
ROW_FILL = "#F0F0F0"
DATE_FORMAT = "%m-%d-%y"
WRAP_FROM_TEXT = "BBG Code 4"  # 使用者指定：比這欄寬的欄位才換行，每格最多兩行
PX_PER_UNIT = 7.5  # Calibri 11pt 粗體每個半形字的約略寬度（全形算 2 個）
FONT_FAMILY = 'font-family: Calibri; mso-fareast-font-family: "Microsoft JhengHei";'
TEMPLATE_COLUMN_UNITS = (
    14.140625,
    11.5703125,
    13.0,
    13.0,
    12.0,
    12.0,
    13.0,
    13.0,
    8.5703125,
    16.140625,
    9.28515625,
    10.42578125,
    10.0,
    7.7109375,
    13.0,
    13.0,
)

LIST_COLUMNS = (("TDCC", 15), ("狀態", 10), ("收件人", 40), ("副本", 50), ("原因", 80))


class MailDraftError(Exception):
    """整批無法建立草稿（資料庫、信箱或路徑有問題）；訊息是給作業人員看的中文。"""


@dataclass(frozen=True)
class Draft:
    tdcc: str
    to: tuple[str, ...]  # 信箱
    cc: tuple[str, ...]
    subject: str
    html_body: str
    attachments: tuple[Path, ...]
    template_path: Path | None = None
    template_product: dict[str, Any] | None = None
    template_branches: tuple[str, ...] = ()


Drafter = Callable[[Draft], None]


@dataclass(frozen=True)
class DraftOutcome:
    tdcc: str
    status: str
    to: tuple[str, ...] = ()  # 人名
    cc: tuple[str, ...] = ()
    reason: str = ""


@dataclass(frozen=True)
class DraftReceipt:
    list_path: Path
    outcomes: tuple[DraftOutcome, ...]
    last_import: ImportInfo | None
    unreadable_dirs: tuple[Path, ...] = ()


MAIL_TEMPLATE_ENV = "SPW_MAIL_TEMPLATE"


def mail_template_from_env(env_path: Path) -> Path | None:
    try:
        values = load_env_file(env_path)
    except DatabaseError as exc:
        raise MailDraftError(str(exc)) from exc
    raw = values.get(MAIL_TEMPLATE_ENV)
    if not raw:
        return None
    path = Path(raw)
    if not path.is_file():
        raise MailDraftError(f"MAIL 草稿格式範本不存在：{path}")
    return path


def create_drafts(
    text: str,
    root: Path,
    out_parent: Path,
    conn: psycopg.Connection,
    *,
    drafter: Drafter | None = None,
    now: dt.datetime | None = None,
    template_path: Path | None = None,
) -> DraftReceipt:
    """依輸入順序處理每個 TDCC；回傳結果清單路徑、逐檔結果與最近一次每日匯入。

    drafter 預設用 Outlook 建草稿（只有真的要建時才會啟動 Outlook）。
    """
    tokens = parse_tdccs(text)
    if not tokens:
        raise MailDraftError("沒有輸入 TDCC")
    if not root.is_dir():
        raise MailDraftError(f"根目錄不存在：{root}")
    if not out_parent.is_dir():
        raise MailDraftError(f"結果清單的位置不存在：{out_parent}")
    try:
        require_current_schema(conn)
    except DatabaseError as exc:
        raise MailDraftError(str(exc)) from exc
    with conn.transaction():
        mailbox = load_mailbox(conn)
    last_import = latest_import(conn)
    fixed_cc = _order_fixed_cc(tuple(p for p in mailbox.people if p.fixed_cc))
    if not fixed_cc:
        raise MailDraftError("信箱沒有任何固定副本（mailbox.person 的 fixed_cc），請先設定")
    stamp = f"{(now or dt.datetime.now()):%Y%m%d-%H%M%S}"
    list_path = out_parent / f"寄信草稿清單_{stamp}.xlsx"
    if list_path.exists():
        raise MailDraftError(f"結果清單已存在：{list_path}")

    valid = {t for t in tokens if is_tdcc(t)}
    folders, unreadable = find_product_folders(root, valid)
    make_draft = drafter or OutlookDrafter()
    outcomes = []
    for tdcc in tokens:
        if tdcc not in valid:
            outcomes.append(DraftOutcome(tdcc, PENDING, reason="不是 12 碼數字"))
            continue
        try:
            outcome = _draft_one(
                tdcc,
                conn,
                mailbox,
                fixed_cc,
                folders.get(tdcc, []),
                unreadable,
                make_draft,
                template_path,
            )
        except Exception as exc:  # 一檔出錯（資料、內文、資料庫）只影響這一檔，結果清單照樣寫出
            outcome = DraftOutcome(tdcc, FAILED, reason=f"處理失敗：{exc}")
        outcomes.append(outcome)
    _write_list(list_path, outcomes, last_import, unreadable)
    return DraftReceipt(list_path, tuple(outcomes), last_import, tuple(unreadable))


def _draft_one(
    tdcc: str,
    conn: psycopg.Connection,
    mailbox: Mailbox,
    fixed_cc: Sequence[Any],
    folders: Sequence[Path],
    unreadable: Sequence[Path],
    make_draft: Drafter,
    template_path: Path | None,
) -> DraftOutcome:
    with conn.transaction():
        product = _product(conn, tdcc)
        branches = [b for (b,) in conn.execute(
            "SELECT branch FROM workflow.branch_allocation WHERE tdcc = %s ORDER BY seq", (tdcc,)
        )]  # fmt: skip
        fc_names = [n for (n,) in conn.execute(
            "SELECT fc_name FROM workflow.trade_row WHERE tdcc = %s ORDER BY row_seq", (tdcc,)
        )]  # fmt: skip
        other = conn.execute("SELECT product FROM workflow.other_product WHERE tdcc = %s", (tdcc,)).fetchone()
    if product is None and not fc_names and other is not None:
        return DraftOutcome(tdcc, PENDING, reason=f"不支援的商品類型：{other[0]}（目前只做 FCN）")
    problems = []
    if product is None:
        problems.append("資料庫沒有這檔的商品條件（還沒匯入 Order 條件，或每日匯入的清單更新還沒有這檔）")
    elif product["issue_date"] is None:
        problems.append("商品資料還沒有發行日（等核對回填）")
    if not fc_names:
        problems.append("資料庫沒有這檔的交易列")
    names = list(dict.fromkeys(n for n in fc_names if not is_pb(n)))
    if fc_names and not names:
        return DraftOutcome(tdcc, NOT_NEEDED, reason="只有 PB 承作，不需寄送")
    people = [mailbox.match(n) for n in names]
    unmatched = [n for n, p in zip(names, people, strict=True) if p is None]
    if unmatched:
        problems.append("名字對不到信箱：" + "、".join(unmatched))
    attachments, folder_problem = _attachments(tdcc, folders, unreadable)
    if folder_problem:
        problems.append(folder_problem)
    if problems:
        return DraftOutcome(tdcc, PENDING, reason="；".join(problems))

    to = list(dict.fromkeys(p for p in people if p is not None))
    cc = [p for p in fixed_cc if p not in to]
    draft = Draft(
        tdcc=tdcc,
        to=tuple(p.email for p in to),
        cc=tuple(p.email for p in cc),
        subject=SUBJECT_FORMAT.format(product=product["product"], tdcc=tdcc),
        html_body=render_body(tdcc, product, branches),
        attachments=attachments,
        template_path=template_path,
        template_product=product,
        template_branches=tuple(branches),
    )
    to_names, cc_names = tuple(p.name for p in to), tuple(p.name for p in cc)
    try:
        make_draft(draft)
    except Exception as exc:  # 一檔失敗不影響其他檔；原因寫進結果
        return DraftOutcome(tdcc, FAILED, to_names, cc_names, f"建立草稿失敗：{exc}")
    return DraftOutcome(tdcc, DRAFTED, to_names, cc_names)


def _order_fixed_cc(people: Sequence[Any]) -> tuple[Any, ...]:
    def key(person: Any) -> tuple[int, str]:
        name = person.name.strip().casefold()
        if name == "max":
            return 0, name
        if name == "bob":
            return 1, name
        if (person.unit or "").strip().casefold() == "pm":
            return 2, name
        return 3, name

    return tuple(sorted(people, key=key))


_PRODUCT_FIELDS = (
    "trade_date", "issue_date", "final_valuation_date", "maturity_date", "initial_pricing", "product", "currency",
    "non_call_months", "underlying_1", "underlying_2", "underlying_3", "underlying_4", "underlying_5", "strike_pct",
    "ko_type", "ko_pct", "coupon_pct", "upfront_pct", "tenor_months", "barrier_type", "ki_pct",
)  # fmt: skip


def _product(conn: psycopg.Connection, tdcc: str) -> dict[str, Any] | None:
    row = conn.execute(f"SELECT {', '.join(_PRODUCT_FIELDS)} FROM workflow.product WHERE tdcc = %s", (tdcc,)).fetchone()
    return None if row is None else dict(zip(_PRODUCT_FIELDS, row, strict=True))


def _attachments(tdcc: str, folders: Sequence[Path], unreadable: Sequence[Path]) -> tuple[tuple[Path, ...], str]:
    if not folders:
        hint = f"（搜尋時有 {len(unreadable)} 個資料夾讀不到，可能在裡面）" if unreadable else ""
        return (), "根目錄底下找不到商品資料夾" + hint
    if len(folders) > 1:
        return (), f"有 {len(folders)} 個同名商品資料夾：" + "、".join(map(str, folders))
    files = tuple(folders[0] / f"{tdcc}_{kind}.pdf" for kind in ("TS", "IIS"))
    missing = [f.name for f in files if not f.is_file()]
    if missing:
        return (), "商品資料夾沒有 " + "、".join(missing) + "（還沒核對通過或還沒放回）"
    return files, ""


# ---------- 信件內容 ----------


def render_body(tdcc: str, product: dict[str, Any], branches: Sequence[str]) -> str:
    """`FYI` 加三個表格（版面與配色依使用者提供的寄送格式）。"""
    p = product
    terms_1 = [
        ("TDCC", tdcc),
        ("Trade date", _date(p["trade_date"])),
        ("Issue date", _date(p["issue_date"])),
        ("Final valuation date", _date(p["final_valuation_date"])),
        ("Maturity date", _date(p["maturity_date"])),
        ("期初價格", p["initial_pricing"] or ""),
        ("風險等級", RISK_LEVEL),
    ]
    terms_2 = [
        ("Product", p["product"]),
        ("Currency", p["currency"]),
        ("Guaranteed Periods (m)", str(p["non_call_months"])),
        *((f"BBG Code {i}", p[f"underlying_{i}"] or "") for i in range(1, 6)),
        ("Strike (%)", _pct(p["strike_pct"])),
        ("KO Type", p["ko_type"]),
        ("KO Barrier (%)", _pct(p["ko_pct"])),
        ("Coupon p.a. (%)", _pct(p["coupon_pct"])),
        ("Upfront / NotePrice (%)", _pct(Decimal(100) - p["upfront_pct"])),
        ("Tenor (m)", str(p["tenor_months"])),
        ("Barrier Type", p["barrier_type"]),
        ("KI Barrier (%)", _pct(p["ki_pct"]) if p["ki_pct"] is not None else "-"),
    ]
    branch_cells = [(f"分公司 {i}", b) for i, b in enumerate(branches, 1)]
    parts = [
        f"<p>{BODY_INTRO}</p>",
        _table(terms_1, NAVY, TEMPLATE_COLUMN_UNITS[: len(terms_1)]),
        _table(terms_2, GREEN, TEMPLATE_COLUMN_UNITS[: len(terms_2)]),
    ]
    if branch_cells:
        parts.append(_table(branch_cells, ORANGE, TEMPLATE_COLUMN_UNITS[: len(branch_cells)]))
    return f'<html><body style="{FONT_FAMILY} font-size: 11pt;">' + "<br>".join(parts) + "</body></html>"


def _table(cells: Sequence[tuple[str, str]], header_color: str, template_units: Sequence[float]) -> str:
    """一列表頭、一列資料。比「BBG Code 4」寬的欄位換行，欄寬取能讓表頭與內容都在兩行內的最小寬度。

    Outlook 不支援 max-width，所以用 width 屬性加 white-space: normal；不需要換行的欄維持一行。
    """
    border = "padding: 2px 6px;"
    head, body = [], []
    widths = []
    for index, (k, v) in enumerate(cells):
        required = _wrap_width(k, v) or 0
        if required == 0:
            required = math.ceil(_text_units(k) * PX_PER_UNIT + 12)
        px = max(math.ceil(template_units[index] * PX_PER_UNIT), math.ceil(required))
        widths.append(px)
        width = f' width="{px}"'
        wrap = f"width: {px}px; white-space: normal;"
        cell_style = f"{border} {wrap} {FONT_FAMILY} text-align: center; vertical-align: middle;"
        head.append(
            f'<th bgcolor="{header_color}" color="#FFFFFF"{width} '
            f'style="{cell_style} background: {header_color}; color: #FFFFFF; font-weight: bold;">'
            f"{_font_html(k, '#FFFFFF')}</th>"
        )
        body.append(
            f'<td bgcolor="{ROW_FILL}"{width} style="{cell_style} background: {ROW_FILL};">{_font_html(v)}</td>'
        )
    table_width = sum(widths)
    return (
        f'<table border="0" cellpadding="2" cellspacing="0" width="{table_width}" '
        f'style="border: 0; border-collapse: collapse; table-layout: fixed; width: {table_width}px; '
        f"mso-table-lspace: 0pt; "
        f'mso-table-rspace: 0pt; {FONT_FAMILY} text-align: center;">'
        f"<tr>{''.join(head)}</tr><tr>{''.join(body)}</tr></table>"
    )


def _wrap_width(*texts: str) -> int | None:
    """需要換行時回傳欄寬（px），不需要回傳 None。欄寬至少是「BBG Code 4」，並讓每段文字最多兩行。"""
    limit = _text_units(WRAP_FROM_TEXT)
    if max(_text_units(t) for t in texts) <= limit:
        return None
    units = max(limit, *(_two_line_units(t) for t in texts))
    return math.ceil(units * PX_PER_UNIT)


def _two_line_units(text: str) -> int:
    """在空白處斷成兩行時，較長那行最短是多少；沒有空白就只能一行。"""
    words = text.split(" ")
    if len(words) == 1:
        return _text_units(text)
    return min(max(_text_units(" ".join(words[:i])), _text_units(" ".join(words[i:]))) for i in range(1, len(words)))


def _font_html(text: str, color: str | None = None) -> str:
    """Use explicit font tags because Outlook's Word renderer ignores some CSS font fallbacks."""
    if not text:
        return ""
    runs: list[tuple[str, str]] = []
    current_font = ""
    current_text = []
    for char in text:
        font = "Microsoft JhengHei" if unicodedata.east_asian_width(char) in "WF" else "Calibri"
        if font != current_font and current_text:
            runs.append((current_font, "".join(current_text)))
            current_text = []
        current_font = font
        current_text.append(char)
    runs.append((current_font, "".join(current_text)))
    color_attr = f' color="{color}"' if color else ""
    return "".join(
        f'<font face="{font}"{color_attr}><span style="font-family: {font};">{html.escape(value)}</span></font>'
        for font, value in runs
    )


def _text_units(text: str) -> int:
    """估計顯示寬度：全形字（中文）算 2，其他算 1。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _date(value: dt.date) -> str:
    return value.strftime(DATE_FORMAT)


def _pct(value: Decimal) -> str:
    return f"{value:.2f}"


# ---------- Outlook ----------


OUTLOOK_FAILURE = "Outlook 無法建立郵件（可能還沒設定 Outlook 帳號、Outlook 啟動時跳出對話框，或附件打不開）"


class OutlookError(Exception):
    """Outlook 建草稿失敗；訊息是給作業人員看的中文，附上 COM 錯誤碼方便維護人員查。"""


class OutlookDrafter:
    """用傳統版 Outlook（COM）在執行者的預設信箱建草稿；只呼叫 Save()，不會寄出。"""

    def __call__(self, draft: Draft) -> None:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()  # 視窗在背景執行緒呼叫，COM 要在這個執行緒初始化
        excel = None
        workbook = None
        temp_dir = None
        try:
            mail = win32com.client.Dispatch("Outlook.Application").CreateItem(0)  # 0 = olMailItem
            mail.To = "; ".join(draft.to)
            mail.CC = "; ".join(draft.cc)
            mail.Subject = draft.subject
            if draft.template_path is None:
                mail.HTMLBody = draft.html_body
            else:
                if draft.template_product is None:
                    raise OutlookError("MAIL 草稿缺少 Excel 範本資料")
                excel, workbook, sheet, temp_dir = self._open_template(
                    win32com.client,
                    draft.template_path,
                    draft.tdcc,
                    draft.template_product,
                    draft.template_branches,
                )
                inspector = mail.GetInspector
                inspector.Display()
                inspector.Activate()
                time.sleep(1)
                self._paste_template(sheet, inspector.WordEditor, len(draft.template_branches))
            for path in draft.attachments:
                mail.Attachments.Add(str(path))
            mail.Save()
        except pythoncom.com_error as exc:  # 原始訊息常是系統語系的亂碼，改成中文說明
            code = exc.args[0] if exc.args else ""
            hresult = f"0x{code & 0xFFFFFFFF:08X}" if isinstance(code, int) else str(code)
            raise OutlookError(f"{OUTLOOK_FAILURE}，錯誤碼 {hresult}") from exc
        finally:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
            if excel is not None:
                excel.Quit()
            if temp_dir is not None:
                shutil.rmtree(temp_dir, ignore_errors=True)
            pythoncom.CoUninitialize()

    @staticmethod
    def _open_template(
        win32com: Any,
        template_path: Path,
        tdcc: str,
        product: dict[str, Any],
        branches: Sequence[str],
    ) -> tuple[Any, Any, Any, str]:
        if len(branches) > 16:
            raise OutlookError("MAIL 格式範本最多支援 16 間分公司")
        excel = win32com.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        temp_dir = tempfile.mkdtemp(prefix="spw-mail-")
        temp_path = Path(temp_dir) / template_path.name
        shutil.copy2(template_path, temp_path)
        workbook = excel.Workbooks.Open(str(temp_path), ReadOnly=False)
        sheet = workbook.Worksheets(1)
        sheet.Columns(5).ColumnWidth = 12.0  # BBG Code 2：比範本略寬，避免表頭擠壓
        values = {
            2: _excel_date(product["trade_date"]),
            3: _excel_date(product["issue_date"]),
            4: _excel_date(product["final_valuation_date"]),
            5: _excel_date(product["maturity_date"]),
            6: product["initial_pricing"] or "",
            7: RISK_LEVEL,
        }
        sheet.Cells(2, 1).NumberFormat = "@"
        sheet.Cells(2, 1).Value = str(tdcc)
        for column, value in values.items():
            sheet.Cells(2, column).Value = value
        terms = (
            product["product"],
            product["currency"],
            product["non_call_months"],
            product["underlying_1"] or "",
            product["underlying_2"] or "",
            product["underlying_3"] or "",
            product["underlying_4"] or "",
            product["underlying_5"] or "",
            product["strike_pct"],
            product["ko_type"],
            product["ko_pct"],
            product["coupon_pct"],
            Decimal(100) - product["upfront_pct"],
            product["tenor_months"],
            product["barrier_type"],
            product["ki_pct"] if product["ki_pct"] is not None else "-",
        )
        for column, value in enumerate(terms, 1):
            sheet.Cells(5, column).Value = float(value) if isinstance(value, Decimal) else value
        sheet.Range("A7:P8").ClearContents()
        for column, branch in enumerate(branches, 1):
            if column > 4:
                letter = get_column_letter(column)
                sheet.Range("D7:D8").Copy(sheet.Range(f"{letter}7"))
            sheet.Cells(7, column).Value = f"分公司 {column}"
            sheet.Cells(8, column).Value = branch
        workbook.Save()
        return excel, workbook, sheet, temp_dir

    @staticmethod
    def _paste_template(sheet: Any, word_editor: Any, branch_count: int) -> None:
        selection = word_editor.Application.Selection
        selection.WholeStory()
        selection.Delete()
        selection.TypeText("FYI")
        selection.TypeParagraph()
        selection.TypeParagraph()
        ranges = ("A1:G2", "A4:P5")
        if branch_count:
            ranges += (f"A7:{get_column_letter(branch_count)}8",)
        for address in ranges:
            sheet.Range(address).Copy()
            selection.PasteExcelTable(False, False, False)
            selection.TypeParagraph()


def _excel_date(value: dt.date) -> dt.datetime:
    return dt.datetime.combine(value, dt.time.min)


# ---------- 結果清單 ----------


def _write_list(
    path: Path, outcomes: Sequence[DraftOutcome], last_import: ImportInfo | None, unreadable: Sequence[Path]
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "寄信草稿清單"
    ws.append([name for name, _ in LIST_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for o in outcomes:
        ws.append([o.tdcc, o.status, "、".join(o.to) or None, "、".join(o.cc) or None, o.reason or None])
        ws.cell(ws.max_row, 1).number_format = "@"  # TDCC 是文字，保留前導零
    for i, (_, width) in enumerate(LIST_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    info = wb.create_sheet("資料來源")
    if last_import is None:
        info.append(["最近一次每日匯入", "沒有匯入紀錄"])
    else:
        info.append(["最近一次每日匯入", f"{last_import.imported_at:%Y-%m-%d %H:%M:%S}"])
        info.append(["匯入編號", last_import.run_id])
        info.append(["每日更新檔", last_import.file_name])
        info.append(["SHA-256", last_import.sha256])
    info.append(["讀不到的資料夾", len(unreadable)])
    for folder in unreadable:
        info.append([None, str(folder)])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 70
    wb.save(path)
